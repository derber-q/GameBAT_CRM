"""Партнёры, подписанный прайс, взаимные долги и единый финансовый отчёт."""
from decimal import Decimal
from io import BytesIO
import uuid
from unittest.mock import patch
from django.contrib.auth.models import Permission
from django.core import signing
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import IntegrityError, transaction
from django.test import TestCase
from django.urls import reverse
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from accounts.models import User
from cash.models import CashTransaction
from catalog.models import CD, Platform
from creditors.models import Creditor, CreditorTransaction
from creditors.services import adjust_creditor_debt
from warehouse.models import Warehouse, CDWarehouseStock
from price.excel import generate_wholesale_price_xlsx, import_wholesale_price_to_sale, META_SHEET, MAIN_SHEET, HEADER_ROW
from sales.models import Sale
from sales.services import create_sale, mark_sale_paid, cancel_sale, advance_order_status
from sales.partner_payments import quote, SALT
from sales.statistics_service import build_sales_statistics_report
from sales.statistics_excel import render_statistics_workbook
from .models import Partner, SalesPlatform


class PartnerTests(TestCase):
    def setUp(self):
        self.actor = User.objects.create_superuser('partner-admin')
        self.warehouse = Warehouse.objects.create(name='Склад')
        self.partner = Partner.objects.create(name='Игромир', contact_person='Иван', phone_1='00123', phone_2='00456', address='Адрес')
        self.key = uuid.uuid4()
        self.product = CD.objects.create(name='Игра', platform=Platform.objects.create(name='PS5'), cost=6000, wholesale_price=10000, avito_price=11000)
        self.stock = CDWarehouseStock.objects.create(warehouse=self.warehouse, cd=self.product, quantity=10)

    def debt(self, amount):
        return adjust_creditor_debt(actor=self.actor, creditor_id=self.partner.creditor_id,
            warehouse_id=self.warehouse.pk, money_source_type='offset', new_debt=amount, comment='Исходный долг')

    def token(self, received, *, sale=None):
        self.partner.creditor.refresh_from_db()
        return signing.dumps(quote(total=Decimal('10000'), received=received,
            debt=self.partner.creditor.current_debt, partner_id=self.partner.pk,
            sale_id=sale.pk if sale else None, creation_key=None if sale else self.key), salt=SALT)

    def sale(self, **kwargs):
        values = dict(actor=self.actor, warehouse_id=self.warehouse.pk, price_type='wholesale', sale_type='wholesale',
            payment_method='cash', partner_id=self.partner.pk, creation_key=self.key,
            note='Расчёт с партнёром', lines=[{'product_type': 'cd', 'product_id': self.product.pk, 'quantity': 1}])
        values.update(kwargs)
        return create_sale(**values)

    def test_partner_account_atomic_unique_and_platform_untouched(self):
        self.assertEqual(self.partner.creditor.current_debt, 0)
        account_id = self.partner.creditor_id
        self.partner.name = 'Новое имя'; self.partner.save(); self.partner.save()
        self.assertEqual(self.partner.creditor_id, account_id)
        self.assertEqual(Creditor.objects.count(), 1)
        platform = SalesPlatform.objects.create(name='Площадка', address='Адрес', legal_entity='ИП', phone_1='123')
        self.partner.platform = platform; self.partner.save()
        with self.assertRaises(IntegrityError), transaction.atomic():
            Partner.objects.create(name='Второй', platform=platform)
        self.assertEqual(Creditor.objects.count(), 1)
        with self.assertRaises(ValidationError): Partner.objects.create(name='   ')
        with patch('creditors.models.Creditor.objects.create', side_effect=RuntimeError('audit')):
            with self.assertRaises(RuntimeError): Partner.objects.create(name='Ошибка')
        self.assertEqual(Partner.objects.count(), 1)
        self.stock.refresh_from_db(); self.assertEqual(self.stock.quantity, 10)

    def test_contact_name_labels_use_partner_ids_and_preserve_document_name(self):
        from .customers import CustomerSelectionForm, resolve_customer
        from sales.statistics_forms import SaleStatisticsFilterForm
        from sales.partner_payments import payment_partner_choices
        from orders.trip_addition_views import AdditionPartnersForm

        label = 'Игромир (Иван)'
        self.assertEqual(self.partner.display_name, label)
        self.assertEqual(str(self.partner), 'Игромир')
        self.assertEqual(self.partner.creditor.display_name, label)
        self.assertEqual(resolve_customer(partner_id=self.partner.pk)[1], 'Игромир')
        self.assertIn((str(self.partner.pk), label), CustomerSelectionForm().fields['customer'].choices)
        for form, field in [(SaleStatisticsFilterForm(), 'partner'), (AdditionPartnersForm(), 'partners')]:
            self.assertIn((str(self.partner.pk), label), [(str(value), text) for value, text in form.fields[field].choices])
        self.assertEqual(list(payment_partner_choices())[0].display_name, label)
        sale = self.sale(payment_method='cash_postpay')
        self.assertEqual(sale.partner_name_snapshot, 'Игромир')

        self.client.force_login(self.actor)
        for endpoint in ['creditors:list', 'sales:create', 'sales:incomplete', 'resource_storefront:contacts']:
            with self.subTest(endpoint=endpoint):
                self.assertContains(self.client.get(reverse(endpoint)), label)
        self.partner.is_archived = True
        self.partner.save()
        self.assertNotIn(str(self.partner.pk), dict(CustomerSelectionForm().fields['customer'].choices))
        archived = CustomerSelectionForm(allow_archived_initial=True, initial={'customer': str(self.partner.pk)})
        self.assertEqual(dict(archived.fields['customer'].choices)[str(self.partner.pk)], label + ' (в архиве)')

    def test_identifying_name_is_escaped_and_does_not_expose_other_contacts(self):
        from .customers import CustomerSelectionForm

        self.assertEqual(Partner(name='Без контакта', contact_person='  ').display_name, 'Без контакта')
        self.partner.contact_person = '<Иван & Пётр>'
        self.partner.save()
        self.assertIn('Игромир (&lt;Иван &amp; Пётр&gt;)', str(CustomerSelectionForm()['customer']))
        viewer = User.objects.create_user('creditor-name-viewer')
        viewer.user_permissions.add(Permission.objects.get(codename='view_creditors', content_type__app_label='creditors'))
        self.client.force_login(viewer)
        response = self.client.get(reverse('creditors:list'))
        self.assertContains(response, 'Игромир (&lt;Иван &amp; Пётр&gt;)')
        self.assertNotContains(response, self.partner.phone_1)
        self.assertNotContains(response, self.partner.phone_2)

    def test_wholesale_requires_explicit_customer_and_custom_has_no_account(self):
        with self.assertRaisesMessage(ValidationError, 'заказчика'): self.sale(partner_id=None)
        sale = self.sale(partner_id=None, partner_name='Вручную')
        self.assertIsNone(sale.partner_id); self.assertEqual(sale.partner_name_snapshot, 'Вручную')
        self.assertEqual(Partner.objects.count(), 1); self.assertEqual(Creditor.objects.count(), 1)

    def test_repayment_and_overpayment_do_not_duplicate_cash_or_profit(self):
        self.debt('2000')
        sale = self.sale(cash_received_amount='13000', settlement_confirmation=self.token('13000'))
        self.partner.creditor.refresh_from_db(); self.warehouse.cash_register.refresh_from_db()
        self.assertEqual((sale.total_amount, sale.cash_received_amount, sale.debt_repayment_amount, sale.extra_cash_amount),
            (Decimal('10000'), Decimal('13000'), Decimal('2000'), Decimal('1000')))
        self.assertEqual(self.partner.creditor.current_debt, 0)
        self.assertEqual(self.warehouse.cash_register.balance, 13000)
        self.assertEqual(sale.cash_transactions.count(), 1)
        report = build_sales_statistics_report({'partner': self.partner})
        self.assertEqual(report.totals.profit, 5000)
        self.assertEqual(report.totals.actual_revenue, 11000)
        self.assertEqual(report.totals.debt_repayment, 2000)
        self.assertEqual(report.totals.actual_money, 13000)
        self.assertEqual(self.sale(cash_received_amount='13000', settlement_confirmation='expired').pk, sale.pk)
        self.assertEqual(sale.creditor_movements.count(), 1)

    def test_offset_and_cancellation_restore_debt_and_real_cash_once(self):
        self.debt('-3000')
        sale = self.sale(cash_received_amount='8000', settlement_confirmation=self.token('8000'))
        self.partner.creditor.refresh_from_db(); self.warehouse.cash_register.refresh_from_db()
        self.assertEqual((self.partner.creditor.current_debt, self.warehouse.cash_register.balance), (-1000, 8000))
        self.assertEqual(sale.debt_offset_amount, 2000)
        cancelled = cancel_sale(actor=self.actor, sale_id=sale.pk, comment='Возврат')
        self.assertTrue(cancelled.cancelled)
        self.assertFalse(cancel_sale(actor=self.actor, sale_id=sale.pk, comment='Повтор').cancelled)
        self.partner.creditor.refresh_from_db(); self.warehouse.cash_register.refresh_from_db(); self.stock.refresh_from_db()
        self.assertEqual((self.partner.creditor.current_debt, self.warehouse.cash_register.balance, self.stock.quantity), (-3000, 0, 10))
        self.assertEqual(sale.creditor_movements.count(), 2)
        report = build_sales_statistics_report({'partner': self.partner})
        self.assertEqual((report.totals.sales_count, report.totals.profit, report.totals.debt_offset), (0, 0, 0))

    def test_confirmation_note_and_stale_balance_are_rejected(self):
        self.debt('-1000')
        with self.assertRaisesMessage(ValidationError, 'Подтвердите'): self.sale(cash_received_amount='8000')
        self.debt('-3000')
        token = self.token('8000')
        with self.assertRaisesMessage(ValidationError, 'Подтвердите'): self.sale(cash_received_amount='8000')
        with self.assertRaisesMessage(ValidationError, 'примечание'): self.sale(cash_received_amount='8000', settlement_confirmation=token, note='')
        self.debt('-2500')
        with self.assertRaisesMessage(ValidationError, 'изменились'): self.sale(cash_received_amount='8000', settlement_confirmation=token)
        self.assertFalse(Sale.objects.exists()); self.stock.refresh_from_db(); self.assertEqual(self.stock.quantity, 10)

    def test_postpay_uses_current_debt_only_when_payment_is_confirmed(self):
        self.debt('3000')
        sale = self.sale(payment_method='bank_transfer_postpay', cash_received_amount='12000')
        self.assertEqual(sale.creditor_movements.count(), 0)
        mark_sale_paid(actor=self.actor, sale_id=sale.pk, cash_received_amount='12000', settlement_confirmation=self.token('12000', sale=sale))
        with self.assertRaises(ValidationError): mark_sale_paid(actor=self.actor, sale_id=sale.pk)
        self.assertFalse(sale.cash_transactions.exists())
        self.partner.creditor.refresh_from_db(); self.assertEqual(self.partner.creditor.current_debt, 1000)
        report = build_sales_statistics_report({'partner': self.partner})
        self.assertEqual(report.totals.sales_count, 0)
        self.assertEqual(report.totals.debt_repayment, 2000)

    def test_permissions_and_preview(self):
        self.debt('3000')
        user = User.objects.create_user('seller')
        user.user_permissions.add(Permission.objects.get(codename='create_sale', content_type__app_label='sales'))
        with self.assertRaises(PermissionDenied): self.sale(actor=user, cash_received_amount='12000', settlement_confirmation=self.token('12000'))
        self.client.force_login(self.actor)
        response = self.client.post(reverse('sales:settlement_preview'), {'partner_id': self.partner.pk, 'total': '10000', 'received': '12000', 'creation_key': str(self.key)})
        self.assertEqual(response.status_code, 200); self.assertTrue(response.json()['required'])
        self.assertEqual(self.client.post(reverse('sales:settlement_preview'), {'partner_id': 'wrong'}).status_code, 400)
        self.client.force_login(user)
        self.assertEqual(self.client.get(reverse('partners:partner_detail', args=[self.partner.pk])).status_code, 403)

    def workbook(self):
        content = generate_wholesale_price_xlsx(warehouse_id=self.warehouse.pk, actor=self.actor, partner_id=self.partner.pk)
        book = load_workbook(BytesIO(content))
        sheet = book[MAIN_SHEET]
        for row in range(HEADER_ROW + 1, sheet.max_row + 1):
            if sheet.cell(row, 2).value == 10000: sheet.cell(row, 4, 1)
        return book

    def upload(self, book):
        output = BytesIO(); book.save(output); book.close()
        return SimpleUploadedFile('renamed.xlsx', output.getvalue())

    def test_excel_round_trip_tamper_legacy_and_name_snapshot(self):
        book = self.workbook()
        _, lines, customer = import_wholesale_price_to_sale(self.upload(book), with_customer=True)
        self.assertEqual(customer['partner_id'], self.partner.pk); self.assertEqual(lines[0]['quantity'], 1)
        book = self.workbook()
        for row in book[META_SHEET]:
            if row[0].value == 'customer_token': row[1].value = str(row[1].value) + 'bad'
        with self.assertRaisesMessage(ValidationError, 'повреждены'): import_wholesale_price_to_sale(self.upload(book), with_customer=True)
        book = self.workbook()
        for row in reversed(list(book[META_SHEET])):
            if row[0].value in ('customer_token', 'customer_required'): book[META_SHEET].delete_rows(row[0].row)
        _, _, customer = import_wholesale_price_to_sale(self.upload(book), with_customer=True)
        self.assertIsNone(customer)
        sale = self.sale()
        self.partner.name = 'Переименован'; self.partner.save()
        sale.refresh_from_db(); self.assertEqual(sale.partner_name_snapshot, 'Игромир')
        report = build_sales_statistics_report({'partner': self.partner})
        workbook = load_workbook(BytesIO(render_statistics_workbook(report)))
        self.assertIn('Все операции', workbook.sheetnames)
        self.assertIn('Игромир', [cell.value for row in workbook['Все операции'] for cell in row])

    def test_partner_form_repeated_post_and_contact_visibility(self):
        self.client.force_login(self.actor)
        response = self.client.get(reverse('partners:partner_create'))
        key = response.context['form']['request_key'].value()
        payload = {'request_key': str(key), 'name': 'Новый партнёр', 'contact_person': 'Скрытый контакт', 'phone_1': '0011'}
        first = self.client.post(reverse('partners:partner_create'), payload)
        second = self.client.post(reverse('partners:partner_create'), payload)
        self.assertEqual((first.status_code, second.status_code), (302, 302))
        self.assertEqual(Partner.objects.filter(name='Новый партнёр').count(), 1)
        self.assertEqual(Creditor.objects.count(), 2)
        viewer = User.objects.create_user('partner-viewer')
        viewer.user_permissions.add(Permission.objects.get(codename='view_partner', content_type__app_label='partners'))
        self.client.force_login(viewer)
        response = self.client.get(reverse('partners:partner_list'))
        self.assertContains(response, 'Новый партнёр'); self.assertNotContains(response, 'Скрытый контакт')

    def test_full_offset_has_no_zero_cash_entry_and_bank_never_changes_cash_or_safe(self):
        self.debt('-10000')
        sale = self.sale(cash_received_amount='0', settlement_confirmation=self.token('0'))
        self.assertFalse(sale.cash_transactions.exists())
        cancel_sale(actor=self.actor, sale_id=sale.pk, comment='Отмена полного зачёта')
        self.assertFalse(sale.cash_transactions.exists())
        self.partner.creditor.refresh_from_db(); self.assertEqual(self.partner.creditor.current_debt, -10000)
        adjust_creditor_debt(actor=self.actor, creditor_id=self.partner.creditor_id,
            warehouse_id=self.warehouse.pk, money_source_type='bank', new_debt='2000', comment='Банк')
        self.warehouse.cash_register.refresh_from_db(); self.warehouse.safe.refresh_from_db()
        self.assertEqual((self.warehouse.cash_register.balance, self.warehouse.safe.balance), (0, 0))
        self.assertEqual(CashTransaction.objects.count(), 0)

    def test_legacy_import_retains_lines_until_manual_customer_choice(self):
        book = self.workbook()
        for row in reversed(list(book[META_SHEET])):
            if row[0].value in ('customer_token', 'customer_required'): book[META_SHEET].delete_rows(row[0].row)
        self.client.force_login(self.actor)
        response = self.client.post(reverse('sales:import_wholesale'), {'file': self.upload(book)})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'без повторной загрузки')
        response = self.client.post(reverse('sales:import_wholesale'), {'customer': str(self.partner.pk)})
        self.assertRedirects(response, reverse('sales:create'))
        response = self.client.post(reverse('sales:create'), {'payment_method': 'bank_account',
            'product_type': ['cd'], 'product_id': [self.product.pk], 'quantity': [1]})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Sale.objects.get().partner_id, self.partner.pk)

    def test_arbitrary_excel_and_production_filename(self):
        content = generate_wholesale_price_xlsx(warehouse_id=self.warehouse.pk, actor=self.actor, partner_name='Без счёта')
        book = load_workbook(BytesIO(content))
        sheet = book[MAIN_SHEET]
        for row in range(HEADER_ROW + 1, sheet.max_row + 1):
            if sheet.cell(row, 2).value == 10000: sheet.cell(row, 4, 1)
        _, _, customer = import_wholesale_price_to_sale(self.upload(book), with_customer=True)
        self.assertEqual(customer, {'partner_id': None, 'partner_name': 'Без счёта'})
        self.assertEqual(Partner.objects.count(), 1)
        self.client.force_login(self.actor)
        response = self.client.post(reverse('price:wholesale_export'), {'warehouse': self.warehouse.pk,
            'customer': str(self.partner.pk)})
        from urllib.parse import quote as urlquote
        self.assertIn(urlquote(self.partner.name), response['Content-Disposition'])

    def test_anonymous_wholesale_export_keeps_signed_empty_customer(self):
        from partners.excel import read_customer_metadata
        self.client.force_login(self.actor)
        response = self.client.post(reverse('price:wholesale_export'), {'warehouse': self.warehouse.pk})
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content))
        values = {
            row[0].value: str(row[1].value or '')
            for row in book[META_SHEET] if isinstance(row[0].value, str)
        }
        self.assertTrue(values.get('customer_token'))
        self.assertIsNone(read_customer_metadata(values))
        book.close()
        from urllib.parse import quote as urlquote
        self.assertNotIn(urlquote(self.partner.name), response['Content-Disposition'])
        self.assertNotIn(urlquote('Без_заказчика'), response['Content-Disposition'])
        values['customer_token'] += 'broken'
        with self.assertRaisesMessage(ValidationError, 'повреждены'):
            read_customer_metadata(values)

    def test_anonymous_wholesale_import_requires_customer_then_continues_without_reupload(self):
        content = generate_wholesale_price_xlsx(warehouse_id=self.warehouse.pk, actor=self.actor)
        book = load_workbook(BytesIO(content))
        sheet = book[MAIN_SHEET]
        for row in range(HEADER_ROW + 1, sheet.max_row + 1):
            if sheet.cell(row, 2).value == 10000:
                sheet.cell(row, 4, 2)
        self.client.force_login(self.actor)
        session = self.client.session
        session['sale_wholesale_import_draft'] = {
            'warehouse_id': self.warehouse.pk, 'partner_id': self.partner.pk, 'lines': [],
        }
        session.save()
        response = self.client.post(reverse('sales:import_wholesale'), {'file': self.upload(book)})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Выберите заказчика')
        self.assertIn('wholesale_customer_pending', self.client.session)
        self.assertNotIn('sale_wholesale_import_draft', self.client.session)
        self.assertFalse(Sale.objects.exists())
        response = self.client.post(reverse('sales:import_wholesale'), {})
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Sale.objects.exists())
        response = self.client.post(reverse('sales:import_wholesale'), {'customer': str(self.partner.pk)})
        self.assertRedirects(response, reverse('sales:create'))
        self.assertNotIn('wholesale_customer_pending', self.client.session)
        draft = self.client.session['sale_wholesale_import_draft']
        self.assertEqual(draft['partner_id'], self.partner.pk)
        self.assertEqual(draft['lines'][0]['quantity'], 2)
        page = self.client.get(reverse('sales:create'))
        self.assertEqual(page.context['form'].initial['customer'], str(self.partner.pk))
        self.assertTrue(page.context['form'].fields['customer'].disabled)
        self.stock.refresh_from_db()
        self.assertEqual(self.stock.quantity, 10)
        self.assertFalse(Sale.objects.exists())

    def test_standalone_operations_filter_date_and_excel_share_one_report(self):
        entry = self.debt('3000')
        adjust_creditor_debt(actor=self.actor, creditor_id=self.partner.creditor_id, warehouse_id=self.warehouse.pk,
            money_source_type='bank', new_debt='2000', comment='Получен банковский перевод')
        report = build_sales_statistics_report({'partner': self.partner})
        self.assertEqual((report.totals.profit, report.totals.sales_count, report.totals.actual_money, report.totals.debt_repayment), (0, 0, 1000, 1000))
        self.assertEqual(len(report.operations), 2)
        from datetime import timedelta
        from django.utils import timezone
        report = build_sales_statistics_report({'partner': self.partner, 'date_from': timezone.localdate() + timedelta(days=1)})
        self.assertEqual(report.operations, [])

    def test_import_order_partner_and_invoice_filename_preserve_snapshots(self):
        from partners.models import Supplier
        from pricing.models import SupplierCDPrice
        from price.services import create_procurement_price_list
        from orders.services import create_customer_procurement_order
        supplier = Supplier.objects.create(name='Поставщик', letter='A', highlight_color='#123456', legal_entity='ИП', phone_1='123')
        SupplierCDPrice.objects.create(supplier=supplier, cd=self.product, price=100)
        price = create_procurement_price_list(actor=self.actor, exchange_rate_usdt_aed=4, exchange_rate_usdt_rub=80,
            delivery_general_usdt_per_kg=0, delivery_devices_usdt_per_kg=0, delivery_storage_usdt_per_kg=0, delivery_cd_rub=0)
        order = create_customer_procurement_order(actor=self.actor, price_list_id=price.pk,
            recipient='Получатель', comment='', partner_id=self.partner.pk,
            lines=[{'price_list_item_id': price.items.get().pk, 'prepayment_quantity': 1, 'postpayment_quantity': 0}])
        self.assertEqual((order.partner_id, order.partner_name_snapshot), (self.partner.pk, 'Игромир'))
        self.client.force_login(self.actor)
        response = self.client.get(reverse('orders:export', args=[order.pk]))
        from urllib.parse import quote as urlquote
        self.assertIn(urlquote(self.partner.name), response['Content-Disposition'])
        self.assertIn(str(order.pk), response['Content-Disposition'])
        self.assertIn('Игромир', self.client.get(reverse('orders:detail', args=[order.pk])).content.decode())

    def test_cash_verification_counts_creditor_cash_and_safe_once(self):
        from cash.services import deposit_cash, deposit_safe
        from django.core.management import call_command
        deposit_cash(actor=self.actor, warehouse_id=self.warehouse.pk, amount='5000', comment='Начальная касса')
        deposit_safe(actor=self.actor, warehouse_id=self.warehouse.pk, amount='5000', comment='Начальный сейф')
        adjust_creditor_debt(actor=self.actor, creditor_id=self.partner.creditor_id, warehouse_id=self.warehouse.pk,
            money_source_type='cash', new_debt='1000', comment='Выдано из кассы')
        adjust_creditor_debt(actor=self.actor, creditor_id=self.partner.creditor_id, warehouse_id=self.warehouse.pk,
            money_source_type='safe', new_debt='2000', comment='Выдано из сейфа')
        adjust_creditor_debt(actor=self.actor, creditor_id=self.partner.creditor_id, warehouse_id=self.warehouse.pk,
            money_source_type='safe', new_debt='1500', comment='Возврат в сейф')
        from io import StringIO
        call_command('verify_cash', stdout=StringIO())

    def test_new_renamed_excel_assigns_partner_through_import_and_sale_form(self):
        self.client.force_login(self.actor)
        response = self.client.post(reverse('sales:import_wholesale'), {'file': self.upload(self.workbook())})
        self.assertRedirects(response, reverse('sales:create'))
        response = self.client.post(reverse('sales:create'), {'payment_method': 'bank_account',
            'product_type': ['cd'], 'product_id': [self.product.pk], 'quantity': [1]})
        self.assertEqual(response.status_code, 302)
        sale = Sale.objects.get()
        self.assertEqual((sale.partner_id, sale.partner_name_snapshot, sale.total_amount), (self.partner.pk, self.partner.name, 10000))
        self.assertEqual((Partner.objects.count(), Creditor.objects.count()), (1, 1))
