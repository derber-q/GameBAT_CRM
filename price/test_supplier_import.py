from io import BytesIO
from decimal import Decimal

from django.contrib.auth.models import Permission
from django.core.exceptions import ValidationError
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import User
from catalog.models import Brand, CD, Platform, ProductType, Tech
from partners.models import Supplier
from pricing.models import SupplierCDPrice, SupplierTechPrice

from .excel import (
    MAIN_SHEET, META_SHEET, generate_supplier_template, import_supplier_price,
    import_supplier_price_files,
)
from .tests import uploaded, workbook_bytes


class SupplierTemplateImportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_superuser("supplier-admin", password="StrongAdmin!123")
        cls.supplier = Supplier.objects.create(
            name="M", letter="M", highlight_color="#37A7BA", legal_entity="M", phone_1="1",
        )
        cls.cd = CD.objects.create(platform=Platform.objects.create(name="PS5"), name="Игра", sku="CD-1")
        cls.tech = Tech.objects.create(
            brand=Brand.objects.create(name="Sony"),
            product_type=ProductType.objects.create(name="Консоли"), name="Консоль", sku="TECH-1",
        )

    def template(self, kind):
        return load_workbook(BytesIO(generate_supplier_template(actor=self.user, product_kind=kind)))

    def fill(self, workbook, kind, product_id, price):
        sheet = workbook[MAIN_SHEET]
        row = next(
            row for row in range(2, sheet.max_row + 1)
            if sheet.cell(row, 1).value == kind.upper() and sheet.cell(row, 2).value == product_id
        )
        sheet.cell(row, 6, price)
        return row

    def import_file(self, workbook, kind):
        return import_supplier_price(
            supplier_id=self.supplier.pk, upload=uploaded(workbook_bytes(workbook)),
            actor=self.user, product_kind=kind,
        )

    def test_old_files_update_both_types_and_preserve_missing_new_products(self):
        workbooks = {kind: self.template(kind) for kind in ("cd", "tech")}
        self.fill(workbooks["cd"], "cd", self.cd.pk, 100.25)
        self.fill(workbooks["tech"], "tech", self.tech.pk, 200)
        new_cd = CD.objects.create(platform=self.cd.platform, name="Новая игра", sku="CD-2")
        new_tech = Tech.objects.create(
            brand=self.tech.brand, product_type=self.tech.product_type, name="Новая консоль", sku="TECH-2",
        )
        SupplierCDPrice.objects.create(supplier=self.supplier, cd=new_cd, price=77)
        SupplierTechPrice.objects.create(supplier=self.supplier, tech=new_tech, price=88)
        self.client.force_login(self.user)
        response = self.client.post(reverse("price:supplier_upload"), {
            "supplier": self.supplier.pk,
            **{f"{kind}_file": uploaded(workbook_bytes(book), f"M-{kind}.xlsx") for kind, book in workbooks.items()},
        })
        self.assertRedirects(response, reverse("price:imports"))
        self.assertEqual(SupplierCDPrice.objects.get(supplier=self.supplier, cd=self.cd).price, Decimal("100.25"))
        self.assertEqual(SupplierTechPrice.objects.get(supplier=self.supplier, tech=self.tech).price, Decimal("200"))
        self.assertEqual(SupplierCDPrice.objects.get(supplier=self.supplier, cd=new_cd).price, Decimal("77"))
        self.assertEqual(SupplierTechPrice.objects.get(supplier=self.supplier, tech=new_tech).price, Decimal("88"))

    def test_archived_rows_with_or_without_price_are_skipped(self):
        for kind, product, price_model, field in (
            ("cd", self.cd, SupplierCDPrice, "cd"), ("tech", self.tech, SupplierTechPrice, "tech"),
        ):
            for price in (None, 99):
                with self.subTest(kind=kind, price=price):
                    product.is_archived = False
                    product.save(update_fields=("is_archived",))
                    workbook = self.template(kind)
                    if price is not None:
                        self.fill(workbook, kind, product.pk, price)
                    offer, _ = price_model.objects.get_or_create(
                        supplier=self.supplier, **{field: product}, defaults={"price": 77},
                    )
                    product.is_archived = True
                    product.save(update_fields=("is_archived",))
                    self.assertEqual(self.import_file(workbook, kind), 0)
                    offer.refresh_from_db()
                    self.assertEqual(offer.price, Decimal("77"))

    def test_blank_price_removes_only_offer_present_in_file(self):
        workbook = self.template("cd")
        SupplierCDPrice.objects.create(supplier=self.supplier, cd=self.cd, price=77)
        SupplierTechPrice.objects.create(supplier=self.supplier, tech=self.tech, price=88)
        self.assertEqual(self.import_file(workbook, "cd"), 0)
        self.assertFalse(SupplierCDPrice.objects.filter(supplier=self.supplier).exists())
        self.assertEqual(SupplierTechPrice.objects.get(supplier=self.supplier).price, Decimal("88"))

    def price_cells(self, workbook):
        sheet = workbook[MAIN_SHEET]
        return {
            (sheet.cell(row, 1).value.lower(), sheet.cell(row, 2).value): sheet.cell(row, 6)
            for row in range(2, sheet.max_row + 1)
            if sheet.cell(row, 1).value in ("CD", "TECH")
        }

    def test_prefilled_templates_use_only_selected_supplier_and_include_new_products(self):
        other = Supplier.objects.create(
            name="F", letter="F", highlight_color="#37A7BA", legal_entity="F", phone_1="2",
        )
        for kind, product, model, field in (
            ("cd", self.cd, SupplierCDPrice, "cd"),
            ("tech", self.tech, SupplierTechPrice, "tech"),
        ):
            model.objects.create(supplier=self.supplier, **{field: product}, price=Decimal("77.123456"))
            model.objects.create(supplier=other, **{field: product}, price=Decimal("99.654321"))
            new = type(product).objects.create(**{
                **({"platform": product.platform} if kind == "cd" else {
                    "brand": product.brand, "product_type": product.product_type,
                }), "name": "Новый товар", "sku": f"{kind}-new",
            })
            for supplier, expected in ((self.supplier, "77.123456"), (other, "99.654321")):
                with self.subTest(kind=kind, supplier=supplier.letter):
                    content = generate_supplier_template(
                        actor=self.user, product_kind=kind, supplier_id=supplier.pk,
                    )
                    workbook = load_workbook(BytesIO(content))
                    cells = self.price_cells(workbook)
                    self.assertEqual(set(cells), {(kind, product.pk), (kind, new.pk)})
                    self.assertEqual(Decimal(str(cells[kind, product.pk].value)), Decimal(expected))
                    self.assertFalse(cells[kind, product.pk].protection.locked)
                    self.assertIsNone(cells[kind, new.pk].value)
                    workbook.close()
            blank = self.template(kind)
            self.assertTrue(all(cell.value is None for cell in self.price_cells(blank).values()))
            blank.close()

    def test_cleared_prefilled_prices_stay_blank_in_next_download_for_both_types(self):
        books = {}
        for kind, product, model, field in (
            ("cd", self.cd, SupplierCDPrice, "cd"),
            ("tech", self.tech, SupplierTechPrice, "tech"),
        ):
            model.objects.create(supplier=self.supplier, **{field: product}, price=Decimal("88.25"))
            books[kind] = load_workbook(BytesIO(generate_supplier_template(
                actor=self.user, product_kind=kind, supplier_id=self.supplier.pk,
            )))
            cell = self.price_cells(books[kind])[kind, product.pk]
            self.assertEqual(cell.value, 88.25)
            cell.value = None
        self.client.force_login(self.user)
        response = self.client.post(reverse("price:supplier_upload"), {
            "supplier": self.supplier.pk,
            **{f"{kind}_file": uploaded(workbook_bytes(book)) for kind, book in books.items()},
        })
        self.assertRedirects(response, reverse("price:imports"))
        self.assertFalse(SupplierCDPrice.objects.filter(supplier=self.supplier).exists())
        self.assertFalse(SupplierTechPrice.objects.filter(supplier=self.supplier).exists())
        for kind in ("cd", "tech"):
            workbook = load_workbook(BytesIO(generate_supplier_template(
                actor=self.user, product_kind=kind, supplier_id=self.supplier.pk,
            )))
            self.assertTrue(all(cell.value is None for cell in self.price_cells(workbook).values()))
            self.assertTrue(self.price_cells(workbook))
            workbook.close()

    def test_prefilled_download_requires_both_permissions_and_masks_supplier_details(self):
        user = User.objects.create_user("template-reader", password="StrongReader!123")
        self.supplier.name = "Закрытое название поставщика"
        self.supplier.save(update_fields=("name",))
        url = reverse("price:supplier_template_prefilled", args=(self.supplier.pk, "cd"))
        permissions = {
            code: Permission.objects.get(content_type__app_label="price", codename=code)
            for code in ("view_price_page", "download_supplier_template", "view_supplier_prices")
        }
        for granted in ((), ("download_supplier_template",), ("view_supplier_prices",),
                        ("download_supplier_template", "view_supplier_prices")):
            with self.subTest(granted=granted):
                user.user_permissions.set([permissions["view_price_page"], *[permissions[p] for p in granted]])
                self.client.force_login(user)
                allowed = len(granted) == 2
                response = self.client.get(url)
                self.assertEqual(response.status_code, 200 if allowed else 403)
                page = self.client.get(reverse("price:imports"))
                self.assertEqual(page.status_code, 200)
                self.assertNotContains(page, self.supplier.name)
                if allowed:
                    self.assertContains(page, url)
                    self.assertContains(page, self.supplier.safe_label)
                else:
                    self.assertNotContains(page, url)

    def test_prefilled_download_is_read_only_and_uses_supplier_in_filename(self):
        SupplierCDPrice.objects.create(supplier=self.supplier, cd=self.cd, price=Decimal("123.456789"))
        self.client.force_login(self.user)
        url = reverse("price:supplier_template_prefilled", args=(self.supplier.pk, "cd"))
        with CaptureQueriesContext(connection) as queries:
            response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertIn("resource-supplier-M-cd-template.xlsx", response["Content-Disposition"])
        self.assertIn("no-store", response["Cache-Control"])
        self.assertFalse(any(
            query["sql"].lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE"))
            for query in queries.captured_queries
        ))
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(self.price_cells(workbook)["cd", self.cd.pk].value, 123.456789)
        workbook.close()
        self.assertEqual(self.client.post(url).status_code, 405)
        self.assertEqual(self.client.get(reverse(
            "price:supplier_template_prefilled", args=(999999, "cd"),
        )).status_code, 404)
        self.assertEqual(self.client.get(reverse(
            "price:supplier_template_prefilled", args=(self.supplier.pk, "other"),
        )).status_code, 404)

    def test_prefilled_templates_skip_excluded_archived_and_zero_prices(self):
        for kind, product, model, field in (
            ("cd", self.cd, SupplierCDPrice, "cd"),
            ("tech", self.tech, SupplierTechPrice, "tech"),
        ):
            model.objects.create(supplier=self.supplier, **{field: product}, price=0)
            base = {"platform": product.platform} if kind == "cd" else {
                "brand": product.brand, "product_type": product.product_type,
            }
            for flag in ("is_archived", "exclude_from_supplier_template"):
                excluded = type(product).objects.create(**base, name=flag, **{flag: True})
                model.objects.create(supplier=self.supplier, **{field: excluded}, price=77)
            workbook = load_workbook(BytesIO(generate_supplier_template(
                actor=self.user, product_kind=kind, supplier_id=self.supplier.pk,
            )))
            cells = self.price_cells(workbook)
            self.assertEqual(set(cells), {(kind, product.pk)})
            self.assertIsNone(cells[kind, product.pk].value)
            workbook.close()

    def test_prefilled_import_changes_only_selected_type_and_supplier_and_exports_latest_price(self):
        other = Supplier.objects.create(
            name="F", letter="F", highlight_color="#37A7BA", legal_entity="F", phone_1="2",
        )
        SupplierTechPrice.objects.create(supplier=self.supplier, tech=self.tech, price=88)
        SupplierCDPrice.objects.create(supplier=self.supplier, cd=self.cd, price=77)
        SupplierCDPrice.objects.create(supplier=other, cd=self.cd, price=66)
        workbook = load_workbook(BytesIO(generate_supplier_template(
            actor=self.user, product_kind="cd", supplier_id=self.supplier.pk,
        )))
        self.price_cells(workbook)["cd", self.cd.pk].value = 99.5
        self.assertEqual(self.import_file(workbook, "cd"), 1)
        self.assertEqual(SupplierTechPrice.objects.get(supplier=self.supplier).price, Decimal("88"))
        self.assertEqual(SupplierCDPrice.objects.get(supplier=other).price, Decimal("66"))
        next_book = load_workbook(BytesIO(generate_supplier_template(
            actor=self.user, product_kind="cd", supplier_id=self.supplier.pk,
        )))
        self.assertEqual(self.price_cells(next_book)["cd", self.cd.pk].value, 99.5)
        next_book.close()

    def test_hidden_product_absent_from_file_keeps_price(self):
        self.cd.exclude_from_supplier_template = True
        self.cd.save(update_fields=("exclude_from_supplier_template",))
        workbook = self.template("cd")
        offer = SupplierCDPrice.objects.create(supplier=self.supplier, cd=self.cd, price=77)
        self.assertEqual(self.import_file(workbook, "cd"), 0)
        offer.refresh_from_db()
        self.assertEqual(offer.price, Decimal("77"))

    def test_hidden_product_present_in_old_file_is_updated(self):
        workbook = self.template("cd")
        self.fill(workbook, "cd", self.cd.pk, 99)
        self.cd.exclude_from_supplier_template = True
        self.cd.save(update_fields=("exclude_from_supplier_template",))
        self.assertEqual(self.import_file(workbook, "cd"), 1)
        self.assertEqual(SupplierCDPrice.objects.get(supplier=self.supplier).price, Decimal("99"))

    def test_row_shift_or_identity_change_is_rejected_without_changing_price(self):
        offer = SupplierCDPrice.objects.create(supplier=self.supplier, cd=self.cd, price=77)
        for changed_column, changed_value in ((1, "TECH"), (2, self.cd.pk + 1), (3, "OTHER-SKU")):
            with self.subTest(column=changed_column):
                workbook = self.template("cd")
                row = self.fill(workbook, "cd", self.cd.pk, 99)
                workbook[MAIN_SHEET].cell(row, changed_column, changed_value)
                with self.assertRaisesMessage(ValidationError, "не соответствует техническому листу"):
                    self.import_file(workbook, "cd")
                offer.refresh_from_db()
                self.assertEqual(offer.price, Decimal("77"))

    def test_removed_metadata_or_added_unmapped_price_is_rejected(self):
        for mutation in ("metadata", "unmapped_price"):
            with self.subTest(mutation=mutation):
                workbook = self.template("cd")
                if mutation == "metadata":
                    meta = workbook[META_SHEET]
                    meta.delete_rows(meta.max_row)
                else:
                    sheet = workbook[MAIN_SHEET]
                    sheet.cell(sheet.max_row + 1, 6, 99)
                with self.assertRaisesMessage(ValidationError, "неполна"):
                    self.import_file(workbook, "cd")

    def test_unknown_product_is_rejected_even_with_blank_price(self):
        workbook = self.template("cd")
        sheet = workbook[MAIN_SHEET]
        row = self.fill(workbook, "cd", self.cd.pk, None)
        sheet.cell(row, 2, 999999)
        meta = workbook[META_SHEET]
        meta.cell(meta.max_row, 3, 999999)
        with self.assertRaisesMessage(ValidationError, "товар не найден"):
            self.import_file(workbook, "cd")

    def test_second_file_error_rolls_back_first_file(self):
        offer = SupplierCDPrice.objects.create(supplier=self.supplier, cd=self.cd, price=77)
        cd_book, tech_book = self.template("cd"), self.template("tech")
        self.fill(cd_book, "cd", self.cd.pk, 99)
        self.fill(tech_book, "tech", self.tech.pk, -1)
        with self.assertRaisesMessage(ValidationError, "больше нуля"):
            import_supplier_price_files(supplier_id=self.supplier.pk, actor=self.user, uploads={
                "cd": uploaded(workbook_bytes(cd_book)), "tech": uploaded(workbook_bytes(tech_book)),
            })
        offer.refresh_from_db()
        self.assertEqual(offer.price, Decimal("77"))
        self.assertFalse(SupplierTechPrice.objects.filter(supplier=self.supplier).exists())

    def test_legacy_combined_template_updates_only_selected_type(self):
        workbook = self.template(None)
        self.fill(workbook, "cd", self.cd.pk, 99)
        self.fill(workbook, "tech", self.tech.pk, 100)
        offer = SupplierTechPrice.objects.create(supplier=self.supplier, tech=self.tech, price=88)
        CD.objects.create(platform=self.cd.platform, name="Новая игра", sku="CD-2")
        self.assertEqual(self.import_file(workbook, "cd"), 1)
        self.assertEqual(SupplierCDPrice.objects.get(supplier=self.supplier).price, Decimal("99"))
        offer.refresh_from_db()
        self.assertEqual(offer.price, Decimal("88"))
