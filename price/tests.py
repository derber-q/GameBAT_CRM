from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import Permission
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import User
from catalog.models import Brand, CD, Platform, ProductType, Tech
from partners.models import Supplier
from pricing.models import SupplierCDPrice, SupplierTechPrice
from warehouse.models import CDWarehouseStock, TechWarehouseStock, Warehouse
from sales.models import Sale

from .excel import (
    HEADER_ROW,
    MAIN_SHEET,
    META_SHEET,
    generate_procurement_customer_xlsx,
    generate_retail_price_xlsx,
    generate_supplier_template,
    generate_wholesale_price_xlsx,
    import_procurement_order_xlsx,
    import_supplier_price,
    import_wholesale_price_to_sale,
)
from .models import ProcurementPriceListItem
from .services import calculate_postpayment, create_procurement_price_list, resolve_best_suppliers


def uploaded(content, name="price.xlsx"):
    return SimpleUploadedFile(
        name, content,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def workbook_bytes(workbook):
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class PriceExcelTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser("admin", password="StrongAdmin!123")
        self.moscow = Warehouse.objects.create(name="Москва")
        self.spb = Warehouse.objects.create(name="СПб")
        platform = Platform.objects.create(name="PS5")
        brand = Brand.objects.create(name="Sony")
        product_type = ProductType.objects.create(name="Консоли")
        self.cd = CD.objects.create(
            platform=platform, name="Игра", sku="CD-1", barcode="1", cost=100,
            retail_price=200, wholesale_price=150,
        )
        self.no_price_cd = CD.objects.create(
            platform=platform, name="Без цены", sku="CD-2", barcode="2", cost=100,
        )
        self.tech = Tech.objects.create(
            brand=brand, product_type=product_type, name="Консоль", sku="T-1", barcode="3", cost=100,
            retail_price=300, wholesale_price=250,
        )
        CDWarehouseStock.objects.create(warehouse=self.moscow, cd=self.cd, quantity=5)
        CDWarehouseStock.objects.create(warehouse=self.moscow, cd=self.no_price_cd, quantity=3)
        CDWarehouseStock.objects.create(warehouse=self.spb, cd=self.cd, quantity=0)
        TechWarehouseStock.objects.create(warehouse=self.moscow, tech=self.tech, quantity=0)
        self.alpha = Supplier.objects.create(
            name="Alpha", letter="A", highlight_color="#37A7BA", legal_entity="A", phone_1="1"
        )
        self.beta = Supplier.objects.create(
            name="Beta", letter="B", highlight_color="#37A7BA", legal_entity="B", phone_1="2"
        )

    def test_retail_uses_only_local_positive_stock_and_skips_missing_price(self):
        content, skipped = generate_retail_price_xlsx(warehouse_id=self.moscow.pk, actor=self.user)
        workbook = load_workbook(BytesIO(content), data_only=False)
        sheet = workbook[MAIN_SHEET]
        names = [sheet.cell(row, 1).value for row in range(HEADER_ROW + 1, sheet.max_row + 1)]
        self.assertIn("Игра", names)
        self.assertNotIn("Без цены", names)
        self.assertEqual(skipped, 1)
        self.assertIn("1", sheet["A7"].value)
        workbook.close()

        other, skipped = generate_retail_price_xlsx(warehouse_id=self.spb.pk, actor=self.user)
        workbook = load_workbook(BytesIO(other))
        self.assertEqual(workbook[MAIN_SHEET].max_row, HEADER_ROW)
        self.assertEqual(skipped, 0)
        workbook.close()

    def test_wholesale_file_has_availability_formulas_ids_and_import_ignores_price(self):
        content = generate_wholesale_price_xlsx(warehouse_id=self.moscow.pk, actor=self.user)
        workbook = load_workbook(BytesIO(content), data_only=False)
        sheet = workbook[MAIN_SHEET]
        self.assertEqual(sheet.cell(HEADER_ROW + 1, 3).value, 5)
        self.assertTrue(str(sheet.cell(sheet.max_row, 4).value).startswith("=SUMPRODUCT"))
        self.assertTrue(sheet.cell(sheet.max_row, 4).protection.locked)
        self.assertFalse(sheet.cell(HEADER_ROW + 1, 4).protection.locked)
        self.assertEqual(workbook[META_SHEET].sheet_state, "veryHidden")
        sheet.cell(HEADER_ROW + 1, 2, 1)
        sheet.cell(HEADER_ROW + 1, 4, 2)
        changed = workbook_bytes(workbook)

        warehouse, lines = import_wholesale_price_to_sale(uploaded(changed))
        self.assertEqual(warehouse, self.moscow)
        self.assertEqual(lines[0]["product_id"], self.cd.pk)
        self.assertEqual(lines[0]["quantity"], 2)
        self.cd.refresh_from_db()
        self.assertEqual(self.cd.wholesale_price, Decimal("150.00"))

        CDWarehouseStock.objects.filter(warehouse=self.moscow, cd=self.cd).update(quantity=1)
        with self.assertRaisesMessage(ValidationError, "запрошено 2 шт., доступно 1 шт"):
            import_wholesale_price_to_sale(uploaded(changed))

    def test_wholesale_formula_quantity_is_rejected(self):
        content = generate_wholesale_price_xlsx(warehouse_id=self.moscow.pk, actor=self.user)
        workbook = load_workbook(BytesIO(content), data_only=False)
        workbook[MAIN_SHEET].cell(HEADER_ROW + 1, 4, "=1+1")
        with self.assertRaisesMessage(ValidationError, "не формулой"):
            import_wholesale_price_to_sale(uploaded(workbook_bytes(workbook)))

    def test_wholesale_upload_prefills_normal_sale_without_automatic_write(self):
        self.client.force_login(self.user)
        content = generate_wholesale_price_xlsx(warehouse_id=self.moscow.pk, actor=self.user)
        workbook = load_workbook(BytesIO(content), data_only=False)
        workbook[MAIN_SHEET].cell(HEADER_ROW + 1, 4, 2)
        response = self.client.post(reverse("sales:import_wholesale"), {
            "file": uploaded(workbook_bytes(workbook)),
        })
        self.assertRedirects(response, reverse("sales:create"))
        self.assertFalse(Sale.objects.exists())
        self.assertEqual(
            CDWarehouseStock.objects.get(warehouse=self.moscow, cd=self.cd).quantity, 5
        )
        preview = self.client.get(reverse("sales:create"))
        self.assertEqual(preview.context["initial_items"][0]["quantity"], 2)
        response = self.client.post(reverse("sales:create"), {
            "warehouse": self.spb.pk,
            "price_type": Sale.PriceType.RETAIL,
            "sale_type": Sale.SaleType.WHOLESALE_PICKUP,
            "payment_method": Sale.PaymentMethod.BANK_ACCOUNT,
            "product_type": ["cd"], "product_id": [self.cd.pk], "quantity": [2],
        })
        self.assertEqual(response.status_code, 302)
        sale = Sale.objects.get()
        self.assertEqual(sale.warehouse, self.moscow)
        self.assertEqual(sale.price_type, Sale.PriceType.WHOLESALE)
        self.assertEqual(sale.cd_items.get().unit_price, Decimal("150.00"))
        self.assertEqual(CDWarehouseStock.objects.get(warehouse=self.moscow, cd=self.cd).quantity, 3)

    def test_supplier_template_contains_all_products_and_replaces_full_snapshot(self):
        SupplierCDPrice.objects.create(supplier=self.alpha, cd=self.no_price_cd, price=50)
        SupplierTechPrice.objects.create(supplier=self.alpha, tech=self.tech, price=60)
        content = generate_supplier_template(actor=self.user)
        workbook = load_workbook(BytesIO(content), data_only=False)
        sheet = workbook[MAIN_SHEET]
        names = [sheet.cell(row, 4).value for row in range(HEADER_ROW + 1, sheet.max_row + 1)]
        self.assertCountEqual(names, [self.cd.name, self.no_price_cd.name, self.tech.name])
        for row in range(HEADER_ROW + 1, sheet.max_row + 1):
            if sheet.cell(row, 1).value == "CD" and sheet.cell(row, 2).value == self.cd.pk:
                sheet.cell(row, 5, 100)
        count = import_supplier_price(
            supplier_id=self.alpha.pk, upload=uploaded(workbook_bytes(workbook)), actor=self.user
        )
        self.assertEqual(count, 1)
        self.assertEqual(SupplierCDPrice.objects.get(supplier=self.alpha).cd, self.cd)
        self.assertFalse(SupplierTechPrice.objects.filter(supplier=self.alpha).exists())

    def test_failed_supplier_import_keeps_previous_snapshot(self):
        original = SupplierCDPrice.objects.create(supplier=self.alpha, cd=self.cd, price=77)
        workbook = load_workbook(BytesIO(generate_supplier_template(actor=self.user)), data_only=False)
        workbook[MAIN_SHEET].cell(HEADER_ROW + 1, 5, -1)
        with self.assertRaisesMessage(ValidationError, "больше нуля"):
            import_supplier_price(
                supplier_id=self.alpha.pk, upload=uploaded(workbook_bytes(workbook)), actor=self.user
            )
        original.refresh_from_db()
        self.assertEqual(original.price, Decimal("77"))

    def test_incomplete_supplier_template_is_rejected_without_deleting_prices(self):
        original = SupplierCDPrice.objects.create(supplier=self.alpha, cd=self.cd, price=77)
        workbook = load_workbook(BytesIO(generate_supplier_template(actor=self.user)), data_only=False)
        meta = workbook[META_SHEET]
        meta.delete_rows(meta.max_row)
        with self.assertRaisesMessage(ValidationError, "неполна"):
            import_supplier_price(
                supplier_id=self.alpha.pk, upload=uploaded(workbook_bytes(workbook)), actor=self.user
            )
        original.refresh_from_db()
        self.assertEqual(original.price, Decimal("77"))

    def test_best_supplier_tie_uses_unique_wins_then_alphabetical(self):
        SupplierCDPrice.objects.create(supplier=self.alpha, cd=self.cd, price=100)
        SupplierCDPrice.objects.create(supplier=self.beta, cd=self.cd, price=100)
        SupplierCDPrice.objects.create(supplier=self.alpha, cd=self.no_price_cd, price=90)
        SupplierCDPrice.objects.create(supplier=self.beta, cd=self.no_price_cd, price=95)
        winners, counts = resolve_best_suppliers()
        self.assertEqual(winners[("cd", self.cd.pk)][0], self.alpha)
        self.assertEqual(counts[self.alpha.pk], 1)

        tie_offers = {("tech", self.tech.pk): [
            (self.beta, Decimal("10"), self.tech),
            (self.alpha, Decimal("10"), self.tech),
        ]}
        winners, counts = resolve_best_suppliers(tie_offers)
        self.assertEqual(counts[self.alpha.pk], 0)
        self.assertEqual(winners[("tech", self.tech.pk)][0], self.alpha)

    def test_postpayment_boundaries(self):
        self.assertEqual(calculate_postpayment("1999.99"), Decimal("2059.99"))
        self.assertEqual(calculate_postpayment("2000"), Decimal("2080.00"))
        self.assertEqual(calculate_postpayment("2999.99"), Decimal("3079.99"))
        self.assertEqual(calculate_postpayment("3000"), Decimal("3100.00"))

    def test_procurement_snapshot_client_privacy_and_imported_price_is_ignored(self):
        SupplierCDPrice.objects.create(supplier=self.alpha, cd=self.cd, price=100)
        price_list = create_procurement_price_list(actor=self.user, exchange_rate="20")
        item = price_list.items.get()
        SupplierCDPrice.objects.filter(supplier=self.alpha, cd=self.cd).update(price=1)
        item.refresh_from_db()
        self.assertEqual(item.selected_supplier, self.alpha)
        self.assertEqual(item.supplier_price_aed, Decimal("100"))
        self.assertEqual(item.base_price_rub, Decimal("2000.00"))

        content = generate_procurement_customer_xlsx(price_list_id=price_list.pk, actor=self.user)
        workbook = load_workbook(BytesIO(content), data_only=False)
        all_values = " ".join(
            str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row
            if cell.value is not None
        )
        self.assertNotIn(self.alpha.name, all_values)
        self.assertFalse(any(
            cell.value == self.alpha.letter
            for worksheet in workbook.worksheets for row in worksheet.iter_rows() for cell in row
        ))
        sheet = workbook[MAIN_SHEET]
        self.assertFalse(sheet.cell(HEADER_ROW + 1, 3).protection.locked)
        self.assertTrue(sheet.cell(sheet.max_row, 5).protection.locked)
        sheet.cell(HEADER_ROW + 1, 2, 1)
        sheet.cell(HEADER_ROW + 1, 3, 2)
        saved = workbook_bytes(workbook)
        imported_list, lines = import_procurement_order_xlsx(uploaded(saved))
        self.assertEqual(imported_list, price_list)
        self.assertEqual(lines, [{
            "price_list_item_id": item.pk, "prepayment_quantity": 2, "postpayment_quantity": 0,
        }])
        item.refresh_from_db()
        self.assertEqual(item.prepayment_price_rub, Decimal("2000.00"))

    def test_internal_supplier_prices_and_identity_follow_separate_permissions(self):
        SupplierCDPrice.objects.create(supplier=self.alpha, cd=self.cd, price=100)
        price_list = create_procurement_price_list(actor=self.user, exchange_rate=20)
        worker = User.objects.create_user("viewer", password="StrongWorker!123")
        worker.user_permissions.add(Permission.objects.get(
            content_type__app_label="price", codename="view_price_page"
        ))
        self.client.force_login(worker)
        response = self.client.get(reverse("price:procurement_detail", args=(price_list.pk,)))
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "Поставщик")
        self.assertNotContains(response, self.alpha.name)
        self.assertNotContains(response, "Цена поставщика")

        worker.user_permissions.add(Permission.objects.get(
            content_type__app_label="price", codename="view_supplier_prices"
        ))
        response = self.client.get(reverse("price:procurement_detail", args=(price_list.pk,)))
        self.assertContains(response, self.alpha.safe_label)
        self.assertNotContains(response, self.alpha.name)
        self.assertContains(response, "AED")

        worker.user_permissions.add(Permission.objects.get(
            content_type__app_label="partners", codename="view_supplier_details"
        ))
        response = self.client.get(reverse("price:procurement_detail", args=(price_list.pk,)))
        self.assertContains(response, self.alpha.name)

    def test_price_endpoints_require_individual_permissions(self):
        worker = User.objects.create_user("worker", password="StrongWorker!123")
        self.client.force_login(worker)
        self.assertEqual(self.client.get(reverse("price:index")).status_code, 403)
        worker.user_permissions.add(Permission.objects.get(
            content_type__app_label="price", codename="view_price_page"
        ))
        self.assertEqual(self.client.get(reverse("price:index")).status_code, 200)
        self.assertEqual(self.client.post(reverse("price:retail_export"), {
            "warehouse": self.moscow.pk,
        }).status_code, 403)

    def test_all_price_mutation_and_export_endpoints_reject_missing_permissions(self):
        worker = User.objects.create_user("restricted", password="StrongWorker!123")
        self.client.force_login(worker)
        checks = (
            ("post", reverse("price:settings_update"), {}),
            ("post", reverse("price:retail_export"), {"warehouse": self.moscow.pk}),
            ("post", reverse("price:wholesale_export"), {"warehouse": self.moscow.pk}),
            ("get", reverse("price:supplier_template"), {}),
            ("post", reverse("price:supplier_upload"), {}),
            ("post", reverse("price:procurement_create"), {"exchange_rate": 20}),
            ("get", reverse("price:procurement_detail", args=(1,)), {}),
            ("post", reverse("price:procurement_update", args=(1,)), {}),
            ("get", reverse("price:procurement_export", args=(1,)), {}),
            ("get", reverse("sales:import_wholesale"), {}),
        )
        for method, url, data in checks:
            with self.subTest(url=url):
                self.assertEqual(getattr(self.client, method)(url, data).status_code, 403)

    def test_malformed_or_wrong_document_is_rejected(self):
        content = generate_wholesale_price_xlsx(warehouse_id=self.moscow.pk, actor=self.user)
        with self.assertRaisesMessage(ValidationError, "другого типа"):
            import_procurement_order_xlsx(uploaded(content))
        workbook = load_workbook(BytesIO(content), data_only=False)
        del workbook[META_SHEET]
        with self.assertRaisesMessage(ValidationError, "технический лист"):
            import_wholesale_price_to_sale(uploaded(workbook_bytes(workbook)))

    def test_formula_injection_is_neutralized(self):
        self.cd.name = "=HYPERLINK(\"bad\")"
        self.cd.save(update_fields=("name",))
        content = generate_wholesale_price_xlsx(warehouse_id=self.moscow.pk, actor=self.user)
        workbook = load_workbook(BytesIO(content), data_only=False)
        self.assertTrue(workbook[MAIN_SHEET].cell(HEADER_ROW + 1, 1).value.startswith("'="))
        workbook.close()
