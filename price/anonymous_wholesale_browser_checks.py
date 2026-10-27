"""Безымянный оптовый прайс: скачивание, импорт и выбор заказчика в браузере."""
import os
from contextlib import closing
from pathlib import Path
from unittest.mock import patch

from django.conf import settings
from django.contrib.staticfiles.testing import StaticLiveServerTestCase
from django.db import connections
from django.urls import reverse
from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

from accounts.models import User
from catalog.models import CD, Platform
from partners.models import Partner
from sales.models import Sale
from warehouse.models import CDWarehouseStock, Warehouse
from .excel import HEADER_ROW, MAIN_SHEET


class AnonymousWholesaleBrowserChecks(StaticLiveServerTestCase):
    @patch.dict(os.environ, {'DJANGO_ALLOW_ASYNC_UNSAFE': 'true'})
    def test_download_and_import_without_customer_requires_selection(self):
        actor = User.objects.create_superuser('anonymous-price-browser')
        warehouse = Warehouse.objects.create(name='Склад браузера')
        partner = Partner.objects.create(name='Партнёр браузера')
        product = CD.objects.create(
            name='Игра браузера', platform=Platform.objects.create(name='PS5'),
            wholesale_price=10000, cost=6000,
        )
        stock = CDWarehouseStock.objects.create(warehouse=warehouse, cd=product, quantity=10)
        self.client.force_login(actor)
        output = Path(settings.BASE_DIR) / 'outputs/trip_operations_20261025/browser'
        output.mkdir(parents=True, exist_ok=True)
        with sync_playwright() as pw, closing(connections['default']):
            browser = pw.chromium.launch()
            context = browser.new_context(viewport={'width': 1440, 'height': 950})
            context.add_cookies([{
                'name': settings.SESSION_COOKIE_NAME,
                'value': self.client.cookies[settings.SESSION_COOKIE_NAME].value,
                'url': self.live_server_url,
            }])
            context.route('**/crm/api/exchange-rates/', lambda route: route.fulfill(json={'rates': {}, 'available': False}))
            page = context.new_page()
            errors = []
            page.on('pageerror', lambda error: errors.append(str(error)))
            page.goto(self.live_server_url + reverse('price:index'))
            form = page.locator(f'form[action="{reverse("price:wholesale_export")}"]')
            expect(form.locator('[name="customer"]')).to_have_value('')
            with page.expect_download() as download_info:
                form.get_by_role('button', name='Скачать оптовый XLSX', exact=True).click()
            download = download_info.value
            self.assertNotIn(partner.name, download.suggested_filename)
            path = output / 'anonymous-wholesale.xlsx'
            download.save_as(path)
            book = load_workbook(path)
            sheet = book[MAIN_SHEET]
            for row in range(HEADER_ROW + 1, sheet.max_row + 1):
                if sheet.cell(row, 2).value == 10000:
                    sheet.cell(row, 4, 2)
            book.save(path)
            book.close()
            page.goto(self.live_server_url + reverse('sales:import_wholesale'))
            page.locator('[name="file"]').set_input_files(path)
            page.get_by_role('button', name='Продолжить', exact=True).click()
            expect(page.get_by_text('В прайсе не указан заказчик.', exact=False)).to_be_visible()
            self.assertFalse(Sale.objects.exists())
            page.screenshot(path=str(output / 'anonymous-wholesale-customer.png'), full_page=True)
            page.locator('[name="customer"]').select_option(str(partner.pk))
            page.get_by_role('button', name='Продолжить', exact=True).click()
            page.wait_for_url(self.live_server_url + reverse('sales:create'))
            expect(page.locator('[name="customer"]')).to_have_value(str(partner.pk))
            expect(page.locator('[name="customer"]')).to_be_disabled()
            expect(page.locator('.stock-lines [name="quantity"]')).to_have_value('2')
            stock.refresh_from_db()
            self.assertEqual(stock.quantity, 10)
            self.assertFalse(Sale.objects.exists())
            self.assertEqual(errors, [])
            browser.close()
