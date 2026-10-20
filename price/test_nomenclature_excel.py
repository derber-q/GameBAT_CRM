"""Группировка закупочных файлов и совместимость обратного импорта."""
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

from django.core.exceptions import ValidationError
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from openpyxl import Workbook, load_workbook

from accounts.models import User
from catalog.models import Brand, CD, GameSeries, Platform, ProductType, Tech
from partners.models import Supplier
from pricing.models import SupplierCDPrice, SupplierTechPrice
from orders.models import CustomerProcurementOrder
from orders.services import advance_order_status, create_customer_procurement_order, receive_stock_order
from orders.trips import create_trip
from warehouse.models import Warehouse

from .excel import (
    HEADER_ROW, LEGACY_STOCK_HEADERS, MAIN_SHEET, META_SHEET, _meta_sheet, generate_procurement_customer_xlsx,
    generate_procurement_stock_xlsx, generate_supplier_template,
    import_procurement_order_files, import_procurement_order_xlsx, import_supplier_price,
)
from .services import create_procurement_price_list
from .tests import uploaded, workbook_bytes


class ProcurementNomenclatureExcelTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_superuser("grouping-admin", password="StrongAdmin!123")
        cls.supplier = Supplier.objects.create(
            name="Поставщик A", letter="A", highlight_color="#37A7BA", legal_entity="A", phone_1="1",
        )
        sony = Brand.objects.create(name="Sony")
        gamepad = ProductType.objects.create(name="Геймпад")
        cls.white, cls.edge = [Tech.objects.create(
            name=name, brand=sony, product_type=gamepad, weight_grams=1000, sku=f"000-tech-{index}",
        ) for index, name in enumerate(("Геймпад PS5 Белый (Original White)", "Геймпад Dualsense Edge"))]
        ns, ns2, ps5 = [Platform.objects.create(name=name) for name in (
            "Nintendo Switch", "Nintendo Switch 2", "PlayStation 5",
        )]
        zelda = GameSeries.objects.create(name="Zelda")
        assassin = GameSeries.objects.create(name="Assassin's Creed")
        cls.ns = CD.objects.create(name="Zulu Zelda", platform=ns, game_series=zelda, sku="000-cd-1")
        cls.ns2 = CD.objects.create(name="Alpha Assassin", platform=ns2, game_series=assassin, sku="000-cd-2")
        cls.ps5 = CD.objects.create(name="Beta без серии", platform=ps5, sku="000-cd-3")
        for index, product in enumerate((cls.ns, cls.ns2, cls.ps5)):
            SupplierCDPrice.objects.create(supplier=cls.supplier, cd=product, price=Decimal("41.123456") + index)
        for index, product in enumerate((cls.white, cls.edge)):
            SupplierTechPrice.objects.create(supplier=cls.supplier, tech=product, price=Decimal("151.654321") + index)
        cls.price = create_procurement_price_list(
            actor=cls.user, exchange_rate_usdt_aed=4, exchange_rate_usdt_rub=80,
            delivery_general_usdt_per_kg=0, delivery_devices_usdt_per_kg=0,
            delivery_storage_usdt_per_kg=0, delivery_cd_rub=0,
        )

    @staticmethod
    def metadata(workbook):
        return [row for row in workbook[META_SHEET].iter_rows(values_only=True) if isinstance(row[0], int)]

    def test_all_price_generators_follow_page_order_and_keep_identity_without_lazy_series_queries(self):
        expected = [("tech", self.white.pk), ("tech", self.edge.pk),
                    ("cd", self.ns2.pk), ("cd", self.ns.pk), ("cd", self.ps5.pk)]
        titles = ["Tech · Sony · Геймпад", "CD · Платформа: Nintendo Switch 2",
                  "CD · Платформа: Nintendo Switch",
                  "CD · Платформа: PlayStation 5"]
        cases = [
            ("supplier", None, lambda: generate_supplier_template(actor=self.user)),
            ("prefilled", None, lambda: generate_supplier_template(actor=self.user, supplier_id=self.supplier.pk)),
            ("customer", None, lambda: generate_procurement_customer_xlsx(price_list_id=self.price.pk, actor=self.user)),
            ("stock", None, lambda: generate_procurement_stock_xlsx(price_list_id=self.price.pk, actor=self.user)),
        ]
        for kind in ("cd", "tech"):
            cases.extend((
                ("supplier", kind, lambda kind=kind: generate_supplier_template(actor=self.user, product_kind=kind)),
                ("prefilled", kind, lambda kind=kind: generate_supplier_template(actor=self.user, product_kind=kind, supplier_id=self.supplier.pk)),
                ("customer", kind, lambda kind=kind: generate_procurement_customer_xlsx(price_list_id=self.price.pk, actor=self.user, product_kind=kind)),
            ))
        for label, kind, generate in cases:
            with self.subTest(document=label, kind=kind):
                with CaptureQueriesContext(connection) as queries:
                    content = generate()
                self.assertFalse(any('FROM "catalog_gameseries"' in q["sql"] for q in queries))
                self.assertLessEqual(len(queries), 8)
                workbook = load_workbook(BytesIO(content))
                sheet = workbook[MAIN_SHEET]
                column = 4 if label in ("supplier", "prefilled") else 1
                actual_titles = [row[column - 1] for row in sheet.iter_rows(values_only=True)
                                 if isinstance(row[column - 1], str) and row[column - 1].startswith(("Tech ·", "CD ·"))]
                self.assertEqual(actual_titles, [title for title in titles if kind is None or title.lower().startswith(kind)])
                keys = [(row[1], row[2]) for row in self.metadata(workbook)]
                self.assertEqual(keys, [key for key in expected if kind is None or key[0] == kind])
                self.assertEqual(len(keys), len(set(keys)))
                workbook.close()

    def test_supplier_round_trip_preserves_exact_prices_and_cleared_cells(self):
        snapshots = list(self.price.items.values_list("id", "supplier_price_aed"))
        for prefilled in (False, True):
            workbook = load_workbook(BytesIO(generate_supplier_template(
                actor=self.user, supplier_id=self.supplier.pk if prefilled else None,
            )))
            expected = {}
            rows = self.metadata(workbook)
            for index, (row, kind, pk, article) in enumerate(rows):
                value = None if prefilled and index == len(rows) - 1 else Decimal("97.123456") + index
                workbook[MAIN_SHEET].cell(row, 6).value = value
                if value is not None:
                    expected[kind, pk] = value
            self.assertEqual(import_supplier_price(
                supplier_id=self.supplier.pk, actor=self.user, upload=uploaded(workbook_bytes(workbook)),
            ), len(expected))
            actual = {("cd", pk): price for pk, price in SupplierCDPrice.objects.filter(supplier=self.supplier).values_list("cd_id", "price")}
            actual.update({("tech", pk): price for pk, price in SupplierTechPrice.objects.filter(supplier=self.supplier).values_list("tech_id", "price")})
            self.assertEqual(actual, expected)
        self.assertEqual(list(self.price.items.values_list("id", "supplier_price_aed")), snapshots)
        next_price = create_procurement_price_list(
            actor=self.user, exchange_rate_usdt_aed=4, exchange_rate_usdt_rub=80,
            delivery_general_usdt_per_kg=0, delivery_devices_usdt_per_kg=0,
            delivery_storage_usdt_per_kg=0, delivery_cd_rub=0,
        )
        self.assertEqual({
            (item.product_kind, item.product_id): item.supplier_price_aed
            for item in next_price.items.all()
        }, expected)

    def test_customer_round_trip_creates_orders_and_trips_with_saved_prices(self):
        for kind in (None, "cd", "tech"):
            with self.subTest(kind=kind):
                workbook = load_workbook(BytesIO(generate_procurement_customer_xlsx(
                    actor=self.user, price_list_id=self.price.pk, product_kind=kind,
                )))
                expected = {}
                for index, (row, row_kind, pk, article) in enumerate(self.metadata(workbook), 1):
                    sheet = workbook[MAIN_SHEET]
                    sheet.cell(row, 3, index)
                    sheet.cell(row, 5, index + 1)
                    sheet.cell(row, 2, 1)
                    sheet.cell(row, 4, 1)
                    expected[row_kind, pk] = (index, index + 1)
                source, lines = import_procurement_order_xlsx(uploaded(workbook_bytes(workbook)), product_kind=kind)
                order = create_customer_procurement_order(
                    actor=self.user, price_list_id=source.pk, recipient=f"Проверка {kind}", comment="", lines=lines,
                )
                self.assertEqual(len(lines), len(expected))
                for item in order.items.select_related("price_list_item"):
                    self.assertEqual((item.prepayment_quantity, item.postpayment_quantity), expected[item.product_kind, item.product_id])
                    self.assertEqual(item.prepayment_unit_price, item.price_list_item.prepayment_price_rub)
                    self.assertEqual(item.postpayment_unit_price, item.price_list_item.postpayment_price_rub)
                    self.assertEqual(item.selected_supplier_id, self.supplier.pk)
                trip = create_trip(actor=self.user, order_ids=[order.pk], creation_key=uuid4())
                self.assertEqual(sum(line.quantity for line in trip.lines.all()), order.total_units)
                self.assertEqual(trip.lines.count(), len(expected))

    def test_stock_round_trip_formulas_create_trip_and_receipt_at_saved_cost(self):
        workbook = load_workbook(BytesIO(generate_procurement_stock_xlsx(actor=self.user, price_list_id=self.price.pk)))
        expected = {}
        for index, (row, kind, pk, article) in enumerate(self.metadata(workbook), 1):
            sheet = workbook[MAIN_SHEET]
            self.assertEqual(sheet.cell(row, 9).value, f"=G{row}*H{row}")
            self.assertEqual(sheet.cell(row, 4).value, article)
            self.assertFalse(sheet.cell(row, 8).protection.locked)
            sheet.cell(row, 8, index)
            sheet.cell(row, 7, 1)
            expected[kind, pk] = index
        source, lines, purpose = import_procurement_order_files({"stock": uploaded(workbook_bytes(workbook))}, with_purpose=True)
        self.assertEqual(purpose, CustomerProcurementOrder.Purpose.STOCK)
        order = create_customer_procurement_order(
            actor=self.user, price_list_id=source.pk, recipient="На склад", comment="", lines=lines, purpose=purpose,
        )
        for item in order.items.select_related("price_list_item"):
            self.assertEqual(item.quantity, expected[item.product_kind, item.product_id])
            self.assertEqual(item.prepayment_unit_price, item.price_list_item.stock_unit_cost)
        trip = create_trip(actor=self.user, order_ids=[order.pk], creation_key=uuid4())
        self.assertEqual(sum(line.quantity for line in trip.lines.all()), sum(expected.values()))
        advance_order_status(actor=self.user, order_id=order.pk, next_status=order.Status.CONFIRMED)
        advance_order_status(actor=self.user, order_id=order.pk, next_status=order.Status.RECEIVING)
        received_order = receive_stock_order(actor=self.user, order_id=order.pk, warehouse_id=Warehouse.objects.first().pk)
        supply = received_order.supply
        self.assertEqual(supply.cd_items.count() + supply.tech_items.count(), len(expected))
        order.refresh_from_db()
        self.assertEqual(order.status, order.Status.COMPLETED)

    def test_customer_import_rejects_shifted_product_rows(self):
        workbook = load_workbook(BytesIO(generate_procurement_customer_xlsx(actor=self.user, price_list_id=self.price.pk)))
        sheet = workbook[MAIN_SHEET]
        first, second = [row[0] for row in self.metadata(workbook)[:2]]
        sheet.cell(first, 3, 2)
        sheet.cell(second, 3, 7)
        first_values = [sheet.cell(first, col).value for col in range(1, 6)]
        second_values = [sheet.cell(second, col).value for col in range(1, 6)]
        for col, (one, two) in enumerate(zip(first_values, second_values), 1):
            sheet.cell(first, col).value = two
            sheet.cell(second, col).value = one
        with self.assertRaisesMessage(ValidationError, "не соответствует"):
            import_procurement_order_xlsx(uploaded(workbook_bytes(workbook)))

    def legacy_workbook(self, document_type):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = MAIN_SHEET
        headers = {
            "resource_supplier": ("Тип", "ID", "Артикул", "Название", "Цена AED"),
            "resource_procurement": ("Товар", "Цена предоплата", "Количество предоплата", "Цена постоплата", "Количество постоплата"),
            "resource_stock_procurement": LEGACY_STOCK_HEADERS,
        }[document_type]
        for column, value in enumerate(headers, 1):
            sheet.cell(HEADER_ROW, column, value)
        items = [self.price.items.get(cd=self.ns), self.price.items.get(tech=self.white)]
        metadata = []
        for index, item in enumerate(items, 1):
            row = HEADER_ROW + index
            if document_type == "resource_supplier":
                values = (item.product_kind.upper(), item.product_id, item.article_snapshot, item.product_name_snapshot, Decimal("90.123456") + index)
            elif document_type == "resource_procurement":
                values = (item.product_name_snapshot, item.prepayment_price_rub, index, item.postpayment_price_rub, 3 - index)
            else:
                values = (item.product_name_snapshot, item.article_snapshot, item.base_price_rub,
                          item.delivery_rub, item.stock_unit_cost, index, f"=E{row}*F{row}")
            for column, value in enumerate(values, 1):
                sheet.cell(row, column, value)
            metadata.append((row, item.product_kind, item.product_id, item.article_snapshot))
        _meta_sheet(workbook, document_type=document_type, values={
            "main_sheet": MAIN_SHEET, "price_list_token": self.price.public_token,
        }, rows=metadata)
        return workbook, items

    def test_legacy_supplier_customer_and_stock_formats_are_still_imported(self):
        workbook, items = self.legacy_workbook("resource_supplier")
        self.assertEqual(import_supplier_price(
            supplier_id=self.supplier.pk, actor=self.user, upload=uploaded(workbook_bytes(workbook)),
        ), 2)
        self.assertEqual(SupplierCDPrice.objects.get(supplier=self.supplier, cd=self.ns).price, Decimal("91.123456"))
        self.assertEqual(SupplierTechPrice.objects.get(supplier=self.supplier, tech=self.white).price, Decimal("92.123456"))
        self.assertEqual(SupplierCDPrice.objects.filter(supplier=self.supplier).count(), 3)
        workbook, items = self.legacy_workbook("resource_procurement")
        source, lines = import_procurement_order_xlsx(uploaded(workbook_bytes(workbook)))
        self.assertEqual([(row["price_list_item_id"], row["prepayment_quantity"], row["postpayment_quantity"]) for row in lines],
                         [(items[0].pk, 1, 2), (items[1].pk, 2, 1)])
        workbook, items = self.legacy_workbook("resource_stock_procurement")
        source, lines, purpose = import_procurement_order_files({"stock": uploaded(workbook_bytes(workbook))}, with_purpose=True)
        self.assertEqual(purpose, CustomerProcurementOrder.Purpose.STOCK)
        self.assertEqual([(row["price_list_item_id"], row["prepayment_quantity"]) for row in lines], [(items[0].pk, 1), (items[1].pk, 2)])
