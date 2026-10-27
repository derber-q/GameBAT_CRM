from decimal import Decimal
from io import BytesIO
from uuid import uuid4
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.test import TestCase
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import User
from catalog.models import Brand, CD, Platform, ProductType, Tech
from catalog.supplier_exclusions import set_supplier_exclusions
from orders.services import create_customer_procurement_order
from orders.stock_order_form import stock_order_items
from orders.trips import change_trip_line, create_trip
from partners.models import Supplier
from pricing.models import SupplierCDPrice, SupplierTechPrice
from warehouse.models import CDWarehouseStock, Warehouse

from .excel import (MAIN_SHEET, META_SHEET, generate_procurement_customer_xlsx, generate_retail_price_xlsx,
                    generate_supplier_template, generate_wholesale_price_xlsx, import_procurement_order_xlsx,
                    import_supplier_price, import_wholesale_price_to_sale)
from .models import ProcurementPriceList, ProcurementPriceListItem, ProcurementPricePreference
from .services import create_procurement_price_list, resolve_best_suppliers, round_customer_price, update_procurement_pricing
from .tests import find_sheet_row, uploaded, workbook_bytes


class SupplierNewnessRoundingTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser('rounding-owner')
        self.platform = Platform.objects.create(name='Nintendo Switch')
        self.platform2 = Platform.objects.create(name='Nintendo Switch 2')
        self.cd = CD.objects.create(platform=self.platform, name='Первая игра', sku='000001', wholesale_price=3011)
        self.cd2 = CD.objects.create(platform=self.platform2, name='Вторая игра', sku='000002', wholesale_price=3021)
        self.suppliers = {}
        for number, (letter, amount) in enumerate((('A', 95), ('F', 80), ('M', 70)), 1):
            supplier = Supplier.objects.create(name=letter, letter=letter, priority=number, highlight_color='#37A7BA', legal_entity=letter, phone_1='1')
            self.suppliers[letter] = supplier
            for product in (self.cd, self.cd2):
                SupplierCDPrice.objects.create(supplier=supplier, cd=product, price=amount)

    def create_price(self):
        return create_procurement_price_list(actor=self.user, exchange_rate_usdt_aed=1, exchange_rate_usdt_rub=1,
            delivery_general_usdt_per_kg=0, delivery_devices_usdt_per_kg=0, delivery_storage_usdt_per_kg=0, delivery_cd_rub=0)

    def order(self, price, item, quantity=3):
        return create_customer_procurement_order(actor=self.user, price_list_id=price.pk, recipient='Покупатель', comment='',
            lines=[{'price_list_item_id': item.pk, 'prepayment_quantity': quantity, 'postpayment_quantity': 0}])

    def test_supplier_filter_selects_f_and_keeps_m_offer_in_template_and_import(self):
        set_supplier_exclusions(product=self.cd, supplier_ids=[self.suppliers['M'].pk], actor=self.user)
        winners, _ = resolve_best_suppliers()
        self.assertEqual(winners[('cd', self.cd.pk)][:2], (self.suppliers['F'], Decimal('80')))
        self.assertEqual(winners[('cd', self.cd2.pk)][0], self.suppliers['M'])
        content = generate_supplier_template(actor=self.user, product_kind='cd', supplier_id=self.suppliers['M'].pk)
        workbook = load_workbook(BytesIO(content))
        meta = workbook[META_SHEET]
        rows = [r for r in meta.iter_rows(values_only=True) if len(r) >= 3 and r[1] == 'cd' and r[2] == self.cd.pk]
        self.assertEqual(len(rows), 1)
        row_number = int(rows[0][0])
        sheet = workbook[MAIN_SHEET]
        # Шаблон имеет служебные идентификаторы; цена AED — последняя колонка.
        sheet.cell(row_number, sheet.max_column, 65)
        import_supplier_price(supplier_id=self.suppliers['M'].pk, upload=uploaded(workbook_bytes(workbook)), actor=self.user, product_kind='cd')
        self.assertEqual(SupplierCDPrice.objects.get(cd=self.cd, supplier=self.suppliers['M']).price, Decimal('65'))
        price = self.create_price()
        item = price.items.get(cd=self.cd)
        self.assertEqual(item.selected_supplier, self.suppliers['F'])
        self.assertEqual(item.supplier_price_aed, Decimal('80'))
        self.assertEqual(item.supplier_offers.get(supplier=self.suppliers['M']).price_aed, Decimal('65'))
        set_supplier_exclusions(product=self.cd, supplier_ids=[], actor=self.user)
        self.assertEqual(resolve_best_suppliers()[0][('cd', self.cd.pk)][0], self.suppliers['M'])

    def test_all_excluded_has_no_winner_or_zero_price_fallback(self):
        for product in (self.cd, self.cd2):
            set_supplier_exclusions(product=product, supplier_ids=[s.pk for s in self.suppliers.values()], actor=self.user)
        self.assertEqual(resolve_best_suppliers()[0], {})
        with self.assertRaisesMessage(ValidationError, 'Нет доступных предложений'):
            self.create_price()
        self.assertFalse(ProcurementPriceList.objects.exists())

    def test_tech_supplier_selection_uses_same_product_specific_filter(self):
        product = Tech.objects.create(name='Техника', brand=Brand.objects.create(name='Sony'),
                                      product_type=ProductType.objects.create(name='Игровая консоль'), weight_grams=1000)
        for letter, price in (('A', 95), ('F', 80), ('M', 70)):
            SupplierTechPrice.objects.create(tech=product, supplier=self.suppliers[letter], price=price)
        set_supplier_exclusions(product=product, supplier_ids=[self.suppliers['M'].pk], actor=self.user)
        self.assertEqual(resolve_best_suppliers()[0][('tech', product.pk)][:2], (self.suppliers['F'], Decimal('80')))
        self.assertEqual(resolve_best_suppliers()[0][('cd', self.cd.pk)][0], self.suppliers['M'])

    def test_supplier_exclusion_selection_does_not_add_query_per_product(self):
        products = CD.objects.bulk_create([CD(platform=self.platform, name=f'Товар {index}') for index in range(100)])
        SupplierCDPrice.objects.bulk_create([SupplierCDPrice(cd=product, supplier=supplier, price=50)
                                            for product in products for supplier in self.suppliers.values()])
        with CaptureQueriesContext(connection) as queries:
            winners, _ = resolve_best_suppliers()
        self.assertEqual(len(winners), 102)
        self.assertLessEqual(len(queries), 4)

    def test_old_price_and_order_unchanged_but_new_assignment_is_blocked(self):
        price = self.create_price()
        item = price.items.get(cd=self.cd)
        order = self.order(price, item)
        before = list(order.items.values())
        old_price = item.supplier_price_aed
        set_supplier_exclusions(product=self.cd, supplier_ids=[item.selected_supplier_id], actor=self.user)
        self.assertEqual(list(order.items.values()), before)
        item.refresh_from_db()
        self.assertEqual(item.supplier_price_aed, old_price)
        self.assertNotIn(item.pk, [r.pk for r in stock_order_items(price)])
        with self.assertRaisesMessage(ValidationError, 'исключён'):
            self.order(price, item)
        with self.assertRaisesMessage(ValidationError, 'исключён'):
            create_trip(actor=self.user, order_ids=[order.pk], creation_key=uuid4())
        order.refresh_from_db()
        self.assertIsNone(order.trip_id)

    def test_automatic_m_to_f_rule_respects_exclusion_and_manual_transfer_cannot_bypass(self):
        price = self.create_price()
        item = price.items.get(cd=self.cd)
        order = self.order(price, item)
        set_supplier_exclusions(product=self.cd, supplier_ids=[self.suppliers['F'].pk], actor=self.user)
        trip = create_trip(actor=self.user, order_ids=[order.pk], creation_key=uuid4())
        line = trip.lines.get()
        self.assertEqual(line.supplier_id, self.suppliers['M'].pk)
        with self.assertRaisesMessage(ValidationError, 'исключён'):
            change_trip_line(actor=self.user, batch_id=trip.pk, line_id=line.pk, revision=trip.revision, quantity=0,
                             action='transfer', target_supplier_id=self.suppliers['F'].pk, price_supplier_id=self.suppliers['A'].pk,
                             _allow_priority_transfer=True)
        line.refresh_from_db()
        self.assertEqual(line.quantity, 3)
        set_supplier_exclusions(product=self.cd, supplier_ids=[self.suppliers['M'].pk], actor=self.user)
        with self.assertRaisesMessage(ValidationError, 'исключён'):
            change_trip_line(actor=self.user, batch_id=trip.pk, line_id=line.pk, revision=trip.revision, quantity=0,
                             action='transfer', target_supplier_id=self.suppliers['A'].pk, price_supplier_id=self.suppliers['M'].pk,
                             _allow_priority_transfer=True)

    def test_rounding_examples(self):
        for before, after in (('3001', '3010'), ('3009', '3010'), ('254', '260'), ('1230', '1230'),
                              ('5891', '5900'), ('3000.50', '3010')):
            with self.subTest(before=before):
                self.assertEqual(round_customer_price(Decimal(before)), Decimal(after))

    def test_independent_rounding_persists_effective_markup_original_setting_and_order_prices(self):
        price = ProcurementPriceList.objects.create(exchange_rate_aed_rub=1, created_by=self.user)
        item = ProcurementPriceListItem.objects.create(price_list=price, product_kind='cd', cd=self.cd,
            product_name_snapshot=self.cd.name, article_snapshot=self.cd.sku, selected_supplier=self.suppliers['A'],
            supplier_price_aed=2800, exchange_rate_aed_rub=1, base_price_rub=2800, delivery_rub=0,
            markup_rub=201, prepayment_price_rub=3001, postpayment_price_rub=3127, postpayment_surcharge_rub=126)
        rows = {item.pk: {'markup': '201', 'delivery': '0', 'postpayment_surcharge': '126', 'surcharge_mode': 'fixed'}}
        for repeat in range(2):
            update_procurement_pricing(actor=self.user, price_list_id=price.pk, rows=rows, can_change_markup=True, can_change_delivery=True)
            item.refresh_from_db()
            self.assertEqual((item.prepayment_price_rub, item.postpayment_price_rub), (Decimal('3010'), Decimal('3130')))
            self.assertEqual((item.prepayment_rounding_extra_rub, item.postpayment_rounding_extra_rub), (Decimal('9'), Decimal('3')))
            self.assertEqual((item.original_markup_rub, item.markup_rub), (Decimal('201'), Decimal('210')))
            self.assertEqual(item.base_price_rub, Decimal('2800'))
        self.assertEqual(ProcurementPricePreference.objects.get(cd=self.cd).markup_rub, Decimal('201'))
        content = generate_procurement_customer_xlsx(price_list_id=price.pk, actor=self.user, product_kind='cd')
        workbook = load_workbook(BytesIO(content))
        sheet = workbook[MAIN_SHEET]
        row = find_sheet_row(sheet, self.cd.name)
        self.assertEqual((sheet.cell(row, 2).value, sheet.cell(row, 4).value), (3010, 3130))
        sheet.cell(row, 3, 2)
        sheet.cell(row, 5, 3)
        imported, lines = import_procurement_order_xlsx(uploaded(workbook_bytes(workbook)), product_kind='cd')
        order = create_customer_procurement_order(actor=self.user, price_list_id=imported.pk, recipient='Итог', comment='', lines=lines)
        self.assertEqual((order.items.get().prepayment_unit_price, order.items.get().postpayment_unit_price), (Decimal('3010'), Decimal('3130')))
        self.assertEqual(order.grand_total, Decimal('15410'))

    def test_newness_and_switch_order_round_trip_with_expired_previous_file(self):
        self.cd2.newness_months = 1
        self.cd2.save(update_fields=('newness_months',))
        price = self.create_price()
        content = generate_procurement_customer_xlsx(price_list_id=price.pk, actor=self.user, product_kind='cd')
        workbook = load_workbook(BytesIO(content))
        sheet = workbook[MAIN_SHEET]
        second = find_sheet_row(sheet, '🌟 ' + self.cd2.name)
        first = find_sheet_row(sheet, self.cd.name)
        self.assertLess(second, first)
        self.assertEqual(sheet.cell(second, 1).fill.fgColor.rgb, '00E2F0D9')
        self.assertFalse(sheet.cell(second, 3).protection.locked)
        sheet.cell(second, 3, 4)
        previous_file = workbook_bytes(workbook)
        with patch('price.excel.timezone.now', return_value=self.cd2.newness_until):
            current = load_workbook(BytesIO(generate_procurement_customer_xlsx(price_list_id=price.pk, actor=self.user, product_kind='cd')))
            self.assertIsNotNone(find_sheet_row(current[MAIN_SHEET], self.cd2.name))
            current.close()
        _price, lines = import_procurement_order_xlsx(uploaded(previous_file), product_kind='cd')
        self.assertEqual(lines[0]['price_list_item_id'], price.items.get(cd=self.cd2).pk)
        self.cd2.refresh_from_db()
        self.assertEqual(self.cd2.name, 'Вторая игра')

    def test_warehouse_customer_wholesale_newness_and_platform_order_do_not_round_stock_prices(self):
        warehouse = Warehouse.objects.create(name='Склад')
        self.cd2.newness_months = 1
        self.cd2.avito_price = 3021
        self.cd2.save()
        self.cd.avito_price = 3011
        self.cd.save()
        for product in (self.cd, self.cd2):
            CDWarehouseStock.objects.create(warehouse=warehouse, cd=product, quantity=5)
        workbook = load_workbook(BytesIO(generate_wholesale_price_xlsx(warehouse_id=warehouse.pk, actor=self.user, partner_name='Тестовый заказчик')))
        sheet = workbook[MAIN_SHEET]
        second = find_sheet_row(sheet, '🌟 ' + self.cd2.name)
        first = find_sheet_row(sheet, self.cd.name)
        self.assertLess(second, first)
        self.assertEqual(sheet.cell(second, 2).value, 3021)
        sheet.cell(second, 4, 2)
        _warehouse, lines = import_wholesale_price_to_sale(uploaded(workbook_bytes(workbook)))
        self.assertEqual((lines[0]['product_id'], lines[0]['unit_price']), (self.cd2.pk, '3021.00'))
        retail, _ = generate_retail_price_xlsx(warehouse_id=warehouse.pk, actor=self.user)
        workbook = load_workbook(BytesIO(retail))
        self.assertLess(find_sheet_row(workbook[MAIN_SHEET], self.cd2.name), find_sheet_row(workbook[MAIN_SHEET], self.cd.name))
        workbook.close()
