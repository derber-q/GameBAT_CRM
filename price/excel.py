import logging
from copy import copy
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from textwrap import wrap

from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Sum
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from catalog.models import CD, Tech
from catalog.nomenclature_grouping import nomenclature_product_groups
from catalog.product_ordering import platform_order_key, tech_price_order
from catalog.newness import customer_product_name
from partners.models import Supplier
from pricing.models import SupplierCDPrice, SupplierTechPrice
from warehouse.models import CDWarehouseStock, TechWarehouseStock, Warehouse
from core.excel_layout import wrapped_text_height

from .models import PriceDocumentSettings, ProcurementPriceList

logger = logging.getLogger("gamebat.business")
TEMPLATE_VERSION = "1"
MAIN_SHEET = "Прайс"
META_SHEET = "_resource_meta"
HEADER_ROW = 8
SUPPLIER_HEADER_ROW = 1
BLUE = "334660"
TEAL = "2C879B"
LIGHT = "EAF3F5"
WHITE = "FFFFFF"
GRAY = "DCE5E7"
GROUP_FILL = "D5E6EA"
MONEY_FORMAT = '#,##0.00 "₽"'
AED_FORMAT = '#,##0.000000 "AED"'


def safe_text(value):
    result = ILLEGAL_CHARACTERS_RE.sub("", str(value or ""))
    if result.startswith(("=", "+", "-", "@")):
        result = "'" + result
    return result


def _base_workbook(*, title, subtitle, columns, settings=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = MAIN_SHEET
    settings = settings or PriceDocumentSettings.get_solo()
    last_column = get_column_letter(columns)
    sheet.merge_cells(f"A1:{last_column}1")
    sheet["A1"] = safe_text(settings.company_name)
    sheet["A1"].font = Font(size=20, bold=True, color=WHITE)
    sheet["A1"].fill = PatternFill("solid", fgColor=BLUE)
    sheet["A1"].alignment = Alignment(vertical="center")
    if settings.logo:
        try:
            from openpyxl.drawing.image import Image as ExcelImage

            logo = ExcelImage(settings.logo.path)
            original_height = max(logo.height, 1)
            logo.height = 28
            logo.width = min(100, int(logo.width * 28 / original_height))
            sheet.add_image(logo, f"{last_column}1")
        except (FileNotFoundError, OSError, ValueError):
            pass
    sheet.row_dimensions[1].height = 34
    sheet.merge_cells(f"A2:{last_column}2")
    sheet["A2"] = safe_text(title)
    sheet["A2"].font = Font(size=15, bold=True, color=TEAL)
    sheet.merge_cells(f"A3:{last_column}3")
    sheet["A3"] = safe_text(subtitle)
    contacts = " · ".join(filter(None, [
        settings.address, settings.phone_1, settings.phone_2,
        settings.email, settings.website, settings.telegram,
    ]))
    sheet.merge_cells(f"A4:{last_column}4")
    sheet["A4"] = safe_text(contacts)
    sheet["A4"].font = Font(color="7B8999")
    sheet.merge_cells(f"A5:{last_column}5")
    sheet["A5"] = safe_text(settings.additional_text)
    sheet["A5"].alignment = Alignment(wrap_text=True)
    sheet.merge_cells(f"A6:{last_column}6")
    sheet["A6"] = f"Дата формирования: {timezone.localtime().strftime('%d.%m.%Y %H:%M')}"
    sheet["A6"].font = Font(size=10, color="7B8999")
    sheet.freeze_panes = f"A{HEADER_ROW + 1}"
    sheet.auto_filter.ref = f"A{HEADER_ROW}:{last_column}{HEADER_ROW}"
    return workbook, sheet


def _style_table(sheet, headers, *, editable_columns=(), editable_end_row=None, header_row=HEADER_ROW):
    thin = Side(style="thin", color=GRAY)
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(header_row, column, header)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    for column in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 22
    sheet.column_dimensions["A"].width = 62
    for row in sheet.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.border = Border(bottom=thin)
            is_editable = cell.column in editable_columns and (
                editable_end_row is None or cell.row <= editable_end_row
            )
            cell.protection = Protection(locked=not is_editable)
        if row and (row[0].row - header_row) % 2 == 0:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="F8FAFB")
        for column in editable_columns:
            if editable_end_row is None or row[0].row <= editable_end_row:
                row[column - 1].fill = PatternFill("solid", fgColor=LIGHT)
    sheet.protection.sheet = True
    sheet.protection.autoFilter = False
    sheet.protection.sort = False


def _meta_sheet(workbook, *, document_type, values, rows=()):
    meta = workbook.create_sheet(META_SHEET)
    pairs = {"document_type": document_type, "template_version": TEMPLATE_VERSION, **values}
    for index, (key, value) in enumerate(pairs.items(), 1):
        meta.cell(index, 1, key)
        meta.cell(index, 2, str(value))
    start = len(pairs) + 2
    meta.cell(start, 1, "row_number")
    meta.cell(start, 2, "product_type")
    meta.cell(start, 3, "product_id")
    meta.cell(start, 4, "article")
    for offset, row in enumerate(rows, 1):
        for column, value in enumerate(row, 1):
            meta.cell(start + offset, column, value)
    meta.sheet_state = "veryHidden"
    meta.protection.sheet = True


def _bytes(workbook):
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _norm(value):
    return " ".join(str(value or "").casefold().replace("ё", "е").split())


def _tech_rows(warehouse, *, price_field):
    stocks = TechWarehouseStock.objects.filter(
        warehouse=warehouse, quantity__gt=0, tech__is_archived=False,
    ).select_related("tech__brand", "tech__product_type")
    rows, skipped = [], 0
    for stock in stocks:
        product = stock.tech
        price = getattr(product, price_field)
        if price is None:
            skipped += 1
            continue
        sort_key, group = tech_price_order(product, for_excel=True)
        rows.append((sort_key, group, ("tech", product, stock.quantity, price)))
    rows.sort(key=lambda row: row[0])
    groups = []
    for _key, title, item in rows:
        if not groups or groups[-1][0] != title: groups.append((title, []))
        groups[-1][1].append(item)
    return groups, skipped


def _cd_rows(warehouse, *, price_field):
    stocks = CDWarehouseStock.objects.filter(
        warehouse=warehouse, quantity__gt=0, cd__is_archived=False,
    ).select_related("cd__platform")
    entries, skipped = [], 0
    for stock in stocks:
        product = stock.cd
        price = getattr(product, price_field)
        if price is None:
            skipped += 1
            continue
        entries.append((platform_order_key(product.platform if product.platform_id else ''), product, stock.quantity, price))
    entries.sort(key=lambda row: (row[0], _norm(row[1].name), row[1].pk))
    groups = []
    for platform, product, quantity, price in entries:
        title = f"CD · Платформа: {product.platform.name if product.platform_id else 'Без платформы'}"
        if not groups or groups[-1][0] != title:
            groups.append((title, []))
        groups[-1][1].append(("cd", product, quantity, price))
    return groups, skipped


def _warehouse_price_groups(warehouse, *, price_field):
    """Единая структура для розничного и оптового прайса: Tech перед CD."""
    tech_groups, tech_skipped = _tech_rows(warehouse, price_field=price_field)
    cd_groups, cd_skipped = _cd_rows(warehouse, price_field=price_field)
    return tech_groups + cd_groups, tech_skipped + cd_skipped


DOT = "\u00b7"
SECTION_FILLS = {
    f"Tech {DOT} VR-шлемы": "DDEEEA",
    f"Tech {DOT} Sony": "D9EAF7",
    f"Tech {DOT} Nintendo": "FCE4E4",
    f"Tech {DOT} \u041f\u043e\u0440\u0442\u0430\u0442\u0438\u0432\u043d\u044b\u0435 \u043a\u043e\u043d\u0441\u043e\u043b\u0438": "E8DDF5",
    f"Tech {DOT} \u041f\u0440\u043e\u0447\u0438\u0435 \u0442\u043e\u0432\u0430\u0440\u044b": "F9EAD8",
    f"CD {DOT} \u041f\u043b\u0430\u0442\u0444\u043e\u0440\u043c\u0430: Nintendo Switch 2": "FCE4E4",
    f"CD {DOT} \u041f\u043b\u0430\u0442\u0444\u043e\u0440\u043c\u0430: PlayStation 4": "E3F0FA",
    f"CD {DOT} \u041f\u043b\u0430\u0442\u0444\u043e\u0440\u043c\u0430: PlayStation 5": "D9EAF7",
    f"CD {DOT} ": "EEEAF7",
}


def _group_fill(title):
    for prefix, color in SECTION_FILLS.items():
        if title.startswith(prefix):
            return color
    return GROUP_FILL


def _style_group_rows(sheet, rows, *, columns):
    thin = Side(style="thin", color="8A9AA8")
    strong = Side(style="medium", color="334660")
    for row_number in rows:
        title = str(sheet.cell(row_number, 1).value or "")
        fill = _group_fill(title)
        for column in range(1, columns + 1):
            cell = sheet.cell(row_number, column)
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.border = Border(top=strong, bottom=thin)
            cell.protection = Protection(locked=True)
        sheet.cell(row_number, 1).font = Font(bold=True, color=BLUE, size=11)


def generate_retail_price_xlsx(*, warehouse_id, actor):
    try:
        warehouse = Warehouse.objects.get(pk=warehouse_id)
    except Warehouse.DoesNotExist as exc:
        raise ValidationError("Выберите существующий склад.") from exc
    groups, skipped = _warehouse_price_groups(warehouse, price_field="avito_price")
    workbook, sheet = _base_workbook(
        title="Розничный прайс", subtitle=f"Склад: {warehouse.name}", columns=2
    )
    if skipped:
        sheet["A7"] = f"Не включено товаров без розничной цены: {skipped}"
        sheet["A7"].font = Font(color="C94F55", italic=True)
    excel_row = HEADER_ROW + 1
    group_rows = []
    product_count = 0
    for group_title, products in groups:
        group_rows.append(excel_row)
        sheet.cell(excel_row, 1, safe_text(group_title))
        excel_row += 1
        for _kind, product, _quantity, price in products:
            sheet.cell(excel_row, 1, safe_text(product.name))
            sheet.cell(excel_row, 2, price).number_format = MONEY_FORMAT
            product_count += 1
            excel_row += 1
    _style_table(sheet, ("Товар", "Цена"))
    _style_group_rows(sheet, group_rows, columns=2)
    logger.info(
        "Розничный прайс сформирован: user_id=%s warehouse_id=%s rows=%s skipped=%s",
        actor.pk, warehouse.pk, product_count, skipped,
    )
    return _bytes(workbook), skipped


def generate_wholesale_price_xlsx(*, warehouse_id, actor):
    try:
        warehouse = Warehouse.objects.get(pk=warehouse_id)
    except Warehouse.DoesNotExist as exc:
        raise ValidationError("Выберите существующий склад.") from exc
    groups, _skipped = _warehouse_price_groups(warehouse, price_field="wholesale_price")
    workbook, sheet = _base_workbook(
        title="Оптовый прайс", subtitle=f"Склад: {warehouse.name}", columns=4
    )
    metadata = []
    excel_row = HEADER_ROW + 1
    group_rows = []
    product_rows = []
    new_product_rows = []
    generated_at = timezone.now()
    for group_title, products in groups:
        group_rows.append(excel_row)
        sheet.cell(excel_row, 1, safe_text(group_title))
        excel_row += 1
        for kind, product, quantity, price in products:
            product_rows.append(excel_row)
            sheet.cell(excel_row, 1, safe_text(customer_product_name(product.name, product, at=generated_at)))
            sheet.cell(excel_row, 2, price).number_format = MONEY_FORMAT
            sheet.cell(excel_row, 3, quantity)
            sheet.cell(excel_row, 4, 0)
            metadata.append((excel_row, kind, product.pk, safe_text(product.sku)))
            if product.created_at and product.created_at <= generated_at and product.newness_until and generated_at < product.newness_until:
                new_product_rows.append(excel_row)
            excel_row += 1
    end = max(HEADER_ROW + 1, excel_row - 1)
    totals = end + 2
    sheet.cell(totals, 3, "Выбрано позиций")
    sheet.cell(totals, 4, f'=COUNTIF(D{HEADER_ROW + 1}:D{end},">0")')
    sheet.cell(totals + 1, 3, "Заказано единиц")
    sheet.cell(totals + 1, 4, f"=SUM(D{HEADER_ROW + 1}:D{end})")
    sheet.cell(totals + 2, 3, "Итоговая стоимость")
    sheet.cell(totals + 2, 4, f"=SUMPRODUCT(B{HEADER_ROW + 1}:B{end},D{HEADER_ROW + 1}:D{end})")
    sheet.cell(totals + 2, 4).number_format = MONEY_FORMAT
    validation = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    sheet.add_data_validation(validation)
    for product_row in product_rows:
        validation.add(sheet.cell(product_row, 4))
    _style_table(
        sheet,
        ("Товар", "Оптовая цена", "Доступно для заказа", "Количество к заказу"),
        editable_columns=(4,),
        editable_end_row=end,
    )
    _style_group_rows(sheet, group_rows, columns=4)
    for row in new_product_rows:
        for column in range(1, 5):
            sheet.cell(row, column).fill = PatternFill('solid', fgColor='E2F0D9')
    _meta_sheet(workbook, document_type="resource_wholesale", values={
        "warehouse_id": warehouse.pk, "generated_at": timezone.now().isoformat(), "main_sheet": MAIN_SHEET,
    }, rows=metadata)
    logger.info(
        "Оптовый прайс сформирован: user_id=%s warehouse_id=%s rows=%s",
        actor.pk, warehouse.pk, len(product_rows),
    )
    return _bytes(workbook)


def _supplier_text_height(value, *, width, minimum=28, line_height=16):
    lines = sum(
        max(1, len(wrap(line, width=width, break_on_hyphens=False)))
        for line in str(value or "").split("\n")
    )
    return max(minimum, lines * line_height + 10)


def _style_supplier_template(sheet, *, group_rows, product_kind=None):
    """Показываем название, CUSA/PPSA и цену, сохраняя служебные колонки."""
    for column, width in (("A", 7), ("B", 9), ("C", 22)):
        sheet.column_dimensions[column].width = width
        sheet.column_dimensions[column].hidden = True
    sheet.column_dimensions["D"].width = 78
    sheet.column_dimensions["E"].width = 18
    sheet.column_dimensions["E"].hidden = product_kind == "tech"
    sheet.column_dimensions["F"].width = 22

    sheet.row_dimensions[SUPPLIER_HEADER_ROW].height = 30
    for row in range(SUPPLIER_HEADER_ROW + 1, sheet.max_row + 1):
        if row in group_rows:
            title = group_rows[row]
            for cell in sheet[row]:
                cell.fill = PatternFill("solid", fgColor=_group_fill(title))
                cell.protection = Protection(locked=True)
                cell.border = Border(top=Side(style="thin", color=TEAL), bottom=Side(style="thin", color=GRAY))
            sheet.merge_cells(start_row=row, start_column=4, end_row=row, end_column=6)
            sheet.cell(row, 4).font = Font(name="Calibri", size=11, bold=True, color=BLUE)
            sheet.cell(row, 4).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            sheet.row_dimensions[row].height = _supplier_text_height(title, width=100, minimum=32)
            continue
        name = sheet.cell(row, 4)
        name.font = Font(name="Calibri", size=11, color=BLUE)
        name.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        code = sheet.cell(row, 5)
        code.font = Font(name="Calibri", size=11, color=BLUE)
        code.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        code.number_format = "@"
        sheet.row_dimensions[row].height = max(
            _supplier_text_height(name.value, width=64),
            _supplier_text_height(code.value, width=16),
        )
        price = sheet.cell(row, 6)
        price.font = Font(name="Calibri", size=11, bold=True, color=TEAL)
        price.alignment = Alignment(horizontal="right", vertical="center", wrap_text=False)
        price.number_format = "#,##0.00####"
        price.border = Border(
            left=Side(style="thin", color=TEAL), bottom=Side(style="thin", color=GRAY),
        )

    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.topLeftCell = "D1"
    sheet.freeze_panes = f"A{SUPPLIER_HEADER_ROW + 1}"
    first_product_row = next(
        (row for row in range(SUPPLIER_HEADER_ROW + 1, sheet.max_row + 1) if row not in group_rows),
        SUPPLIER_HEADER_ROW + 1,
    )
    for selection in sheet.sheet_view.selection:
        selection.activeCell = f"F{first_product_row}"
        selection.sqref = f"F{first_product_row}"
    sheet.auto_filter.ref = f"A{SUPPLIER_HEADER_ROW}:F{sheet.max_row}"
    sheet.print_area = f"D1:F{sheet.max_row}"
    sheet.print_title_rows = f"1:{SUPPLIER_HEADER_ROW}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0


def generate_supplier_template(*, actor, product_kind=None, supplier_id=None):
    if product_kind not in (None, "cd", "tech"):
        raise ValidationError("Выберите диски или технику.")
    prices = {"cd": {}, "tech": {}}
    if supplier_id is not None:
        for kind, model, field in (
            ("cd", SupplierCDPrice, "cd_id"), ("tech", SupplierTechPrice, "tech_id"),
        ):
            if product_kind is None or product_kind == kind:
                prices[kind] = dict(model.objects.filter(
                    supplier_id=supplier_id, price__gt=0,
                ).values_list(field, "price"))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = MAIN_SHEET
    groups = nomenclature_product_groups(
        Tech.objects.active().filter(exclude_from_supplier_template=False).select_related("brand", "product_type")
        if product_kind != "cd" else [],
        CD.objects.active().filter(exclude_from_supplier_template=False).select_related("platform", "game_series")
        if product_kind != "tech" else [],
    )
    metadata = []
    group_rows = {}
    excel_row = SUPPLIER_HEADER_ROW + 1
    for kind, title, products in groups:
        group_rows[excel_row] = title
        sheet.cell(excel_row, 4, safe_text(title))
        excel_row += 1
        for product in products:
            values = (
                kind.upper(), product.pk, safe_text(product.sku), safe_text(product.name),
                safe_text(product.cusa_ppsa_code) if kind == "cd" else None,
                prices[kind].get(product.pk),
            )
            for column, value in enumerate(values, 1):
                sheet.cell(excel_row, column, value)
            sheet.cell(excel_row, 6).number_format = AED_FORMAT
            metadata.append((excel_row, kind, product.pk, safe_text(product.sku)))
            excel_row += 1
    _style_table(
        sheet, ("Тип", "ID", "Артикул", "Название", "CUSA/PPSA", "Цена AED"),
        editable_columns=(6,), header_row=SUPPLIER_HEADER_ROW,
    )
    _style_supplier_template(sheet, group_rows=group_rows, product_kind=product_kind)
    _meta_sheet(workbook, document_type="resource_supplier", values={
        "generated_at": timezone.now().isoformat(), "main_sheet": MAIN_SHEET,
        "header_row": SUPPLIER_HEADER_ROW,
        "price_column": 6,
        "product_kind": product_kind or "",
    }, rows=metadata)
    logger.info(
        "Шаблон поставщика сформирован: user_id=%s supplier_id=%s rows=%s",
        actor.pk, supplier_id, len(metadata),
    )
    return _bytes(workbook)


def generate_procurement_customer_xlsx(*, price_list_id, actor, product_kind=None):
    if product_kind not in (None, "cd", "tech"):
        raise ValidationError("Выберите диски или технику.")
    try:
        price_list = ProcurementPriceList.objects.prefetch_related("items").get(pk=price_list_id)
    except ProcurementPriceList.DoesNotExist as exc:
        raise ValidationError("Закупочный прайс не найден.") from exc
    workbook, sheet = _base_workbook(
        title="Закупочный прайс" + ({"cd": " · Диски", "tech": " · Техника"}.get(product_kind, "")),
        subtitle="Укажите количество по предоплате или постоплате", columns=5,
    )
    metadata = []
    rows = [
        item for item in price_list.items.select_related("cd__platform", "cd__game_series", "tech__brand", "tech__product_type")
        if not item.product.is_archived and (product_kind is None or item.product_kind == product_kind)
    ]
    groups = nomenclature_product_groups(
        [item for item in rows if item.product_kind == "tech"],
        [item for item in rows if item.product_kind == "cd"],
        product_of=lambda item: item.product,
    )
    excel_row = HEADER_ROW + 1
    group_rows, product_rows, new_product_rows = [], [], []
    generated_at = timezone.now()
    for _kind, title, group_items in groups:
        if title:
            sheet.cell(excel_row, 1, safe_text(title))
            group_rows.append(excel_row)
            excel_row += 1
        for item in group_items:
            values = (
                safe_text(customer_product_name(item.product_name_snapshot, item.product, at=generated_at)), item.prepayment_price_rub, 0,
                item.postpayment_price_rub, 0,
            )
            for column, value in enumerate(values, 1):
                sheet.cell(excel_row, column, value)
            sheet.cell(excel_row, 2).number_format = MONEY_FORMAT
            sheet.cell(excel_row, 4).number_format = MONEY_FORMAT
            metadata.append((excel_row, item.product_kind, item.product_id, safe_text(item.article_snapshot)))
            product_rows.append(excel_row)
            if item.product.created_at and item.product.created_at <= generated_at and item.product.newness_until and generated_at < item.product.newness_until:
                new_product_rows.append(excel_row)
            excel_row += 1
    end = max(HEADER_ROW + 1, excel_row - 1)
    totals = end + 2
    formulas = (
        ("Количество позиций", f'=SUMPRODUCT(--((C{HEADER_ROW + 1}:C{end}+E{HEADER_ROW + 1}:E{end})>0))'),
        ("Единиц предоплата", f"=SUM(C{HEADER_ROW + 1}:C{end})"),
        ("Единиц постоплата", f"=SUM(E{HEADER_ROW + 1}:E{end})"),
        ("Всего единиц", f"=SUM(C{HEADER_ROW + 1}:C{end})+SUM(E{HEADER_ROW + 1}:E{end})"),
        ("Сумма предоплаты", f"=SUMPRODUCT(B{HEADER_ROW + 1}:B{end},C{HEADER_ROW + 1}:C{end})"),
        ("Сумма постоплаты", f"=SUMPRODUCT(D{HEADER_ROW + 1}:D{end},E{HEADER_ROW + 1}:E{end})"),
        ("Итого", f"=SUMPRODUCT(B{HEADER_ROW + 1}:B{end},C{HEADER_ROW + 1}:C{end})+SUMPRODUCT(D{HEADER_ROW + 1}:D{end},E{HEADER_ROW + 1}:E{end})"),
    )
    for offset, (label, formula) in enumerate(formulas):
        sheet.cell(totals + offset, 4, label)
        sheet.cell(totals + offset, 5, formula)
        if offset >= 4:
            sheet.cell(totals + offset, 5).number_format = MONEY_FORMAT
    validation = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    sheet.add_data_validation(validation)
    for row in product_rows:
        validation.add(f"C{row}")
        validation.add(f"E{row}")
    _style_table(
        sheet,
        ("Товар", "Цена предоплата", "Количество предоплата", "Цена постоплата", "Количество постоплата"),
        editable_columns=(3, 5),
        editable_end_row=end,
    )
    _style_group_rows(sheet, group_rows, columns=5)
    for row in new_product_rows:
        for column in range(1, 6):
            sheet.cell(row, column).fill = PatternFill('solid', fgColor='E2F0D9')
    for row in group_rows:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        sheet.cell(row, 1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        sheet.row_dimensions[row].height = wrapped_text_height(
            sheet.cell(row, 1).value, 130, font_size=11, min_height=28,
        )
    sheet.row_dimensions[HEADER_ROW].height = 38
    sheet.column_dimensions["A"].width = 70
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["E"].width = 16
    for row in product_rows:
        name_cell = sheet.cell(row, 1)
        name_cell.alignment = Alignment(vertical="center", wrap_text=True)
        name_font = copy(name_cell.font)
        name_font.sz = 13
        name_cell.font = name_font
        sheet.row_dimensions[row].height = wrapped_text_height(
            name_cell.value, 70, font_name=name_font.name or "Calibri", font_size=13,
            padding_points=4, min_height=24,
        )
        for column in (2, 3, 4, 5):
            sheet.cell(row, column).alignment = Alignment(horizontal="right", vertical="center")
    sheet.sheet_view.showGridLines = False
    first_product_row = product_rows[0] if product_rows else HEADER_ROW + 1
    for selection in sheet.sheet_view.selection:
        selection.activeCell = f"C{first_product_row}"
        selection.sqref = f"C{first_product_row}"
    _meta_sheet(workbook, document_type="resource_procurement", values={
        "price_list_token": price_list.public_token, "generated_at": timezone.now().isoformat(),
        "main_sheet": MAIN_SHEET,
        "product_kind": product_kind or "",
    }, rows=metadata)
    logger.info(
        "Клиентский закупочный прайс сформирован: user_id=%s price_list_id=%s rows=%s",
        actor.pk, price_list.pk, len(rows),
    )
    return _bytes(workbook)


LEGACY_STOCK_HEADERS = (
    "Товар", "Артикул", "Закупочная база", "Доставка",
    "Себестоимость с доставкой", "Количество к заказу", "Сумма",
)
STOCK_HEADERS = (
    "Товар", "Остаток всех складов", "Текущая себестоимость", "Артикул", "Закупочная база", "Доставка",
    "Себестоимость с доставкой", "Количество к заказу", "Сумма",
)


def generate_procurement_stock_xlsx(*, price_list_id, actor):
    try:
        price_list = ProcurementPriceList.objects.get(pk=price_list_id)
    except ProcurementPriceList.DoesNotExist as exc:
        raise ValidationError("Закупочный прайс не найден.") from exc
    workbook, sheet = _base_workbook(
        title=f"Заказ на склад · Прайс №{price_list.pk}",
        subtitle="Заполните количество к заказу. Остаток и текущая себестоимость — на дату формирования.", columns=9,
    )
    items = {(item.product_kind, item.product_id): item for item in price_list.items.select_related(
        "cd__platform", "cd__game_series", "tech__brand", "tech__product_type"
    ) if not item.product.is_archived}
    stock_quantities = {}
    for kind, model, field in (("cd", CDWarehouseStock, "cd_id"), ("tech", TechWarehouseStock, "tech_id")):
        product_ids = [product_id for product_kind, product_id in items if product_kind == kind]
        stock_quantities[kind] = dict(model.objects.filter(**{f"{field}__in": product_ids}).values(field).annotate(
            total=Sum("quantity"),
        ).values_list(field, "total"))
    groups = nomenclature_product_groups(
        [item.tech for item in items.values() if item.product_kind == "tech"],
        [item.cd for item in items.values() if item.product_kind == "cd"],
    )
    metadata, group_rows, product_rows = [], [], []
    row = HEADER_ROW + 1
    for kind, title, products in groups:
        sheet.cell(row, 1, safe_text(title))
        group_rows.append(row)
        row += 1
        for product in products:
            item = items[(kind, product.pk)]
            values = (safe_text(item.product_name_snapshot), stock_quantities[kind].get(product.pk, 0), product.cost,
                      safe_text(item.article_snapshot), item.base_price_rub, item.delivery_rub,
                      item.stock_unit_cost, 0, f"=G{row}*H{row}")
            for column, value in enumerate(values, 1):
                sheet.cell(row, column, value)
            sheet.cell(row, 2).number_format = "0"
            sheet.cell(row, 4).number_format = "@"
            for column in (3, 5, 6, 7, 9):
                sheet.cell(row, column).number_format = MONEY_FORMAT
            metadata.append((row, kind, product.pk, safe_text(item.article_snapshot)))
            product_rows.append(row)
            row += 1
    end = max(HEADER_ROW + 1, row - 1)
    sheet.cell(end + 2, 7, "Всего единиц")
    sheet.cell(end + 2, 8, f"=SUM(H{HEADER_ROW + 1}:H{end})")
    sheet.cell(end + 3, 7, "Итого с доставкой")
    sheet.cell(end + 3, 9, f"=SUM(I{HEADER_ROW + 1}:I{end})").number_format = MONEY_FORMAT
    _style_table(sheet, STOCK_HEADERS, editable_columns=(8,), editable_end_row=end)
    _style_group_rows(sheet, group_rows, columns=9)
    for group_row in group_rows:
        sheet.merge_cells(start_row=group_row, start_column=1, end_row=group_row, end_column=9)
        sheet.row_dimensions[group_row].height = wrapped_text_height(
            sheet.cell(group_row, 1).value, 130, font_size=11, min_height=30,
        )
    validation = DataValidation(type="whole", operator="between", formula1="0", formula2="2147483647",
                                allow_blank=True)
    validation.showErrorMessage = True
    validation.error = "Введите целое неотрицательное количество."
    sheet.add_data_validation(validation)
    for product_row in product_rows:
        validation.add(f"H{product_row}")
        sheet.cell(product_row, 1).alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[product_row].height = max(
            30, (len(wrap(str(sheet.cell(product_row, 1).value), width=48)) or 1) * 16 + 8,
        )
    for column, width in (("A", 52), ("B", 14), ("C", 20), ("D", 20), ("E", 21), ("F", 18),
                          ("G", 25), ("H", 22), ("I", 22)):
        sheet.column_dimensions[column].width = width
    for column in ("D", "E", "F"):
        sheet.column_dimensions[column].hidden = True
    sheet.row_dimensions[HEADER_ROW].height = 42
    sheet.sheet_view.showGridLines = False
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = f"1:{HEADER_ROW}"
    sheet.print_area = f"A1:I{end + 3}"
    _meta_sheet(workbook, document_type="resource_stock_procurement", values={
        "price_list_token": price_list.public_token, "generated_at": timezone.now().isoformat(),
        "main_sheet": MAIN_SHEET,
    }, rows=metadata)
    logger.info("Прайс для заказа на склад сформирован: user_id=%s price_list_id=%s rows=%s",
                actor.pk, price_list.pk, len(items))
    return _bytes(workbook)


def import_stock_procurement_order_xlsx(upload):
    workbook, sheet, values, technical_rows = _load_uploaded_workbook(upload, "resource_stock_procurement")
    try:
        is_legacy = tuple(sheet.cell(HEADER_ROW, column).value for column in range(1, 8)) == LEGACY_STOCK_HEADERS
        _require_headers(sheet, LEGACY_STOCK_HEADERS if is_legacy else STOCK_HEADERS)
        article_column, quantity_column = (2, 6) if is_legacy else (4, 8)
        try:
            price_list = ProcurementPriceList.objects.prefetch_related("items").get(
                public_token=values.get("price_list_token"),
            )
        except (ProcurementPriceList.DoesNotExist, ValueError, ValidationError) as exc:
            raise ValidationError("Версия закупочного прайса не найдена.") from exc
        items = {(item.product_kind, item.product_id): item for item in price_list.items.all()}
        lines, seen, seen_rows = [], set(), set()
        for excel_row, kind, product_id, _article in technical_rows:
            try:
                excel_row, product_id = int(excel_row), int(product_id)
            except (TypeError, ValueError) as exc:
                raise ValidationError("Технические идентификаторы файла повреждены.") from exc
            key = (str(kind), product_id)
            if key in seen or excel_row in seen_rows or not HEADER_ROW < excel_row <= sheet.max_row:
                raise ValidationError("Технические строки файла повреждены или повторяются.")
            seen.add(key)
            seen_rows.add(excel_row)
            item = items.get(key)
            if item is None:
                raise ValidationError(f"Строка {excel_row}: товар отсутствует в сохранённом прайсе.")
            if (sheet.cell(excel_row, 1).value != safe_text(item.product_name_snapshot)
                    or str(sheet.cell(excel_row, article_column).value or "") != safe_text(item.article_snapshot)
                    or str(_article or "") != safe_text(item.article_snapshot)):
                raise ValidationError(f"Строка {excel_row}: название, артикул или порядок товаров изменён. Заполняйте только количество.")
            quantity = _quantity(sheet.cell(excel_row, quantity_column).value, excel_row=excel_row)
            if quantity > 2147483647:
                raise ValidationError(f"Строка {excel_row}: слишком большое количество.")
            if quantity:
                lines.append({"price_list_item_id": item.pk, "prepayment_quantity": quantity,
                              "postpayment_quantity": 0})
        visible_rows = {row for row in range(HEADER_ROW + 1, sheet.max_row + 1)
                        if sheet.cell(row, 1).value not in (None, "")
                        and (sheet.cell(row, article_column).value not in (None, "")
                             or sheet.cell(row, quantity_column).value is not None)}
        if visible_rows != seen_rows:
            raise ValidationError("Состав видимого и технического листов не совпадает. Сохраните строки исходного шаблона.")
        if not lines:
            raise ValidationError("В файле не выбрано ни одной товарной позиции.")
        # Цены из Excel не принимаются: заказ копирует стоимость из сохранённого прайса.
        return price_list, lines
    finally:
        workbook.close()


def _load_uploaded_workbook(upload, expected_type):
    filename = Path(getattr(upload, "name", "")).name
    if Path(filename).suffix.lower() != ".xlsx":
        raise ValidationError("Разрешены только файлы .xlsx.")
    try:
        workbook = load_workbook(
            upload, data_only=False, read_only=False, keep_vba=False, keep_links=False
        )
    except Exception as exc:
        raise ValidationError("Не удалось прочитать XLSX-файл.") from exc
    if META_SHEET not in workbook.sheetnames or MAIN_SHEET not in workbook.sheetnames:
        workbook.close()
        raise ValidationError("Файл не является документом ReSOURCE: отсутствует технический лист.")
    meta = workbook[META_SHEET]
    values = {}
    row = 1
    while meta.cell(row, 1).value not in (None, ""):
        values[str(meta.cell(row, 1).value)] = str(meta.cell(row, 2).value or "")
        row += 1
    if values.get("document_type") != expected_type:
        workbook.close()
        raise ValidationError("Загружен документ другого типа.")
    if values.get("template_version") != TEMPLATE_VERSION:
        workbook.close()
        raise ValidationError("Версия шаблона не поддерживается.")
    header_row = row + 1
    technical_rows = []
    row = header_row + 1
    while meta.cell(row, 1).value not in (None, ""):
        technical_rows.append((
            meta.cell(row, 1).value, meta.cell(row, 2).value,
            meta.cell(row, 3).value, meta.cell(row, 4).value,
        ))
        row += 1
    return workbook, workbook[MAIN_SHEET], values, technical_rows


def _require_headers(sheet, expected, *, header_row=HEADER_ROW):
    actual = tuple(sheet.cell(header_row, column).value for column in range(1, len(expected) + 1))
    if actual != tuple(expected):
        raise ValidationError("Структура видимого листа изменена: заголовки не соответствуют шаблону.")


def _quantity(value, *, excel_row):
    if value in (None, ""):
        return 0
    if isinstance(value, bool) or isinstance(value, str):
        raise ValidationError(f"Строка {excel_row}: количество должно быть целым числом, не формулой или текстом.")
    try:
        number = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise ValidationError(f"Строка {excel_row}: некорректное количество.") from exc
    if number < 0 or number != number.to_integral_value():
        raise ValidationError(f"Строка {excel_row}: количество должно быть целым числом не меньше нуля.")
    return int(number)


def import_wholesale_price_to_sale(upload):
    workbook, sheet, values, technical_rows = _load_uploaded_workbook(upload, "resource_wholesale")
    try:
        _require_headers(sheet, ("Товар", "Оптовая цена", "Доступно для заказа", "Количество к заказу"))
        warehouse = Warehouse.objects.get(pk=int(values.get("warehouse_id", "")))
        lines = []
        seen = set()
        for excel_row, kind, product_id, _article in technical_rows:
            try:
                excel_row, product_id = int(excel_row), int(product_id)
            except (TypeError, ValueError) as exc:
                raise ValidationError("Технические идентификаторы файла повреждены.") from exc
            key = (str(kind), product_id)
            if key in seen:
                raise ValidationError("В техническом листе обнаружен повтор товара.")
            seen.add(key)
            quantity = _quantity(sheet.cell(excel_row, 4).value, excel_row=excel_row)
            if not quantity:
                continue
            if kind == "cd":
                product = CD.objects.active().filter(pk=product_id).first()
                stock = CDWarehouseStock.objects.filter(warehouse=warehouse, cd_id=product_id).first()
            elif kind == "tech":
                product = Tech.objects.active().filter(pk=product_id).first()
                stock = TechWarehouseStock.objects.filter(warehouse=warehouse, tech_id=product_id).first()
            else:
                raise ValidationError(f"Строка {excel_row}: неизвестный тип товара.")
            if product is None:
                raise ValidationError(f"Строка {excel_row}: товар больше не существует.")
            available = stock.quantity if stock else 0
            if quantity > available:
                raise ValidationError(
                    f"{product.name}: запрошено {quantity} шт., доступно {available} шт."
                )
            if product.wholesale_price is None:
                raise ValidationError(f"Для товара «{product.name}» больше не задана оптовая цена.")
            lines.append({
                "product_type": kind, "product_id": product_id, "quantity": quantity,
                "label": f"{kind.upper()} — {product.name}",
                "unit_price": f"{product.wholesale_price:.2f}",
            })
        if not lines:
            raise ValidationError("В файле не выбрано ни одной товарной позиции.")
        return warehouse, lines
    except (Warehouse.DoesNotExist, TypeError, ValueError) as exc:
        raise ValidationError("Склад из документа не найден.") from exc
    finally:
        workbook.close()


@transaction.atomic
def import_supplier_price(*, supplier_id, upload, actor, product_kind=None):
    if product_kind not in (None, "cd", "tech"):
        raise ValidationError("Выберите диски или технику.")
    workbook, sheet, values, technical_rows = _load_uploaded_workbook(upload, "resource_supplier")
    try:
        file_kind = values.get("product_kind", "")
        if file_kind not in ("", "cd", "tech") or (file_kind and product_kind and file_kind != product_kind):
            raise ValidationError("Тип загруженного прайса не соответствует полю: диски или техника.")
        scope = product_kind or file_kind or None
        # Ранее скачанные шаблоны имели титульную часть и заголовки в строке 8.
        header_row = values.get("header_row", str(HEADER_ROW))
        if header_row not in (str(SUPPLIER_HEADER_ROW), str(HEADER_ROW)):
            raise ValidationError("Структура шаблона поставщика не поддерживается.")
        price_column = values.get("price_column", "5")
        if price_column not in ("5", "6"):
            raise ValidationError("Структура шаблона поставщика не поддерживается.")
        headers = ("Тип", "ID", "Артикул", "Название")
        if price_column == "6":
            headers += ("CUSA/PPSA",)
        _require_headers(
            sheet, headers + ("Цена AED",), header_row=int(header_row),
        )
        supplier = Supplier.objects.select_for_update().get(pk=supplier_id)
        prepared = {"cd": [], "tech": []}
        products = {
            (kind, product_id): archived
            for kind, model in (("cd", CD), ("tech", Tech))
            if not file_kind or kind == file_kind
            for product_id, archived in model.objects.values_list("id", "is_archived")
        }
        seen = set()
        seen_rows = set()
        included = {"cd": [], "tech": []}
        archived_count = 0
        for excel_row, kind, product_id, article in technical_rows:
            try:
                excel_row, product_id = int(excel_row), int(product_id)
            except (TypeError, ValueError) as exc:
                raise ValidationError("Технические идентификаторы файла повреждены.") from exc
            key = (str(kind), product_id)
            if key in seen:
                raise ValidationError("В техническом листе обнаружен повтор товара.")
            seen.add(key)
            if kind not in ("cd", "tech") or (file_kind and kind != file_kind):
                raise ValidationError("В шаблоне обнаружен товар другого типа.")
            if excel_row <= int(header_row) or excel_row > sheet.max_row or excel_row in seen_rows:
                raise ValidationError("Технические номера строк файла повреждены.")
            seen_rows.add(excel_row)
            if (
                sheet.cell(excel_row, 1).value != kind.upper()
                or sheet.cell(excel_row, 2).value != product_id
                or str(sheet.cell(excel_row, 3).value or "") != str(article or "")
            ):
                raise ValidationError(
                    f"Строка {excel_row}: товар не соответствует техническому листу. "
                    "Сохраните строки шаблона и заполните только колонку «Цена AED»."
                )
            if key not in products:
                raise ValidationError(f"Строка {excel_row}: товар не найден.")
            if scope and kind != scope:
                continue
            # Архивирование после скачивания шаблона не мешает загрузке остальных строк.
            if products[key]:
                archived_count += 1
                continue
            included[kind].append(product_id)
            raw_price = sheet.cell(excel_row, int(price_column)).value
            if raw_price in (None, ""):
                continue
            if isinstance(raw_price, bool) or isinstance(raw_price, str):
                raise ValidationError(f"Строка {excel_row}: цена AED должна быть числом, не формулой или текстом.")
            try:
                price = Decimal(str(raw_price)).quantize(Decimal("0.000001"))
            except (InvalidOperation, TypeError, ValueError) as exc:
                raise ValidationError(f"Строка {excel_row}: некорректная цена AED.") from exc
            if price <= 0:
                raise ValidationError(f"Строка {excel_row}: цена AED должна быть больше нуля.")
            prepared[kind].append((product_id, price))
        # Полноту проверяем внутри скачанного документа, а не по изменяемому каталогу.
        visible_rows = {
            row for row in range(int(header_row) + 1, sheet.max_row + 1)
            if any(sheet.cell(row, column).value not in (None, "") for column in (1, 2, 3, int(price_column)))
        }
        if seen_rows != visible_rows:
            raise ValidationError(
                "Структура шаблона поставщика неполна: товарные строки и технический лист не совпадают. "
                "Сохраните строки шаблона и заполните только колонку «Цена AED»."
            )
        for kind, model, field in (("cd", SupplierCDPrice, "cd_id"), ("tech", SupplierTechPrice, "tech_id")):
            if scope is None or kind == scope:
                # Отсутствующие в файле и архивные товары сохраняют прежние предложения.
                model.objects.filter(supplier=supplier, **{f"{field}__in": included[kind]}).delete()
                model.objects.bulk_create([
                    model(supplier=supplier, **{field: product_id}, price=price)
                    for product_id, price in prepared[kind]
                ])
        count = len(prepared["cd"]) + len(prepared["tech"])
        logger.info(
            "Прайс поставщика обновлён: user_id=%s supplier_id=%s kind=%s prices=%s archived_skipped=%s",
            actor.pk, supplier.pk, scope or "all", count, archived_count,
        )
        return count
    except Supplier.DoesNotExist as exc:
        raise ValidationError("Поставщик не найден.") from exc
    finally:
        workbook.close()


@transaction.atomic
def import_supplier_price_files(*, supplier_id, uploads, actor):
    if not uploads or not set(uploads).issubset({"cd", "tech"}):
        raise ValidationError("Выберите прайс дисков, техники или оба файла.")
    return {
        kind: import_supplier_price(supplier_id=supplier_id, upload=upload, actor=actor, product_kind=kind)
        for kind, upload in uploads.items()
    }


def import_procurement_order_xlsx(upload, *, product_kind=None, allow_empty=False):
    if product_kind not in (None, "cd", "tech"):
        raise ValidationError("Выберите диски или технику.")
    workbook, sheet, values, technical_rows = _load_uploaded_workbook(upload, "resource_procurement")
    try:
        file_kind = values.get("product_kind", "")
        if file_kind not in ("", "cd", "tech") or (file_kind and product_kind and file_kind != product_kind):
            raise ValidationError("Тип клиентского прайса не соответствует выбранному полю.")
        scope = product_kind or file_kind or None
        _require_headers(sheet, (
            "Товар", "Цена предоплата", "Количество предоплата",
            "Цена постоплата", "Количество постоплата",
        ))
        try:
            price_list = ProcurementPriceList.objects.prefetch_related("items").get(
                public_token=values.get("price_list_token")
            )
        except (ProcurementPriceList.DoesNotExist, ValueError, ValidationError) as exc:
            raise ValidationError("Версия закупочного прайса не найдена.") from exc
        items = {(item.product_kind, item.product_id): item for item in price_list.items.all()}
        lines = []
        seen = set()
        seen_rows = set()
        for excel_row, kind, product_id, _article in technical_rows:
            try:
                excel_row, product_id = int(excel_row), int(product_id)
            except (TypeError, ValueError) as exc:
                raise ValidationError("Технические идентификаторы файла повреждены.") from exc
            key = (str(kind), product_id)
            if key in seen:
                raise ValidationError("В техническом листе обнаружен повтор товара.")
            seen.add(key)
            if excel_row <= HEADER_ROW or excel_row > sheet.max_row or excel_row in seen_rows:
                raise ValidationError("Технические номера строк файла повреждены.")
            seen_rows.add(excel_row)
            if kind not in ("cd", "tech") or (file_kind and kind != file_kind):
                raise ValidationError("В клиентском прайсе обнаружен товар другого типа.")
            item = items.get(key)
            if item is None:
                raise ValidationError(f"Строка {excel_row}: товар отсутствует в сохранённом прайсе.")
            if (sheet.cell(excel_row, 1).value not in (safe_text(item.product_name_snapshot), safe_text('🌟 ' + item.product_name_snapshot))
                    or str(_article or "") != safe_text(item.article_snapshot)):
                raise ValidationError(f"Строка {excel_row}: товар не соответствует сохранённому прайсу. Заполняйте только количество.")
            if scope and kind != scope:
                continue
            prepayment_quantity = _quantity(sheet.cell(excel_row, 3).value, excel_row=excel_row)
            postpayment_quantity = _quantity(sheet.cell(excel_row, 5).value, excel_row=excel_row)
            if prepayment_quantity or postpayment_quantity:
                lines.append({
                    "price_list_item_id": item.pk,
                    "prepayment_quantity": prepayment_quantity,
                    "postpayment_quantity": postpayment_quantity,
                })
        if not lines and not allow_empty:
            raise ValidationError("В файле не выбрано ни одной товарной позиции.")
        return price_list, lines
    finally:
        workbook.close()


def import_procurement_order_files(uploads, *, with_purpose=False):
    if uploads and "stock" in uploads:
        if len(uploads) != 1:
            raise ValidationError("Заказ на склад загружается отдельно от клиентских файлов.")
        if not with_purpose:
            raise ValidationError("Для складского прайса требуется создание заказа на склад.")
        price_list, lines = import_stock_procurement_order_xlsx(uploads["stock"])
        return price_list, lines, "stock"
    if not uploads or not set(uploads).issubset({"cd", "tech", None}):
        raise ValidationError("Загрузите заполненный клиентский прайс.")
    result_list, result_lines = None, []
    seen_items = set()
    for kind, upload in uploads.items():
        price_list, lines = import_procurement_order_xlsx(upload, product_kind=kind, allow_empty=True)
        if result_list is not None and price_list.pk != result_list.pk:
            raise ValidationError("Прайсы дисков и техники должны относиться к одной версии закупочного прайса.")
        result_list = price_list
        for line in lines:
            if line["price_list_item_id"] in seen_items:
                raise ValidationError("Товар повторяется в загруженных файлах.")
            seen_items.add(line["price_list_item_id"])
            result_lines.append(line)
    if not result_lines:
        raise ValidationError("В файлах не выбрано ни одной товарной позиции.")
    if with_purpose:
        return result_list, result_lines, "customer"
    return result_list, result_lines
