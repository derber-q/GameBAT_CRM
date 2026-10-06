"""Порядок прайса в таблицах CRM и сохранность строк при группировке."""
from html import escape
from io import BytesIO
from types import SimpleNamespace

from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import User
from price.excel import HEADER_ROW, MAIN_SHEET, generate_wholesale_price_xlsx
from warehouse.models import CDWarehouseStock, TechWarehouseStock, Warehouse

from .models import Brand, CD, Platform, ProductType, Tech
from .product_ordering import tech_brand_groups


def create_ordering_fixture():
    warehouse = Warehouse.objects.create(name="Склад проверки порядка")
    sony_types = [
        "Игровая консоль", "Внешний привод", "Зарядная станция", "Геймпад",
        "Запасная часть", "Игровая гарнитура", "VR-шлем", "VR-аксессуар",
        "Стриминговая приставка", "Адаптер", "Кабель",
    ]
    nintendo_types = ["Портативная консоль", "Чехол", "Геймпад", "Игровая гарнитура"]
    definitions = [
        ("Sony", kind, f"Sony item {index}") for index, kind in enumerate(sony_types)
    ] + [
        ("Nintendo", kind, f"Nintendo item {index}") for index, kind in enumerate(nintendo_types)
    ] + [
        ("Anbernic", "Портативная консоль", "Anbernic console"),
        ("Steam Deck", "Портативная консоль", "Steam Deck console"),
        ("Apple", "Ноутбук", "Apple laptop"),
    ]
    # Создание в обратном порядке исключает случайное совпадение сортировки с ID.
    products = {}
    for brand_name, type_name, name in reversed(definitions):
        brand, _ = Brand.objects.get_or_create(name=brand_name)
        kind, _ = ProductType.objects.get_or_create(name=type_name)
        product = Tech.objects.create(
            brand=brand, product_type=kind, name=name,
            cost=100, wholesale_price=200, avito_price=250,
        )
        TechWarehouseStock.objects.create(warehouse=warehouse, tech=product, quantity=2)
        products[name] = product
    ordered = [products[name] for _, _, name in definitions]
    for platform_name in ("PS5", "NS2"):
        platform = Platform.objects.create(name=platform_name)
        game = CD.objects.create(
            platform=platform, name=f"{platform_name} game", cost=100, wholesale_price=200,
        )
        CDWarehouseStock.objects.create(warehouse=warehouse, cd=game, quantity=2)
    cds = list(CD.objects.select_related("platform").order_by("platform__name", "name"))
    return warehouse, ordered, cds


def ordering_pages(warehouse):
    return [
        reverse("warehouse:global_stock"),
        reverse("warehouse:detail", args=[warehouse.pk]),
        reverse("nomenclature:list"),
        reverse("pricing:list"),
    ]


class ProductOrderingTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_superuser("ordering-admin", password="test-password")
        cls.warehouse, cls.ordered, cls.cds = create_ordering_fixture()

    def setUp(self):
        self.client.force_login(self.user)

    def test_all_pages_show_brands_types_and_products_in_price_order(self):
        expected_brands = ["Sony", "Nintendo", "Anbernic", "Steam Deck", "Apple"]
        for url in ordering_pages(self.warehouse):
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 200)
                groups = response.context["tech_brand_groups"]
                self.assertEqual([group["brand"].name for group in groups], expected_brands)
                rows = [row for brand in groups for _, items in brand["type_groups"] for row in items]
                actual = [row["product"] if isinstance(row, dict) else row for row in rows]
                self.assertEqual([p.pk for p in actual], [p.pk for p in self.ordered])
                self.assertEqual(sum(group["count"] for group in groups), len(self.ordered))
                html = response.content.decode()
                positions = [html.index(escape(product.name) + "</a>") for product in self.ordered + self.cds]
                self.assertEqual(positions, sorted(positions))

    def test_export_keeps_the_same_flat_order_and_group_headers(self):
        content = generate_wholesale_price_xlsx(warehouse_id=self.warehouse.pk, actor=self.user)
        sheet = load_workbook(BytesIO(content))[MAIN_SHEET]
        values = [row[0] for row in sheet.iter_rows(min_row=HEADER_ROW + 1, values_only=True)]
        expected = [p.name for p in self.ordered + self.cds]
        self.assertEqual([value for value in values if value in expected], expected)
        self.assertIn("Tech · Sony · Игровая консоль", values)
        self.assertIn("Tech · Nintendo · Портативная консоль", values)
        self.assertIn("Tech · Портативные консоли · Anbernic", values)
        self.assertIn("Tech · Прочие товары", values)

    def test_sony_gamepads_keep_model_and_color_order_on_every_page(self):
        template = self.ordered[3]
        expected = ["PS4 DualShock Black", "PS5 Original White", "PS5 Midnight Black"]
        for name in reversed(expected):
            product = Tech.objects.create(brand=template.brand, product_type=template.product_type, name=name)
            TechWarehouseStock.objects.create(warehouse=self.warehouse, tech=product, quantity=1)
        for url in ordering_pages(self.warehouse):
            with self.subTest(url=url):
                response = self.client.get(url)
                sony = response.context["tech_brand_groups"][0]
                rows = next(items for kind, items in sony["type_groups"] if kind.pk == template.product_type_id)
                products = [row["product"] if isinstance(row, dict) else row for row in rows]
                self.assertEqual([p.name for p in products], expected + [template.name])


class ProductGroupingTests(SimpleTestCase):
    def test_brand_spanning_sections_keeps_each_row_once(self):
        brand = SimpleNamespace(pk=1, name="Other")
        products = [
            SimpleNamespace(pk=1, name="Laptop", brand=brand, brand_id=1,
                            product_type=SimpleNamespace(pk=1, name="Ноутбук"), product_type_id=1),
            SimpleNamespace(pk=2, name="Handheld", brand=brand, brand_id=1,
                            product_type=SimpleNamespace(pk=2, name="Портативная консоль"), product_type_id=2),
        ]
        rows = [{"product": p, "quantity": 12, "extra": object()} for p in products]
        groups = tech_brand_groups(rows, product_of=lambda row: row["product"])
        self.assertEqual(len(groups), 1)
        self.assertEqual(groups[0]["count"], 2)
        self.assertEqual([kind.pk for kind, _ in groups[0]["type_groups"]], [2, 1])
        actual = [row for _, items in groups[0]["type_groups"] for row in items]
        self.assertIs(actual[0], rows[1])
        self.assertIs(actual[1], rows[0])
