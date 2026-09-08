import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.db.models import Sum
from openpyxl import load_workbook

from accounts.models import User
from catalog.models import Brand, CD, Platform, ProductChangeEvent, ProductType, Tech
from consignment.models import (
    CDConsignmentStock,
    ConsignmentMovement,
    TechConsignmentStock,
)
from consignment.services import transfer_many_to_consignment
from partners.models import SalesPlatform
from warehouse.models import CDWarehouseStock, TechWarehouseStock, Warehouse


CENT = Decimal("0.01")
MAIN_BLOCKS = (
    ("tech", 34, 116),
    ("ns2", 118, 157),
    ("ps4", 159, 358),
    ("ps5", 360, 578),
    ("new", 580, 587),
)
CONS_PRODUCT_ROWS = {
    3: 103,
    4: 94,
    5: 40,
    6: 68,
    7: 86,
    8: 87,
    9: 44,
    10: 97,
    11: 42,
    12: 100,
    13: 100,
    14: 94,
    15: 86,
    16: 87,
    17: 86,
    18: 40,
    19: 42,
    20: 44,
    21: 47,
    22: 47,
    23: 100,
    24: 105,
    25: 119,
    26: 125,
    27: 131,
    28: 138,
    29: 139,
    30: 143,
    31: 135,
    32: 124,
}
SALES_PLATFORM_ALIASES = {
    "Саня PS": "PlayStore",
    "Саня ПС": "PlayStore",
    "Джой": "Джойстик",
    "ПлБокс": "PlayBox",
}
CD_PLATFORM_BY_BLOCK = {
    "ns2": "Nintendo Switch 2",
    "ps4": "PlayStation 4",
    "ps5": "PlayStation 5",
}
NEW_NINTENDO_ROWS = {582, 584}


@dataclass(frozen=True)
class MainRow:
    row: int
    kind: str
    name: str
    quantity: int
    cost: Decimal
    brand: str | None = None
    product_type: str | None = None
    platform: str | None = None


@dataclass(frozen=True)
class ConsignmentRow:
    row: int
    main_row: int
    platform: str
    quantity: int
    receivable: Decimal


def _decimal(value, *, row, label, allow_blank=False):
    if value in (None, ""):
        if allow_blank:
            return Decimal("0.00")
        raise CommandError(f"Строка {row}: не заполнено поле «{label}».")
    try:
        result = Decimal(str(value)).quantize(CENT, rounding=ROUND_HALF_UP)
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise CommandError(f"Строка {row}: некорректное значение «{label}»: {value!r}.") from exc
    if result < 0:
        raise CommandError(f"Строка {row}: «{label}» не может быть отрицательным.")
    return result


def _quantity(value, *, row, label="Наличие"):
    if value in (None, ""):
        return 0
    try:
        decimal = Decimal(str(value))
        result = int(decimal)
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise CommandError(f"Строка {row}: некорректное количество в «{label}»: {value!r}.") from exc
    if decimal != result or result < 0:
        raise CommandError(f"Строка {row}: «{label}» должно быть целым неотрицательным числом.")
    return result


def _tech_reference(row, name):
    if 34 <= row <= 39:
        return "Sony", "Игровая консоль"
    if row == 40:
        return "Sony", "Внешний привод для консоли"
    if 41 <= row <= 42:
        return "Sony", "VR-шлем"
    if row == 43:
        return "Sony", "Зарядная станция"
    if row == 44:
        return "Sony", "VR-аксессуар"
    if row == 45:
        return "Sony", "Геймпад"
    if row in {46, 47}:
        return "Sony", "Зарядная станция для геймпадов"
    if row == 48:
        return "Без бренда", "Зарядная станция для геймпадов"
    if row == 49:
        return "PowerA", "Зарядная станция для геймпадов"
    if 50 <= row <= 51 or 53 <= row <= 80:
        return "Sony", "Геймпад"
    if row == 52:
        return "Sony", "Запасная часть"
    if 81 <= row <= 82:
        return "Sony", "Стриминговая приставка"
    if 83 <= row <= 87:
        return "Sony", "Игровая гарнитура"
    if row == 88:
        return "Без бренда", "Подставка для игровой консоли"
    if row == 89:
        return "PlayX", "Геймпад"
    if row == 90:
        return "Без бренда", "Охлаждающая подставка для консоли"
    if 91 <= row <= 92:
        return "Valve", "Портативная игровая консоль"
    if 93 <= row <= 94:
        return "ASUS ROG", "Портативная игровая консоль"
    if row == 95:
        return "Lenovo", "Портативная игровая консоль"
    if 96 <= row <= 98:
        return "Meta", "VR-шлем"
    if 99 <= row <= 103:
        return "Nintendo", "Портативная игровая консоль"
    if row == 104:
        return "Nintendo", "Чехол для игровой консоли"
    if 105 <= row <= 108:
        return "Nintendo", "Геймпад"
    if row == 109:
        return "Nintendo", "Игровая гарнитура"
    if row == 110:
        return "Nintendo", "MicroSD-карта"
    if row == 111:
        return "Nintendo", "Аксессуар для игровой консоли"
    if row == 112:
        return "Logitech G", "Ручная коробка передач для симрейсинга"
    if row in {113, 114}:
        return "Без бренда", "Защитное стекло"
    if row == 115:
        return "Без бренда", "Кабель питания"
    if row == 116:
        return "Без бренда", "Переходник"
    raise CommandError(f"Строка {row}: не определены бренд и тип техники «{name}».")


def _parse_workbook(path):
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise CommandError(f"Не удалось прочитать Excel-файл: {exc}") from exc
    try:
        if "Лист1" not in workbook.sheetnames:
            raise CommandError("В книге отсутствует лист «Лист1».")
        sheet = workbook["Лист1"]
        main_rows = []
        for block, start, end in MAIN_BLOCKS:
            for row in range(start, end + 1):
                raw_name = sheet.cell(row, 1).value
                if raw_name in (None, ""):
                    continue
                name = str(raw_name).strip()
                if not name:
                    raise CommandError(f"Строка {row}: пустое наименование товара.")
                if len(name) > 255:
                    raise CommandError(f"Строка {row}: наименование длиннее 255 символов.")
                quantity = _quantity(sheet.cell(row, 3).value, row=row)
                cost = _decimal(sheet.cell(row, 7).value, row=row, label="Закуп", allow_blank=True)
                if block == "tech":
                    brand, product_type = _tech_reference(row, name)
                    main_rows.append(MainRow(row, "tech", name, quantity, cost, brand, product_type))
                else:
                    platform = CD_PLATFORM_BY_BLOCK.get(block)
                    if block == "new":
                        platform = "Nintendo Switch 2" if row in NEW_NINTENDO_ROWS else "PlayStation 5"
                    main_rows.append(MainRow(row, "cd", name, quantity, cost, platform=platform))

        by_row = {item.row: item for item in main_rows}
        if len(by_row) != len(main_rows):
            raise CommandError("В основной части обнаружены повторяющиеся номера строк.")
        names = Counter((item.kind, item.name.casefold()) for item in main_rows)
        duplicates = [name for name, count in names.items() if count > 1]
        if duplicates:
            raise CommandError(f"В основной части обнаружены дубли товаров: {duplicates!r}.")

        consignment_rows = []
        presence_pattern = re.compile(r"^(.*?)\s+(\d+)$")
        for row in range(3, 33):
            if row not in CONS_PRODUCT_ROWS:
                raise CommandError(f"Для строки реализации {row} не задано сопоставление с основной базой.")
            raw_presence = str(sheet.cell(row, 3).value or "").strip()
            match = presence_pattern.match(raw_presence)
            if not match:
                raise CommandError(f"Строка {row}: не удалось разобрать наличие «{raw_presence}».")
            raw_platform, raw_quantity = match.groups()
            try:
                platform = SALES_PLATFORM_ALIASES[raw_platform]
            except KeyError as exc:
                raise CommandError(f"Строка {row}: неизвестная площадка «{raw_platform}».") from exc
            quantity = _quantity(raw_quantity, row=row, label="Наличие на реализации")
            receivable = _decimal(sheet.cell(row, 2).value, row=row, label="Цена")
            main_row = CONS_PRODUCT_ROWS[row]
            if main_row not in by_row:
                raise CommandError(f"Строка {row}: основная позиция в строке {main_row} не найдена.")
            consignment_rows.append(ConsignmentRow(row, main_row, platform, quantity, receivable))

        # Строки 15 и 17 относятся к одной основной карточке. Средняя сумма
        # 12 000 ₽ сохраняет общий итог 13 000 + 11 000 без создания дубля товара.
        # Пользователь подтвердил объединение строк 21–22: 7 шт. по 3 500 ₽.
        normalized = []
        for item in consignment_rows:
            if item.row in {15, 21}:
                continue
            if item.row == 17:
                normalized.append(ConsignmentRow(17, item.main_row, item.platform, 2, Decimal("12000.00")))
                continue
            if item.row == 22:
                normalized.append(ConsignmentRow(22, item.main_row, item.platform, 7, Decimal("3500.00")))
            else:
                normalized.append(item)

        grouped = {}
        for item in normalized:
            key = (item.main_row, item.platform)
            if key in grouped:
                previous = grouped[key]
                if previous.receivable != item.receivable:
                    raise CommandError(
                        f"Строки {previous.row} и {item.row}: разные суммы к получению для одного товара и площадки."
                    )
                grouped[key] = ConsignmentRow(
                    item.row, item.main_row, item.platform,
                    previous.quantity + item.quantity, item.receivable,
                )
            else:
                grouped[key] = item
        return main_rows, list(grouped.values())
    finally:
        workbook.close()


def _require_exact(model, names, label):
    found = {item.name: item for item in model.objects.filter(name__in=names)}
    missing = sorted(set(names) - set(found))
    if missing:
        raise CommandError(f"В справочнике «{label}» отсутствуют значения: {', '.join(missing)}.")
    return found


def _extract_code(name):
    match = re.search(r"\b(?:CUSA|PPSA|PSSA)\s*\d+\b", name, flags=re.IGNORECASE)
    return match.group(0) if match else ""


class Command(BaseCommand):
    help = "Проверяет и импортирует текущие остатки из согласованного Excel-файла."

    def add_arguments(self, parser):
        parser.add_argument("xlsx", type=Path)
        parser.add_argument("--apply", action="store_true", help="Записать проверенные данные в базу.")
        parser.add_argument("--actor", default="derber", help="Пользователь для документов реализации.")

    def handle(self, *args, **options):
        path = options["xlsx"].resolve()
        if not path.is_file():
            raise CommandError(f"Файл не найден: {path}")
        main_rows, consignment_rows = _parse_workbook(path)

        brand_names = {item.brand for item in main_rows if item.brand}
        type_names = {item.product_type for item in main_rows if item.product_type}
        cd_platform_names = {item.platform for item in main_rows if item.platform}
        sales_platform_names = {item.platform for item in consignment_rows}
        brands = _require_exact(Brand, brand_names, "Бренды")
        product_types = _require_exact(ProductType, type_names, "Типы товаров")
        cd_platforms = _require_exact(Platform, cd_platform_names, "Платформы")
        sales_platforms = _require_exact(SalesPlatform, sales_platform_names, "Площадки реализации")

        if Warehouse.objects.count() != 1:
            raise CommandError("Для импорта должен существовать ровно один склад.")
        warehouse = Warehouse.objects.get()
        try:
            actor = User.objects.get(username=options["actor"], is_active=True)
        except User.DoesNotExist as exc:
            raise CommandError(f"Активный пользователь «{options['actor']}» не найден.") from exc

        tech_rows = [item for item in main_rows if item.kind == "tech"]
        cd_rows = [item for item in main_rows if item.kind == "cd"]
        consignment_units = sum(item.quantity for item in consignment_rows)
        self.stdout.write(f"Источник: {path}")
        self.stdout.write(f"Техника: {len(tech_rows)} поз., {sum(item.quantity for item in tech_rows)} шт.")
        self.stdout.write(f"CD: {len(cd_rows)} поз., {sum(item.quantity for item in cd_rows)} шт.")
        self.stdout.write(f"Реализация: {len(consignment_rows)} поз., {consignment_units} шт.")
        self.stdout.write("Площадки реализации: " + ", ".join(
            f"{name} — {sum(item.quantity for item in consignment_rows if item.platform == name)} шт."
            for name in sorted(sales_platform_names)
        ))
        self.stdout.write("Платформы CD: " + ", ".join(
            f"{name} — {sum(1 for item in cd_rows if item.platform == name)} поз."
            for name in sorted(cd_platform_names)
        ))
        self.stdout.write("Классификация техники: " + ", ".join(
            f"{name} — {sum(1 for item in tech_rows if item.product_type == name)} поз."
            for name in sorted(type_names)
        ))
        if not options["apply"]:
            self.stdout.write(self.style.WARNING("Проверка завершена. База не изменена; для импорта добавьте --apply."))
            return

        nonempty = {
            "CD": CD.objects.count(),
            "Tech": Tech.objects.count(),
            "остатки CD": CDWarehouseStock.objects.count(),
            "остатки техники": TechWarehouseStock.objects.count(),
            "реализация CD": CDConsignmentStock.objects.count(),
            "реализация техники": TechConsignmentStock.objects.count(),
            "документы реализации": ConsignmentMovement.objects.count(),
            "история товаров": ProductChangeEvent.objects.count(),
        }
        occupied = {label: count for label, count in nonempty.items() if count}
        if occupied:
            raise CommandError(f"Целевые таблицы не пусты: {occupied}. Импорт отменён.")

        extra_by_main_row = defaultdict(int)
        for item in consignment_rows:
            extra_by_main_row[item.main_row] += item.quantity

        with transaction.atomic():
            products_by_row = {}
            for item in main_rows:
                common = {
                    "name": item.name,
                    "sku": f"{'TECH' if item.kind == 'tech' else 'CD'}-XLSX-{item.row:04d}",
                    "cost": item.cost,
                    "quantity_on_consignment": 0,
                    "retail_price": None,
                    "wholesale_price": None,
                    "yandex_market_price": None,
                    "comment": "",
                }
                if item.kind == "tech":
                    product = Tech.objects.create(
                        brand=brands[item.brand],
                        product_type=product_types[item.product_type],
                        **common,
                    )
                    TechWarehouseStock.objects.create(
                        warehouse=warehouse,
                        tech=product,
                        quantity=item.quantity + extra_by_main_row[item.row],
                    )
                else:
                    product = CD.objects.create(
                        platform=cd_platforms[item.platform],
                        cusa_ppsa_code=_extract_code(item.name),
                        **common,
                    )
                    CDWarehouseStock.objects.create(
                        warehouse=warehouse,
                        cd=product,
                        quantity=item.quantity + extra_by_main_row[item.row],
                    )
                products_by_row[item.row] = product

            for platform_name in sorted(sales_platform_names):
                lines = []
                for item in consignment_rows:
                    if item.platform != platform_name:
                        continue
                    product = products_by_row[item.main_row]
                    lines.append({
                        "product_type": "tech" if isinstance(product, Tech) else "cd",
                        "product_id": product.pk,
                        "quantity": item.quantity,
                        "receivable_per_unit": item.receivable,
                    })
                transfer_many_to_consignment(
                    actor=actor,
                    warehouse_id=warehouse.pk,
                    platform_id=sales_platforms[platform_name].pk,
                    lines=lines,
                )

            cd_warehouse_total = CDWarehouseStock.objects.aggregate(total=Sum("quantity"))["total"] or 0
            tech_warehouse_total = TechWarehouseStock.objects.aggregate(total=Sum("quantity"))["total"] or 0
            actual_consignment = (
                (CDConsignmentStock.objects.aggregate(total=Sum("quantity"))["total"] or 0)
                + (TechConsignmentStock.objects.aggregate(total=Sum("quantity"))["total"] or 0)
            )
            expected_cd = sum(item.quantity for item in cd_rows)
            expected_tech = sum(item.quantity for item in tech_rows)
            if (cd_warehouse_total, tech_warehouse_total, actual_consignment) != (
                expected_cd, expected_tech, consignment_units
            ):
                raise CommandError(
                    "Контрольные суммы не сошлись: "
                    f"склад CD {cd_warehouse_total}/{expected_cd}, "
                    f"склад техники {tech_warehouse_total}/{expected_tech}, "
                    f"реализация {actual_consignment}/{consignment_units}."
                )

        self.stdout.write(self.style.SUCCESS(
            f"Импорт завершён: {len(main_rows)} товаров, "
            f"{sum(item.quantity for item in main_rows)} шт. на складе, "
            f"{consignment_units} шт. на реализации."
        ))
