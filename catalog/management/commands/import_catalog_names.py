import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from difflib import SequenceMatcher
from functools import lru_cache
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from catalog.models import Brand, CD, Platform, ProductType, Tech
from warehouse.models import CDWarehouseStock, TechWarehouseStock, Warehouse


CD_SECTIONS = (
    (2, 97, "Nintendo Switch 2"),
    (98, 540, "Nintendo Switch"),
    (541, 1093, "PlayStation 4"),
    (1094, 1861, "PlayStation 5"),
    (1862, 1901, "Xbox Series X"),
)
CD_HEADER_ROWS = {2, 98, 541, 1094, 1862}
TECH_HEADER_ROWS = {2, 6, 20, 132, 228, 263}
TECH_MATCH_SKU_BY_ROW = {
    29: "TECH-XLSX-0039",   # Ghost of Yotei с физическим диском
    53: "TECH-XLSX-0044",   # оригинальный PS VR2 PC Adapter
    60: "TECH-XLSX-0040",   # дисковод PS5
    61: "TECH-XLSX-0046",   # оригинальная док-станция DualSense
    62: "TECH-XLSX-0048",   # копия док-станции DualSense
    86: "TECH-XLSX-0071",   # Astro Bot Joyful Original
    89: "TECH-XLSX-0072",   # Death Stranding 2 Original
    98: "TECH-XLSX-0076",   # God of War 20th Anniversary Original
    120: "TECH-XLSX-0083",  # Pulse 3D Black
    121: "TECH-XLSX-0084",  # Pulse 3D Camo
    122: "TECH-XLSX-0085",  # Pulse 3D White
    130: "TECH-XLSX-0088",  # оригинальная вертикальная подставка
    134: "TECH-XLSX-0099",  # Switch 2 + физический Mario Kart
    137: "TECH-XLSX-0110",  # microSD 256 GB
    141: "TECH-XLSX-0108",  # стандартный комплект Joy-Con 2
    177: "TECH-XLSX-0102",  # Switch Lite Gray
    179: "TECH-XLSX-0103",  # Switch Lite Turquoise
    188: "TECH-XLSX-0107",  # стандартный Pro Controller (Black)
}
MONEY = Decimal("0.01")


@dataclass(frozen=True)
class SourceItem:
    row: int
    name: str
    cost: Decimal
    classification: str
    secondary: str = ""


CONFUSABLES = str.maketrans({
    "а": "a", "А": "a", "в": "b", "В": "b", "е": "e", "Е": "e",
    "к": "k", "К": "k", "м": "m", "М": "m", "н": "h", "Н": "h",
    "о": "o", "О": "o", "р": "p", "Р": "p", "с": "c", "С": "c",
    "т": "t", "Т": "t", "х": "x", "Х": "x", "у": "y", "У": "y",
})


@lru_cache(maxsize=None)
def _base_normalize(value):
    value = unicodedata.normalize("NFKC", str(value or "")).strip().casefold().replace("ё", "е")
    value = value.translate(CONFUSABLES)
    value = value.replace("&", " and ").replace("+", " plus ")
    value = re.sub(r"\bplay\s*station\s*([45])\b", r"ps\1", value)
    value = re.sub(r"\bnintendo\s+switch\s*2\b", "ns2", value)
    value = re.sub(r"\bnintendo\s+switch\b", "ns", value)
    value = re.sub(r"\bxbox\s+series\s+[xs]\b", "xbox", value)
    value = re.sub(r"[^0-9a-zа-я]+", " ", value)
    return " ".join(value.split())


@lru_cache(maxsize=None)
def _cd_key(value):
    value = _base_normalize(value)
    value = re.sub(r"^(?:ns2|ns|ps4|ps5|xbox)\s+", "", value)
    return value


@lru_cache(maxsize=None)
def _soft_key(value, *, cd=False):
    value = _cd_key(value) if cd else _base_normalize(value)
    value = re.sub(r"\b(?:cusa|ppsa|pssa)\s*\d+(?:\s+\d+)*\b", " ", value)
    value = re.sub(r"\b(?:rus|eng)\s+(?:lang|sub)\b", " ", value)
    value = re.sub(r"\b(?:rus|eng)\b", " ", value)
    return " ".join(value.split())


@lru_cache(maxsize=None)
def _numbers(value):
    value = re.sub(r"\b(?:cusa|ppsa|pssa)\s*\d+(?:\s+\d+)*\b", " ", value)
    return tuple(token for token in value.split() if any(char.isdigit() for char in token))


EDITION_WORDS = {
    "anniversary", "bundle", "collection", "complete", "deluxe", "digital", "edition",
    "gold", "goty", "launch", "legacy", "lenticular", "limited", "new", "original",
    "physical", "pro", "ref", "remake", "remastered", "slim", "special", "ultimate",
}
STRICT_WORDS = EDITION_WORDS | {
    "black", "blue", "camo", "camouflage", "coral", "europe", "gold", "gray", "green",
    "grey", "japan", "l", "lcd", "light", "neon", "oled", "orange", "pastel", "pink",
    "purple", "r", "red", "silver", "turquoise", "usa", "vr", "white", "yellow",
    "i", "ii", "iii", "iv", "v",
}


@lru_cache(maxsize=None)
def _edition_words(value):
    return {word for word in _soft_key(value).split() if word in EDITION_WORDS}


def _strict_signature(value, *, cd=False, soft=False):
    normalized = _soft_key(value, cd=cd) if soft else (_cd_key(value) if cd else _base_normalize(value))
    words = normalized.split()
    signature = {word for word in words if word in STRICT_WORDS or any(char.isdigit() for char in word)}
    if not cd and "usb" in words:
        signature.update(word for word in words if word in {"a", "b", "c"})
    raw = unicodedata.normalize("NFKC", str(value)).casefold().replace("ё", "е")
    for marker in ("rus lang", "rus sub", "eng lang", "eng sub", "русская обложка", "уценка"):
        if marker in raw:
            signature.add(marker)
    return frozenset(signature)


def _catalog_codes(value):
    normalized = _base_normalize(value)
    return frozenset(re.findall(r"\b(?:cusa|ppsa|pssa)\s*(\d+)", normalized))


def _similarity(left, right, *, cd=False):
    a = _soft_key(left, cd=cd)
    b = _soft_key(right, cd=cd)
    if not a or not b:
        return 0.0
    return SequenceMatcher(None, a, b).ratio()


def _is_same_product(left, right, *, cd=False, existing=False):
    a = _cd_key(left) if cd else _base_normalize(left)
    b = _cd_key(right) if cd else _base_normalize(right)
    if a == b:
        return True
    if not existing and _strict_signature(left, cd=cd) != _strict_signature(right, cd=cd):
        return False
    soft_a = _soft_key(left, cd=cd) if existing else a
    soft_b = _soft_key(right, cd=cd) if existing else b
    if existing:
        sig_a = _strict_signature(left, cd=cd, soft=True)
        sig_b = _strict_signature(right, cd=cd, soft=True)
        language_markers = {"rus lang", "rus sub", "eng lang", "eng sub", "русская обложка"}
        langs_a, langs_b = sig_a & language_markers, sig_b & language_markers
        if langs_a and langs_b and langs_a != langs_b:
            return False
        codes_a, codes_b = _catalog_codes(left), _catalog_codes(right)
        if codes_a and codes_b and not (codes_a & codes_b):
            return False
        neutral = language_markers | {"original"}
        if (sig_a - neutral) != (sig_b - neutral):
            return False
    if soft_a == soft_b:
        return True
    tokens_a, tokens_b = set(soft_a.split()), set(soft_b.split())
    overlap = len(tokens_a & tokens_b) / max(1, min(len(tokens_a), len(tokens_b)))
    if overlap < 0.80:
        return False
    score = max(
        SequenceMatcher(None, soft_a, soft_b).ratio(),
        SequenceMatcher(None, " ".join(sorted(tokens_a)), " ".join(sorted(tokens_b))).ratio(),
    )
    return score >= (0.94 if cd else 0.93)


def _money(value, row):
    if value is None or value == "":
        return None
    try:
        result = Decimal(str(value)).quantize(MONEY, rounding=ROUND_HALF_UP)
    except Exception as exc:
        raise CommandError(f"Строка {row}: некорректная себестоимость {value!r}.") from exc
    if result < 0:
        raise CommandError(f"Строка {row}: отрицательная себестоимость.")
    return result


def _read_cd(path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook.active
        items = []
        missing_cost = []
        for row, values in enumerate(sheet.iter_rows(min_col=1, max_col=2, values_only=True), start=1):
            if row in CD_HEADER_ROWS:
                continue
            raw_name, raw_cost = values
            if raw_name in (None, ""):
                continue
            platform = next((name for start, end, name in CD_SECTIONS if start <= row <= end), None)
            if platform is None:
                continue
            name = " ".join(str(raw_name).split())
            cost = _money(raw_cost, row)
            if cost is None:
                missing_cost.append((row, name))
                continue
            items.append(SourceItem(row, name, cost, platform))
        return items, missing_cost
    finally:
        workbook.close()


def _tech_brand(row, name):
    low = _base_normalize(name)
    if "copy" in low:
        return "Без бренда"
    if "powera" in low:
        return "PowerA"
    if "victrix" in low:
        return "Victrix"
    if "nacon" in low or "revolution 5" in low:
        return "Nacon"
    if "turtle beach" in low:
        return "Turtle Beach"
    if "airlite" in low:
        return "PDP"
    if row <= 131:
        if "wd sn850p" in low:
            return "WD_BLACK"
        return "Sony"
    if 133 <= row <= 227:
        if "sega" in low:
            return "Sega"
        return "Nintendo"
    if 229 <= row <= 262:
        return "Microsoft"
    if "asus rog" in low:
        return "ASUS ROG"
    if low.startswith("lenovo"):
        return "Lenovo"
    if low.startswith("msi"):
        return "MSI"
    if low.startswith("valve"):
        return "Valve"
    if "logitech" in low:
        return "Logitech G"
    if "marshall" in low or any(model in low for model in ("major iv", "monitor ii", "minor iii", "motif", "minor ii")):
        return "Marshall"
    if "razer" in low:
        return "Razer"
    if "astro" in low:
        return "Astro Gaming"
    if "hyperx" in low:
        return "HyperX"
    if "sony" in low:
        return "Sony"
    if "xbox" in low:
        return "Microsoft"
    if low.startswith("meta") or low.startswith("oculus"):
        return "Meta"
    if low.startswith("pico"):
        return "Pico"
    return "Без бренда"


def _tech_type(row, name):
    low = _base_normalize(name)
    raw = unicodedata.normalize("NFKC", str(name)).casefold().replace("ё", "е")
    if row == 3:
        return "Ретро-консоль"
    if row in {4, 5}:
        return "Портативная игровая консоль"
    if ("oculus" in low or "quest" in low) and "charging station" in low:
        return "VR-аксессуар"
    if "playstation vr2 charging station" in low:
        return "Зарядная станция"
    if "charging station" in low or "док станция" in raw:
        return "Зарядная станция для геймпадов"
    if "stick module" in low:
        return "Запасная часть"
    if "disc drive" in low or "дисковод" in raw:
        return "Внешний привод для консоли"
    if "console covers" in low or "cable cover" in low:
        return "Аксессуар для игровой консоли"
    if "vr2 pc adapter" in low:
        return "VR-аксессуар"
    if "playstation vr" in low or "oculus quest" in low or low.startswith("meta quest") or low.startswith("pico 4"):
        return "VR-шлем"
    if "remote portal" in low:
        return "Стриминговая приставка"
    if "hd camera" in low or low.endswith(" camera") or "camera piranha" in low:
        return "Камера для игровой консоли"
    if "link usb adapter" in low or "lan adapter" in low or "ac adapter" in low:
        return "Сетевое зарядное устройство" if "ac adapter" in low else "Адаптер"
    if "micro sd" in low or "microsd" in low:
        return "MicroSD-карта"
    if "ssd" in low:
        return "SSD"
    if "headset" in low or "наушники" in raw:
        return "Игровая гарнитура"
    if "racing wheel" in low:
        return "Руль"
    if "игровой руль" in raw:
        return "Руль"
    if "shifter" in low:
        return "Ручная коробка передач для симрейсинга"
    if "carrying case all in one" in low or "crossbody bag" in low or "messenger bag" in low:
        return "Сумка для игровой консоли"
    if "carrying case" in low or "travel case" in low:
        return "Чехол для игровой консоли"
    if "screen protective" in low:
        return "Защитная плёнка"
    if "hdmi cable" in low:
        return "Кабель HDMI"
    if "usb a to usb c cable" in low or "usb c to usb c cable" in low:
        return "Кабель USB-C"
    if "battery" in low:
        return "Аксессуар для игровой консоли"
    if "alarmo" in low:
        return "Умный будильник"
    if "talking flower" in low or "virtual boy" in low or "ring fit" in low or "playstand" in low:
        return "Аксессуар для игровой консоли"
    if "vertical stand" in low:
        return "Подставка для игровой консоли"
    if "joy con" in low and any(word in low for word in ("strap", "wheel", "grip")):
        return "Аксессуар для игровой консоли"
    if "dock" in low:
        return "Док-станция для игровой консоли"
    if "steam machine" in low:
        return "Игровая консоль"
    if "joy con" in low or "controller" in low or "control pad" in low or "геймпад" in raw or "split pad" in low:
        return "Геймпад"
    if row in range(7, 16) or row in range(21, 46) or row in range(229, 234):
        return "Игровая консоль"
    if row in range(133, 137) or row in range(170, 183) or row in range(264, 282):
        return "Портативная игровая консоль"
    if row == 16:
        return "VR-шлем"
    if row in range(17, 20) or row in range(64, 120) or row in range(139, 151) or row in range(188, 208) or row in range(219, 224) or row in range(234, 261):
        return "Геймпад"
    return "Аксессуар для игровой консоли"


def _read_tech(path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook.active
        items = []
        missing_cost = []
        for row, values in enumerate(sheet.iter_rows(min_col=1, max_col=2, values_only=True), start=1):
            if row < 2:
                continue
            if row in TECH_HEADER_ROWS:
                continue
            raw_name, raw_cost = values
            if raw_name in (None, ""):
                continue
            name = " ".join(str(raw_name).split())
            cost = _money(raw_cost, row)
            if cost is None:
                missing_cost.append((row, name))
                continue
            items.append(SourceItem(row, name, cost, _tech_brand(row, name), _tech_type(row, name)))
        return items, missing_cost
    finally:
        workbook.close()


def _dedupe_source(items, *, cd):
    groups = []
    duplicate_rows = []
    token_index = defaultdict(set)
    exact_index = {}
    for item in items:
        match_index = None
        class_key = (item.classification, item.secondary)
        base = _cd_key(item.name) if cd else _base_normalize(item.name)
        soft = _soft_key(item.name, cd=cd)
        candidates = set()
        for exact in (base, soft):
            index = exact_index.get((class_key, exact))
            if index is not None:
                candidates.add(index)
        tokens = {token for token in soft.split() if len(token) >= 3}
        for token in tokens:
            candidates.update(token_index[(class_key, token)])
        if len(tokens) <= 1:
            candidates.update(
                index for index, candidate in enumerate(groups)
                if (candidate.classification, candidate.secondary) == class_key
                and abs(len(_soft_key(candidate.name, cd=cd)) - len(soft)) <= 3
                and _soft_key(candidate.name, cd=cd)[:1] == soft[:1]
            )
        for index in sorted(candidates):
            candidate = groups[index]
            if _is_same_product(item.name, candidate.name, cd=cd):
                match_index = index
                break
        if match_index is None:
            index = len(groups)
            groups.append(item)
            exact_index[(class_key, base)] = index
            exact_index[(class_key, soft)] = index
            for token in tokens:
                token_index[(class_key, token)].add(index)
            continue
        chosen = groups[match_index]
        duplicate_rows.append((item.row, item.name, chosen.row, chosen.name))
        if item.cost < chosen.cost:
            groups[match_index] = SourceItem(chosen.row, chosen.name, item.cost, chosen.classification, chosen.secondary)
    return groups, duplicate_rows


def _match_existing(items, existing, *, cd):
    proposals = []
    ambiguous = []
    pools = defaultdict(list)
    token_index = defaultdict(set)
    exact_index = defaultdict(set)
    for obj in existing:
        class_key = obj.platform.name if cd else "tech"
        index = len(pools[class_key])
        pools[class_key].append(obj)
        base = _cd_key(obj.name) if cd else _base_normalize(obj.name)
        soft = _soft_key(obj.name, cd=cd)
        exact_index[(class_key, base)].add(index)
        exact_index[(class_key, soft)].add(index)
        for token in {token for token in soft.split() if len(token) >= 3}:
            token_index[(class_key, token)].add(index)
    tech_by_sku = {obj.sku: obj for obj in existing} if not cd else {}
    for item_index, item in enumerate(items):
        class_key = item.classification if cd else "tech"
        pool = pools[class_key]
        base = _cd_key(item.name) if cd else _base_normalize(item.name)
        soft = _soft_key(item.name, cd=cd)
        candidate_indices = set(exact_index[(class_key, base)]) | set(exact_index[(class_key, soft)])
        tokens = {token for token in soft.split() if len(token) >= 3}
        for token in tokens:
            candidate_indices.update(token_index[(class_key, token)])
        if len(tokens) <= 1:
            candidate_indices.update(
                index for index, obj in enumerate(pool)
                if abs(len(_soft_key(obj.name, cd=cd)) - len(soft)) <= 3
                and _soft_key(obj.name, cd=cd)[:1] == soft[:1]
            )
        candidate_pool = [pool[index] for index in sorted(candidate_indices)]
        exact = [obj for obj in candidate_pool if _is_same_product(item.name, obj.name, cd=cd, existing=True)]
        override = tech_by_sku.get(TECH_MATCH_SKU_BY_ROW.get(item.row)) if not cd else None
        if override is not None and override not in exact:
            exact.append(override)
        if exact:
            def quality(obj):
                override_bonus = 1000 if obj is override else 0
                exact_bonus = 100 if (_cd_key(item.name) if cd else _base_normalize(item.name)) == (_cd_key(obj.name) if cd else _base_normalize(obj.name)) else 0
                code_bonus = 10 if _catalog_codes(item.name) & _catalog_codes(obj.name) else 0
                return override_bonus + exact_bonus + code_bonus + _similarity(item.name, obj.name, cd=cd)
            exact.sort(key=lambda obj: (-quality(obj), obj.pk))
            for obj in exact:
                proposals.append((quality(obj), item_index, item, obj, _similarity(item.name, obj.name, cd=cd)))
            if len(exact) > 1:
                ambiguous.append((item, [(obj, _similarity(item.name, obj.name, cd=cd)) for obj in exact[:5]]))
    assigned_items = set()
    assigned_objects = set()
    matches = []
    for _quality, item_index, item, obj, score in sorted(proposals, key=lambda row: (-row[0], row[1], row[3].pk)):
        if item_index in assigned_items or obj.pk in assigned_objects:
            continue
        assigned_items.add(item_index)
        assigned_objects.add(obj.pk)
        matches.append((item, obj, score))
    new_items = [item for index, item in enumerate(items) if index not in assigned_items]
    return new_items, matches, ambiguous


def _require_references(model, names, label):
    found = {obj.name: obj for obj in model.objects.filter(name__in=names)}
    missing = sorted(set(names) - set(found))
    if missing:
        raise CommandError(f"В справочнике «{label}» отсутствуют: {', '.join(missing)}.")
    return found


class Command(BaseCommand):
    help = "Добавляет новые наименования и себестоимость из файлов Наименования cd/tech.xlsx без изменения существующих товаров."

    def add_arguments(self, parser):
        parser.add_argument("cd_xlsx", type=Path)
        parser.add_argument("tech_xlsx", type=Path)
        parser.add_argument("--apply", action="store_true")
        parser.add_argument("--show-matches", action="store_true")

    def handle(self, *args, **options):
        cd_path = options["cd_xlsx"].resolve()
        tech_path = options["tech_xlsx"].resolve()
        for path in (cd_path, tech_path):
            if not path.is_file():
                raise CommandError(f"Файл не найден: {path}")

        cd_source, cd_missing = _read_cd(cd_path)
        tech_source, tech_missing = _read_tech(tech_path)
        cd_unique, cd_source_dupes = _dedupe_source(cd_source, cd=True)
        tech_unique, tech_source_dupes = _dedupe_source(tech_source, cd=False)
        new_cd, cd_matches, cd_ambiguous = _match_existing(cd_unique, list(CD.objects.select_related("platform")), cd=True)
        new_tech, tech_matches, tech_ambiguous = _match_existing(tech_unique, list(Tech.objects.select_related("brand", "product_type")), cd=False)

        self.stdout.write(f"CD: строк с ценой {len(cd_source)}, дублей внутри файла {len(cd_source_dupes)}, уже в базе {len(cd_matches)}, новых {len(new_cd)}.")
        self.stdout.write(f"Техника: строк с ценой {len(tech_source)}, дублей внутри файла {len(tech_source_dupes)}, уже в базе {len(tech_matches)}, новых {len(new_tech)}.")
        self.stdout.write(f"Без себестоимости: CD {len(cd_missing)}, техника {len(tech_missing)}.")
        self.stdout.write("Новые CD по платформам: " + ", ".join(f"{name}: {count}" for name, count in sorted(Counter(x.classification for x in new_cd).items())))
        self.stdout.write("Новая техника по типам: " + ", ".join(f"{name}: {count}" for name, count in sorted(Counter(x.secondary for x in new_tech).items())))
        self.stdout.write("Новая техника по брендам: " + ", ".join(f"{name}: {count}" for name, count in sorted(Counter(x.classification for x in new_tech).items())))

        for row, name in cd_missing + tech_missing:
            self.stdout.write(self.style.WARNING(f"Пропущена строка {row} без себестоимости: {name}"))
        if cd_ambiguous or tech_ambiguous:
            self.stdout.write(self.style.WARNING(f"Неоднозначных совпадений: CD {len(cd_ambiguous)}, техника {len(tech_ambiguous)}."))
        if options["show_matches"]:
            for label, rows in (("CD", cd_matches), ("TECH", tech_matches)):
                for item, obj, score in rows:
                    self.stdout.write(f"MATCH {label} row={item.row} score={score:.3f}: {item.name} -> #{obj.pk} {obj.name}")
            for label, rows in (("CD-SOURCE", cd_source_dupes), ("TECH-SOURCE", tech_source_dupes)):
                for row, name, kept_row, kept_name in rows:
                    self.stdout.write(f"DUP {label} row={row}: {name} -> row={kept_row} {kept_name}")

        if not options["apply"]:
            self.stdout.write(self.style.WARNING("Проверка завершена, база не изменена. Для записи добавьте --apply."))
            return

        platforms = _require_references(Platform, {x.classification for x in new_cd}, "Платформы")
        brands = _require_references(Brand, {x.classification for x in new_tech}, "Бренды")
        product_types = _require_references(ProductType, {x.secondary for x in new_tech}, "Типы товаров")
        warehouses = list(Warehouse.objects.all())
        if not warehouses:
            raise CommandError("В базе нет склада.")

        with transaction.atomic():
            for item in new_cd:
                product = CD.objects.create(
                    platform=platforms[item.classification], name=item.name,
                    sku=f"CD-NAMES-{item.row:04d}", cost=item.cost,
                )
                CDWarehouseStock.objects.bulk_create([
                    CDWarehouseStock(warehouse=warehouse, cd=product, quantity=0) for warehouse in warehouses
                ])
            for item in new_tech:
                product = Tech.objects.create(
                    brand=brands[item.classification], product_type=product_types[item.secondary],
                    name=item.name, sku=f"TECH-NAMES-{item.row:04d}", cost=item.cost,
                )
                TechWarehouseStock.objects.bulk_create([
                    TechWarehouseStock(warehouse=warehouse, tech=product, quantity=0) for warehouse in warehouses
                ])

        self.stdout.write(self.style.SUCCESS(f"Добавлено: CD {len(new_cd)}, техника {len(new_tech)}."))
