import re
import unicodedata
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from accounts.models import User
from catalog.models import CD, Tech
from pricing.services import update_product_prices


CENT = Decimal("0.01")
SOURCE_BLOCKS = (
    ("tech", 11, 59),
    ("cd", 61, 93),
    ("cd", 95, 239),
    ("cd", 241, 383),
)

# В этих строках источник использует сокращённые названия. SKU ссылаются на
# карточки, созданные из исходной ведомости текущих остатков.
MANUAL_SKU_BY_SOURCE_ROW = {
    20: "TECH-XLSX-0054",  # Белый (Original White)
    21: "TECH-XLSX-0055",  # Чёрный (Midnight Black)
    23: "TECH-XLSX-0057",  # Камуфляж (Grey Camouflage)
    25: "TECH-XLSX-0060",  # Фиолетовый (Galactic Purple)
    26: "TECH-XLSX-0061",  # Volcanic/Vulcanic Red
    34: "TECH-XLSX-0071",  # Astro Bot, пометка Original
    35: "TECH-XLSX-0072",  # Death Stranding 2, пометка Original
    39: "TECH-XLSX-0076",  # God of War, пометка Original
}


@dataclass(frozen=True)
class PriceRow:
    row: int
    product_kind: str
    source_name: str
    price: Decimal
    product: object


def _normalize(value):
    value = unicodedata.normalize("NFKC", str(value)).casefold().replace("ё", "е")
    return " ".join(re.findall(r"[a-zа-я0-9]+", value))


def _price(value, row):
    if value in (None, ""):
        raise CommandError(f"Строка {row}: оптовая цена не заполнена.")
    try:
        result = Decimal(str(value)).quantize(CENT, rounding=ROUND_HALF_UP)
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise CommandError(f"Строка {row}: некорректная оптовая цена {value!r}.") from exc
    if result <= 0:
        raise CommandError(f"Строка {row}: оптовая цена должна быть больше нуля.")
    return result


def _load_source(path):
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise CommandError(f"Не удалось прочитать Excel-файл: {exc}") from exc
    try:
        sheet = workbook.active
        rows = []
        for product_kind, start, end in SOURCE_BLOCKS:
            for row in range(start, end + 1):
                raw_name = sheet.cell(row, 1).value
                raw_price = sheet.cell(row, 2).value
                if raw_name in (None, "") and raw_price in (None, ""):
                    continue
                name = str(raw_name or "").strip()
                if not name:
                    raise CommandError(f"Строка {row}: отсутствует наименование товара.")
                rows.append((row, product_kind, name, _price(raw_price, row)))
        return sheet.title, rows
    finally:
        workbook.close()


def _resolve_rows(source_rows):
    models = {"cd": CD, "tech": Tech}
    products = {kind: list(model.objects.all()) for kind, model in models.items()}
    by_normalized_name = {kind: {} for kind in models}
    for kind, items in products.items():
        for product in items:
            by_normalized_name[kind].setdefault(_normalize(product.name), []).append(product)

    resolved = []
    used_targets = {}
    for row, kind, source_name, price in source_rows:
        manual_sku = MANUAL_SKU_BY_SOURCE_ROW.get(row)
        if manual_sku:
            try:
                product = models[kind].objects.get(sku=manual_sku)
            except models[kind].DoesNotExist as exc:
                raise CommandError(f"Строка {row}: карточка с SKU {manual_sku} не найдена.") from exc
        else:
            matches = by_normalized_name[kind].get(_normalize(source_name), [])
            if len(matches) != 1:
                raise CommandError(
                    f"Строка {row}: для «{source_name}» найдено карточек: {len(matches)}."
                )
            product = matches[0]
        target = (kind, product.pk)
        if target in used_targets:
            previous = used_targets[target]
            raise CommandError(
                f"Строки {previous} и {row} сопоставлены с одной карточкой «{product.name}»."
            )
        used_targets[target] = row
        resolved.append(PriceRow(row, kind, source_name, price, product))
    return resolved


class Command(BaseCommand):
    help = "Проверяет и импортирует оптовые цены из согласованного Excel-файла."

    def add_arguments(self, parser):
        parser.add_argument("xlsx", type=Path)
        parser.add_argument("--apply", action="store_true", help="Записать проверенные цены в базу.")
        parser.add_argument("--actor", default="derber", help="Пользователь для истории изменений.")

    def handle(self, *args, **options):
        path = options["xlsx"].resolve()
        if not path.is_file():
            raise CommandError(f"Файл не найден: {path}")
        sheet_name, source_rows = _load_source(path)
        rows = _resolve_rows(source_rows)
        try:
            actor = User.objects.get(username=options["actor"], is_active=True)
        except User.DoesNotExist as exc:
            raise CommandError(f"Активный пользователь «{options['actor']}» не найден.") from exc

        cd_rows = [item for item in rows if item.product_kind == "cd"]
        tech_rows = [item for item in rows if item.product_kind == "tech"]
        self.stdout.write(f"Источник: {path}, лист «{sheet_name}»")
        self.stdout.write(f"Оптовые цены: {len(rows)} карточек (CD: {len(cd_rows)}, техника: {len(tech_rows)}).")
        self.stdout.write(
            f"Диапазон цен: {min(item.price for item in rows):.2f}–{max(item.price for item in rows):.2f} руб."
        )
        if not options["apply"]:
            self.stdout.write(self.style.WARNING("Проверка завершена. База не изменена; для импорта добавьте --apply."))
            return

        already_priced = [
            item for item in rows if item.product.wholesale_price is not None
        ]
        if already_priced:
            raise CommandError(
                f"У {len(already_priced)} сопоставленных карточек оптовая цена уже заполнена. Импорт отменён."
            )

        with transaction.atomic():
            for item in rows:
                update_product_prices(
                    actor=actor,
                    product_type=item.product_kind,
                    product_id=item.product.pk,
                    changes={"wholesale_price": item.price},
                )

            errors = []
            for item in rows:
                model = CD if item.product_kind == "cd" else Tech
                actual = model.objects.values_list("wholesale_price", flat=True).get(pk=item.product.pk)
                if actual != item.price:
                    errors.append(f"строка {item.row}: {actual}/{item.price}")
            if errors:
                raise CommandError("Контроль цен не сошёлся: " + "; ".join(errors[:10]))

        self.stdout.write(self.style.SUCCESS(f"Импортировано {len(rows)} оптовых цен."))
