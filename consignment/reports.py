"""Проверка подписанного отчёта; продажи проводятся штатным сервисом Sale."""
from decimal import Decimal, InvalidOperation
from uuid import UUID
from zipfile import BadZipFile

from django.core import signing
from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook

from partners.models import Partner, SalesPlatform
from sales.models import Sale
from .models import CDConsignmentStock, TechConsignmentStock

SALT = 'consignment.reconciliation.v1'
DRAFT_KEY = 'consignment_report_sale'


def sign_report(value):
    return signing.dumps(value, salt=SALT, compress=True)


def unsign_report(value):
    try:
        return signing.loads(value, salt=SALT)
    except (signing.BadSignature, TypeError, ValueError) as exc:
        raise ValidationError('Служебные данные отчёта повреждены. Скачайте новый Excel площадки.') from exc


def _quantity(value):
    try:
        if isinstance(value, bool):
            raise ValueError
        result = Decimal(str(value or 0))
        if not result.is_finite() or result < 0 or result != result.to_integral_value() or result > 2147483647:
            raise ValueError
        return int(result)
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise ValidationError('Количество реализованного товара должно быть целым неотрицательным числом.') from exc


def current_report_rows(*, platform_id, rows, actor):
    result, seen = [], set()
    for source in rows:
        kind, stock_id = source['kind'], source['stock_id']
        if kind not in ('cd', 'tech') or (kind, stock_id) in seen:
            raise ValidationError('В отчёте повторяется или повреждена партия товара.')
        seen.add((kind, stock_id))
        if not actor.has_perm(f'consignment.view_{kind}consignmentstock'):
            raise ValidationError('Нет права просмотра партии реализации из отчёта.')
        model = CDConsignmentStock if kind == 'cd' else TechConsignmentStock
        try:
            stock = model.objects.select_for_update().select_related(kind).get(pk=stock_id, platform_id=platform_id)
        except model.DoesNotExist as exc:
            raise ValidationError('Партия из отчёта не найдена на этой площадке.') from exc
        product = getattr(stock, kind)
        quantity = _quantity(source['sold'])
        if quantity > stock.quantity:
            raise ValidationError(f'{product.name}: в отчёте {quantity} шт., на площадке осталось {stock.quantity}.')
        if quantity and stock.receivable_per_unit <= 0:
            raise ValidationError(f'{product.name}: не задана цена реализации.')
        result.append({**source, 'sold': quantity, 'name': product.name, 'warehouse_id': stock.warehouse_id,
                       'current_quantity': stock.quantity, 'current_price': str(stock.receivable_per_unit),
                       'amount': str(stock.receivable_per_unit * quantity),
                       'changed': stock.quantity != source['quantity'] or stock.receivable_per_unit != Decimal(source['price'])})
    return result


@transaction.atomic
def import_report(*, upload, platform_id, actor):
    if upload.size > 10 * 1024 * 1024:
        raise ValidationError('Размер отчёта превышает 10 МБ.')
    try:
        book = load_workbook(upload, read_only=True, data_only=False)
        meta = book['_resource_report']
        token = ''.join(str(row[0]) for row in meta.iter_rows(values_only=True) if row[0])
        manifest = unsign_report(token)
        if manifest.get('version') != 1 or manifest['platform_id'] != platform_id:
            raise ValidationError('Отчёт относится к другой площадке или имеет неподдерживаемый формат.')
        UUID(manifest['report_id'])
        if Sale.objects.filter(consignment_report_id=manifest['report_id']).exists():
            raise ValidationError('По этому отчёту уже создана продажа. Повторное списание запрещено.')
        sheet = book['Реализация']
        if sheet.max_row > 10000:
            raise ValidationError('В отчёте слишком много строк.')
        rows, seen = [], set()
        for values in sheet.iter_rows(min_row=5, max_col=6, values_only=True):
            if not values[5]:
                if values[3] not in (None, '', 0) and not str(values[3]).startswith('='):
                    raise ValidationError('У заполненной строки отсутствует идентификатор партии.')
                continue
            source = unsign_report(values[5])
            identity = f"{source['kind']}:{source['stock_id']}"
            if (source['report_id'] != manifest['report_id'] or source['platform_id'] != platform_id
                    or identity not in manifest['identities'] or identity in seen):
                raise ValidationError('В отчёте повторные строки, посторонние партии или неверная площадка.')
            seen.add(identity)
            rows.append({**source, 'sold': _quantity(values[3])})
        if seen != set(manifest['identities']):
            raise ValidationError('Состав отчёта повреждён: отсутствуют служебные строки партий.')
    except (BadZipFile, KeyError, ValueError, TypeError, OSError) as exc:
        raise ValidationError('Не удалось прочитать отчёт. Загрузите заполненный Excel, скачанный у этой площадки.') from exc
    finally:
        if 'book' in locals():
            book.close()
    rows = current_report_rows(platform_id=platform_id, rows=rows, actor=actor)
    if not any(row['sold'] for row in rows):
        raise ValidationError('В отчёте не указано ни одного проданного товара.')
    partner = Partner.objects.filter(platform_id=platform_id, is_archived=False).first()
    return {'report_id': manifest['report_id'], 'platform_id': platform_id,
            'partner_id': partner.pk if partner else None, 'rows': rows,
            'warehouse_id': next(row['warehouse_id'] for row in rows if row['sold'])}


@transaction.atomic
def create_report_sale(*, actor, draft, **kwargs):
    from sales.services import create_sale
    platform = SalesPlatform.objects.select_for_update().get(pk=draft['platform_id'])
    existing = Sale.objects.filter(consignment_report_id=draft['report_id']).first()
    if existing:
        if existing.created_by_id != actor.pk:
            raise ValidationError('По этому отчёту уже создана продажа.')
        return existing
    rows = current_report_rows(platform_id=platform.pk, rows=draft['rows'], actor=actor)
    if any(row['current_price'] != old['current_price'] or row['current_quantity'] != old['current_quantity']
           for row, old in zip(rows, draft['rows'])):
        raise ValidationError('Остаток или цена изменились после проверки отчёта. Загрузите его заново и проверьте различия.')
    # Состав формируется исключительно из проверенного отчёта, а не произвольного POST.
    lines = [{'product_type': 'consignment', 'product_kind': row['kind'], 'stock_id': row['stock_id'],
              'quantity': row['sold']} for row in rows if row['sold']]
    kwargs.update(warehouse_id=draft['warehouse_id'], lines=lines, sale_type=Sale.SaleType.WHOLESALE,
                  price_type=Sale.PriceType.WHOLESALE, payment_method=Sale.PaymentMethod.CASH_POSTPAY,
                  creation_key=draft['report_id'])
    linked = Partner.objects.filter(platform=platform, is_archived=False).first()
    if linked:
        kwargs.update(partner_id=linked.pk, partner_name='')
    sale = create_sale(actor=actor, **kwargs)
    sale.consignment_platform, sale.consignment_report_id = platform, draft['report_id']
    sale.save(update_fields=('consignment_platform', 'consignment_report_id'))
    return sale
