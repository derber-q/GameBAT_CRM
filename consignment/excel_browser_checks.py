"""Явно запускаемая браузерная проверка скачивания на отдельной тестовой БД."""
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.contrib.staticfiles.testing import StaticLiveServerTestCase
from django.urls import reverse
from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

from accounts.models import User
from catalog.models import CD, Platform
from partners.models import SalesPlatform
from warehouse.models import Warehouse
from .models import CDConsignmentStock


class ConsignmentExcelBrowserChecks(StaticLiveServerTestCase):
    def test_desktop_and_mobile_download_from_collapsed_filtered_panel_and_platform_page(self):
        user = User.objects.create_superuser("consignment-browser", password="StrongAdmin!123")
        self.client.force_login(user)
        shop = SalesPlatform.objects.create(name="Площадка браузерной проверки", address="А", legal_entity="А", phone_1="1")
        warehouse = Warehouse.objects.create(name="Тестовый склад")
        cd = CD.objects.create(platform=Platform.objects.create(name="PS5"), name="Игра для экспорта", sku="BROWSER-CD")
        CDConsignmentStock.objects.create(platform=shop, warehouse=warehouse, cd=cd, quantity=3, receivable_per_unit="123.45")
        output = Path(settings.BASE_DIR) / "outputs" / "consignment-excel-browser"
        output.mkdir(parents=True, exist_ok=True)
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch()
            for width in (1280, 390):
                context = browser.new_context(viewport={"width": width, "height": 900}, accept_downloads=True)
                context.add_cookies([{"name": settings.SESSION_COOKIE_NAME,
                                     "value": self.client.cookies[settings.SESSION_COOKIE_NAME].value,
                                     "url": self.live_server_url}])
                context.route("**/crm/api/exchange-rates/", lambda route: route.fulfill(json={"rates": {}, "available": False}))
                page = context.new_page()
                errors = []
                page.on("pageerror", lambda error: errors.append(str(error)))
                page.goto(self.live_server_url + reverse("consignment:list") + "?q=ничего-нет")
                toggle = page.locator(f'[aria-controls="consignment-platform-{shop.pk}"]')
                toggle.click()
                expect(page.locator(f"#consignment-platform-{shop.pk}")).to_be_hidden()
                button = page.get_by_role("link", name="Скачать Excel", exact=True)
                expect(button).to_be_visible()
                with page.expect_download() as downloaded:
                    button.click()
                download = downloaded.value
                self.assertIsNone(download.failure())
                workbook = load_workbook(BytesIO(Path(download.path()).read_bytes()))
                self.assertTrue(any(row[0] == cd.name and row[1] == 3 and row[2] == 123.45 and row[3] == 0
                                    for row in workbook.active.iter_rows(values_only=True)))
                page.screenshot(path=str(output / f"collapsed-{width}.png"), full_page=True)
                page.goto(self.live_server_url + reverse("consignment:platform", args=(shop.pk,)))
                expect(page.get_by_role("link", name="Скачать Excel", exact=True)).to_be_visible()
                with page.expect_download() as downloaded:
                    page.get_by_role("link", name="Скачать Excel", exact=True).click()
                self.assertIsNone(downloaded.value.failure())
                page.screenshot(path=str(output / f"platform-{width}.png"), full_page=True)
                self.assertEqual(errors, [])
                context.close()
            browser.close()
