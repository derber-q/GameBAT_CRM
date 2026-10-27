from decimal import Decimal
from io import BytesIO
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from accounts.models import User
from catalog.models import CD, Platform
from partners.models import SalesPlatform, Partner
from warehouse.models import Warehouse, CDWarehouseStock
from sales.models import Sale
from sales.services import cancel_sale, mark_sale_paid
from .excel import generate_consignment_xlsx
from .reports import import_report, create_report_sale, DRAFT_KEY
from .services import transfer_to_consignment, update_consignment_reward


class ConsignmentReportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser('report', password='test')
        self.platform = SalesPlatform.objects.create(name='Площадка', address='Адрес', legal_entity='ООО', phone_1='1')
        self.warehouse = Warehouse.objects.create(name='Склад')
        self.cd = CD.objects.create(name='Игра', sku='REPORT', platform=Platform.objects.create(name='PS5'), cost=100)
        CDWarehouseStock.objects.create(cd=self.cd, warehouse=self.warehouse, quantity=10)
        transfer_to_consignment(actor=self.user, warehouse_id=self.warehouse.pk, platform_id=self.platform.pk,
            product_type='cd', product_id=self.cd.pk, quantity=5, receivable_per_unit=150)
        self.stock = self.cd.consignment_stocks.get()
        self.data = generate_consignment_xlsx(platform=self.platform, user=self.user)
        self.client.force_login(self.user)

    def upload(self, sold=2, mutate=None):
        book = load_workbook(BytesIO(self.data))
        sheet = book.active
        row = next(row for row in sheet if row[5].value)
        sheet.cell(row[0].row,4,sold)
        if mutate: mutate(sheet,row[0].row)
        stream=BytesIO(); book.save(stream); book.close()
        return SimpleUploadedFile('report.xlsx',stream.getvalue())

    def draft(self, **kwargs):
        return import_report(upload=self.upload(**kwargs),platform_id=self.platform.pk,actor=self.user)

    def sale(self, draft=None):
        return create_report_sale(actor=self.user,draft=draft or self.draft(),partner_name='Заказчик',note='Отчёт площадки')

    def test_import_is_read_only_and_uses_signed_stock_not_edited_name_and_price(self):
        draft=self.draft(mutate=lambda sheet,row:(sheet.cell(row,1,'Подмена'),sheet.cell(row,3,1),sheet.cell(row,5,1)))
        self.assertEqual((draft['rows'][0]['name'],draft['rows'][0]['amount']),('Игра','300.00'))
        self.stock.refresh_from_db();self.assertEqual(self.stock.quantity,5)
        self.assertFalse(Sale.objects.exists()); self.assertFalse(Partner.objects.exists())

    def test_negative_fraction_formula_excess_and_duplicate_rows_rejected(self):
        for sold in (-1,'1.5','=1+1',6,True):
            with self.subTest(sold=sold),self.assertRaises(ValidationError): self.draft(sold=sold)
        def duplicate(sheet,row): sheet.append([cell.value for cell in sheet[row]])
        with self.assertRaises(ValidationError):self.draft(mutate=duplicate)
        def damage(sheet,row):sheet.cell(row,6,'invalid')
        with self.assertRaises(ValidationError):self.draft(mutate=damage)

    def test_other_platform_and_deleted_identity_rejected(self):
        other=SalesPlatform.objects.create(name='Другая',address='А',legal_entity='ООО',phone_1='2')
        with self.assertRaises(ValidationError):import_report(upload=self.upload(),platform_id=other.pk,actor=self.user)
        with self.assertRaises(ValidationError):self.draft(mutate=lambda sheet,row:sheet.delete_rows(row))

    def test_changed_price_preview_and_stale_submit_rejected(self):
        update_consignment_reward(actor=self.user,product_type='cd',stock_id=self.stock.pk,value=160)
        draft=self.draft();self.assertTrue(draft['rows'][0]['changed']);self.assertEqual(draft['rows'][0]['amount'],'320.00')
        update_consignment_reward(actor=self.user,product_type='cd',stock_id=self.stock.pk,value=170)
        with self.assertRaises(ValidationError):self.sale(draft)
        self.assertFalse(Sale.objects.exists())

    def test_standard_sale_report_repeat_and_cancellation_restore_exact_lot(self):
        draft=self.draft();sale=self.sale(draft)
        self.assertEqual((sale.sale_type,sale.payment_method,sale.total_amount),('wholesale','cash_postpay',300))
        self.assertEqual(sale.consignment_items.get().source_stock_id,self.stock.pk)
        self.stock.refresh_from_db();self.assertEqual(self.stock.quantity,3)
        self.assertEqual(self.sale(draft).pk,sale.pk)
        self.stock.refresh_from_db();self.assertEqual(self.stock.quantity,3)
        with self.assertRaises(ValidationError):self.draft()
        cancel_sale(actor=self.user,sale_id=sale.pk,comment='Отмена отчёта')
        self.stock.refresh_from_db();self.assertEqual(self.stock.quantity,5)

    def test_linked_partner_is_selected_automatically(self):
        partner=Partner.objects.create(name='Партнёр',platform=self.platform)
        draft=self.draft();self.assertEqual(draft['partner_id'],partner.pk)
        self.assertEqual(self.sale(draft).partner_id,partner.pk)

    def test_upload_preview_standard_form_and_submit(self):
        response=self.client.post(reverse('consignment:report_import',args=(self.platform.pk,)),{'file':self.upload(),'customer':'custom','custom_partner_name':'Заказчик'})
        self.assertContains(response,'Открыть обычную форму продажи')
        response=self.client.post(reverse('consignment:report_import',args=(self.platform.pk,)),{'confirm_report':'1','customer':'custom','custom_partner_name':'Заказчик'})
        self.assertEqual(response.status_code,302)
        response=self.client.get(response.url);self.assertContains(response,'consignment_report')
        key=response.context['creation_key']
        response=self.client.post(reverse('sales:create'),{'consignment_report':'1','creation_key':key,'warehouse':self.warehouse.pk,
            'price_type':'wholesale','sale_type':'wholesale','payment_method':'cash_postpay','customer':'custom','custom_partner_name':'Заказчик','note':'Отчёт'})
        self.assertEqual(response.status_code,302)
        self.assertEqual(Sale.objects.get().total_amount,300)
        response=self.client.post(reverse('sales:create'),{'consignment_report':'1','creation_key':key})
        self.assertEqual(response.status_code,302)
        self.assertEqual(Sale.objects.count(),1)
        self.stock.refresh_from_db();self.assertEqual(self.stock.quantity,3)

    def test_partial_payment_prompts_partner_only_at_payment_and_cancel_keeps_unpaid(self):
        sale=self.sale();self.assertFalse(Partner.objects.exists())
        response=self.client.post(reverse('sales:mark_paid',args=(sale.pk,)),{'cash_received_amount':'200'})
        self.assertRedirects(response,reverse('consignment:report_partner_payment',args=(sale.pk,)))
        self.assertFalse(Partner.objects.exists());sale.refresh_from_db();self.assertEqual(sale.payment_status,'unpaid')
        page=self.client.get(response.url)
        self.assertContains(page,'100,00')
        form=page.context['form']
        response=self.client.post(response.url,{'name':'Новый партнёр','request_key':str(form['request_key'].value()),'payment_note':'Недоплата по отчёту'})
        self.assertEqual(response.status_code,302)
        sale.refresh_from_db();partner=sale.partner
        self.assertEqual(partner.platform_id,self.platform.pk);self.assertEqual(partner.creditor.current_debt,100)
        self.assertEqual(sale.cash_transactions.filter(operation_type='sale_payment').count(),1)
        cancel_sale(actor=self.user,sale_id=sale.pk,comment='Отмена')
        partner.creditor.refresh_from_db();self.assertEqual(partner.creditor.current_debt,0)
        self.stock.refresh_from_db();self.assertEqual(self.stock.quantity,5)

    def test_full_payment_does_not_create_partner(self):
        sale=self.sale();mark_sale_paid(actor=self.user,sale_id=sale.pk,cash_received_amount=300)
        sale.refresh_from_db()
        self.assertFalse(Partner.objects.exists());self.assertEqual(sale.cash_received_amount,300)

    def test_reward_ajax_errors_and_success_keep_row_protocol(self):
        url=reverse('consignment:reward_update',args=('cd',self.stock.pk))
        response=self.client.post(url,{'receivable_per_unit':'-1'},HTTP_ACCEPT='application/json')
        self.assertEqual(response.status_code,400);self.assertFalse(response.json()['success']);self.assertNotIn('row_html',response.json())
        response=self.client.post(url,{'receivable_per_unit':'175'},HTTP_ACCEPT='application/json')
        self.assertTrue(response.json()['success']);self.assertIn('175',response.json()['row_html'])
