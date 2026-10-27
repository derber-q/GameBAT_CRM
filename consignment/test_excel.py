"""Проверки выгрузки реализации: права, реальные остатки, партии и отсутствие затрат."""
from decimal import Decimal
from io import BytesIO
from zipfile import ZipFile

from django.contrib.auth.models import Group, Permission
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import User
from catalog.models import Brand, CD, Platform, ProductType, Tech
from partners.models import SalesPlatform
from warehouse.models import CDWarehouseStock, TechWarehouseStock, Warehouse
from .models import CDConsignmentStock, TechConsignmentStock
from .services import record_consignment_sale, return_from_consignment, transfer_to_consignment, update_consignment_reward


class ConsignmentExcelTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser("excel-admin", password="StrongAdmin!123")
        cls.viewer = User.objects.create_user("excel-viewer")
        cls.warehouse = Warehouse.objects.create(name="Склад")
        cls.shop = SalesPlatform.objects.create(name="Магазин А", address="А", legal_entity="А", phone_1="1")
        cls.other = SalesPlatform.objects.create(name="Магазин Б", address="Б", legal_entity="Б", phone_1="2")
        cls.empty = SalesPlatform.objects.create(name="Пустая площадка", address="В", legal_entity="В", phone_1="3")
        cls.ps5 = Platform.objects.create(name="PlayStation 5")
        cls.ps4 = Platform.objects.create(name="PlayStation 4")
        cls.cd = CD.objects.create(platform=cls.ps5, name="Зета", sku="000123", barcode="123", cost="9876543.21")
        cls.alpha = CD.objects.create(platform=cls.ps5, name="Альфа", sku="CD-A", cost="8765432.10")
        cls.old = CD.objects.create(platform=cls.ps4, name="Игра PS4", sku="CD-4", cost="7654321.09")
        brand = Brand.objects.create(name="Sony")
        cls.tech = Tech.objects.create(brand=brand, product_type=ProductType.objects.create(name="Геймпад"),
                                      name="Контроллер", sku="TECH-C", cost="6543210.98")
        cls.console = Tech.objects.create(brand=brand, product_type=ProductType.objects.create(name="Консоль"),
                                         name="Консоль", sku="TECH-S", cost="5432109.87")
        for item in (cls.cd, cls.alpha, cls.old):
            CDWarehouseStock.objects.create(warehouse=cls.warehouse, cd=item, quantity=100)
        for item in (cls.tech, cls.console):
            TechWarehouseStock.objects.create(warehouse=cls.warehouse, tech=item, quantity=100)
        cls.cd_stock = cls.transfer(cls.cd, 4, "123.45")
        alpha_stock = cls.transfer(cls.alpha, 2, "1")
        update_consignment_reward(actor=cls.admin, product_type="cd", stock_id=alpha_stock.pk, value="0")
        cls.transfer(cls.old, 1, "7.89")
        cls.tech_stock = cls.transfer(cls.tech, 3, "678.91")
        cls.transfer(cls.console, 1, "99.99")
        cls.transfer(cls.cd, 20, "777.77", platform=cls.other)

    @classmethod
    def transfer(cls, product, quantity, reward, platform=None):
        return transfer_to_consignment(
            actor=cls.admin, warehouse_id=cls.warehouse.pk, platform_id=(platform or cls.shop).pk,
            product_type="cd" if isinstance(product, CD) else "tech", product_id=product.pk,
            quantity=quantity, reward_per_unit=reward,
        )

    def setUp(self):
        self.client.force_login(self.admin)

    def export(self, sales_platform=None, **params):
        return self.client.get(reverse("consignment:platform_export", args=((sales_platform or self.shop).pk,)), params)

    def workbook(self, response):
        self.assertEqual(response.status_code, 200)
        return load_workbook(BytesIO(response.content), data_only=False)

    def products(self, sheet):
        return [row[:5] for row in sheet.iter_rows(min_row=5, values_only=True) if len(row) > 5 and row[5]]

    def test_buttons_platform_scope_mime_and_get_only(self):
        response=self.client.get(reverse("consignment:list"),{"q":"Не найдено"})
        self.assertContains(response,"Скачать Excel",count=3)
        response=self.export()
        self.assertEqual(response["Content-Type"],"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("no-store",response["Cache-Control"])
        self.assertEqual(self.client.post(reverse("consignment:platform_export",args=(self.shop.pk,))).status_code,405)
        self.assertEqual(self.client.get(reverse("consignment:platform_export",args=(999999,))).status_code,404)

    def test_visible_columns_current_stock_and_sold_formulas(self):
        book=self.workbook(self.export());sheet=book.active
        self.assertEqual([cell.value for cell in sheet[4]][:5],["Наименование","Остаток на реализации","Цена за единицу, ₽","Количество реализованного","Сумма реализованного, ₽"])
        rows=self.products(sheet)
        self.assertEqual([row[:4] for row in rows],[("Игра PS4",1,7.89,0),("Альфа",2,0,0),("Зета",4,123.45,0),("Контроллер",3,678.91,0),("Консоль",1,99.99,0)])
        self.assertTrue(all(row[4].startswith('=IF(') for row in rows))
        self.assertEqual(sheet.cell(sheet.max_row,1).value,"Итого площадка должна за реализованные товары")
        self.assertTrue(sheet.cell(sheet.max_row,5).value.startswith('=SUM('))
        self.assertEqual([row[:4] for row in self.products(self.workbook(self.export(self.other)).active)],[("Зета",20,777.77,0)])
        self.assertEqual(sheet.freeze_panes,"A5")
        self.assertTrue(sheet.column_dimensions['F'].hidden)
        self.assertEqual(book['_resource_report'].sheet_state,'veryHidden')
        self.assertIsNotNone(sheet.data_validations)
        book.close()

    def test_grouping_format_and_filters_do_not_limit_export(self):
        sheet=self.workbook(self.export(q='Нет',page=99,platform=999999,kind='tech')).active
        self.assertEqual(len(self.products(sheet)),5)
        labels=[sheet.cell(area.min_row,1).value for area in sorted(sheet.merged_cells.ranges,key=lambda area:area.min_row)]
        for title in ('CD · Диски','PlayStation 4','PlayStation 5','Tech · Техника','Геймпад','Консоль'):self.assertIn(title,labels)
        for row in sheet:
            if row[0].value=='Зета':
                self.assertEqual(row[1].number_format,'0');self.assertIn('₽',row[2].number_format);self.assertTrue(row[0].alignment.wrap_text)

    def test_distinct_lots_are_not_merged_or_multiplied_by_barcodes(self):
        from catalog.product_identifiers import add_product_barcode
        add_product_barcode(actor=self.admin,product_kind='cd',product_id=self.cd.pk,value='001999')
        CDConsignmentStock.objects.create(platform=self.shop,warehouse=self.warehouse,cd=self.cd,lot_key='historical:1',quantity=2,receivable_per_unit='222.22')
        other=Warehouse.objects.create(name='Другой склад')
        CDConsignmentStock.objects.create(platform=self.shop,warehouse=other,cd=self.cd,quantity=1,receivable_per_unit='123.45')
        rows=[row[:4] for row in self.products(self.workbook(self.export()).active) if row[0]=='Зета']
        self.assertCountEqual(rows,[('Зета',4,123.45,0),('Зета',1,123.45,0),('Зета',2,222.22,0)])

    def test_new_download_reflects_sale_return_reward_and_excludes_zero(self):
        record_consignment_sale(actor=self.admin,product_type='cd',stock_id=self.cd_stock.pk,quantity=1,payment_method='bank_account')
        return_from_consignment(actor=self.admin,warehouse_id=self.warehouse.pk,platform_id=self.shop.pk,product_type='tech',product_id=self.tech.pk,quantity=1,stock_id=self.tech_stock.pk)
        update_consignment_reward(actor=self.admin,product_type='cd',stock_id=self.cd_stock.pk,value='111.11')
        rows={row[0]:row for row in self.products(self.workbook(self.export()).active)}
        self.assertEqual(rows['Зета'][1:4],(3,111.11,0));self.assertEqual(rows['Контроллер'][1:4],(2,678.91,0))
        record_consignment_sale(actor=self.admin,product_type='cd',stock_id=self.cd_stock.pk,quantity=3,payment_method='bank_account')
        self.assertNotIn('Зета',[row[0] for row in self.products(self.workbook(self.export()).active)])

    def test_no_cost_profit_or_writes_including_hidden_metadata(self):
        with CaptureQueriesContext(connection) as queries:response=self.export()
        for query in queries:
            sql=query['sql'].lower();self.assertNotRegex(sql,r'\b(insert|update|delete|replace)\b');self.assertNotIn('"unit_cost"',sql);self.assertNotIn('unit_cost_snapshot',sql)
        with ZipFile(BytesIO(response.content)) as archive:payload='\n'.join(archive.read(name).decode('utf-8') for name in archive.namelist()).casefold()
        for forbidden in ('себестоимость','закупочная','прибыль','маржинальность','наценка','unit_cost','9876543.21','8765432.1','7654321.09','6543210.98','5432109.87'):self.assertNotIn(forbidden,payload)

    def test_empty_platform_is_valid_workbook(self):
        sheet=self.workbook(self.export(self.empty)).active
        self.assertEqual(sheet['A5'].value,'Нет товаров на реализации');self.assertEqual(self.products(sheet),[])
        self.assertEqual(sheet.cell(sheet.max_row,5).value,'=SUM(E5:E5)')

    def test_permissions_scope_each_kind(self):
        self.client.logout();self.assertEqual(self.export().status_code,302)
        self.client.force_login(self.viewer);self.assertEqual(self.export().status_code,403)
        for kind,expected in (('cd',{'Игра PS4','Альфа','Зета'}),('tech',{'Контроллер','Консоль'})):
            group=Group.objects.create(name=f'Excel {kind}')
            group.permissions.add(Permission.objects.get(content_type__app_label='consignment',codename=f'view_{kind}consignmentstock'))
            self.viewer.groups.set([group]);self.client.force_login(self.viewer)
            self.assertEqual({row[0] for row in self.products(self.workbook(self.export()).active)},expected)

    def test_user_text_is_literal_not_formula(self):
        self.alpha.name='=HYPERLINK("https://example.com","Игра")';self.alpha.save(update_fields=('name',))
        rows=self.products(self.workbook(self.export()).active)
        self.assertIn("'"+self.alpha.name,[row[0] for row in rows])
