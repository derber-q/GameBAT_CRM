from collections import defaultdict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from catalog.management.commands.import_catalog_names import _base_normalize, _is_same_product
from catalog.models import CD, Tech
from consignment.models import CDConsignmentStock, TechConsignmentStock
from consignment.services import return_from_consignment, transfer_many_to_consignment
from partners.models import SalesPlatform
from warehouse.models import CDWarehouseStock, TechWarehouseStock, Warehouse


CENT = Decimal("0.01")
EXPECTED_SOURCE_ROWS = 62
EXPECTED_SOURCE_UNITS = 77


@dataclass(frozen=True)
class SourceLine:
    row: int
    kind: str
    name: str
    quantity: int
    receivable: Decimal
    platform: str = ""


@dataclass(frozen=True)
class PreparedLine:
    kind: str
    product_id: int
    product_name: str
    quantity: int
    receivable: Decimal
    source_rows: tuple[int, ...]
    source_total: Decimal


def _money(value, *, row, label):
    try:
        result = Decimal(str(value)).quantize(CENT, rounding=ROUND_HALF_UP)
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise CommandError(f"Строка {row}: некорректное значение «{label}».") from exc
    if result <= 0:
        raise CommandError(f"Строка {row}: значение «{label}» должно быть больше нуля.")
    return result


def _quantity(value, *, row):
    try:
        result = int(value)
    except (TypeError, ValueError) as exc:
        raise CommandError(f"Строка {row}: некорректное количество.") from exc
    if result <= 0 or result != value:
        raise CommandError(f"Строка {row}: количество должно быть положительным целым числом.")
    return result


def _read_source(path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook.active
        lines = []
        for row in range(3, 65):
            name = sheet.cell(row, 1).value
            if name in (None, ""):
                continue
            kind = "tech" if row <= 9 else "cd"
            platform = ""
            if kind == "cd":
                platform = "Nintendo Switch 2" if row <= 30 else "Nintendo Switch"
            lines.append(SourceLine(
                row=row,
                kind=kind,
                name=" ".join(str(name).split()),
                quantity=_quantity(sheet.cell(row, 3).value, row=row),
                receivable=_money(sheet.cell(row, 2).value, row=row, label="вознаграждение"),
                platform=platform,
            ))
        return lines
    finally:
        workbook.close()


def _match_line(line):
    if line.kind == "tech":
        candidates = list(Tech.objects.all())
    else:
        candidates = list(CD.objects.select_related("platform").filter(platform__name=line.platform))

    exact = [obj for obj in candidates if _base_normalize(obj.name) == _base_normalize(line.name)]
    fuzzy = [
        obj for obj in candidates
        if _is_same_product(line.name, obj.name, cd=line.kind == "cd", existing=True)
    ]
    matches = {obj.pk: obj for obj in exact + fuzzy}

    # В источнике сокращено название существующей позиции Master Collection.
    if not matches and line.row == 16:
        matches = {
            obj.pk: obj for obj in candidates
            if "metal gear solid vol 2" in _base_normalize(obj.name)
        }

    if len(matches) == 1:
        return next(iter(matches.values()))
    if len(matches) > 1:
        details = ", ".join(f"#{obj.pk} {obj.name}" for obj in matches.values())
        raise CommandError(f"Строка {line.row}: неоднозначное совпадение «{line.name}»: {details}")
    raise CommandError(f"Строка {line.row}: товар не найден в номенклатуре: «{line.name}».")


def _prepare(lines):
    grouped = defaultdict(list)
    products = {}
    for line in lines:
        product = _match_line(line)
        key = (line.kind, product.pk)
        grouped[key].append(line)
        products[key] = product

    prepared = []
    weighted = []
    for key, source_lines in grouped.items():
        quantity = sum(line.quantity for line in source_lines)
        source_total = sum(line.receivable * line.quantity for line in source_lines)
        receivable = (source_total / quantity).quantize(CENT, rounding=ROUND_HALF_UP)
        rates = {line.receivable for line in source_lines}
        if len(rates) > 1:
            weighted.append((products[key].name, tuple(line.row for line in source_lines), receivable))
        prepared.append(PreparedLine(
            kind=key[0],
            product_id=key[1],
            product_name=products[key].name,
            quantity=quantity,
            receivable=receivable,
            source_rows=tuple(line.row for line in source_lines),
            source_total=source_total,
        ))
    prepared.sort(key=lambda item: (item.kind, item.product_name.casefold(), item.product_id))
    return prepared, weighted


def _stock_config(kind):
    if kind == "cd":
        return CDConsignmentStock, CDWarehouseStock, "cd"
    return TechConsignmentStock, TechWarehouseStock, "tech"


def _current_lines(platform, warehouse):
    result = []
    for kind in ("cd", "tech"):
        consignment_model, _warehouse_model, field = _stock_config(kind)
        for stock in consignment_model.objects.filter(
            platform=platform, warehouse=warehouse, quantity__gt=0
        ).select_related(field):
            product = getattr(stock, field)
            result.append(PreparedLine(
                kind=kind,
                product_id=product.pk,
                product_name=product.name,
                quantity=stock.quantity,
                receivable=stock.receivable_per_unit,
                source_rows=(),
                source_total=stock.receivable_per_unit * stock.quantity,
            ))
    return result


def _signature(lines):
    return {
        (line.kind, line.product_id): (line.quantity, line.receivable)
        for line in lines
    }


class Command(BaseCommand):
    help = "Полностью заменяет остатки площадки PlayBox на данные из Excel."

    def add_arguments(self, parser):
        parser.add_argument("xlsx", type=Path)
        parser.add_argument("--apply", action="store_true")
        parser.add_argument("--actor", default="derber")

    def handle(self, *args, **options):
        path = options["xlsx"].resolve()
        if not path.is_file():
            raise CommandError(f"Файл не найден: {path}")

        lines = _read_source(path)
        units = sum(line.quantity for line in lines)
        if len(lines) != EXPECTED_SOURCE_ROWS or units != EXPECTED_SOURCE_UNITS:
            raise CommandError(
                f"Контрольный итог источника не сошёлся: {len(lines)} строк, {units} шт.; "
                f"ожидалось {EXPECTED_SOURCE_ROWS} строк и {EXPECTED_SOURCE_UNITS} шт."
            )
        prepared, weighted = _prepare(lines)

        if Warehouse.objects.count() != 1:
            raise CommandError("Для импорта должен существовать ровно один склад.")
        warehouse = Warehouse.objects.get()
        try:
            platform = SalesPlatform.objects.get(name__iexact="PlayBox")
        except SalesPlatform.DoesNotExist as exc:
            raise CommandError("Площадка PlayBox не найдена.") from exc

        current = _current_lines(platform, warehouse)
        source_total = sum(line.receivable * line.quantity for line in lines)
        stored_total = sum(line.receivable * line.quantity for line in prepared)
        self.stdout.write(
            f"Источник: {len(lines)} строк, {len(prepared)} уникальных позиций, {units} шт., "
            f"сумма вознаграждений {source_total:.2f} ₽."
        )
        self.stdout.write(
            f"Сейчас на PlayBox: {len(current)} позиций, {sum(line.quantity for line in current)} шт."
        )
        for name, rows, rate in weighted:
            self.stdout.write(self.style.WARNING(
                f"Разные ставки объединены для «{name}» (строки {', '.join(map(str, rows))}): "
                f"средневзвешенная ставка {rate:.2f} ₽."
            ))
        if stored_total != source_total:
            self.stdout.write(self.style.WARNING(
                f"Из-за округления ставок итог в модели составит {stored_total:.2f} ₽ "
                f"(отклонение {stored_total - source_total:+.2f} ₽)."
            ))

        if _signature(current) == _signature(prepared):
            self.stdout.write(self.style.SUCCESS("PlayBox уже полностью соответствует источнику."))
            return
        if not options["apply"]:
            self.stdout.write(self.style.WARNING(
                "Проверка завершена, база не изменена. Для записи добавьте --apply."
            ))
            return

        User = get_user_model()
        try:
            actor = User.objects.get(username=options["actor"], is_active=True)
        except User.DoesNotExist as exc:
            raise CommandError(f"Активный пользователь «{options['actor']}» не найден.") from exc

        expected_by_key = {(line.kind, line.product_id): line.quantity for line in prepared}
        current_by_key = {(line.kind, line.product_id): line.quantity for line in current}
        affected_keys = set(expected_by_key) | set(current_by_key)

        with transaction.atomic():
            baseline = {}
            for kind, product_id in sorted(affected_keys):
                _consignment_model, warehouse_model, field = _stock_config(kind)
                warehouse_stock, _created = warehouse_model.objects.select_for_update().get_or_create(
                    warehouse=warehouse, **{f"{field}_id": product_id}
                )
                baseline[(kind, product_id)] = warehouse_stock.quantity

            # Штатно возвращаем весь старый список, чтобы сохранить историю операций.
            for line in current:
                return_from_consignment(
                    actor=actor,
                    warehouse_id=warehouse.pk,
                    platform_id=platform.pk,
                    product_type=line.kind,
                    product_id=line.product_id,
                    quantity=line.quantity,
                )

            # Это корректирующий импорт фактической реализации, поэтому основной склад
            # должен остаться ровно таким, каким был до замены списка PlayBox.
            for kind, product_id in sorted(affected_keys):
                _consignment_model, warehouse_model, field = _stock_config(kind)
                warehouse_stock = warehouse_model.objects.select_for_update().get(
                    warehouse=warehouse, **{f"{field}_id": product_id}
                )
                warehouse_stock.quantity = baseline[(kind, product_id)] + expected_by_key.get((kind, product_id), 0)
                warehouse_stock.full_clean()
                warehouse_stock.save(update_fields=("quantity",))

            movement = transfer_many_to_consignment(
                actor=actor,
                warehouse_id=warehouse.pk,
                platform_id=platform.pk,
                lines=[{
                    "product_type": line.kind,
                    "product_id": line.product_id,
                    "quantity": line.quantity,
                    "receivable_per_unit": line.receivable,
                } for line in prepared],
            )

            # Нулевые строки старого списка не должны оставаться в справочнике площадки.
            for kind in ("cd", "tech"):
                consignment_model, _warehouse_model, _field = _stock_config(kind)
                consignment_model.objects.filter(platform=platform, warehouse=warehouse, quantity=0).delete()

            final = _current_lines(platform, warehouse)
            if _signature(final) != _signature(prepared):
                raise CommandError("Итоговый список PlayBox не совпал с подготовленными данными.")
            for key, quantity in baseline.items():
                kind, product_id = key
                _consignment_model, warehouse_model, field = _stock_config(kind)
                actual = warehouse_model.objects.get(
                    warehouse=warehouse, **{f"{field}_id": product_id}
                ).quantity
                if actual != quantity:
                    raise CommandError(
                        f"Основной склад изменился для {kind} #{product_id}: было {quantity}, стало {actual}."
                    )

        self.stdout.write(self.style.SUCCESS(
            f"PlayBox заменён: операция передачи №{movement.pk}, {len(prepared)} позиций, {units} шт. "
            "Основной склад сохранён без изменений."
        ))
