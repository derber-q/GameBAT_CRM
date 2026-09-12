from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from catalog.management.commands.import_catalog_names import _base_normalize, _is_same_product
from catalog.models import Brand, CD, Platform, ProductType, Tech
from consignment.models import CDConsignmentStock, TechConsignmentStock
from consignment.services import transfer_many_to_consignment
from partners.models import SalesPlatform
from warehouse.models import CDWarehouseStock, TechWarehouseStock, Warehouse


CENT = Decimal("0.01")
HEADER_ROWS = {1, 2, 8, 55}


@dataclass(frozen=True)
class SourceLine:
    row: int
    kind: str
    name: str
    quantity: int
    receivable: Decimal
    cost: Decimal
    platform: str = ""


def _decimal(value, *, row, label):
    try:
        result = Decimal(str(value)).quantize(CENT, rounding=ROUND_HALF_UP)
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise CommandError(f"Строка {row}: некорректное значение «{label}». ") from exc
    if result < 0:
        raise CommandError(f"Строка {row}: значение «{label}» не может быть отрицательным.")
    return result


def _read_source(path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook.active
        result = []
        for row, values in enumerate(sheet.iter_rows(min_col=1, max_col=7, values_only=True), start=1):
            if row in HEADER_ROWS:
                continue
            name, receivable, _days, _profit, quantity, _total, cost = values
            if name in (None, ""):
                continue
            try:
                quantity = int(quantity)
            except (TypeError, ValueError) as exc:
                raise CommandError(f"Строка {row}: некорректное количество.") from exc
            if quantity <= 0:
                raise CommandError(f"Строка {row}: количество должно быть больше нуля.")
            kind = "cd" if row < 55 else "tech"
            platform = "PlayStation 4" if 3 <= row <= 7 else ("PlayStation 5" if 9 <= row <= 54 else "")
            result.append(SourceLine(
                row=row,
                kind=kind,
                name=" ".join(str(name).split()),
                quantity=quantity,
                receivable=_decimal(receivable, row=row, label="вознаграждение"),
                cost=_decimal(cost, row=row, label="себестоимость"),
                platform=platform,
            ))
        return result
    finally:
        workbook.close()


def _match_line(line):
    if line.kind == "cd":
        candidates = list(CD.objects.select_related("platform").filter(platform__name=line.platform))
    else:
        candidates = list(Tech.objects.select_related("brand", "product_type"))
    exact = [obj for obj in candidates if _base_normalize(obj.name) == _base_normalize(line.name)]
    if len(exact) == 1:
        return exact[0]
    fuzzy = [obj for obj in candidates if _is_same_product(line.name, obj.name, cd=line.kind == "cd", existing=True)]
    unique = {obj.pk: obj for obj in exact + fuzzy}
    if len(unique) == 1:
        return next(iter(unique.values()))
    if len(unique) > 1:
        names = ", ".join(f"#{obj.pk} {obj.name}" for obj in unique.values())
        raise CommandError(f"Строка {line.row}: неоднозначное совпадение «{line.name}»: {names}")
    return None


class Command(BaseCommand):
    help = "Импортирует фактические остатки и вознаграждения площадки iTech из Excel."

    def add_arguments(self, parser):
        parser.add_argument("xlsx", type=Path)
        parser.add_argument("--apply", action="store_true")
        parser.add_argument("--actor", default="derber")

    def handle(self, *args, **options):
        path = options["xlsx"].resolve()
        if not path.is_file():
            raise CommandError(f"Файл не найден: {path}")
        lines = _read_source(path)
        if len(lines) != 53 or sum(line.quantity for line in lines) != 57:
            raise CommandError(
                f"Контрольный итог источника не сошёлся: {len(lines)} позиций, "
                f"{sum(line.quantity for line in lines)} шт.; ожидалось 53 позиции и 57 шт."
            )

        if Warehouse.objects.count() != 1:
            raise CommandError("Для импорта должен существовать ровно один склад.")
        warehouse = Warehouse.objects.get()
        try:
            platform = SalesPlatform.objects.get(name="iTech")
        except SalesPlatform.DoesNotExist as exc:
            raise CommandError("Площадка iTech не найдена.") from exc

        matched = {}
        missing = []
        already_present = []
        transfer_lines = []
        for line in lines:
            product = _match_line(line)
            if product is None:
                missing.append(line)
                continue
            matched[line.row] = product
            stock_model = CDConsignmentStock if line.kind == "cd" else TechConsignmentStock
            product_field = "cd" if line.kind == "cd" else "tech"
            current = stock_model.objects.filter(
                warehouse=warehouse, platform=platform, **{product_field: product}
            ).first()
            if current is not None and current.quantity == line.quantity and current.receivable_per_unit == line.receivable:
                already_present.append(line)
                continue
            if current is not None and current.quantity:
                raise CommandError(
                    f"Строка {line.row}: на iTech уже есть «{product.name}» в количестве {current.quantity} "
                    f"с вознаграждением {current.receivable_per_unit}."
                )
            transfer_lines.append({
                "product_type": line.kind,
                "product_id": product.pk,
                "quantity": line.quantity,
                "receivable_per_unit": line.receivable,
            })

        allowed_missing = {"Labubu Have a Seat", "Labubu Energy"}
        unexpected_missing = [line for line in missing if line.name not in allowed_missing]
        if unexpected_missing:
            details = ", ".join(f"строка {line.row} «{line.name}»" for line in unexpected_missing)
            raise CommandError(f"Не найдены товары в основной базе: {details}.")

        self.stdout.write(
            f"Источник: {len(lines)} позиций, {sum(line.quantity for line in lines)} шт., "
            f"вознаграждение всего {sum(line.receivable * line.quantity for line in lines):.2f} ₽."
        )
        self.stdout.write(
            f"Найдено в базе: {len(matched)}; требуется создать: {len(missing)}; "
            f"уже корректно на iTech: {len(already_present)}; к передаче: {len(transfer_lines)}."
        )
        if transfer_lines:
            self.stdout.write("Текущие количества основного склада будут сохранены; единицы iTech учитываются дополнительно.")
        for line in missing:
            self.stdout.write(self.style.WARNING(
                f"Будет создана техника: строка {line.row}, {line.name}, себестоимость {line.cost:.2f} ₽, "
                f"количество {line.quantity}."
            ))

        if not options["apply"]:
            self.stdout.write(self.style.WARNING("Проверка завершена, база не изменена. Для записи добавьте --apply."))
            return

        if already_present and len(already_present) != len(lines):
            raise CommandError("Площадка iTech заполнена частично; автоматическое дополнение отменено.")
        if len(already_present) == len(lines):
            self.stdout.write(self.style.SUCCESS("Площадка iTech уже полностью соответствует источнику."))
            return

        User = get_user_model()
        try:
            actor = User.objects.get(username=options["actor"], is_active=True)
        except User.DoesNotExist as exc:
            raise CommandError(f"Активный пользователь «{options['actor']}» не найден.") from exc
        tech_brand = Brand.objects.get(name="POP MART")
        tech_type = ProductType.objects.get(name="Коллекционная фигурка")

        with transaction.atomic():
            for line in lines:
                product = matched.get(line.row)
                if product is None:
                    continue
                stock_model = CDWarehouseStock if line.kind == "cd" else TechWarehouseStock
                product_field = "cd" if line.kind == "cd" else "tech"
                warehouse_stock, _created = stock_model.objects.select_for_update().get_or_create(
                    warehouse=warehouse,
                    **{product_field: product},
                )
                warehouse_stock.quantity += line.quantity
                warehouse_stock.full_clean()
                warehouse_stock.save(update_fields=("quantity",))
            for line in missing:
                product = Tech.objects.create(
                    brand=tech_brand,
                    product_type=tech_type,
                    name=line.name,
                    sku=f"TECH-ITECH-{line.row:04d}",
                    cost=line.cost,
                )
                TechWarehouseStock.objects.create(
                    warehouse=warehouse,
                    tech=product,
                    quantity=line.quantity,
                )
                matched[line.row] = product
                transfer_lines.append({
                    "product_type": "tech",
                    "product_id": product.pk,
                    "quantity": line.quantity,
                    "receivable_per_unit": line.receivable,
                })
            movement = transfer_many_to_consignment(
                actor=actor,
                warehouse_id=warehouse.pk,
                platform_id=platform.pk,
                lines=transfer_lines,
            )

        self.stdout.write(self.style.SUCCESS(
            f"iTech заполнена: операция №{movement.pk}, {len(lines)} позиций, "
            f"{sum(line.quantity for line in lines)} шт."
        ))
