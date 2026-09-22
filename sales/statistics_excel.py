"""XLSX rendering from the exact same StatisticsReport used by the web page."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from django.utils import timezone

from .models import Sale
from .statistics_service import CHANNEL_LABELS


def _append(sheet, values):
    sheet.append([value if value is not None else "Недоступно" for value in values])


def _header(sheet, values):
    _append(sheet, values)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="287F98")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def render_statistics_workbook(report):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Сводка"
    _header(summary, ("Показатель", "Значение"))
    totals = report.totals
    filters = report.filters
    period = f"{filters.get('date_from') or 'Все даты'} — {filters.get('date_to') or 'Все даты'}"
    selected_channels = ", ".join(CHANNEL_LABELS.get(value, value) for value in filters.get("channels", [])) or "Все каналы"
    payment_labels = dict(Sale.PaymentMethod.choices)
    payments = ", ".join(payment_labels.get(value, value) for value in filters.get("payment_methods", [])) or "Все способы"
    for label, value in (
        ("Период завершения", period), ("Каналы", selected_channels),
        ("Склад", filters["warehouse"].name if filters.get("warehouse") else "Все склады"),
        ("Оплата", payments), ("Товар", filters.get("product_query") or "Все товары"),
        ("Платформа CD", filters["platform"].name if filters.get("platform") else "Все"),
        ("Игровая серия", filters["game_series"].name if filters.get("game_series") else "Все"),
        ("Бренд Tech", filters["brand"].name if filters.get("brand") else "Все"),
        ("Тип Tech", filters["product_type"].name if filters.get("product_type") else "Все"),
        ("Выбранные ID продаж", ", ".join(map(str, sorted(filters.get("sale_ids", [])))) or "Все продажи"),
        ("Минимальная прибыль продажи", filters.get("min_profit") if filters.get("min_profit") is not None else "Не задана"),
        ("Максимальная прибыль продажи", filters.get("max_profit") if filters.get("max_profit") is not None else "Не задана"),
        ("Группировка", dict((("day", "По дням"), ("week", "По неделям"), ("month", "По месяцам"))).get(filters.get("grouping"), "По дням")),
        ("Завершённых продаж", totals.sales_count),
        ("Продано единиц", totals.units), ("Стоимость товаров", totals.goods_total),
        ("Переплаты", totals.overpayment), ("Фактическая выручка", totals.actual_revenue),
        ("Себестоимость", totals.cost), ("Чистая прибыль", totals.profit),
        ("Маржа, %", totals.margin), ("Наценка, %", totals.markup),
        ("Средний чек", totals.average_check),
        ("Продаж без полной себестоимости", totals.missing_sales),
        ("Позиций без себестоимости", totals.missing_lines),
        ("Завершённых продаж без даты", report.undated_completed_count),
    ):
        _append(summary, (label, value))

    sales = workbook.create_sheet("Продажи")
    _header(sales, ("ID", "Дата завершения", "Склад", "Канал", "Оплата", "Стоимость товаров",
                    "Переплата", "Фактическая выручка", "Себестоимость", "Чистая прибыль",
                    "Маржа, %", "Наценка, %", "Единиц", "Позиций", "Нет себестоимости"))
    items = workbook.create_sheet("Позиции продаж")
    _header(items, ("Продажа ID", "Дата завершения", "Тип", "Товар ID", "Артикул", "Товар",
                    "Количество", "Цена единицы", "Себестоимость единицы", "Выручка",
                    "Себестоимость", "Прибыль", "Маржа, %", "Наценка, %"))
    for row in report.sales:
        sale = row.sale
        date = timezone.localtime(sale.completed_at).replace(tzinfo=None)
        _append(sales, (sale.pk, date, sale.warehouse.name, row.channel_label,
                        sale.get_payment_method_display(), row.goods_total, row.overpayment,
                        row.actual_revenue, row.cost, row.profit, row.margin, row.markup,
                        row.units, row.positions, row.missing_lines))
        for line in row.lines:
            _append(items, (sale.pk, date, line.kind.upper(), line.product_id, line.sku, line.name,
                            line.quantity, line.unit_price, line.unit_cost, line.revenue,
                            line.cost, line.profit, line.margin, line.markup))

    products = workbook.create_sheet("Товары")
    _header(products, ("Тип", "ID", "Артикул", "Товар", "Категория", "Продано единиц",
                       "Количество продаж", "Выручка", "Себестоимость", "Прибыль", "Маржа, %",
                       "Наценка, %", "Средняя цена единицы", "Доля в прибыли товаров, %"))
    for row in report.products:
        _append(products, (row.kind.upper(), row.product_id, row.sku, row.name, row.category,
                           row.units, row.sales_count, row.revenue, row.cost, row.profit,
                           row.margin, row.markup, row.average_unit_price, row.profit_share))

    channels = workbook.create_sheet("Каналы")
    _header(channels, ("Канал", "Продаж", "Единиц", "Стоимость товаров", "Переплаты",
                       "Фактическая выручка", "Себестоимость", "Прибыль", "Маржа, %",
                       "Наценка, %", "Средний чек"))
    for row in report.channels:
        t = row.totals
        _append(channels, (row.label, t.sales_count, t.units, t.goods_total, t.overpayment,
                           t.actual_revenue, t.cost, t.profit, t.margin, t.markup, t.average_check))

    series = workbook.create_sheet("Динамика")
    _header(series, ("Период", "Продаж", "Единиц", "Фактическая выручка", "Себестоимость", "Прибыль"))
    for row in report.time_series:
        t = row.totals
        _append(series, (row.label, t.sales_count, t.units, t.actual_revenue, t.cost, t.profit))

    for sheet in workbook:
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(48, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
