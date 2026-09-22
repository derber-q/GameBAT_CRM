from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import Permission
from django.core.exceptions import ValidationError
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import User
from catalog.models import CD, GameSeries, Platform
from partners.models import SalesPlatform
from consignment.services import record_consignment_sale, transfer_to_consignment
from warehouse.models import CDWarehouseStock, Warehouse

from .models import Sale
from .services import advance_order_status, create_sale, edit_postpay_sale_items, mark_sale_paid
from .statistics_service import build_sales_statistics_report


class SalesStatisticsTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser("stats-admin", password="TestPassword123!")
        self.client.force_login(self.user)
        self.warehouse = Warehouse.objects.create(name="Склад отчёта")
        self.other_warehouse = Warehouse.objects.create(name="Другой склад")
        self.platform = Platform.objects.create(name="PS5")
        self.product = CD.objects.create(
            platform=self.platform, name="Игра", sku="STAT-CD", cost=1000,
            avito_price=1500, wholesale_price=1500, yandex_market_price=1500,
        )
        CDWarehouseStock.objects.create(warehouse=self.warehouse, cd=self.product, quantity=30)
        CDWarehouseStock.objects.create(warehouse=self.other_warehouse, cd=self.product, quantity=30)

    def sale(self, *, sale_type=Sale.SaleType.AVITO, payment=Sale.PaymentMethod.CASH,
             quantity=1, warehouse=None, received=None):
        return create_sale(
            actor=self.user, warehouse_id=(warehouse or self.warehouse).pk,
            price_type=Sale.PriceType.RETAIL, sale_type=sale_type, payment_method=payment,
            lines=[{"product_type": "cd", "product_id": self.product.pk, "quantity": quantity}],
            cash_received_amount=received,
        )

    def report(self, **filters):
        return build_sales_statistics_report(filters)

    def test_snapshot_profit_overpayment_margin_markup_and_cost_immutability(self):
        sale = self.sale(quantity=3, received="5000")
        item = sale.cd_items.get()
        self.assertEqual(item.unit_cost_snapshot, Decimal("1000"))
        self.product.cost = Decimal("1400")
        self.product.save(update_fields=("cost",))
        report = self.report()
        row = report.sales[0]
        self.assertEqual((row.lines[0].revenue, row.lines[0].cost, row.lines[0].profit),
                         (Decimal("4500"), Decimal("3000"), Decimal("1500")))
        self.assertEqual((row.overpayment, row.actual_revenue, row.profit),
                         (Decimal("500"), Decimal("5000"), Decimal("2000")))
        self.assertEqual(row.margin, Decimal("40.00"))
        self.assertEqual(row.markup, Decimal("66.67"))
        self.assertEqual(report.totals.average_check, Decimal("5000.00"))

    def test_only_completed_not_cancelled_and_completed_date_not_created_date(self):
        completed = self.sale()
        unfinished = self.sale(payment=Sale.PaymentMethod.BANK_ACCOUNT)
        cancelled = self.sale()
        from .services import cancel_sale
        cancel_sale(actor=self.user, sale_id=cancelled.pk, comment="Тест отмены")
        created = timezone.make_aware(datetime(2026, 9, 1, 12))
        finished = timezone.make_aware(datetime(2026, 9, 5, 12))
        Sale.objects.filter(pk=completed.pk).update(created_at=created, completed_at=finished)
        self.assertEqual([row.sale.pk for row in self.report(date_from=finished.date(), date_to=finished.date()).sales],
                         [completed.pk])
        self.assertFalse(self.report(date_from=created.date(), date_to=created.date()).sales)
        self.assertEqual(self.report().totals.sales_count, 1)
        self.assertIsNone(unfinished.completed_at)

    def test_deferred_snapshot_fixed_at_creation_and_new_postpay_line_uses_current_cost(self):
        sale = self.sale(payment=Sale.PaymentMethod.CASH_POSTPAY)
        original = sale.cd_items.get()
        self.assertEqual(original.unit_cost_snapshot, Decimal("1000"))
        second = CD.objects.create(
            platform=self.platform, name="Новая игра", sku="STAT-NEW", cost=700, avito_price=900,
        )
        CDWarehouseStock.objects.create(warehouse=self.warehouse, cd=second, quantity=5)
        self.product.cost = Decimal("1400")
        self.product.save(update_fields=("cost",))
        edit_postpay_sale_items(actor=self.user, sale_id=sale.pk, lines=[
            {"product_type": "cd", "product_id": self.product.pk, "quantity": 2},
            {"product_type": "cd", "product_id": second.pk, "quantity": 1},
        ])
        self.assertEqual(sale.cd_items.get(cd=self.product).unit_cost_snapshot, Decimal("1000"))
        self.assertEqual(sale.cd_items.get(cd=second).unit_cost_snapshot, Decimal("700"))
        mark_sale_paid(actor=self.user, sale_id=sale.pk)
        for status in (Sale.OrderStatus.ASSEMBLED, Sale.OrderStatus.SHIPPED, Sale.OrderStatus.DELIVERED):
            advance_order_status(actor=self.user, sale_id=sale.pk, next_status=status)
        self.assertEqual(self.report().totals.cost, Decimal("2700"))

    def test_missing_historical_cost_never_uses_current_product_cost(self):
        sale = self.sale()
        sale.cd_items.update(unit_cost_snapshot=None)
        report = self.report()
        self.assertEqual(report.totals.actual_revenue, Decimal("1500"))
        self.assertIsNone(report.totals.cost)
        self.assertIsNone(report.totals.profit)
        self.assertEqual((report.totals.missing_sales, report.totals.missing_lines), (1, 1))
        response = self.client.get(reverse("statistics:index"))
        self.assertContains(response, "Себестоимость недоступна")

    def test_multi_channel_wholesale_warehouse_and_sale_ids(self):
        avito = self.sale()
        yandex = self.sale(sale_type=Sale.SaleType.YANDEX_MARKET)
        pickup = self.sale(sale_type=Sale.SaleType.WHOLESALE_PICKUP)
        delivery = self.sale(sale_type=Sale.SaleType.WHOLESALE_DELIVERY, warehouse=self.other_warehouse)
        self.assertEqual(self.report(channels=["avito", "yandex_market"]).totals.sales_count, 2)
        self.assertEqual(self.report(channels=["wholesale"]).totals.sales_count, 2)
        self.assertEqual(self.report(warehouse=self.warehouse).totals.sales_count, 3)
        self.assertEqual({row.sale.pk for row in self.report(sale_ids={avito.pk, delivery.pk}).sales},
                         {avito.pk, delivery.pk})
        self.assertEqual({row.key for row in self.report(channels=["avito", "yandex_market"]).channels},
                         {"avito", "yandex_market"})

    def test_product_aggregation_sort_and_time_groups(self):
        first = self.sale(quantity=2)
        second = self.sale(quantity=3)
        timestamp = timezone.make_aware(datetime(2026, 9, 5, 12))
        Sale.objects.filter(pk__in=(first.pk, second.pk)).update(completed_at=timestamp)
        report = self.report(grouping="day", product_sort="-units")
        product = report.products[0]
        self.assertEqual((product.units, product.sales_count, product.revenue, product.cost, product.profit),
                         (5, 2, Decimal("7500"), Decimal("5000"), Decimal("2500")))
        self.assertEqual(report.time_series[0].totals.profit, Decimal("2500"))
        self.assertEqual(report.time_series[0].label, "05.09.2026")
        self.assertEqual(self.report(grouping="week").time_series[0].label, "31.08.2026–06.09.2026")
        self.assertEqual(self.report(grouping="month").time_series[0].label, "Сентябрь 2026")
        self.assertEqual(len(report.chart["series"]), 3)

    def test_web_excel_same_selection_and_permissions(self):
        chosen = self.sale()
        self.sale(warehouse=self.other_warehouse)
        params = {"warehouse": self.warehouse.pk, "channels": "avito", "sale_ids_text": str(chosen.pk)}
        page = self.client.get(reverse("statistics:index"), params)
        self.assertEqual(page.status_code, 200)
        self.assertEqual(page.context["report"].totals.sales_count, 1)
        excel = self.client.get(reverse("statistics:export"), params)
        self.assertEqual(excel.status_code, 200)
        workbook = load_workbook(BytesIO(excel.content), read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames,
                         ["Сводка", "Продажи", "Позиции продаж", "Товары", "Каналы", "Динамика"])
        self.assertEqual(workbook["Продажи"].max_row, 2)
        self.assertEqual(workbook["Продажи"]["A2"].value, chosen.pk)
        viewer = User.objects.create_user("stats-viewer", password="TestPassword123!")
        self.client.force_login(viewer)
        self.assertEqual(self.client.get(reverse("statistics:index")).status_code, 403)
        viewer.user_permissions.add(Permission.objects.get(codename="view_sales_statistics", content_type__app_label="sales"))
        self.assertEqual(self.client.get(reverse("statistics:index")).status_code, 200)
        self.assertEqual(self.client.get(reverse("statistics:export")).status_code, 403)

    def test_product_category_barcode_payment_and_profit_filters(self):
        series = GameSeries.objects.create(name="Серия отчёта")
        self.product.game_series = series
        self.product.barcodes.create(value="001234-STATS", product_kind="cd")
        self.product.save(update_fields=("game_series",))
        cash_sale = self.sale()
        bank_sale = self.sale(payment=Sale.PaymentMethod.BANK_ACCOUNT)
        mark_sale_paid(actor=self.user, sale_id=bank_sale.pk)
        for status in (Sale.OrderStatus.ASSEMBLED, Sale.OrderStatus.SHIPPED, Sale.OrderStatus.DELIVERED):
            advance_order_status(actor=self.user, sale_id=bank_sale.pk, next_status=status)
        self.assertEqual(self.report(product_query="001234-STATS").totals.sales_count, 2)
        self.assertEqual(self.report(platform=self.platform, game_series=series).totals.sales_count, 2)
        self.assertEqual(self.report(payment_methods=[Sale.PaymentMethod.CASH]).totals.sales_count, 1)
        self.assertEqual(self.report(min_profit=Decimal("501")).totals.sales_count, 0)
        self.assertEqual(self.report(min_profit=Decimal("500"), max_profit=Decimal("500")).totals.sales_count, 2)

    def test_negative_profit_and_zero_cost_percentage_rules(self):
        self.product.cost = Decimal("1600")
        self.product.save(update_fields=("cost",))
        loss = self.sale()
        loss_row = self.report().sales[0]
        self.assertEqual(loss_row.profit, Decimal("-100"))
        self.assertEqual(loss_row.margin, Decimal("-6.67"))
        self.assertEqual(loss_row.markup, Decimal("-6.25"))
        self.product.cost = Decimal("0")
        self.product.save(update_fields=("cost",))
        zero = self.sale()
        zero_row = next(row for row in self.report().sales if row.sale.pk == zero.pk)
        self.assertEqual(zero_row.profit, Decimal("1500"))
        self.assertIsNone(zero_row.markup)

    def test_report_queries_do_not_grow_per_sale_or_item(self):
        for _ in range(10):
            self.sale()
        with CaptureQueriesContext(connection) as captured:
            report = self.report()
        self.assertEqual(report.totals.sales_count, 10)
        self.assertLessEqual(len(captured), 7)

    def test_consignment_sale_captures_cost_when_recorded(self):
        platform = SalesPlatform.objects.create(
            name="Площадка отчёта", address="Адрес", legal_entity="Юрлицо", phone_1="1",
        )
        stock = transfer_to_consignment(
            actor=self.user, warehouse_id=self.warehouse.pk, platform_id=platform.pk,
            product_type="cd", product_id=self.product.pk, quantity=2, receivable_per_unit="1500",
        )
        sale = record_consignment_sale(
            actor=self.user, product_type="cd", stock_id=stock.pk, quantity=1,
            payment_method=Sale.PaymentMethod.BANK_ACCOUNT,
        )
        self.assertEqual(sale.cd_items.get().unit_cost_snapshot, Decimal("1000"))
        self.assertEqual(self.report(channels=["consignment"]).totals.profit, Decimal("500"))

    def test_completed_at_and_historical_cost_cannot_be_edited_normally(self):
        sale = self.sale()
        original_date = sale.completed_at
        sale.completed_at = original_date + timedelta(days=1)
        with self.assertRaisesMessage(ValidationError, "Дату завершения"):
            sale.save(update_fields=("completed_at",))
        item = sale.cd_items.get()
        item.unit_cost_snapshot = Decimal("1")
        with self.assertRaisesMessage(ValidationError, "Историческую себестоимость"):
            item.save(update_fields=("unit_cost_snapshot",))

    def test_sales_layout_shows_every_field_and_action_without_horizontal_table(self):
        sale = self.sale()
        response = self.client.get(reverse("statistics:index"))
        self.assertEqual(response.status_code, 200)
        section = response.content.decode("utf-8").split("stats-sales-section", 1)[1].split(
            "Статистика по товарам", 1,
        )[0]
        self.assertNotIn('class="table-scroll"', section)
        self.assertNotIn("stats-sales-table", section)
        for label in (
            "Склад", "Канал", "Оплата", "Стоимость товаров", "Переплата", "Выручка",
            "Себестоимость", "Прибыль", "Маржа %", "Наценка %", "Единиц / позиций",
            "Состав", "Количество", "Цена / ед.",
        ):
            self.assertIn(label, section)
        self.assertIn(f'data-stats-expand="stats-sale-{sale.pk}"', section)
        self.assertIn('form="statistics-filter-form"', section)
