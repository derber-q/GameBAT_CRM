from decimal import Decimal
from io import BytesIO
from zipfile import is_zipfile

from django.contrib.auth.models import Permission
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import User
from catalog.models import CD, Platform
from sales.models import Sale, SaleCDItem, SaleConsignmentItem, SaleCustomItem, SaleTechItem
from sales.services import create_sale
from sales.invoice_excel import HEADER_ROW, HEADERS, render_sale_invoice_xlsx
from cash.models import CashRegister, CashTransaction
from partners.models import SalesPlatform
from consignment.services import transfer_to_consignment, record_consignment_sale
from warehouse.models import CDWarehouseStock, Warehouse


class SaleInvoiceTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser("invoice-admin", password="StrongAdmin!123")
        self.warehouse = Warehouse.objects.create(name="Основной склад")
        self.platform = Platform.objects.create(name="PS5")
        self.cd = CD.objects.create(
            platform=self.platform, name="Тестовый диск", sku="CD-INV-1", barcode="INV-1",
            cost=Decimal("100.00"), avito_price=Decimal("250.00"),
        )
        CDWarehouseStock.objects.create(warehouse=self.warehouse, cd=self.cd, quantity=3)
        self.client.force_login(self.user)

    def test_invoice_download_is_available_for_unfinished_sale(self):
        sale = create_sale(
            actor=self.user, warehouse_id=self.warehouse.pk,
            price_type=Sale.PriceType.RETAIL, sale_type=Sale.SaleType.RETAIL,
            payment_method=Sale.PaymentMethod.CASH_POSTPAY,
            lines=[{"product_type": "cd", "product_id": self.cd.pk, "quantity": 2}],
        )
        response = self.client.get(reverse("sales:invoice", args=(sale.pk,)))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Disposition"], f'attachment; filename="nakladnaya_{sale.visible_id}.xlsx"')
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertTrue(is_zipfile(BytesIO(response.content)))
        book = load_workbook(BytesIO(response.content))
        sheet = book.active
        self.assertEqual(sheet["A1"].value, f"Накладная по продаже {sale.visible_id}")
        self.assertIn(self.warehouse.name, sheet["A2"].value)
        self.assertIn(timezone.localtime(sale.created_at).strftime("%d.%m.%Y %H:%M"), sheet["A2"].value)
        self.assertEqual(tuple(cell.value for cell in sheet[HEADER_ROW]), HEADERS)
        self.assertEqual(tuple(cell.value for cell in sheet[HEADER_ROW + 1]),
                         (1, "CD", "Тестовый диск", "CD-INV-1", 2, 250, 500))
        self.assertEqual(sheet.cell(sheet.max_row, 7).value, 500)
        book.close()
        sale.refresh_from_db()
        self.assertEqual(sale.order_status, Sale.OrderStatus.CREATED)
        self.assertEqual(sale.payment_status, Sale.PaymentStatus.UNPAID)

    def sale(self, **overrides):
        values = dict(actor=self.user, warehouse_id=self.warehouse.pk, price_type="retail", sale_type="retail",
            payment_method="cash_postpay", lines=[{"product_type": "cd", "product_id": self.cd.pk, "quantity": 2}])
        values.update(overrides)
        return create_sale(**values)

    def sheet(self, sale):
        response = self.client.get(reverse("sales:invoice", args=(sale.pk,)))
        self.assertEqual(response.status_code, 200)
        return load_workbook(BytesIO(response.content)).active

    def test_historical_names_prices_discounts_and_surcharges_are_preserved(self):
        for discount in (Decimal("25.50"), Decimal("-10.25")):
            with self.subTest(discount=discount):
                # Отдельная номенклатура для каждого варианта без повторного расхода склада.
                product = CD.objects.create(platform=self.platform, name=f"Исходная игра {discount}",
                    sku=f"OLD-{discount}", cost=100, avito_price=250)
                CDWarehouseStock.objects.create(warehouse=self.warehouse, cd=product, quantity=2)
                sale = self.sale(note="Проверка сохранённой скидки или доплаты", lines=[{"product_type": "cd", "product_id": product.pk,
                    "quantity": 2, "unit_discount": str(discount)}])
                item = sale.cd_items.get()
                original_name, original_sku = product.name, product.sku
                product.name = "Новое название"
                product.sku = "NEW-" + str(product.pk)
                product.avito_price = 99999
                product.save(update_fields=("name", "sku", "avito_price"))
                sheet = self.sheet(sale)
                self.assertEqual(sheet["C5"].value, original_name)
                self.assertEqual(sheet["D5"].value, original_sku)
                self.assertEqual(Decimal(str(sheet["F5"].value)), item.effective_unit_price)
                self.assertEqual(Decimal(str(sheet["G5"].value)), item.line_total)
                self.assertEqual(Decimal(str(sheet.cell(sheet.max_row, 7).value)), sale.total_amount)

    def test_invoice_is_read_only_and_uses_snapshot_total_not_received_cash(self):
        sale = self.sale(payment_method="cash", cash_received_amount="550")
        models = (Sale, SaleCDItem, SaleTechItem, SaleCustomItem, SaleConsignmentItem,
                  CDWarehouseStock, CashRegister, CashTransaction, CD)
        before = {model: list(model.objects.order_by("pk").values()) for model in models}
        with CaptureQueriesContext(connection) as queries:
            sheet = self.sheet(sale)
        self.assertEqual(sheet.cell(sheet.max_row, 7).value, 500)
        self.assertFalse(any(q["sql"].lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE")) for q in queries))
        self.assertEqual(before, {model: list(model.objects.order_by("pk").values()) for model in models})

    def test_consignment_invoice_does_not_duplicate_legacy_mirror(self):
        platform = SalesPlatform.objects.create(name="Площадка накладной", address="Адрес", legal_entity="ИП", phone_1="1")
        stock = transfer_to_consignment(actor=self.user, warehouse_id=self.warehouse.pk, platform_id=platform.pk,
            product_type="cd", product_id=self.cd.pk, quantity=2, receivable_per_unit="310.25")
        sale = record_consignment_sale(actor=self.user, product_type="cd", stock_id=stock.pk,
                                       quantity=2, payment_method="bank_account")
        sheet = self.sheet(sale)
        self.assertEqual(sheet["B5"].value, "Реализация · Площадка накладной")
        self.assertEqual((sheet["E5"].value, sheet["F5"].value, sheet["G5"].value), (2, 310.25, 620.5))
        self.assertEqual(sheet.max_row, 6)

    def test_custom_digital_and_expense_rows_use_existing_invoice_rules(self):
        for kind in ("custom", "DIGITAL_PRODUCT", "EXPENSE"):
            with self.subTest(kind=kind):
                line = {"product_type": "custom", "name": "Особая позиция", "unit_cost": "10",
                        "unit_price": "25.50", "quantity": 2}
                if kind != "custom":
                    line.update(product_type="system", system_kind=kind)
                if kind == "EXPENSE":
                    line["unit_price"] = "0"
                sale = self.sale(lines=[line], note="Описание операции для накладной")
                item = sale.custom_items.get()
                sheet = self.sheet(sale)
                self.assertEqual(sheet["C5"].value, item.product_name_snapshot)
                self.assertEqual(sheet["D5"].value, "—")
                self.assertEqual(Decimal(str(sheet["G5"].value)), item.line_total)
                self.assertEqual(Decimal(str(sheet["G6"].value)), sale.total_amount)

    def test_long_names_print_settings_and_multiple_pages(self):
        sale = self.sale()
        long_name = "PlayStation 5 Slim Digital Edition комплект с двумя геймпадами DualSense " * 3
        rows = [{"kind": "Tech", "name": long_name, "article": "0000123", "quantity": 12,
                 "unit_price": Decimal("99999.99"), "line_total": Decimal("1199999.88")} for _ in range(100)]
        sheet = load_workbook(BytesIO(render_sale_invoice_xlsx(sale=sale, rows=rows))).active
        self.assertEqual(sheet["C5"].value, long_name)
        self.assertGreater(sheet.row_dimensions[5].height, 60)
        self.assertEqual(sheet["D5"].data_type, "s")
        self.assertEqual(sheet["E5"].data_type, "n")
        self.assertEqual(sheet["F5"].data_type, "n")
        self.assertEqual(sheet["F5"].number_format, "0.00")
        self.assertEqual(str(sheet.page_setup.paperSize), sheet.PAPERSIZE_A4)
        self.assertEqual(sheet.page_setup.orientation, "portrait")
        self.assertEqual((sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight), (1, 0))
        self.assertAlmostEqual(sheet.page_margins.left * 25.4, 12)
        self.assertEqual(sheet.print_title_rows, "'Накладная'!$4:$4" if "!" in sheet.print_title_rows else "$4:$4")
        self.assertGreater(len(sheet.row_breaks.brk), 1)
        self.assertEqual(sheet.max_row, 105)
        self.assertEqual(sheet["A1"].font.name, "Arial")
        self.assertEqual(sheet["A1"].font.sz, 16.5)
        self.assertEqual(sheet["A4"].fill.fgColor.rgb[-6:], "EAF2F8")
        self.assertEqual(sheet["G105"].fill.fgColor.rgb[-6:], "F3F6F9")

    def test_literal_excel_text_empty_old_sale_and_missing_cost_snapshot(self):
        sale = self.sale()
        item = sale.cd_items.get()
        # Старый снимок себестоимости может отсутствовать и не нужен в накладной.
        SaleCDItem.objects.filter(pk=item.pk).update(unit_cost_snapshot=None)
        self.assertEqual(self.sheet(sale)["G5"].value, 500)
        rows = [{"kind": "CD", "name": "=SUM(A1:A2)", "article": "000001", "quantity": 1,
                 "unit_price": Decimal("1.25"), "line_total": Decimal("1.25")}]
        sheet = load_workbook(BytesIO(render_sale_invoice_xlsx(sale=sale, rows=rows))).active
        self.assertEqual(sheet["C5"].data_type, "s")
        self.assertEqual(sheet["C5"].value, "=SUM(A1:A2)")
        sheet = load_workbook(BytesIO(render_sale_invoice_xlsx(sale=sale, rows=[]))).active
        self.assertEqual(sheet["A5"].value, "В продаже пока нет товарных позиций.")

    def test_existing_permission_and_download_button(self):
        sale = self.sale()
        self.assertContains(self.client.get(reverse("sales:detail", args=(sale.pk,))), "Скачать накладную Excel")
        employee = User.objects.create_user("invoice-viewer")
        self.client.force_login(employee)
        url = reverse("sales:invoice", args=(sale.pk,))
        self.assertEqual(self.client.get(url).status_code, 403)
        employee.user_permissions.add(Permission.objects.get(content_type__app_label="sales", codename="view_sale_detail"))
        self.assertEqual(self.client.get(url).status_code, 200)
        self.assertEqual(self.client.post(url).status_code, 405)
