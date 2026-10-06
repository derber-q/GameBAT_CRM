from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import Permission
from django.core.exceptions import ValidationError
from django.http import QueryDict
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import User
from catalog.models import CD, Platform
from warehouse.models import CDWarehouseStock, Warehouse

from .commissions import calculate_avito_commission
from .models import Sale
from .services import cancel_sale, create_sale, edit_postpay_sale_items
from .statistics_excel import render_statistics_workbook
from .statistics_service import build_sales_statistics_report
from .views import _parse_lines


class AvitoCommissionTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(
            "avito-commission-admin", password="StrongAdmin!123",
        )
        self.client.force_login(self.user)
        self.warehouse = Warehouse.objects.create(name="Склад комиссии Avito")
        self.platform = Platform.objects.create(name="PS5 комиссия")
        self.product = CD.objects.create(
            platform=self.platform, name="Товар с комиссией", sku="AVITO-FEE-1",
            cost=Decimal("4000.00"), avito_price=Decimal("10000.00"),
            wholesale_price=Decimal("9000.00"),
        )
        CDWarehouseStock.objects.create(
            warehouse=self.warehouse, cd=self.product, quantity=30,
        )

    def line(self, *, quantity=1, discount="0", enabled=True):
        return {
            "product_type": "cd", "product_id": self.product.pk,
            "quantity": quantity, "unit_discount": discount,
            "avito_commission_enabled": enabled,
        }

    def create(self, *, sale_type=Sale.SaleType.AVITO,
               payment_method=Sale.PaymentMethod.BANK_ACCOUNT,
               lines=None, received=None, note=""):
        return create_sale(
            actor=self.user, warehouse_id=self.warehouse.pk,
            price_type=Sale.PriceType.RETAIL, sale_type=sale_type,
            payment_method=payment_method, lines=lines or [self.line()],
            cash_received_amount=received, note=note,
        )

    def test_formula_uses_half_up_rounding_and_one_ruble_minimum(self):
        self.assertEqual(
            calculate_avito_commission(Decimal("10000.00"), enabled=True),
            Decimal("50.00"),
        )
        self.assertEqual(
            calculate_avito_commission(Decimal("1001.00"), enabled=True),
            Decimal("5.01"),
        )
        self.assertEqual(
            calculate_avito_commission(Decimal("100.00"), enabled=True),
            Decimal("1.00"),
        )
        self.assertEqual(
            calculate_avito_commission(Decimal("10000.00"), enabled=False),
            Decimal("0.00"),
        )

    def test_enabled_product_commission_is_stored_without_changing_sale_total(self):
        sale = self.create()
        item = sale.cd_items.get()
        self.assertTrue(item.avito_commission_enabled)
        self.assertEqual(item.avito_commission_amount, Decimal("50.00"))
        self.assertEqual(item.revenue_after_avito_commission, Decimal("9950.00"))
        self.assertEqual(sale.total_amount, Decimal("10000.00"))

    def test_quantity_and_discount_are_included_in_commission_base(self):
        sale = self.create(
            lines=[self.line(quantity=2, discount="500")],
            note="Скидка для теста комиссии",
        )
        item = sale.cd_items.get()
        self.assertEqual(item.line_total, Decimal("19000.00"))
        self.assertEqual(item.avito_commission_amount, Decimal("95.00"))
        self.assertEqual(sale.total_amount, Decimal("19000.00"))

    def test_avito_commission_is_mandatory_and_non_avito_has_zero(self):
        for requested in (False, None, "invalid"):
            item = self.create(lines=[self.line(enabled=requested)]).cd_items.get()
            self.assertTrue(item.avito_commission_enabled)
            self.assertEqual(item.avito_commission_amount, Decimal("50.00"))
        non_avito = self.create(
            sale_type=Sale.SaleType.RETAIL,
            lines=[self.line(enabled=True)],
        )
        item = non_avito.cd_items.get()
        self.assertFalse(item.avito_commission_enabled)
        self.assertEqual(item.avito_commission_amount, Decimal("0.00"))

    def test_custom_line_supports_commission_and_minimum(self):
        sale = self.create(lines=[{
            "product_type": "custom", "name": "Доставка аксессуара",
            "quantity": 2, "unit_price": "80", "unit_discount": "10",
            "unit_cost": "20", "avito_commission_enabled": False,
        }], note="Скидка на произвольную строку")
        item = sale.custom_items.get()
        self.assertEqual(item.line_total, Decimal("140.00"))
        self.assertTrue(item.avito_commission_enabled)
        self.assertEqual(item.avito_commission_amount, Decimal("1.00"))
        self.assertEqual(sale.total_amount, Decimal("140.00"))

    def test_postpay_edit_recalculates_and_cannot_disable_commission(self):
        sale = self.create(
            payment_method=Sale.PaymentMethod.CASH_POSTPAY,
            lines=[self.line(enabled=False)],
        )
        edit_postpay_sale_items(
            actor=self.user, sale_id=sale.pk,
            lines=[self.line(quantity=2, enabled=True)],
        )
        item = sale.cd_items.get()
        self.assertEqual(item.avito_commission_amount, Decimal("100.00"))
        edit_postpay_sale_items(
            actor=self.user, sale_id=sale.pk,
            lines=[self.line(quantity=2, enabled=False)],
        )
        item.refresh_from_db()
        sale.refresh_from_db()
        self.assertTrue(item.avito_commission_enabled)
        self.assertEqual(item.avito_commission_amount, Decimal("100.00"))
        self.assertEqual(sale.total_amount, Decimal("20000.00"))

    def test_completed_commission_snapshot_cannot_be_changed_normally(self):
        sale = self.create(
            payment_method=Sale.PaymentMethod.CASH,
            received="10000.00",
        )
        item = sale.cd_items.get()
        item.avito_commission_enabled = False
        item.avito_commission_amount = Decimal("0.00")
        with self.assertRaisesMessage(ValidationError, "Историческую комиссию"):
            item.save(update_fields=(
                "avito_commission_enabled", "avito_commission_amount",
            ))

    def test_cash_and_refund_use_gross_sale_total(self):
        register = self.warehouse.cash_register
        sale = self.create(
            payment_method=Sale.PaymentMethod.CASH,
            received="10000.00",
        )
        register.refresh_from_db()
        self.assertEqual(register.balance, Decimal("10000.00"))
        cancel_sale(actor=self.user, sale_id=sale.pk, comment="Тест возврата")
        register.refresh_from_db()
        sale.refresh_from_db()
        self.assertEqual(register.balance, Decimal("0.00"))
        self.assertEqual(sale.refunded_amount, Decimal("10000.00"))

    def test_statistics_deduct_commission_at_all_aggregation_levels(self):
        sale = self.create(
            payment_method=Sale.PaymentMethod.CASH,
            received="10100.00",
            lines=[
                self.line(enabled=True),
                {
                    "product_type": "custom", "name": "Дополнение",
                    "quantity": 1, "unit_price": "100", "unit_discount": "0",
                    "unit_cost": "20", "avito_commission_enabled": True,
                },
            ],
        )
        report = build_sales_statistics_report({"grouping": "day"})
        row = report.sales[0]
        self.assertEqual(sale.total_amount, Decimal("10100.00"))
        self.assertEqual(row.commission, Decimal("51.00"))
        self.assertEqual(row.actual_revenue, Decimal("10100.00"))
        self.assertEqual(row.cost, Decimal("4020.00"))
        self.assertEqual(row.profit, Decimal("6029.00"))
        self.assertEqual(report.totals.commission, Decimal("51.00"))
        self.assertEqual(report.products[0].commission, Decimal("50.00"))
        self.assertEqual(report.channels[0].totals.commission, Decimal("51.00"))
        self.assertEqual(report.time_series[0].totals.commission, Decimal("51.00"))
        self.assertEqual(
            sum(line.commission for line in row.lines), Decimal("51.00"),
        )

    def test_excel_contains_commission_on_every_analytical_sheet(self):
        self.create(
            payment_method=Sale.PaymentMethod.CASH,
            received="10000.00",
        )
        report = build_sales_statistics_report({"grouping": "day"})
        workbook = load_workbook(
            BytesIO(render_statistics_workbook(report)), read_only=True, data_only=True,
        )
        summary = {
            row[0]: row[1] for row in workbook["Сводка"].iter_rows(values_only=True)
        }
        self.assertEqual(summary["Комиссия Avito"], 50)
        for sheet_name in ("Продажи", "Позиции продаж", "Товары", "Каналы", "Динамика"):
            headers = next(workbook[sheet_name].iter_rows(values_only=True))
            self.assertIn("Комиссия Avito", headers)

    def test_parser_keeps_aligned_product_and_custom_flags(self):
        post = QueryDict(mutable=True)
        post.setlist("product_type", ["cd"])
        post.setlist("product_id", [str(self.product.pk)])
        post.setlist("quantity", ["1"])
        post.setlist("unit_discount", ["0"])
        post.setlist("avito_commission_enabled", ["1"])
        post.setlist("custom_name", ["Ручная строка"])
        post.setlist("custom_unit_price", ["100"])
        post.setlist("custom_quantity", ["1"])
        post.setlist("custom_unit_cost", ["20"])
        post.setlist("custom_unit_discount", ["0"])
        post.setlist("custom_avito_commission_enabled", ["0"])
        lines = _parse_lines(post)
        self.assertEqual(lines[0]["avito_commission_enabled"], "1")
        self.assertEqual(lines[1]["avito_commission_enabled"], "0")

    def test_parser_rejects_misaligned_flags(self):
        post = QueryDict(mutable=True)
        post.setlist("product_type", ["cd", "cd"])
        post.setlist("product_id", [str(self.product.pk), str(self.product.pk)])
        post.setlist("quantity", ["1", "1"])
        post.setlist("avito_commission_enabled", ["1"])
        with self.assertRaisesMessage(ValidationError, "флаг комиссии"):
            _parse_lines(post)

    def test_detail_and_forms_show_commission_only_in_authorized_financial_area(self):
        sale = self.create(
            payment_method=Sale.PaymentMethod.CASH,
            received="10000.00",
        )
        detail = self.client.get(reverse("sales:detail", args=(sale.pk,)))
        self.assertContains(detail, "Комиссия Avito")
        self.assertContains(detail, "После комиссии")
        create_page = self.client.get(reverse("sales:create"))
        self.assertContains(create_page, 'data-sale-type-select-id="id_sale_type"')
        self.assertContains(create_page, "К оплате покупателем")
        self.assertContains(create_page, "Сумма после комиссии")
        self.assertContains(create_page, "data-sale-commission-total")
        self.assertContains(create_page, "data-sale-net-total")
        with open("static/js/stock-lines.js", encoding="utf-8") as source:
            javascript = source.read()
        self.assertNotIn("avitoCommissionCheckbox", javascript)
        self.assertNotIn("Учитывать комиссию Avito 0,5%", javascript)
        self.assertIn("Math.max(lineTotal * 0.005, 1)", javascript)
        self.assertIn("total - totalCommission", javascript)
        self.assertIn('commissionPreview.textContent = `${money.format(commission)} ₽`', javascript)
        self.assertNotIn("Комиссия не учитывается", javascript)
        self.assertIn('column.hidden = !visible', javascript)
        self.assertIn('data-avito-commission-column hidden>Комиссия Avito', create_page.content.decode("utf-8"))
        viewer = User.objects.create_user("sale-detail-viewer", password="ViewerPass123!")
        viewer.user_permissions.add(Permission.objects.get(
            codename="view_sale_detail", content_type__app_label="sales",
        ))
        self.client.force_login(viewer)
        restricted = self.client.get(reverse("sales:detail", args=(sale.pk,)))
        self.assertNotContains(restricted, "Комиссия Avito")

    def test_consignment_rows_have_no_commission_fields(self):
        from .models import SaleConsignmentItem
        field_names = {field.name for field in SaleConsignmentItem._meta.fields}
        self.assertNotIn("avito_commission_enabled", field_names)
        self.assertNotIn("avito_commission_amount", field_names)
