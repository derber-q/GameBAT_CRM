"""Накладная XLSX по структуре и цветам прежнего sales/invoice.html."""
from io import BytesIO
from math import ceil

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.page import PageMargins

from core.excel_layout import column_width_pixels, text_width_pixels, wrapped_text_height

HEADER_ROW = 4
HEADERS = ("№", "Тип", "Наименование", "Артикул", "Количество", "Цена", "Сумма")
TEXT = "1F2933"
MUTED = "667085"
LINE = "CBD5E1"
HEADER_FILL = "EAF2F8"
TOTAL_FILL = "F3F6F9"


def _text(cell, value):
    # Строковый тип сохраняет буквально и названия, начинающиеся со знака «=».
    cell.value = ILLEGAL_CHARACTERS_RE.sub("", str(value or ""))
    cell.data_type = "s"


def _height(value, width, *, size=12, bold=False, padding=12, minimum=28):
    return wrapped_text_height(value, width, font_name="Arial", font_size=size,
                               bold=bold, padding_pixels=20, padding_points=padding,
                               min_height=minimum)


def _page_breaks(sheet, last_item, total_row, widths):
    # A4 с полями 12 мм; масштаб по ширине учитывается и при расчёте высоты страницы.
    printable_width = (210 - 24) * 72 / 25.4
    scale = min(1, printable_width / sum(column_width_pixels(width) * .75 for width in widths))
    available = (297 - 24) * 72 / 25.4 / scale
    used = sum(sheet.row_dimensions[row].height for row in range(1, HEADER_ROW + 1))
    repeated_header = sheet.row_dimensions[HEADER_ROW].height
    for row in range(HEADER_ROW + 1, total_row + 1):
        height = sheet.row_dimensions[row].height
        # Последнюю товарную строку и итог оставляем на одной странице.
        needed = height + sheet.row_dimensions[total_row].height if row == last_item else height
        if used + needed > available and used > repeated_header:
            sheet.row_breaks.append(Break(id=row - 1))
            used = repeated_header
        used += height


def render_sale_invoice_xlsx(*, sale, rows):
    book = Workbook()
    sheet = book.active
    sheet.title = "Накладная"
    sheet.sheet_view.showGridLines = False
    widths = [5, 15, 24, 12, 15, 12, 13]
    for column, label in enumerate(HEADERS):
        needed = text_width_pixels(label, font_name="Arial", font_size=12, bold=True)
        widths[column] = max(widths[column], ceil((needed + 20 - 5) / 7 * 10) / 10)
    for column, key in ((5, "unit_price"), (6, "line_total")):
        values = [row[key] for row in rows]
        if column == 6:
            values.append(sale.total_amount)
        needed = max((text_width_pixels(f"{value:.2f}", font_name="Arial", font_size=12,
                                        bold=column == 6) for value in values), default=0)
        widths[column] = max(widths[column], (needed + 20 - 5) / 7)
    for letter, width in zip("ABCDEFG", widths):
        sheet.column_dimensions[letter].width = width
    merged_width = sum(widths) + 30 / 7
    sheet.merge_cells("A1:G1")
    _text(sheet["A1"], f"Накладная по продаже {sale.visible_id}")
    sheet["A1"].font = Font(name="Arial", size=16.5, bold=True, color=TEXT)
    sheet["A1"].alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = _height(sheet["A1"].value, merged_width, size=16.5,
                                             bold=True, padding=4, minimum=25.5)
    sheet.merge_cells("A2:G2")
    date = timezone.localtime(sale.created_at).strftime("%d.%m.%Y %H:%M")
    _text(sheet["A2"], f"Дата: {date} · Склад: {sale.warehouse.name}")
    sheet["A2"].font = Font(name="Arial", size=12, color=MUTED)
    sheet["A2"].alignment = Alignment(vertical="top", wrap_text=True)
    sheet.row_dimensions[2].height = _height(sheet["A2"].value, merged_width, padding=0, minimum=16)
    sheet.row_dimensions[3].height = 18
    border = Border(**{side: Side(style="thin", color=LINE) for side in ("top", "bottom", "left", "right")})

    def table_cell(row, column, value, *, bold=False, fill=None):
        cell = sheet.cell(row, column)
        if isinstance(value, str):
            _text(cell, value)
        else:
            cell.value = value
        cell.font = Font(name="Arial", size=12, bold=bold, color=TEXT)
        cell.border = border
        cell.alignment = Alignment(horizontal="right" if column >= 5 else "left",
                                   vertical="center", wrap_text=column < 5, indent=1)
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        if column in (6, 7) and not isinstance(value, str):
            cell.number_format = "0.00"
        elif column == 4:
            cell.number_format = "@"
        return cell

    for column, label in enumerate(HEADERS, 1):
        cell = table_cell(HEADER_ROW, column, label, bold=True, fill=HEADER_FILL)
        cell.alignment = Alignment(horizontal="right" if column >= 5 else "left",
                                   vertical="center", wrap_text=True, indent=1)
    sheet.row_dimensions[HEADER_ROW].height = max(
        _height(label, width, bold=True) for label, width in zip(HEADERS, widths)
    )
    row_number = HEADER_ROW + 1
    for index, row in enumerate(rows, 1):
        values = (index, row["kind"], row["name"], row["article"], row["quantity"], row["unit_price"], row["line_total"])
        for column, value in enumerate(values, 1):
            table_cell(row_number, column, value)
        sheet.cell(row_number, 5).number_format = "0"
        sheet.row_dimensions[row_number].height = max(
            _height(value, width) for value, width in zip(values[:4], widths[:4])
        )
        row_number += 1
    if not rows:
        for column in range(1, 8):
            table_cell(row_number, column, "")
        sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=7)
        _text(sheet.cell(row_number, 1), "В продаже пока нет товарных позиций.")
        sheet.row_dimensions[row_number].height = 28
        row_number += 1
    total_row = row_number
    for column in range(1, 8):
        table_cell(total_row, column, "", bold=True, fill=TOTAL_FILL)
    sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=6)
    _text(sheet.cell(total_row, 1), "Итого по заказу")
    sheet.cell(total_row, 1).alignment = Alignment(horizontal="right", vertical="center", indent=1)
    table_cell(total_row, 7, sale.total_amount, bold=True, fill=TOTAL_FILL)
    sheet.row_dimensions[total_row].height = 28
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = sheet.ORIENTATION_PORTRAIT
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=12 / 25.4, right=12 / 25.4,
                                    top=12 / 25.4, bottom=12 / 25.4, header=0, footer=0)
    sheet.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    sheet.print_area = f"A1:G{total_row}"
    sheet.freeze_panes = f"A{HEADER_ROW + 1}"
    _page_breaks(sheet, total_row - 1, total_row, widths)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()
