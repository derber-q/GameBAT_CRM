"""Экспорт одного заказа по фактическим поставщикам и сохранённым закупочным ценам."""
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

from django.contrib.auth.models import Permission
from django.core.exceptions import ValidationError
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from pricing.models import SupplierCDPrice
from accounts.models import User
from partners.models import Supplier
from price.services import create_procurement_price_list
from .models import CustomerProcurementOrder
from .order_supplier_excel import build_order_supplier_rows, order_suppliers_xlsx
from .test_excel_grouping import create_export_fixture
from .trips import change_trip_line, create_trip
from .services import create_customer_procurement_order


class OrderSupplierExportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        create_export_fixture(cls)

    def rows(self):
        order = CustomerProcurementOrder.objects.get(pk=self.first.pk)
        items = list(order.items.select_related(
            "selected_supplier", "cd__platform", "cd__game_series", "tech__brand", "tech__product_type",
        ).prefetch_related("trip_allocations__line__supplier").order_by("pk"))
        return order, build_order_supplier_rows(order=order, items=items)

    def test_order_without_trip_uses_its_own_snapshots(self):
        order, rows = self.rows()
        self.assertEqual([(r.supplier.pk, r.quantity, r.unit_price) for r in rows], [
            (self.a.pk, 5, Decimal("100")), (self.b.pk, 1, Decimal("140")),
        ])
        self.assertEqual(sum(row.quantity for row in rows), order.total_units)

    def test_split_in_trip_exports_only_this_orders_allocations_at_saved_prices(self):
        trip = create_trip(actor=self.user, order_ids=[self.first.pk, self.second.pk], creation_key=uuid4())
        source = trip.lines.get(cd=self.cd, supplier=self.a)
        change_trip_line(actor=self.user, batch_id=trip.pk, line_id=source.pk, revision=trip.revision,
                         quantity=3, action="transfer", target_supplier_id=self.c.pk, price_supplier_id=self.c.pk)
        SupplierCDPrice.objects.filter(supplier=self.c, cd=self.cd).update(price=999)
        order, rows = self.rows()
        cd_rows = [row for row in rows if row.item.product_kind == "cd"]
        self.assertCountEqual([(r.supplier.pk, r.quantity, r.unit_price) for r in cd_rows], [
            (self.a.pk, 3, Decimal("100")), (self.c.pk, 2, Decimal("120")),
        ])
        self.assertEqual(sum(row.quantity for row in rows), 6)
        self.assertTrue(all(row.item.order_id == self.first.pk for row in rows))
        self.assertEqual(sum(row.quantity * row.unit_price for row in rows), Decimal("680"))

    def test_missing_supplier_and_missing_prices_are_explicit(self):
        item = SimpleNamespace(pk=1, quantity=2, selected_supplier=None, supplier_price_aed_snapshot=None)
        rows = build_order_supplier_rows(order=SimpleNamespace(trip_id=None), items=[item])
        self.assertEqual(len(rows), 1)
        self.assertIsNone(rows[0].supplier)
        self.assertIsNone(rows[0].unit_price)
        self.assertEqual(rows[0].quantity, 2)

    def test_workbook_keeps_real_supplier_splits_groups_snapshots_and_totals(self):
        trip = create_trip(actor=self.user, order_ids=[self.first.pk, self.second.pk], creation_key=uuid4())
        source = trip.lines.get(cd=self.cd, supplier=self.a)
        change_trip_line(actor=self.user, batch_id=trip.pk, line_id=source.pk, revision=trip.revision,
                         quantity=3, action="transfer", target_supplier_id=self.b.pk, price_supplier_id=self.c.pk)
        self.cd.name = "Название после заказа"
        self.cd.is_archived = True
        self.cd.save(update_fields=("name", "is_archived"))
        order, rows = self.rows()
        book = load_workbook(BytesIO(order_suppliers_xlsx(order=order, rows=rows, show_supplier_details=True)))
        sheet = book.active
        titles = [row[0] for row in sheet.values if isinstance(row[0], str) and row[0].startswith(("Tech ·", "CD ·"))]
        self.assertEqual(titles, ["CD · Платформа: PS5", "Tech · Sony · Консоли", "CD · Платформа: PS5"])
        data = [row for row in sheet.iter_rows(min_row=5, values_only=True) if row[5] in ("cd", "tech")]
        self.assertEqual([(r[9], r[5], r[6], r[0], r[1], r[2], r[3], r[4]) for r in data], [
            (self.a.pk, "cd", self.cd.pk, "Игра", "000123", 3, 100, 300),
            (self.b.pk, "tech", self.tech.pk, "Консоль", "000-tech", 1, 140, 140),
            (self.b.pk, "cd", self.cd.pk, "Игра", "000123", 2, 120, 240),
        ])
        self.assertEqual({r[7] for r in data}, set(order.items.values_list("pk", flat=True)))
        self.assertTrue(all(sheet.column_dimensions[column].hidden for column in "FGHIJ"))
        subtotals = [r for r in sheet.values if r[0] == "Итого по поставщику"]
        self.assertEqual([(r[2], r[4]) for r in subtotals], [(3, 300), (3, 380)])
        self.assertEqual((sheet.cell(sheet.max_row, 3).value, sheet.cell(sheet.max_row, 5).value), (6, 680))
        self.assertNotIn(self.second.recipient, str(list(sheet.values)))
        book.close()

    def test_workbook_shows_unassigned_positions_and_unknown_money_without_faking_zero(self):
        item = SimpleNamespace(pk=123, product=self.cd, product_kind="cd", product_id=self.cd.pk,
                               product_name_snapshot='=HYPERLINK("bad")', article_snapshot="0000001",
                               quantity=2, selected_supplier=None, supplier_price_aed_snapshot=None)
        rows = build_order_supplier_rows(order=SimpleNamespace(trip_id=None), items=[item])
        book = load_workbook(BytesIO(order_suppliers_xlsx(order=self.first, rows=rows)))
        sheet = book.active
        self.assertIn("Поставщик не назначен", [row[0] for row in sheet.values])
        row = next(row for row in sheet if row[5].value == "cd")
        self.assertEqual(row[0].value, "'" + item.product_name_snapshot)
        self.assertEqual(row[0].data_type, "s")
        self.assertEqual(row[1].value, "0000001")
        self.assertIsNone(row[3].value)
        self.assertIsNone(row[4].value)
        self.assertEqual(sheet.cell(sheet.max_row, 3).value, 2)
        self.assertIsNone(sheet.cell(sheet.max_row, 5).value)
        book.close()

    def test_endpoint_and_button_require_order_and_supplier_price_permissions(self):
        user = User.objects.create_user("buy-export-reader", password="StrongReader!123")
        url = reverse("orders:export_suppliers", args=(self.first.pk,))
        detail_permission = Permission.objects.get(content_type__app_label="orders", codename="view_order_detail")
        price_permission = Permission.objects.get(content_type__app_label="price", codename="view_supplier_prices")
        for permissions in ((), (detail_permission,), (price_permission,), (detail_permission, price_permission)):
            with self.subTest(permissions=[p.codename for p in permissions]):
                user.user_permissions.set(permissions)
                self.client.force_login(user)
                response = self.client.get(url)
                self.assertEqual(response.status_code, 200 if len(permissions) == 2 else 403)
                if detail_permission in permissions:
                    page = self.client.get(reverse("orders:detail", args=(self.first.pk,)))
                    if price_permission in permissions:
                        self.assertContains(page, url)
                        self.assertContains(page, "Скачать состав по поставщикам")
                        book = load_workbook(BytesIO(response.content))
                        text = str(list(book.active.values))
                        self.assertNotIn(self.a.name, text)
                        self.assertNotIn(self.b.name, text)
                        self.assertIn(self.a.safe_label, text)
                        book.close()
                    else:
                        self.assertNotContains(page, url)

    def test_download_is_read_only_get_only_and_has_no_per_product_queries(self):
        create_trip(actor=self.user, order_ids=[self.first.pk, self.second.pk], creation_key=uuid4())
        self.client.force_login(self.user)
        url = reverse("orders:export_suppliers", args=(self.first.pk,))
        with CaptureQueriesContext(connection) as queries:
            response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertIn(f"zakaz_{self.first.pk}_suppliers.xlsx", response["Content-Disposition"])
        self.assertIn("no-store", response["Cache-Control"])
        self.assertLessEqual(len(queries), 8)
        self.assertFalse(any(q["sql"].lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE")) for q in queries))
        self.assertFalse(any('FROM "catalog_gameseries"' in q["sql"] for q in queries))
        self.assertEqual(self.client.post(url).status_code, 405)
        self.assertEqual(self.client.get(reverse("orders:export_suppliers", args=(999999,))).status_code, 404)

    def test_partial_and_excessive_allocations_do_not_silently_change_order_quantities(self):
        create_trip(actor=self.user, order_ids=[self.first.pk, self.second.pk], creation_key=uuid4())
        allocation = self.first.items.get(cd=self.cd).trip_allocations.get()
        allocation.postpayment_quantity = 1
        allocation.save(update_fields=("postpayment_quantity",))
        order, rows = self.rows()
        unassigned = [row for row in rows if row.supplier is None]
        self.assertEqual(sum(row.quantity for row in unassigned), 2)
        self.assertEqual(sum(row.quantity for row in rows), 6)
        self.assertTrue(all(row.unit_price is None for row in unassigned))
        allocation.postpayment_quantity = 100
        allocation.save(update_fields=("postpayment_quantity",))
        with self.assertRaisesMessage(ValidationError, "превышает количество"):
            self.rows()

    def test_m_to_f_rule_uses_f_as_recipient_and_keeps_m_price_and_zero_allocation(self):
        m, f = [Supplier.objects.create(name=f"Поставщик {letter}", letter=letter,
                                      highlight_color="#37A7BA", legal_entity=letter, phone_1="1")
                for letter in ("M", "F")]
        SupplierCDPrice.objects.create(supplier=m, cd=self.cd, price=80)
        SupplierCDPrice.objects.create(supplier=f, cd=self.cd, price=95)
        price = create_procurement_price_list(
            actor=self.user, exchange_rate_usdt_aed=4, exchange_rate_usdt_rub=80,
            delivery_general_usdt_per_kg=0, delivery_devices_usdt_per_kg=0,
            delivery_storage_usdt_per_kg=0, delivery_cd_rub=0,
        )
        source = price.items.get(cd=self.cd)
        order = create_customer_procurement_order(actor=self.user, price_list_id=price.pk, recipient="Покупатель",
            comment="", lines=[{"price_list_item_id": source.pk, "prepayment_quantity": 1, "postpayment_quantity": 2}])
        create_trip(actor=self.user, order_ids=[order.pk], creation_key=uuid4())
        self.client.force_login(self.user)
        response = self.client.get(reverse("orders:export_suppliers", args=(order.pk,)))
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content))
        sheet = book.active
        data = [row for row in sheet.values if row[5] in ("cd", "tech")]
        self.assertEqual([(row[9], row[2], row[3]) for row in data if row[2]], [(f.pk, 3, 80)])
        self.assertEqual([(row[9], row[2]) for row in data if not row[2]], [(m.pk, 0)])
        self.assertEqual((sheet.cell(sheet.max_row, 3).value, sheet.cell(sheet.max_row, 5).value), (3, 240))
        book.close()

    def test_stock_order_button_exports_purchase_aed_not_customer_ruble_cost(self):
        order = create_customer_procurement_order(
            actor=self.user, price_list_id=self.price.pk, recipient="На склад", comment="",
            purpose=CustomerProcurementOrder.Purpose.STOCK, lines=[
                {"price_list_item_id": self.cd_source.pk, "prepayment_quantity": 2, "postpayment_quantity": 0},
                {"price_list_item_id": self.tech_source.pk, "prepayment_quantity": 3, "postpayment_quantity": 0},
            ],
        )
        self.client.force_login(self.user)
        url = reverse("orders:export_suppliers", args=(order.pk,))
        self.assertContains(self.client.get(reverse("orders:detail", args=(order.pk,))), url)
        book = load_workbook(BytesIO(self.client.get(url).content))
        sheet = book.active
        self.assertEqual((sheet.cell(sheet.max_row, 3).value, sheet.cell(sheet.max_row, 5).value), (5, 620))
        data = [row for row in sheet.values if row[5] in ("cd", "tech")]
        self.assertEqual([row[3] for row in data], [100, 140])
        book.close()
