from decimal import Decimal
from uuid import uuid4
from django.core.exceptions import ValidationError, PermissionDenied
from django.test import TestCase
from django.urls import reverse
from accounts.models import User
from catalog.models import CD, Platform, Tech, Brand, ProductType
from partners.models import Partner, Supplier
from price.models import ProcurementPriceList, ProcurementPriceListItem
from sales.models import Sale
from .models import CustomerProcurementOrder, SupplierTripAllocation, TripAdditionReceipt
from .services import create_customer_procurement_order
from .trips import create_trip, _assert_balances
from .trip_additions import prepare_addition, commit_addition


class TripAdditionTests(TestCase):
    def setUp(self):
        self.user=User.objects.create_superuser('addition',password='test')
        self.supplier=Supplier.objects.create(name='Поставщик',letter='A',legal_entity='ООО',phone_1='1')
        self.cd=CD.objects.create(name='Игра',sku='ADD-CD',platform=Platform.objects.create(name='PS5'),weight_grams=100)
        self.tech=Tech.objects.create(name='Геймпад',sku='ADD-T',brand=Brand.objects.create(name='Sony'),
            product_type=ProductType.objects.create(name='Геймпад',procurement_delivery_group='devices'),weight_grams=500)
        self.partner=Partner.objects.create(name='Партнёр')
        self.other=Partner.objects.create(name='Другой партнёр')
        self.price=ProcurementPriceList.objects.create(created_by=self.user,exchange_rate_usdt_aed=4,exchange_rate_usdt_rub=80,
            exchange_rate_aed_rub=20,delivery_cd_rub=50,delivery_general_usdt_per_kg=2,delivery_devices_usdt_per_kg=3,delivery_storage_usdt_per_kg=4)
        source=ProcurementPriceListItem.objects.create(price_list=self.price,product_kind='cd',cd=self.cd,product_name_snapshot=self.cd.name,
            article_snapshot=self.cd.sku,selected_supplier=self.supplier,supplier_price_aed=100,exchange_rate_aed_rub=20,
            base_price_rub=2000,delivery_rub=50,prepayment_price_rub=2100,postpayment_price_rub=2200)
        self.old_order=create_customer_procurement_order(actor=self.user,price_list_id=self.price.pk,recipient=self.partner.name,
            partner_id=self.partner.pk,comment='',lines=[{'price_list_item_id':source.pk,'prepayment_quantity':1,'postpayment_quantity':1}])
        self.batch=create_trip(actor=self.user,order_ids=[self.old_order.pk],creation_key=uuid4())
        self.lines=[{'product':f'cd:{self.cd.pk}','supplier':self.supplier.pk,'aed':'110','customer_price':'2500','stock_qty':'2','customer_qty':'3'},
                    {'product':f'tech:{self.tech.pk}','supplier':self.supplier.pk,'aed':'150','customer_price':'100','stock_qty':'1','customer_qty':'1'}]
        self.ids=[self.partner.pk,self.other.pk]
        self.allocations={f'0:{self.partner.pk}':1,f'0:{self.other.pk}':2,f'1:{self.partner.pk}':0,f'1:{self.other.pk}':1}
        self.client.force_login(self.user)

    def commit(self, key=None, **overrides):
        _,_,_,fingerprint=prepare_addition(batch=self.batch,lines=self.lines,partner_ids=self.ids,allocations=self.allocations)
        args=dict(actor=self.user,batch_id=self.batch.pk,key=key or uuid4(),lines=self.lines,partner_ids=self.ids,
                  allocations=self.allocations,preview_fingerprint=fingerprint,revision=self.batch.revision)
        args.update(overrides)
        return commit_addition(**args)

    def test_new_orders_for_existing_partner_preserve_all_old_rows_and_prices(self):
        old_order=CustomerProcurementOrder.objects.filter(pk=self.old_order.pk).values().get()
        old_items=list(self.old_order.items.values());old_lines=list(self.batch.lines.values());old_alloc=list(SupplierTripAllocation.objects.values())
        new=self.commit();self.assertEqual(len(new),3)
        self.assertEqual(CustomerProcurementOrder.objects.filter(pk=self.old_order.pk).values().get(),old_order)
        self.assertEqual(list(self.old_order.items.values()),old_items)
        self.assertEqual(list(self.batch.lines.filter(pk__in=[x['id'] for x in old_lines]).values()),old_lines)
        self.assertEqual(list(SupplierTripAllocation.objects.filter(pk__in=[x['id'] for x in old_alloc]).values()),old_alloc)
        self.assertEqual(CustomerProcurementOrder.objects.filter(partner=self.partner).count(),2)
        self.assertEqual(sum(order.total_units for order in new),7)
        self.assertFalse(Sale.objects.exists());_assert_balances(self.batch)
        self.assertEqual(ProcurementPriceList.objects.filter(is_trip_addition=False).first(),self.price)
        stock=next(order for order in new if order.is_stock_order)
        self.assertEqual(stock.items.get(cd=self.cd).prepayment_unit_price,2250)
        self.assertEqual(stock.items.get(tech=self.tech).prepayment_unit_price,3120)
        customer=next(order for order in new if order.partner_id==self.other.pk)
        self.assertEqual(customer.items.get(tech=self.tech).postpayment_unit_price,100)

    def test_repeat_identical_commit_returns_same_orders(self):
        key=uuid4();first=self.commit(key)
        second=self.commit(key)
        self.assertEqual({order.pk for order in first},{order.pk for order in second})
        self.assertEqual(TripAdditionReceipt.objects.count(),1)

    def test_bad_distribution_stale_revision_and_changed_weight_are_atomic(self):
        _,_,_,fingerprint=prepare_addition(batch=self.batch,lines=self.lines,partner_ids=self.ids,allocations=self.allocations)
        bad={**self.allocations,f'0:{self.other.pk}':1}
        with self.assertRaises(ValidationError):prepare_addition(batch=self.batch,lines=self.lines,partner_ids=self.ids,allocations=bad)
        with self.assertRaises(ValidationError):self.commit(revision=99)
        self.tech.weight_grams=1000;self.tech.save(update_fields=('weight_grams',))
        with self.assertRaises(ValidationError):self.commit(preview_fingerprint=fingerprint)
        self.assertEqual(CustomerProcurementOrder.objects.count(),1);self.assertFalse(TripAdditionReceipt.objects.exists())

    def test_permission_and_missing_fixed_tech_tariff(self):
        viewer=User.objects.create_user('no-add')
        with self.assertRaises(PermissionDenied):self.commit(actor=viewer)
        self.price.delivery_devices_usdt_per_kg=None;self.price.save(update_fields=('delivery_devices_usdt_per_kg',))
        self.batch.refresh_from_db()
        with self.assertRaises(ValidationError):self.commit()

    def test_wizard_all_steps_and_partner_distribution(self):
        url=reverse('orders:trip_addition',args=(self.batch.pk,))
        response=self.client.get(url);self.assertEqual(response.status_code,302)
        url=response.url
        data={'lines-TOTAL_FORMS':'2','lines-INITIAL_FORMS':'0','lines-MIN_NUM_FORMS':'1','lines-MAX_NUM_FORMS':'100'}
        for i,line in enumerate(self.lines):data.update({f'lines-{i}-{key}':value for key,value in line.items()})
        response=self.client.post(url,data);self.assertEqual(response.status_code,302)
        response=self.client.post(response.url,{'partners':self.ids});self.assertEqual(response.status_code,302)
        response=self.client.post(response.url,self.allocations);self.assertEqual(response.status_code,302)
        page=self.client.get(response.url);self.assertContains(page,'Ожидаемая прибыль')
        response=self.client.post(response.url,{});self.assertRedirects(response,reverse('orders:trip_detail',args=(self.batch.pk,)))
        self.assertEqual(CustomerProcurementOrder.objects.count(),4)

    def test_only_supplier_excel_contains_game_codes_without_quantity_duplication(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.cd.cusa_ppsa_code='PPSA-00123, CUSA-00001, PPSA-00123'
        self.cd.save(update_fields=('cusa_ppsa_code',))
        response=self.client.get(reverse('orders:trip_export',args=(self.batch.pk,self.supplier.pk)))
        book=load_workbook(BytesIO(response.content));sheet=book.active
        self.assertEqual(sheet.cell(1,2).value,'CUSA / PPSA')
        data=[row for row in sheet.values if row[0]==self.cd.name]
        self.assertEqual(len(data),1);self.assertEqual(data[0][1],'PPSA-00123, CUSA-00001');self.assertEqual(data[0][-1],2)
        book.close()
        response=self.client.get(reverse('orders:trip_full_export',args=(self.batch.pk,)))
        book=load_workbook(BytesIO(response.content));self.assertNotIn('CUSA / PPSA',[cell.value for cell in book.active[1]]);book.close()

    def test_live_calculation_endpoint_is_read_only(self):
        response=self.client.post(reverse('orders:trip_addition_calculation',args=(self.batch.pk,)),{
            'lines-TOTAL_FORMS':'1','lines-INITIAL_FORMS':'0',**{f'lines-0-{key}':value for key,value in self.lines[0].items()}})
        self.assertEqual(response.status_code,200);self.assertEqual(response.json()['rows'][0]['cost'],'2250.00')
        self.assertEqual(CustomerProcurementOrder.objects.count(),1)
