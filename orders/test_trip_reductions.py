"""Уменьшение закупки у последнего поставщика и автоматические однозначные доли."""
from io import BytesIO
from unittest.mock import patch
from uuid import uuid4

from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from cash.models import CashTransaction
from partners.models import Supplier
from .models import CustomerProcurementOrderItem, OrderAdjustment, SupplierTripAllocation, SupplierTripChange
from .services import create_customer_procurement_order, advance_order_status
from . import test_trips as trip_tests
from .trips import create_trip, _assert_balances


class TripReductionTests(TestCase):
    order = classmethod(trip_tests.TripTests.order.__func__)
    distribution = trip_tests.TripTests.distribution
    change = trip_tests.TripTests.change

    @classmethod
    def setUpTestData(cls):
        trip_tests.TripTests.setUpTestData.__func__(cls)
        cls.b.letter, cls.b.name = "F", "Альфа — второй"
        cls.c.letter, cls.c.name = "M", "Бета — последний"
        cls.a.name = "Я — первый"
        for supplier in (cls.a, cls.b, cls.c):
            supplier.save(update_fields=("name", "letter"))

    def setUp(self):
        self.client.force_login(self.user)

    def trip(self, orders=None):
        return create_trip(actor=self.user, order_ids=[order.pk for order in (orders or [self.first, self.second])],
                           creation_key=uuid4())

    def last_line(self, trip, quantity=None):
        source = trip.lines.get(cd=self.cd, supplier=self.a)
        remaining = source.quantity - quantity if quantity is not None else 0
        self.change(trip, source, remaining, target_supplier_id=self.c.pk, price_supplier_id=self.a.pk)
        return trip.lines.get(cd=self.cd, supplier=self.c)

    def editor(self, trip, line):
        response = self.client.get(reverse("orders:trip_detail", args=[trip.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "js/order-trips.js?v=4")
        return response.context["editor"][str(line.pk)]

    def snapshot(self):
        return list(CustomerProcurementOrderItem.objects.order_by("pk").values())

    def test_priority_one_first_three_last_and_unrelated_suppliers_are_ignored(self):
        trip = self.trip()
        last = self.last_line(trip)
        unrelated = Supplier.objects.create(name="Посторонний", letter="X", priority=999,
                                            legal_entity="X", phone_1="1", highlight_color="#37A7BA")
        tech = trip.lines.get(tech=self.tech)
        self.change(trip, tech, 0, target_supplier_id=unrelated.pk, price_supplier_id=self.b.pk)
        self.assertTrue(self.editor(trip, last)["is_last_supplier"])
        self.assertFalse(self.editor(trip, last)["priority_ambiguous"])
        self.assertFalse(self.editor(trip, trip.lines.get(cd=self.cd, supplier=self.a))["is_last_supplier"])
        other = self.order("Другой рейс", 0, 2)
        other_trip = self.trip([other])
        self.change(other_trip, other_trip.lines.get(cd=self.cd), 0,
                    target_supplier_id=unrelated.pk, price_supplier_id=self.a.pk)
        self.assertTrue(self.editor(trip, last)["is_last_supplier"])

    def test_actual_receiver_without_saved_offer_participates_in_its_product_chain(self):
        trip = self.trip()
        extra = Supplier.objects.create(name="Резервный получатель", letter="X", priority=4,
                                        legal_entity="X", phone_1="1", highlight_color="#37A7BA")
        source = trip.lines.get(cd=self.cd)
        self.change(trip, source, 8, target_supplier_id=extra.pk, price_supplier_id=self.a.pk)
        self.change(trip, source, 3, target_supplier_id=self.c.pk, price_supplier_id=self.a.pk)
        self.assertFalse(self.editor(trip, trip.lines.get(cd=self.cd, supplier=self.c))["is_last_supplier"])
        self.assertTrue(self.editor(trip, trip.lines.get(cd=self.cd, supplier=extra))["is_last_supplier"])

    def test_equal_maximum_priorities_keep_explicit_choice_without_id_tiebreak(self):
        self.b.priority = 3
        self.b.save(update_fields=("priority",))
        trip = self.trip()
        line = self.last_line(trip)
        config = self.editor(trip, line)
        self.assertFalse(config["is_last_supplier"])
        self.assertTrue(config["priority_ambiguous"])
        self.change(trip, line, 6, target_supplier_id=self.b.pk, price_supplier_id=self.a.pk)
        self.assertEqual(trip.lines.get(cd=self.cd, supplier=self.b).quantity, 4)
        self.assertTrue(self.editor(trip, trip.lines.get(cd=self.cd, supplier=self.b))["priority_ambiguous"])

    def test_m_without_own_offer_accepts_transfer_at_a_price(self):
        self.cd_source.supplier_offers.exclude(supplier=self.a).delete()
        trip = self.trip()
        line = trip.lines.get(cd=self.cd, supplier=self.a)
        before = self.snapshot()
        config = self.editor(trip, line)
        self.assertFalse(config['is_last_supplier'])
        self.assertEqual(config['targets'][0]['id'], self.c.pk)
        self.assertEqual([offer['id'] for offer in config['prices']], [self.a.pk])
        self.change(trip, line, 0, target_supplier_id=self.c.pk, price_supplier_id=self.a.pk)
        line.refresh_from_db()
        target = trip.lines.get(cd=self.cd, supplier=self.c)
        self.assertEqual((line.quantity, target.quantity, target.price_supplier_id, target.supplier_price_aed_snapshot),
            (0, 10, self.a.pk, self.cd_source.supplier_price_aed))
        self.assertEqual(self.snapshot(), before)
        self.assertTrue(self.editor(trip, target)['is_last_supplier'])
        _assert_balances(trip)

    def test_excluded_m_is_not_added_as_fallback(self):
        from catalog.supplier_exclusions import set_supplier_exclusions
        self.cd_source.supplier_offers.exclude(supplier=self.a).delete()
        trip = self.trip()
        set_supplier_exclusions(product=self.cd, supplier_ids=[self.c.pk], actor=self.user)
        line = trip.lines.get(cd=self.cd, supplier=self.a)
        config = self.editor(trip, line)
        self.assertTrue(config['is_last_supplier'])
        self.assertNotIn(self.c.pk, [target['id'] for target in config['targets']])
        with self.assertRaises(ValidationError):
            self.change(trip, line, 0, target_supplier_id=self.c.pk, price_supplier_id=self.a.pk)
        self.assertEqual(trip.lines.get(cd=self.cd, supplier=self.a).quantity, 10)

    def test_restore_accidentally_zeroed_item_and_transfer_preserves_customer_prices(self):
        self.cd_source.supplier_offers.exclude(supplier=self.a).delete()
        order = self.order('Восстановление', 0, 1)
        trip = self.trip([order])
        initial_changes = trip.changes.count()
        line = trip.lines.get(cd=self.cd, supplier=self.a)
        item = order.items.get()
        prices = (item.prepayment_unit_price, item.postpayment_unit_price, item.supplier_price_aed_snapshot)
        self.change(trip, line, 0, action='adjust')
        allocation = line.allocations.get()
        self.change(trip, line, 1, action='adjust',
            distribution=[{'allocation_id': allocation.pk, 'final_quantity': 1}])
        self.change(trip, line, 0, target_supplier_id=self.c.pk, price_supplier_id=self.a.pk)
        item.refresh_from_db()
        target = trip.lines.get(cd=self.cd, supplier=self.c)
        self.assertEqual((item.prepayment_quantity, item.postpayment_quantity, target.quantity), (0, 1, 1))
        self.assertEqual((item.prepayment_unit_price, item.postpayment_unit_price, item.supplier_price_aed_snapshot), prices)
        self.assertEqual(target.supplier_price_aed_snapshot, prices[2])
        self.assertEqual(trip.changes.count(), initial_changes + 3)
        self.assertFalse(CashTransaction.objects.exists())
        _assert_balances(trip)

    def test_only_supplier_without_m_fallback_is_last(self):
        self.c.letter = 'C'
        self.c.save(update_fields=('letter',))
        self.cd_source.supplier_offers.exclude(supplier=self.a).delete()
        trip = self.trip()
        line = trip.lines.get(cd=self.cd)
        self.assertTrue(self.editor(trip, line)["is_last_supplier"])
        with self.assertRaisesMessage(ValidationError, "последний поставщик"):
            self.change(trip, line, 6, target_supplier_id=self.c.pk, price_supplier_id=self.a.pk)
        self.assertEqual(trip.lines.get(cd=self.cd).quantity, 10)

    def test_last_supplier_single_order_reduces_automatically_and_preserves_prices(self):
        order = self.order("Один заказ", 6, 4)
        trip = self.trip([order])
        line = self.last_line(trip)
        item = order.items.get()
        prices = (item.prepayment_unit_price, item.postpayment_unit_price, item.supplier_price_aed_snapshot)
        self.change(trip, line, 6, action="adjust")
        item.refresh_from_db()
        line.refresh_from_db()
        order.refresh_from_db()
        self.assertEqual((item.prepayment_quantity, item.postpayment_quantity, line.quantity), (6, 0, 6))
        self.assertEqual((item.prepayment_unit_price, item.postpayment_unit_price, item.supplier_price_aed_snapshot), prices)
        self.assertEqual(order.grand_total, item.line_total)
        self.assertEqual(trip.lines.get(cd=self.cd, supplier=self.a).quantity, 0)
        self.assertEqual(line.allocations.get().prepayment_cap, 6)
        _assert_balances(trip)

    def test_last_supplier_multiple_orders_still_requires_exact_manual_distribution(self):
        trip = self.trip()
        line = self.last_line(trip)
        before = self.snapshot()
        with self.assertRaisesMessage(ValidationError, "каждого заказчика"):
            self.change(trip, line, 7, action="adjust")
        for amounts in ((5, 3), (-1, 8), (6, 1)):
            with self.subTest(amounts=amounts), self.assertRaises(ValidationError):
                self.change(trip, line, 7, action="adjust", distribution=self.distribution(line, **{
                    str(self.first.pk): amounts[0], str(self.second.pk): amounts[1],
                }))
            self.assertEqual(self.snapshot(), before)
        self.change(trip, line, 7, action="adjust", distribution=self.distribution(line, **{
            str(self.first.pk): 5, str(self.second.pk): 2,
        }))
        line.refresh_from_db()
        self.assertEqual(line.quantity, 7)
        self.assertEqual(sum(row.quantity for row in line.allocations.all()), 7)
        self.assertEqual((self.first.items.get(cd=self.cd).quantity, self.second.items.get(cd=self.cd).quantity), (5, 2))
        self.assertEqual(trip.lines.get(tech=self.tech).quantity, 3)
        _assert_balances(trip)

    def test_last_supplier_zero_clears_three_orders_without_deleting_history(self):
        orders = [self.order(f"Заказ {index}", pre, post) for index, (pre, post) in enumerate(((2, 3), (1, 2), (0, 2)))]
        trip = self.trip(orders)
        line = self.last_line(trip)
        allocation_ids = list(line.allocations.values_list("pk", flat=True))
        item_ids = [order.items.get().pk for order in orders]
        self.change(trip, line, 0, action="adjust")
        line.refresh_from_db()
        self.assertEqual(line.quantity, 0)
        self.assertEqual(list(line.allocations.values_list("pk", flat=True)), allocation_ids)
        self.assertTrue(all(row.quantity == 0 for row in line.allocations.all()))
        self.assertEqual([CustomerProcurementOrderItem.objects.get(pk=pk).quantity for pk in item_ids], [0, 0, 0])
        for order in orders:
            order.refresh_from_db()
            self.assertEqual(order.grand_total, 0)
            self.assertEqual(order.adjustments.count(), 1)
        self.assertEqual(self.first.items.get(cd=self.cd).quantity, 5)
        self.assertFalse(CashTransaction.objects.exists())
        _assert_balances(trip)

    def test_zero_and_single_receiver_preserve_quantities_at_other_suppliers(self):
        for quantity, customer_total in ((2, 4), (0, 2)):
            with self.subTest(quantity=quantity):
                order = self.order(f"Частично у других {quantity}", 1, 4)
                trip = self.trip([order])
                line = self.last_line(trip, quantity=3)
                self.change(trip, line, quantity, action="adjust")
                self.assertEqual(order.items.get().quantity, customer_total)
                self.assertEqual(trip.lines.get(cd=self.cd, supplier=self.a).quantity, 2)
                self.assertEqual(trip.lines.get(cd=self.cd, supplier=self.c).quantity, quantity)
                _assert_balances(trip)

    def test_non_last_zero_without_transfer_automatically_clears_its_allocations(self):
        trip = self.trip()
        line = trip.lines.get(cd=self.cd)
        self.assertFalse(self.editor(trip, line)["is_last_supplier"])
        self.change(trip, line, 0, action="adjust")
        self.assertEqual(self.first.items.get(cd=self.cd).quantity, 0)
        self.assertEqual(self.second.items.get(cd=self.cd).quantity, 0)
        self.assertFalse(trip.lines.filter(cd=self.cd).exclude(supplier=self.a).exists())
        self.assertEqual(trip.lines.get(tech=self.tech).quantity, 3)

    def test_non_last_transfer_to_zero_keeps_customer_orders(self):
        trip = self.trip()
        before = self.snapshot()
        self.last_line(trip)
        self.assertEqual(self.snapshot(), before)
        self.assertFalse(OrderAdjustment.objects.exists())
        _assert_balances(trip)

    def test_direct_transfer_from_last_supplier_is_rejected_even_for_zero(self):
        trip = self.trip()
        line = self.last_line(trip)
        before = self.snapshot()
        events = trip.changes.count()
        for quantity in (6, 0):
            with self.subTest(quantity=quantity), self.assertRaisesMessage(ValidationError, "последний поставщик"):
                self.change(trip, line, quantity, target_supplier_id=self.a.pk, price_supplier_id=self.a.pk)
        self.assertEqual(self.snapshot(), before)
        self.assertEqual(trip.changes.count(), events)

    def test_failed_second_customer_update_rolls_back_quantities_totals_and_audit(self):
        trip = self.trip()
        line = self.last_line(trip)
        before = self.snapshot()
        allocations = list(SupplierTripAllocation.objects.order_by("pk").values())
        events = list(SupplierTripChange.objects.order_by("pk").values())
        from .trips import _save_totals
        def save_then_fail(order, items):
            if order.pk == self.second.pk:
                raise ValidationError("Ошибка второго заказа")
            return _save_totals(order, items)
        with patch("orders.trips._save_totals", side_effect=save_then_fail), self.assertRaisesMessage(ValidationError, "второго заказа"):
            self.change(trip, line, 0, action="adjust")
        self.assertEqual(self.snapshot(), before)
        self.assertEqual(list(SupplierTripAllocation.objects.order_by("pk").values()), allocations)
        self.assertEqual(list(SupplierTripChange.objects.order_by("pk").values()), events)
        self.assertFalse(OrderAdjustment.objects.exists())
        self.assertEqual(trip.lines.get(cd=self.cd, supplier=self.c).quantity, 10)

    def test_http_automatic_zero_and_repeated_request_do_not_duplicate_changes(self):
        trip = self.trip()
        line = self.last_line(trip)
        trip.refresh_from_db()
        url = reverse("orders:trip_line_update", args=[trip.pk, line.pk])
        payload = {"revision": trip.revision, "quantity": 0, "action": "adjust", "distribution": "null"}
        self.assertEqual(self.client.post(url, payload).status_code, 200)
        changes = OrderAdjustment.objects.count()
        self.assertEqual(self.client.post(url, payload).status_code, 409)
        self.assertEqual(OrderAdjustment.objects.count(), changes)
        self.assertEqual(trip.lines.get(cd=self.cd, supplier=self.c).quantity, 0)

    def test_increase_retains_manual_distribution_even_for_single_order(self):
        order = self.order("Восстановление", 2, 8)
        trip = self.trip([order])
        line = self.last_line(trip)
        self.change(trip, line, 0, action="adjust")
        with self.assertRaisesMessage(ValidationError, "каждого заказчика"):
            self.change(trip, line, 6, action="adjust")
        self.change(trip, line, 6, action="adjust", distribution=self.distribution(line, **{str(order.pk): 6}))
        self.assertEqual((order.items.get().prepayment_quantity, order.items.get().postpayment_quantity), (2, 4))

    def test_zero_is_hidden_in_supplier_excel_kept_for_customer_and_later_statuses_work(self):
        order = self.order("Нулевая закупка", 2, 8)
        trip = self.trip([order])
        line = self.last_line(trip)
        self.change(trip, line, 0, action="adjust")
        response = self.client.get(reverse("orders:trip_export", args=[trip.pk, self.c.pk]))
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content), data_only=True).active
        self.assertEqual(sheet.max_row, 1)
        self.assertFalse(any(row[0] == self.cd.name for row in sheet.values))
        response = self.client.get(reverse("orders:export", args=[order.pk]))
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content), data_only=True)
        customer_rows = [row for row in book.active.values if row[0] == self.cd.name]
        self.assertEqual(customer_rows, [(self.cd.name, "Не отгружено", 0, None, 0)])
        self.assertEqual(book.active.cell(book.active.max_row, 3).value, 0)
        self.assertEqual(book.active.cell(book.active.max_row, 5).value, 0)
        book.close()
        for status in ("confirmed", "receiving", "completed"):
            advance_order_status(actor=self.user, order_id=order.pk, next_status=status)
        order.refresh_from_db()
        self.assertEqual((order.status, order.grand_total), ("completed", 0))
        self.assertFalse(CashTransaction.objects.exists())

    def test_stock_order_keeps_single_quantity_and_its_original_unit_cost(self):
        order = create_customer_procurement_order(actor=self.user, price_list_id=self.price.pk,
            recipient="На склад", comment="", purpose="stock", lines=[{
                "price_list_item_id": self.cd_source.pk, "prepayment_quantity": 10, "postpayment_quantity": 0,
            }])
        trip = self.trip([order])
        line = self.last_line(trip)
        cost = order.items.get().prepayment_unit_price
        self.change(trip, line, 6, action="adjust")
        item = order.items.get()
        self.assertEqual((item.prepayment_quantity, item.postpayment_quantity, item.prepayment_unit_price), (6, 0, cost))
        self.assertFalse(CashTransaction.objects.exists())
