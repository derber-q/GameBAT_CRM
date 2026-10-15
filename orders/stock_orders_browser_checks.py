"""Путь скачивание → складской заказ → рейс → приход в изолированном Chromium."""
from io import BytesIO
import os
from pathlib import Path
from unittest.mock import patch

from django.conf import settings
from django.contrib.staticfiles.testing import StaticLiveServerTestCase
from django.urls import reverse
from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

from accounts.models import User
from catalog.models import CD, Platform
from partners.models import Supplier
from price.excel import MAIN_SHEET, META_SHEET
from price.services import create_procurement_price_list
from pricing.models import SupplierCDPrice
from warehouse.models import Warehouse, CDWarehouseStock
from .models import CustomerProcurementOrder, SupplierOrderBatch
from .services import create_customer_procurement_order


class StockOrdersBrowserChecks(StaticLiveServerTestCase):
    @patch.dict(os.environ, {"DJANGO_ALLOW_ASYNC_UNSAFE": "true"})
    def test_download_upload_mixed_trip_and_receipt_desktop_mobile(self):
        user = User.objects.create_superuser("stock-browser", password="StrongAdmin!123")
        self.client.force_login(user)
        warehouse = Warehouse.objects.create(name="Склад браузерной проверки")
        supplier = Supplier.objects.create(name="Поставщик A", letter="A", priority=1,
            highlight_color="#37A7BA", legal_entity="A", phone_1="1")
        cd = CD.objects.create(platform=Platform.objects.create(name="PS5"), name="Игра для наличия",
                               sku="000-STOCK", weight_grams=100)
        SupplierCDPrice.objects.create(supplier=supplier, cd=cd, price=100)
        price = create_procurement_price_list(actor=user, exchange_rate_usdt_aed=4, exchange_rate_usdt_rub=80,
            delivery_general_usdt_per_kg=0, delivery_devices_usdt_per_kg=0,
            delivery_storage_usdt_per_kg=0, delivery_cd_rub=50)
        source = price.items.get()
        output = Path(settings.BASE_DIR) / "outputs/stock-orders-browser"
        output.mkdir(parents=True, exist_ok=True)
        with sync_playwright() as pw:
            browser = pw.chromium.launch()
            for index, width in enumerate((1280, 390), 1):
                customer = create_customer_procurement_order(actor=user, price_list_id=price.pk,
                    recipient=f"Клиент {width}", comment="", lines=[{"price_list_item_id": source.pk,
                    "prepayment_quantity": 1, "postpayment_quantity": 0}])
                context = browser.new_context(viewport={"width": width, "height": 900}, accept_downloads=True)
                context.add_cookies([{"name": settings.SESSION_COOKIE_NAME,
                    "value": self.client.cookies[settings.SESSION_COOKIE_NAME].value, "url": self.live_server_url}])
                context.route("**/crm/api/exchange-rates/", lambda route: route.fulfill(json={"rates": {}, "available": False}))
                page = context.new_page()
                page.goto(self.live_server_url + reverse("price:imports"))
                with page.expect_download() as download:
                    page.get_by_role("link", name="Скачать прайс для заказа на склад", exact=True).click()
                book = load_workbook(BytesIO(Path(download.value.path()).read_bytes()))
                for row in book[META_SHEET].iter_rows(values_only=True):
                    if isinstance(row[0], int):
                        book[MAIN_SHEET].cell(row[0], 6, 3)
                filled = BytesIO()
                book.save(filled)
                book.close()
                page.goto(self.live_server_url + reverse("orders:upload"))
                page.locator('[name="stock_file"]').set_input_files({"name": "stock.xlsx",
                    "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "buffer": filled.getvalue()})
                page.get_by_role("button", name="Проверить и показать состав", exact=True).click()
                expect(page.get_by_role("columnheader", name="Себестоимость с доставкой", exact=True)).to_be_visible()
                page.get_by_role("button", name="Создать заказ", exact=True).click()
                expect(page.get_by_role("heading", name="Заказ на склад", exact=False)).to_be_visible()
                stock = CustomerProcurementOrder.objects.filter(purpose="stock").latest("pk")
                page.goto(self.live_server_url + reverse("orders:list"))
                for order_id in (stock.pk, customer.pk):
                    page.locator(f'input[name="order_id"][value="{order_id}"]').check()
                page.get_by_role("button", name="Объединить в рейс", exact=True).click()
                trip = SupplierOrderBatch.objects.latest("pk")
                expect(page).to_have_url(self.live_server_url + reverse("orders:trip_detail", args=(trip.pk,)))
                self.assertEqual(trip.lines.get().quantity, 4)
                page.goto(self.live_server_url + reverse("orders:detail", args=(stock.pk,)))
                page.screenshot(path=str(output / f"stock-order-{width}.png"), full_page=True)
                page.get_by_role("button", name="Подтвердить закупку", exact=True).click()
                page.get_by_role("button", name="Начать приёмку", exact=True).click()
                page.get_by_role("link", name="Скорректировать количество", exact=True).click()
                page.locator('input[name="prepayment_quantity"]').fill("2")
                page.locator('[name="comment"]').fill("Приехало две единицы")
                page.get_by_role("button", name="Сохранить корректировку", exact=True).click()
                page.locator('select[name="warehouse"]').select_option(str(warehouse.pk))
                page.get_by_role("button", name="Оформить приход", exact=True).click()
                stock.refresh_from_db()
                expect(page).to_have_url(self.live_server_url + reverse("supplies:price_review", args=(stock.supply_id,)))
                page.get_by_role("button", name="Подтвердить цены и поставить на баланс", exact=True).click()
                expect(page).to_have_url(self.live_server_url + reverse("supplies:detail", args=(stock.supply_id,)))
                self.assertEqual(CDWarehouseStock.objects.get(cd=cd, warehouse=warehouse).quantity, index * 2)
                self.assertEqual(stock.supply.cd_items.get().purchase_unit_cost, source.stock_unit_cost)
                page.screenshot(path=str(output / f"stock-receipt-{width}.png"), full_page=True)
                context.close()
            browser.close()
