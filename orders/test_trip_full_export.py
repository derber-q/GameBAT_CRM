"""Общий Excel рейса: снимки AED, сортировка, объединение и доступ."""
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

from django.contrib.auth.models import Permission
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import User
from catalog.models import CD
from pricing.models import SupplierCDPrice
from .models import SupplierOrderBatch
from .test_excel_grouping import create_export_fixture
from .trip_excel import trip_order_xlsx
from .trips import create_trip, change_trip_line, apply_m_to_f_rule, return_line_to_m
from .services import advance_order_status


class TripFullExportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        create_export_fixture(cls)
        cls.trip = create_trip(actor=cls.user, order_ids=[cls.first.pk, cls.second.pk], creation_key=uuid4())

    def setUp(self):
        self.client.force_login(self.user)
        self.url = reverse("orders:trip_full_export", args=(self.trip.pk,))

    @staticmethod
    def data_rows(sheet):
        return [row for row in sheet.values if isinstance(row[3], int)]

    def transfer(self, quantity, price_supplier):
        self.trip.refresh_from_db()
        source = self.trip.lines.get(cd=self.cd, supplier=self.a)
        change_trip_line(actor=self.user, batch_id=self.trip.pk, line_id=source.pk,
            revision=self.trip.revision, quantity=quantity, action="transfer",
            target_supplier_id=self.c.pk, price_supplier_id=price_supplier.pk)

    def workbook(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        return load_workbook(BytesIO(response.content))

    def test_all_suppliers_in_one_sorted_table_and_equal_prices_are_combined(self):
        self.transfer(7, self.a)
        book = self.workbook()
        self.assertEqual(book.sheetnames, ["Товары рейса"])
        sheet = book.active
        self.assertEqual(self.data_rows(sheet), [("Консоль", "000-tech", 140, 3, self.b.name),
                                                ("Игра", "000123", 100, 12, f"{self.a.name}, {self.c.name}")])
        titles = [row[0] for row in sheet.values if str(row[0]).startswith(("Tech ·", "CD ·"))]
        self.assertEqual(titles, ["Tech · Sony · Консоли", "CD · Платформа: PS5"])
        self.assertEqual(sheet.max_column, 5)
        self.assertEqual(sheet["E1"].value, "Поставщики")
        self.assertFalse(sheet.column_dimensions["E"].hidden)
        self.assertTrue(sheet.column_dimensions["B"].hidden)
        self.assertEqual(sheet["B5"].number_format, "@")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertFalse(any(str(row[0]).startswith("Поставщик") for row in sheet.values))
        self.assertEqual(sheet.auto_filter.ref, "A1:E5")
        self.assertNotIn(self.first.recipient, str(list(sheet.values)))
        book.close()

    def test_different_prices_stay_separate_adjacent_and_use_snapshots(self):
        self.transfer(7, self.c)
        self.cd.name = "Новое имя после заказа"
        self.cd.is_archived = True
        self.cd.save(update_fields=("name", "is_archived"))
        offer = SupplierCDPrice.objects.get(supplier=self.c, cd=self.cd)
        offer.price = Decimal("999")
        offer.save(update_fields=("price",))
        book = self.workbook()
        self.assertEqual(self.data_rows(book.active), [
            ("Консоль", "000-tech", 140, 3, self.b.name), ("Игра", "000123", 100, 7, self.a.name),
            ("Игра", "000123", 120, 5, self.c.name),
        ])
        self.assertEqual(sum(row[2] * row[3] for row in self.data_rows(book.active)), 1720)
        book.close()

    def test_zero_source_does_not_duplicate_transferred_quantity_and_zero_product_is_kept(self):
        self.transfer(0, self.a)
        book = self.workbook()
        self.assertEqual(self.data_rows(book.active)[1:], [("Игра", "000123", 100, 12, self.c.name)])
        book.close()
        self.trip.refresh_from_db()
        line = self.trip.lines.get(cd=self.cd, supplier=self.c)
        change_trip_line(actor=self.user, batch_id=self.trip.pk, line_id=line.pk,
                         revision=self.trip.revision, quantity=0, action="adjust")
        book = self.workbook()
        self.assertEqual(self.data_rows(book.active), [("Консоль", "000-tech", 140, 3, self.b.name),
                                                     ("Игра", "000123", 100, 0, "—")])
        book.close()
        self.assertEqual(self.trip.lines.filter(cd=self.cd).count(), 2)

    def test_sorting_is_by_products_and_prices_and_preserves_text_and_decimal_precision(self):
        earlier = CD.objects.create(platform=self.cd.platform, name="Альфа", sku="000001")
        def row(product, price, quantity, name=None):
            return SimpleNamespace(product=product, product_kind="cd", product_id=product.pk,
                product_name_snapshot=name or product.name, article_snapshot=product.sku,
                supplier_price_aed_snapshot=Decimal(price), quantity=quantity,
                supplier=self.a, supplier_id=self.a.pk)
        book = load_workbook(BytesIO(trip_order_xlsx([
            row(self.cd, "1.123456", 2), row(earlier, "2.5", 1, '=HYPERLINK("bad")'),
            row(self.cd, "1", 4), row(self.cd, "1.123456", 3),
        ])))
        self.assertEqual(self.data_rows(book.active), [
            ('\'=HYPERLINK("bad")', "000001", 2.5, 1, self.a.name), ("Игра", "000123", 1, 4, self.a.name),
            ("Игра", "000123", 1.123456, 5, self.a.name),
        ])
        self.assertEqual(book.active["A3"].data_type, "s")
        self.assertEqual(book.active["C5"].number_format, "#,##0.######")
        book.close()

    def test_current_receiver_f_is_shown_at_m_price_and_refreshes_after_return(self):
        self.a.letter = "M"
        self.a.save(update_fields=("letter",))
        self.b.letter = "F"
        self.b.save(update_fields=("letter",))
        apply_m_to_f_rule(actor=self.user, batch_id=self.trip.pk)
        book = self.workbook()
        self.assertEqual(self.data_rows(book.active)[1:], [("Игра", "000123", 100, 12, self.b.name)])
        book.close()
        self.trip.refresh_from_db()
        line = self.trip.lines.get(cd=self.cd, supplier=self.b)
        return_line_to_m(actor=self.user, batch_id=self.trip.pk, line_id=line.pk, revision=self.trip.revision)
        book = self.workbook()
        self.assertEqual(self.data_rows(book.active)[1:], [("Игра", "000123", 100, 12, self.a.name)])
        book.close()

    def test_each_trip_has_its_own_link_legacy_batches_do_not(self):
        other = SupplierOrderBatch.objects.create(created_by=self.user, price_list=self.price)
        legacy = SupplierOrderBatch.objects.create(created_by=self.user)
        response = self.client.get(reverse("orders:list"))
        self.assertContains(response, self.url)
        self.assertContains(response, reverse("orders:trip_full_export", args=(other.pk,)))
        legacy_url = reverse("orders:trip_full_export", args=(legacy.pk,))
        self.assertNotContains(response, legacy_url)
        self.assertEqual(self.client.get(legacy_url).status_code, 404)
        self.assertEqual(self.client.get(reverse("orders:trip_full_export", args=(999999,))).status_code, 404)
        book = load_workbook(BytesIO(self.client.get(reverse("orders:trip_full_export", args=(other.pk,))).content))
        self.assertEqual(self.data_rows(book.active), [])
        book.close()

    def test_access_requires_all_trip_view_and_price_permissions_and_hides_button(self):
        user = User.objects.create_user("trip-export-reader", password="StrongReader!123")
        names = (("orders", "generate_supplier_orders"), ("partners", "view_supplier_details"),
                 ("price", "view_supplier_prices"))
        required = [Permission.objects.get(content_type__app_label=app, codename=name) for app, name in names]
        list_permission = Permission.objects.get(content_type__app_label="orders", codename="view_orders")
        for missing in (0, 1, 2, None):
            with self.subTest(missing=missing):
                user.user_permissions.set([list_permission] + [permission for index, permission in enumerate(required) if index != missing])
                self.client.force_login(user)
                self.assertEqual(self.client.get(self.url).status_code, 200 if missing is None else 403)
                response = self.client.get(reverse("orders:list"))
                (self.assertContains if missing is None else self.assertNotContains)(response, self.url)
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 302)

    def test_download_is_read_only_get_only_uncached_and_has_no_per_product_queries(self):
        with CaptureQueriesContext(connection) as queries:
            response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertIn(f"trip_{self.trip.pk}_all.xlsx", response["Content-Disposition"])
        self.assertIn("no-store", response["Cache-Control"])
        self.assertEqual(response["X-Content-Type-Options"], "nosniff")
        self.assertLessEqual(len(queries), 7)
        self.assertFalse(any(query["sql"].lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE")) for query in queries))
        self.assertFalse(any('FROM "catalog_gameseries"' in query["sql"] for query in queries))
        self.assertEqual(self.client.post(self.url).status_code, 405)

    def test_confirmed_trip_can_be_downloaded_without_reopening_orders(self):
        advance_order_status(actor=self.user, order_id=self.first.pk, next_status="confirmed", warehouse_id=1)
        book = self.workbook()
        self.assertEqual(sum(row[3] for row in self.data_rows(book.active)), 15)
        book.close()
