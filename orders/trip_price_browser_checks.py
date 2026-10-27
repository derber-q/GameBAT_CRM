"""Предпросмотр цены и сохранение в реальном Chromium на отдельной БД."""
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from io import BytesIO
import json
from pathlib import Path
from uuid import uuid4

from django.conf import settings
from django.contrib.staticfiles.testing import StaticLiveServerTestCase
from django.urls import reverse
from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

from accounts.models import User
from catalog.models import CD, Platform
from partners.models import Supplier
from price.services import create_procurement_price_list
from pricing.models import SupplierCDPrice
from sales.payment_browser_checks import database_call
from warehouse.models import Warehouse
from .models import SupplierOrderBatch, SupplierOrderBatchLine, SupplierTripChange
from .services import advance_order_status, create_customer_procurement_order, receive_stock_order
from .trips import create_trip
from .trip_prices import change_trip_price


class TripPriceBrowserChecks(StaticLiveServerTestCase):
    def test_preview_confirm_stale_tab_loss_mobile_and_downstream_lock(self):
        user = User.objects.create_superuser('trip-price-browser')
        self.client.force_login(user)
        supplier = Supplier.objects.create(name='Поставщик B', letter='B', priority=1,
            legal_entity='B', phone_1='1', highlight_color='#37A7BA')
        product = CD.objects.create(platform=Platform.objects.create(name='PS5'), name='Игра для проверки цены', sku='000PRICE')
        SupplierCDPrice.objects.create(supplier=supplier, cd=product, price=100)
        price = create_procurement_price_list(actor=user, exchange_rate_usdt_aed=4, exchange_rate_usdt_rub=100,
            delivery_general_usdt_per_kg=0, delivery_devices_usdt_per_kg=0,
            delivery_storage_usdt_per_kg=0, delivery_cd_rub=17)
        source = price.items.get()
        source.prepayment_price_rub, source.postpayment_price_rub = Decimal('3500'), Decimal('3600')
        source.save(update_fields=('prepayment_price_rub', 'postpayment_price_rub'))
        orders = [create_customer_procurement_order(actor=user, price_list_id=price.pk, recipient=name,
            comment='', purpose=purpose, lines=[{'price_list_item_id': source.pk,
                'prepayment_quantity': pre, 'postpayment_quantity': post}])
            for name, purpose, pre, post in [('Покупатель 1', 'customer', 4, 0),
                ('Покупатель 2', 'customer', 0, 3), ('На склад', 'stock', 3, 0)]]
        trip = create_trip(actor=user, order_ids=[row.pk for row in orders], creation_key=uuid4())
        line = trip.lines.get()
        warehouse = Warehouse.objects.create(name='Браузерный склад')
        output = Path(settings.BASE_DIR) / 'outputs/trip_supplier_price_20261022/browser'
        output.mkdir(parents=True, exist_ok=True)
        errors = []
        with ThreadPoolExecutor(max_workers=1) as database, sync_playwright() as pw:
            run = lambda callback: database.submit(database_call, callback).result()
            browser = pw.chromium.launch()
            context = browser.new_context(viewport={'width': 1440, 'height': 1000})
            context.add_cookies([{'name': settings.SESSION_COOKIE_NAME,
                'value': self.client.cookies[settings.SESSION_COOKIE_NAME].value, 'url': self.live_server_url}])
            context.route('**/crm/api/exchange-rates/', lambda route: route.fulfill(json={'rates': {}, 'available': False}))
            page, stale = context.new_page(), context.new_page()
            for tab in (page, stale):
                tab.on('pageerror', lambda error: errors.append(str(error)))
                tab.goto(self.live_server_url + reverse('orders:trip_detail', args=(trip.pk,)))
            sale_navigation = page.locator('.nav-sale')
            sale_navigation.hover()
            expect(sale_navigation.locator('..').locator('.nav-dropdown-menu')).to_be_visible()
            page.locator('.brand').hover()
            dialog = page.locator('#trip-price-dialog')
            button = page.locator(f'[data-trip-price-edit="{line.pk}"]')
            button.click()
            page.locator('#trip-price-input').fill('110')
            expect(page.locator('[data-trip-price-cost-delta]')).to_have_text('2 500,00')
            expect(page.locator('[data-trip-price-profit-delta]')).to_have_text('-1 750,00')
            expect(page.locator('[data-trip-price-stock-delta]')).to_have_text('750,00')
            expect(page.locator('[data-trip-price-cost-delta]')).to_have_css('color', 'rgb(172, 36, 63)')
            expect(page.locator('[data-trip-price-profit-delta]')).to_have_css('color', 'rgb(172, 36, 63)')
            self.assertEqual(run(lambda: SupplierOrderBatchLine.objects.get(pk=line.pk).supplier_price_aed_snapshot), 100)
            page.screenshot(path=str(output / 'preview-desktop.png'))
            page.locator('[data-trip-price-cancel]').click()
            expect(dialog).not_to_be_visible()
            button.click()
            page.locator('#trip-price-input').fill('0')
            expect(page.locator('[data-trip-price-error]')).to_contain_text('больше нуля')
            expect(page.locator('[data-trip-price-save]')).to_be_disabled()
            page.locator('#trip-price-input').fill('110')
            expect(page.get_by_label('Причина изменения (необязательно)')).to_have_attribute('name', 'reason')
            self.assertFalse(page.locator('#trip-price-reason').evaluate('(node) => node.required'))
            expect(page.locator('[data-trip-price-save]')).to_be_enabled()
            page.locator('[data-trip-price-save]').click()
            expect(page.locator(f'[data-trip-line="{line.pk}"] [data-trip-current-price]')).to_contain_text('110')
            expect(dialog).not_to_be_visible()
            self.assertEqual(run(lambda: SupplierTripChange.objects.get(
                batch=trip, details__action='purchase_price').details['reason']), '')
            stale.locator(f'[data-trip-price-edit="{line.pk}"]').click()
            stale.locator('#trip-price-input').fill('120')
            expect(stale.locator('[data-trip-price-error]')).to_contain_text('Рейс уже изменён')
            expect(stale.locator('[data-trip-price-save]')).to_be_disabled()
            stale.close()
            page.set_viewport_size({'width': 390, 'height': 844})
            button.click()
            page.locator('#trip-price-input').fill('200')
            expect(page.locator('[data-trip-price-save]')).to_be_enabled()
            expect(page.locator('.trip-price-loss')).to_have_count(2)
            bounds = dialog.bounding_box()
            self.assertGreaterEqual(bounds['x'], 0)
            self.assertLessEqual(bounds['x'] + bounds['width'], 391)
            self.assertLessEqual(page.evaluate('document.documentElement.scrollWidth'), 390)
            page.screenshot(path=str(output / 'loss-mobile.png'))
            page.keyboard.press('Escape')
            expect(dialog).not_to_be_visible()
            self.assertEqual(run(lambda: SupplierTripChange.objects.filter(batch=trip, details__action='purchase_price').count()), 1)
            page.goto(self.live_server_url + reverse('orders:detail', args=(orders[0].pk,)))
            expect(page.locator('[data-procurement-profit]')).to_contain_text('2932')
            with page.expect_download() as download:
                page.locator(f'a[href="{reverse("orders:export", args=(orders[0].pk,))}"]').click()
            sheet = load_workbook(BytesIO(Path(download.value.path()).read_bytes())).active
            buyer_prices = [row[3] for row in sheet.iter_rows(values_only=True) if row[0] == product.name]
            self.assertIn(3500, buyer_prices)
            def change_price(value):
                revision = SupplierOrderBatch.objects.get(pk=trip.pk).revision
                return change_trip_price(actor=user, batch_id=trip.pk, line_id=line.pk,
                    revision=revision, price_aed=value)
            run(lambda: change_price('200'))
            page.goto(self.live_server_url + reverse('orders:trip_detail', args=(trip.pk,)))
            expect(page.locator('[data-procurement-profit] span')).to_have_css('color', 'rgb(201, 79, 85)')
            expect(page.locator('[data-procurement-profit]')).to_contain_text('-10319')
            run(lambda: change_price('110'))
            def supply():
                advance_order_status(actor=user, order_id=orders[2].pk, next_status='confirmed')
                advance_order_status(actor=user, order_id=orders[2].pk, next_status='receiving')
                result = receive_stock_order(actor=user, order_id=orders[2].pk, warehouse_id=warehouse.pk)
                return result.supply_id, result.supply.cd_items.get().purchase_unit_cost
            supply_id, supply_cost = run(supply)
            self.assertEqual(supply_cost, Decimal('2767'))
            page.goto(self.live_server_url + reverse('orders:trip_detail', args=(trip.pk,)))
            expect(page.locator('[data-trip-price-edit]')).to_have_count(0)
            expect(page.locator('[data-trip-current-price]')).to_contain_text('110')
            expect(page.get_by_text('Цена зафиксирована', exact=True)).to_have_count(1)
            page.screenshot(path=str(output / 'locked-mobile.png'), full_page=True)
            page.goto(self.live_server_url + reverse('supplies:price_review', args=(supply_id,)))
            expect(page.locator('body')).to_contain_text('2767')
            page.screenshot(path=str(output / 'supply-review-mobile.png'), full_page=True)
            self.assertEqual(errors, [])
            (output / 'results.json').write_text(json.dumps({'desktop': 1440, 'mobile': 390,
                'preview_has_no_writes': True, 'customer_price': 3500, 'stock_unit_cost': str(supply_cost),
                'stale_tab_blocked': True, 'zero_rejected': True, 'loss_allowed': True,
                'locked_after_supply': True, 'page_errors': errors}, ensure_ascii=False, indent=2), 'utf-8')
            browser.close()
