"""Скачивание общего состава рейса из таблицы заказов в Chromium."""
from contextlib import closing
from io import BytesIO
import os
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

from django.conf import settings
from django.contrib.staticfiles.testing import StaticLiveServerTestCase
from django.db import connections
from django.urls import reverse
from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

from .test_excel_grouping import create_export_fixture
from .trips import create_trip, change_trip_line


class TripExportBrowserChecks(StaticLiveServerTestCase):
    @patch.dict(os.environ, {"DJANGO_ALLOW_ASYNC_UNSAFE": "true"})
    def test_download_each_trips_sorted_combined_table_on_two_widths(self):
        create_export_fixture(self)
        self.client.force_login(self.user)
        trips = [create_trip(actor=self.user, order_ids=[order.pk], creation_key=uuid4())
                 for order in (self.first, self.second)]
        line = trips[0].lines.get(cd=self.cd)
        change_trip_line(actor=self.user, batch_id=trips[0].pk, line_id=line.pk,
            revision=trips[0].revision, quantity=2, action="transfer",
            target_supplier_id=self.c.pk, price_supplier_id=self.a.pk)
        output = Path(settings.BASE_DIR) / "outputs" / "trip-export-browser"
        output.mkdir(parents=True, exist_ok=True)
        with sync_playwright() as pw, closing(connections["default"]):
            browser = pw.chromium.launch()
            for width in (1280, 390):
                context = browser.new_context(viewport={"width": width, "height": 900}, accept_downloads=True)
                context.add_cookies([{"name": settings.SESSION_COOKIE_NAME,
                    "value": self.client.cookies[settings.SESSION_COOKIE_NAME].value, "url": self.live_server_url}])
                context.route("**/crm/api/exchange-rates/", lambda route: route.fulfill(json={"rates": {}, "available": False}))
                page = context.new_page()
                errors = []
                page.on("pageerror", lambda error: errors.append(str(error)))
                page.goto(self.live_server_url + reverse("orders:list"))
                for trip, quantities in zip(trips, ((1, 5), (2, 7))):
                    button = page.get_by_role("link", name=f"Скачать Excel всего рейса {trip.pk}", exact=True)
                    expect(button).to_have_count(1)
                    with page.expect_download() as download:
                        button.click()
                    self.assertEqual(download.value.suggested_filename, f"trip_{trip.pk}_all.xlsx")
                    contents = Path(download.value.path()).read_bytes()
                    book = load_workbook(BytesIO(contents))
                    self.assertEqual(book.sheetnames, ["Товары рейса"])
                    self.assertEqual(book.active["E1"].value, "Поставщики")
                    rows = [row for row in book.active.values if isinstance(row[3], int)]
                    cd_suppliers = f"{self.a.name}, {self.c.name}" if trip.pk == trips[0].pk else self.a.name
                    self.assertEqual(rows, [("Консоль", "000-tech", 140, quantities[0], self.b.name),
                                            ("Игра", "000123", 100, quantities[1], cd_suppliers)])
                    book.close()
                    (output / f"trip-{trip.pk}-{width}.xlsx").write_bytes(contents)
                page.screenshot(path=str(output / f"trip-list-{width}.png"), full_page=True)
                self.assertEqual(errors, [])
                context.close()
            browser.close()
