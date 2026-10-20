"""Состав одного заказа по фактическим распределениям рейса и сохранённым ценам."""
from collections import OrderedDict
from dataclasses import dataclass
from decimal import Decimal
from io import BytesIO

from django.core.exceptions import ValidationError
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from catalog.nomenclature_grouping import nomenclature_product_groups
from core.excel_layout import wrapped_text_height
from price.excel import BLUE, GRAY, LIGHT, TEAL, WHITE, safe_text


@dataclass(frozen=True)
class SupplierOrderExportRow:
    item: object
    supplier: object | None
    quantity: int
    unit_price: Decimal | None
    allocation_id: int | None = None

    @property
    def product(self):
        return self.item.product


def build_order_supplier_rows(*, order, items):
    """Не использует общий объём строки рейса или актуальные предложения поставщиков."""
    rows = []
    for item in items:
        if not order.trip_id:
            rows.append(SupplierOrderExportRow(
                item, item.selected_supplier, item.quantity, item.supplier_price_aed_snapshot,
            ))
            continue
        allocations = [allocation for allocation in item.trip_allocations.all()
                       if allocation.line.batch_id == order.trip_id]
        allocated = sum(allocation.quantity for allocation in allocations)
        if allocated > item.quantity:
            raise ValidationError(
                f"Распределение «{item.product_name_snapshot}» превышает количество в заказе.",
            )
        for allocation in allocations:
            rows.append(SupplierOrderExportRow(
                item, allocation.line.supplier, allocation.quantity,
                allocation.line.supplier_price_aed_snapshot, allocation.pk,
            ))
        remaining = item.quantity - allocated
        if remaining or not allocations:
            rows.append(SupplierOrderExportRow(item, None, remaining, None))
    return rows


def order_suppliers_xlsx(*, order, rows, show_supplier_details=False):
    """Закупочные суммы считаются один раз по строкам распределения, без повторного итога."""
    rows = list(rows)
    book = Workbook()
    sheet = book.active
    sheet.title = "По поставщикам"
    sheet.sheet_view.showGridLines = False
    widths = (62, 20, 14, 22, 24)
    headers = ("Наименование", "Артикул", "Количество", "Цена, AED", "Сумма, AED",
               "Тип", "ID товара", "ID позиции заказа", "ID распределения", "ID поставщика")
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + column)].width = width
    for letter in "FGHIJ":
        sheet.column_dimensions[letter].hidden = True

    def text_row(value, *, fill=LIGHT, color=BLUE, size=11):
        row = sheet.max_row + 1 if sheet["A1"].value is not None else 1
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        cell = sheet.cell(row, 1, safe_text(value))
        cell.font = Font(name="Arial", size=size, bold=True, color=color)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[row].height = wrapped_text_height(
            value, sum(widths), font_name="Arial", font_size=size, bold=True, min_height=28,
        )

    def data_row(values, *, total=False):
        sheet.append(values)
        row = sheet.max_row
        for cell in list(sheet[row])[:5]:
            cell.font = Font(name="Arial", size=11, color=BLUE, bold=total)
            cell.alignment = Alignment(vertical="center", wrap_text=cell.column <= 2,
                                       horizontal="right" if cell.column >= 3 else "left")
            cell.border = Border(bottom=Side(style="thin", color=GRAY))
            if total:
                cell.fill = PatternFill("solid", fgColor=LIGHT)
        sheet.cell(row, 2).number_format = "@"
        sheet.cell(row, 3).number_format = "0"
        for column in (4, 5):
            value = values[column - 1]
            sheet.cell(row, column).number_format = (
                "#,##0" if value is not None and value == value.to_integral_value() else "#,##0.######"
            )
        sheet.row_dimensions[row].height = max(
            wrapped_text_height(values[0], widths[0], font_name="Arial", font_size=11, min_height=28),
            wrapped_text_height(values[1] or "", widths[1], font_name="Arial", font_size=11, min_height=28),
        )

    text_row(f"Состав заказа №{order.pk} по поставщикам", fill=WHITE, size=16)
    text_row(f"Получатель: {order.recipient}", fill=WHITE)
    text_row(f"Сформировано: {timezone.localtime():%d.%m.%Y %H:%M:%S} · Закупочные цены в AED", fill=WHITE, size=10)
    for column, label in enumerate(headers, 1):
        cell = sheet.cell(4, column, label)
        cell.font = Font(name="Arial", size=11, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[4].height = 34
    suppliers = OrderedDict()
    for row in sorted(rows, key=lambda row: (
        row.supplier is None,
        row.supplier.letter.casefold() if row.supplier else "",
        row.supplier.pk if row.supplier else 0,
    )):
        if row.item.product_kind not in ("cd", "tech"):
            raise ValidationError("В заказе обнаружен неизвестный тип товара.")
        supplier_id = row.supplier.pk if row.supplier else None
        suppliers.setdefault(supplier_id, (row.supplier, []))[1].append(row)

    total_quantity, total_amount, total_unknown = 0, Decimal("0"), False
    for supplier, supplier_rows in suppliers.values():
        label = (supplier.name if show_supplier_details else f"Поставщик {supplier.safe_label}") if supplier else "Поставщик не назначен"
        text_row(label, fill=TEAL, color=WHITE, size=12)
        groups = nomenclature_product_groups(
            [row for row in supplier_rows if row.item.product_kind == "tech"],
            [row for row in supplier_rows if row.item.product_kind == "cd"],
            product_of=lambda row: row.product,
        )
        quantity, amount, unknown = 0, Decimal("0"), False
        for _kind, title, group in groups:
            text_row(title)
            for row in group:
                line_amount = row.quantity * row.unit_price if row.unit_price is not None else None
                data_row([
                    safe_text(row.item.product_name_snapshot), safe_text(row.item.article_snapshot),
                    row.quantity, row.unit_price, line_amount,
                    row.item.product_kind, row.item.product_id, row.item.pk, row.allocation_id,
                    supplier.pk if supplier else None,
                ])
                quantity += row.quantity
                if line_amount is not None:
                    amount += line_amount
                elif row.quantity:
                    unknown = True
        caption = "Итого по поставщику" + (" (есть позиции без цены)" if unknown else "")
        data_row([caption, None, quantity, None, None if unknown else amount], total=True)
        total_quantity += quantity
        total_amount += amount
        total_unknown |= unknown
    if not rows:
        text_row("В заказе нет товарных позиций.")
    caption = "Итого" + (" (есть позиции без цены)" if total_unknown else "")
    data_row([caption, None, total_quantity, None, None if total_unknown else total_amount], total=True)
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:E{sheet.max_row - 1}"
    sheet.print_title_rows = "4:4"
    sheet.print_area = f"A1:E{sheet.max_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()
