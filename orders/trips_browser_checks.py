"""Полный сценарий рейса в Chromium на изолированной тестовой БД."""
from io import BytesIO
import os
from pathlib import Path
from unittest.mock import patch

from django.conf import settings
from django.contrib.staticfiles.testing import StaticLiveServerTestCase
from django.urls import reverse
from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

from accounts.models import User
from catalog.models import CD, Platform
from partners.models import Supplier
from price.services import create_procurement_price_list
from pricing.models import SupplierCDPrice
from .models import SupplierOrderBatch
from .services import create_customer_procurement_order
from .trips import create_trip
from uuid import uuid4


class TripsBrowserChecks(StaticLiveServerTestCase):
    @patch.dict(os.environ, {"DJANGO_ALLOW_ASYNC_UNSAFE": "true"})
    def test_add_free_stock_order_and_move_whole_order_between_trips(self):
        user = User.objects.create_superuser("move-browser", password="StrongAdmin!123")
        self.client.force_login(user)
        supplier = Supplier.objects.create(name="A для переноса", letter="A", priority=1,
            highlight_color="#37A7BA", legal_entity="A", phone_1="1")
        platform = Platform.objects.create(name="PS5")
        products = [CD.objects.create(platform=platform, name=name, sku=f"MOVE-{index}")
                    for index, name in enumerate(("Общая позиция", "Только складской заказ"), 1)]
        for product in products:
            SupplierCDPrice.objects.create(supplier=supplier, cd=product, price=100)
        price = create_procurement_price_list(actor=user, exchange_rate_usdt_aed=4, exchange_rate_usdt_rub=80,
            delivery_general_usdt_per_kg=0, delivery_devices_usdt_per_kg=0,
            delivery_storage_usdt_per_kg=0, delivery_cd_rub=50)
        sources = [price.items.get(cd=product) for product in products]

        def order(recipient, amounts, purpose="customer"):
            return create_customer_procurement_order(actor=user, price_list_id=price.pk,
                recipient=recipient, comment="", purpose=purpose, lines=[
                    {"price_list_item_id": source.pk, "prepayment_quantity": qty, "postpayment_quantity": 0}
                    for source, qty in zip(sources, amounts)])

        output = Path(settings.BASE_DIR) / "outputs/trips-browser"
        output.mkdir(parents=True, exist_ok=True)
        with sync_playwright() as pw:
            browser = pw.chromium.launch()
            for width in (1280, 390):
                customer = order(f"Клиент {width}", (2,))
                stock = order(f"На склад {width}", (3, 1), purpose="stock")
                first = create_trip(actor=user, order_ids=[customer.pk], creation_key=uuid4())
                other = order(f"Другой клиент {width}", (5,))
                second = create_trip(actor=user, order_ids=[other.pk], creation_key=uuid4())
                context = browser.new_context(viewport={"width": width, "height": 900}, accept_downloads=True)
                context.add_cookies([{"name": settings.SESSION_COOKIE_NAME,
                    "value": self.client.cookies[settings.SESSION_COOKIE_NAME].value, "url": self.live_server_url}])
                context.route("**/crm/api/exchange-rates/", lambda route: route.fulfill(json={"rates": {}, "available": False}))
                page = context.new_page()
                errors = []
                page.on("pageerror", lambda error: errors.append(str(error)))
                page.goto(self.live_server_url + reverse("orders:detail", args=(stock.pk,)))
                page.get_by_role("link", name="Добавить в рейс", exact=True).click()
                page.locator('[name="target_trip"]').select_option(f"{first.pk}:{first.revision}")
                page.screenshot(path=str(output / f"add-to-trip-{width}.png"), full_page=True)
                page.get_by_role("button", name="Добавить в рейс", exact=True).click()
                expect(page).to_have_url(self.live_server_url + reverse("orders:trip_detail", args=(first.pk,)))
                stock.refresh_from_db()
                self.assertEqual(stock.trip_id, first.pk)
                page.locator(f'a[href="{reverse("orders:assign_trip", args=(stock.pk,))}"]').click()
                page.locator('[name="target_trip"]').select_option(f"{second.pk}:{second.revision}")
                page.screenshot(path=str(output / f"move-whole-order-{width}.png"), full_page=True)
                page.get_by_role("button", name="Перенести весь заказ", exact=True).click()
                expect(page).to_have_url(self.live_server_url + reverse("orders:trip_detail", args=(second.pk,)))
                stock.refresh_from_db()
                self.assertEqual(stock.trip_id, second.pk)
                self.assertFalse(first.source_orders.filter(pk=stock.pk).exists())
                self.assertEqual(first.lines.count(), 1)
                self.assertEqual(first.lines.get().quantity, 2)
                self.assertEqual(sum(second.lines.values_list("quantity", flat=True)), 9)
                with page.expect_download() as download:
                    page.locator(f'a[href="{reverse("orders:trip_export", args=(second.pk, supplier.pk))}"]').click()
                book = load_workbook(BytesIO(Path(download.value.path()).read_bytes()))
                self.assertEqual(sorted(row[3] for row in list(book.active.values)[1:]), [1, 8])
                book.close()
                page.screenshot(path=str(output / f"moved-orders-{width}.png"), full_page=True)
                self.assertEqual(errors, [])
                context.close()
            browser.close()

    @patch.dict(os.environ, {"DJANGO_ALLOW_ASYNC_UNSAFE": "true"})
    def test_default_m_to_f_download_and_return_button(self):
        user = User.objects.create_superuser("m-f-browser", password="StrongAdmin!123")
        self.client.force_login(user)
        m = Supplier.objects.create(name="M исходный", letter="M", priority=3, highlight_color="#37A7BA", legal_entity="M", phone_1="1")
        f = Supplier.objects.create(name="F удобный", letter="F", priority=2, highlight_color="#37A7BA", legal_entity="F", phone_1="1")
        cd = CD.objects.create(platform=Platform.objects.create(name="PS5"), name="Товар M для F", sku="MF-001")
        SupplierCDPrice.objects.create(supplier=m, cd=cd, price=100)
        SupplierCDPrice.objects.create(supplier=f, cd=cd, price=120)
        price = create_procurement_price_list(actor=user, exchange_rate_usdt_aed=4, exchange_rate_usdt_rub=80,
                    delivery_general_usdt_per_kg=0, delivery_devices_usdt_per_kg=0,
                    delivery_storage_usdt_per_kg=0, delivery_cd_rub=0)
        output = Path(settings.BASE_DIR) / "outputs/trips-browser"
        output.mkdir(parents=True, exist_ok=True)
        with sync_playwright() as pw:
            browser = pw.chromium.launch()
            for width in (1280, 390):
                order = create_customer_procurement_order(actor=user, price_list_id=price.pk,
                            recipient="Заказчик M/F", comment="", lines=[{"price_list_item_id": price.items.get().pk,
                            "prepayment_quantity": 2, "postpayment_quantity": 3}])
                trip = create_trip(actor=user, order_ids=[order.pk], creation_key=uuid4())
                context = browser.new_context(viewport={"width": width, "height": 900}, accept_downloads=True)
                context.add_cookies([{"name": settings.SESSION_COOKIE_NAME,
                                     "value": self.client.cookies[settings.SESSION_COOKIE_NAME].value,
                                     "url": self.live_server_url}])
                context.route("**/crm/api/exchange-rates/", lambda route: route.fulfill(json={"rates": {}, "available": False}))
                page = context.new_page()
                page.goto(self.live_server_url + reverse("orders:trip_detail", args=(trip.pk,)))
                button = page.get_by_role("button", name="Вернуть в M", exact=True)
                expect(button).to_have_count(1)
                link = page.locator(f'a[href="{reverse("orders:trip_export", args=(trip.pk, f.pk))}"]')
                with page.expect_download() as download:
                    link.click()
                workbook = load_workbook(BytesIO(Path(download.value.path()).read_bytes()))
                self.assertEqual(download.value.suggested_filename, f"trip_{trip.pk}_F.xlsx")
                self.assertEqual(workbook.active['C2'].value, 100)
                self.assertEqual(workbook.active['D2'].value, 5)
                self.assertEqual(workbook.active.max_column, 4)
                self.assertTrue(workbook.active.column_dimensions['B'].hidden)
                page.screenshot(path=str(output / f'm-in-f-{width}.png'), full_page=True)
                button.click()
                expect(page.get_by_role("button", name="Вернуть в M", exact=True)).to_have_count(0)
                self.assertEqual(trip.lines.get(supplier=m).quantity, 5)
                self.assertEqual(trip.lines.get(supplier=f).quantity, 0)
                self.assertEqual((order.items.get().prepayment_quantity, order.items.get().postpayment_quantity), (2, 3))
                page.screenshot(path=str(output / f'returned-to-m-{width}.png'), full_page=True)
                context.close()
            browser.close()

    @patch.dict(os.environ, {"DJANGO_ALLOW_ASYNC_UNSAFE": "true"})
    def test_create_manual_distribution_transfer_zero_restore_and_download(self):
        user = User.objects.create_superuser("trip-browser", password="StrongAdmin!123")
        self.client.force_login(user)
        suppliers = [Supplier.objects.create(name=f"Поставщик {letter}", letter=letter, priority=index,
                     highlight_color="#37A7BA", legal_entity=letter, phone_1="1")
                     for index, letter in enumerate("ABC", 1)]
        cd = CD.objects.create(platform=Platform.objects.create(name="PS5"), name="Игра рейса", sku="000123")
        for supplier, price in zip(suppliers, (100, 110, 120)):
            SupplierCDPrice.objects.create(supplier=supplier, cd=cd, price=price)
        price = create_procurement_price_list(actor=user, exchange_rate_usdt_aed=4, exchange_rate_usdt_rub=80,
                    delivery_general_usdt_per_kg=0, delivery_devices_usdt_per_kg=0,
                    delivery_storage_usdt_per_kg=0, delivery_cd_rub=0)
        source = price.items.get()
        output = Path(settings.BASE_DIR) / "outputs/trips-browser"
        output.mkdir(parents=True, exist_ok=True)
        with sync_playwright() as pw:
            browser = pw.chromium.launch()
            for width in (1280, 390):
                orders = [create_customer_procurement_order(actor=user, price_list_id=price.pk,
                            recipient=name, comment="", lines=[{"price_list_item_id": source.pk,
                            "prepayment_quantity": pre, "postpayment_quantity": 5-pre}])
                          for name, pre in ((f"Первый {width}", 2), (f"Второй {width}", 1))]
                context = browser.new_context(viewport={"width": width, "height": 900}, accept_downloads=True)
                context.add_cookies([{"name": settings.SESSION_COOKIE_NAME,
                                     "value": self.client.cookies[settings.SESSION_COOKIE_NAME].value,
                                     "url": self.live_server_url}])
                context.route("**/crm/api/exchange-rates/", lambda route: route.fulfill(json={"rates": {}, "available": False}))
                page = context.new_page()
                errors = []
                page.on("pageerror", lambda error: errors.append(str(error)))
                page.goto(self.live_server_url + reverse("orders:list"))
                for order in orders:
                    page.locator(f'[name="order_id"][value="{order.pk}"]').check()
                page.locator('#trip-name').fill(f"Браузер {width}")
                page.get_by_role("button", name="Объединить в рейс", exact=True).click()
                expect(page.get_by_role("heading", name=f"Браузер {width}", exact=False)).to_be_visible()
                trip = SupplierOrderBatch.objects.get(name=f"Браузер {width}")
                line = trip.lines.get()
                source_row = page.locator(f'[data-trip-line="{line.pk}"]')
                source_row.locator('[name="quantity"]').fill("7")
                source_row.get_by_role("button", name="Изменить", exact=True).click()
                expect(page.locator('#trip-target-supplier')).to_have_value(str(suppliers[2].pk))
                expect(page.locator('#trip-price-supplier')).to_have_value(str(suppliers[0].pk))
                page.locator('#trip-action').select_option('adjust')
                inputs = page.locator('#trip-customer-rows input')
                expect(inputs).to_have_count(2)
                page.locator('#trip-change-save').click()
                expect(page.locator('#trip-change-error')).to_contain_text("Распределите количество")
                inputs.nth(0).fill("3")
                inputs.nth(1).fill("4")
                page.screenshot(path=str(output / f'distribution-{width}.png'))
                page.locator('#trip-change-save').click()
                expect(page.locator('#trip-change-dialog')).to_have_count(1)
                expect(page.locator('#trip-change-dialog')).not_to_be_visible()
                line.refresh_from_db()
                self.assertEqual(line.quantity, 7)
                self.assertEqual(orders[0].items.get().prepayment_quantity + orders[0].items.get().postpayment_quantity, 3)
                source_row = page.locator(f'[data-trip-line="{line.pk}"]')
                source_row.locator('[name="quantity"]').fill("5")
                source_row.get_by_role("button", name="Изменить", exact=True).click()
                page.locator('#trip-price-supplier').select_option(str(suppliers[1].pk))
                page.locator('#trip-change-save').click()
                expect(page.locator('#trip-change-dialog')).not_to_be_visible()
                target = trip.lines.get(supplier=suppliers[2])
                self.assertEqual(target.quantity, 2)
                self.assertEqual(target.supplier_price_aed_snapshot, 110)
                target_row = page.locator(f'[data-trip-line="{target.pk}"]')
                target_row.locator('[name="quantity"]').fill("0")
                target_row.get_by_role("button", name="Изменить", exact=True).click()
                page.locator('#trip-action').select_option('adjust')
                page.locator('#trip-customer-rows input').fill("2")
                page.locator('#trip-change-save').click()
                expect(page.locator('#trip-change-dialog')).not_to_be_visible()
                expect(page.locator(f'[data-trip-line="{target.pk}"]')).to_have_class('trip-zero-row')
                target_row = page.locator(f'[data-trip-line="{target.pk}"]')
                target_row.locator('[name="quantity"]').fill("2")
                target_row.get_by_role("button", name="Изменить", exact=True).click()
                expect(page.locator('#trip-allocation-fields')).to_be_visible()
                page.locator('#trip-customer-rows input').fill("4")
                page.locator('#trip-change-save').click()
                expect(page.locator('#trip-change-dialog')).not_to_be_visible()
                target.refresh_from_db()
                self.assertEqual(target.quantity, 2)
                link = page.locator(f'a[href="{reverse("orders:trip_export", args=(trip.pk, suppliers[2].pk))}"]')
                with page.expect_download() as download:
                    link.click()
                workbook = load_workbook(BytesIO(Path(download.value.path()).read_bytes()))
                self.assertEqual(workbook.active['C2'].value, 110)
                self.assertEqual(workbook.active['D2'].value, 2)
                page.screenshot(path=str(output / f'trip-{width}.png'), full_page=True)
                self.assertEqual(errors, [])
                context.close()
            browser.close()
