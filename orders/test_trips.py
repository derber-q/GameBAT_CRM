from decimal import Decimal
from io import BytesIO
from zipfile import ZipFile
from uuid import uuid4
from types import SimpleNamespace
from importlib import import_module
from unittest.mock import patch

from django.contrib.auth.models import Permission
from django.apps import apps
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import User
from cash.models import CashTransaction
from catalog.models import Brand, CD, Platform, ProductType, Tech
from partners.models import Supplier
from price.models import ProcurementPriceList, ProcurementPriceListItem, ProcurementSupplierOffer
from price.services import create_procurement_price_list
from pricing.models import SupplierCDPrice, SupplierTechPrice
from .models import CustomerProcurementOrder, OrderAdjustment, SupplierOrderBatch, SupplierTripAllocation, SupplierTripChange
from .services import create_customer_procurement_order, edit_created_order, advance_order_status, reduce_receiving_order
from .trips import MANAGE_PERMISSIONS, create_trip, change_trip_line, apply_m_to_f_rule, return_line_to_m


class TripTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_superuser("trip-admin", password="StrongAdmin!123")
        cls.a, cls.b, cls.c = [Supplier.objects.create(name=f"Поставщик {letter}", letter=letter,
            highlight_color="#37A7BA", legal_entity=letter, phone_1="1", priority=index)
            for index, letter in enumerate("ABC", 1)]
        cls.cd = CD.objects.create(platform=Platform.objects.create(name="PS5"), name="Игра", sku="000123")
        cls.tech = Tech.objects.create(brand=Brand.objects.create(name="Sony"), product_type=ProductType.objects.create(name="Консоли"), name="Консоль", sku="T-1", weight_grams=1000)
        for supplier, cd_price, tech_price in ((cls.a, 100, 150), (cls.b, 110, 140), (cls.c, 120, 160)):
            SupplierCDPrice.objects.create(supplier=supplier, cd=cls.cd, price=cd_price)
            SupplierTechPrice.objects.create(supplier=supplier, tech=cls.tech, price=tech_price)
        cls.price = create_procurement_price_list(actor=cls.user, exchange_rate_usdt_aed=4,
            exchange_rate_usdt_rub=80, delivery_general_usdt_per_kg=0, delivery_devices_usdt_per_kg=0,
            delivery_storage_usdt_per_kg=0, delivery_cd_rub=0)
        cls.cd_source = cls.price.items.get(product_kind="cd")
        cls.tech_source = cls.price.items.get(product_kind="tech")
        cls.first = cls.order("Первый заказчик", 2, 3, tech=1)
        cls.second = cls.order("Второй заказчик", 1, 4, tech=2)

    @classmethod
    def order(cls, recipient, pre, post, tech=0):
        lines = [{"price_list_item_id": cls.cd_source.pk, "prepayment_quantity": pre, "postpayment_quantity": post}]
        if tech:
            lines.append({"price_list_item_id": cls.tech_source.pk, "prepayment_quantity": 0, "postpayment_quantity": tech})
        return create_customer_procurement_order(actor=cls.user, price_list_id=cls.price.pk, recipient=recipient, comment="", lines=lines)

    def setUp(self):
        self.client.force_login(self.user)

    def trip(self):
        return create_trip(actor=self.user, order_ids=[self.first.pk, self.second.pk], creation_key=uuid4())

    def change(self, trip, line, quantity, action="transfer", **kwargs):
        trip.refresh_from_db()
        return change_trip_line(actor=self.user, batch_id=trip.pk, line_id=line.pk, revision=trip.revision,
                                quantity=quantity, action=action, **kwargs)

    def distribution(self, line, **quantities):
        return [{"allocation_id": allocation.pk, "final_quantity": quantities[str(allocation.order_item.order_id)]}
                for allocation in line.allocations.select_related("order_item").all()]

    def test_new_price_preserves_all_supplier_prices_and_later_import_cannot_change_them(self):
        self.assertTrue(self.price.supplier_offers_complete)
        self.assertEqual(self.cd_source.supplier_offers.count(), 3)
        offer = self.cd_source.supplier_offers.get(supplier=self.c)
        SupplierCDPrice.objects.filter(cd=self.cd, supplier=self.c).update(price=1)
        offer.refresh_from_db()
        self.assertEqual(offer.price_aed, Decimal("120"))

    def test_old_price_backfill_uses_only_saved_selected_price_and_is_idempotent(self):
        legacy = ProcurementPriceList.objects.create(exchange_rate_aed_rub=20, created_by=self.user)
        original = ProcurementPriceListItem.objects.get(pk=self.cd_source.pk)
        original.pk = None
        original.price_list = legacy
        original.save()
        SupplierCDPrice.objects.filter(cd=self.cd).update(price=999)
        migration = import_module("price.migrations.0009_saved_selected_supplier_offers")
        migration.preserve_selected_offers(apps, SimpleNamespace(connection=connection))
        migration.preserve_selected_offers(apps, SimpleNamespace(connection=connection))
        self.assertFalse(legacy.supplier_offers_complete)
        self.assertEqual(list(original.supplier_offers.values_list("supplier_id", "price_aed")), [(self.a.pk, Decimal("100"))])

    def test_creation_aggregates_cd_tech_and_tracks_every_customer(self):
        trip = self.trip()
        self.assertEqual(trip.price_list_id, self.price.pk)
        self.assertEqual(trip.lines.get(cd=self.cd).quantity, 10)
        self.assertEqual(trip.lines.get(tech=self.tech).quantity, 3)
        self.assertEqual(SupplierTripAllocation.objects.filter(line__batch=trip).count(), 4)
        self.assertEqual(set(CustomerProcurementOrder.objects.filter(trip=trip).values_list("pk", flat=True)), {self.first.pk, self.second.pk})
        self.assertEqual(trip.changes.count(), 1)
        self.assertEqual(CashTransaction.objects.count(), 0)

    def test_creation_is_idempotent_and_order_cannot_belong_to_two_trips(self):
        key = uuid4()
        first = create_trip(actor=self.user, order_ids=[self.first.pk], creation_key=key)
        duplicate = create_trip(actor=self.user, order_ids=[self.first.pk], creation_key=key)
        self.assertEqual(first.pk, duplicate.pk)
        with self.assertRaisesMessage(ValidationError, "уже включён"):
            create_trip(actor=self.user, order_ids=[self.first.pk], creation_key=uuid4())
        with self.assertRaisesMessage(ValidationError, "другой рейс"):
            create_trip(actor=self.user, order_ids=[self.second.pk], creation_key=key)

    def test_different_versions_and_non_created_orders_rejected(self):
        other_price = ProcurementPriceList.objects.create(exchange_rate_aed_rub=20, created_by=self.user)
        self.second.price_list = other_price
        self.second.save(update_fields=("price_list",))
        with self.assertRaisesMessage(ValidationError, "одной версии"):
            self.trip()
        self.assertFalse(SupplierOrderBatch.objects.exists())
        self.second.price_list = self.price
        self.second.status = "confirmed"
        self.second.save(update_fields=("price_list", "status"))
        with self.assertRaisesMessage(ValidationError, "«Создан»"):
            self.trip()

    def test_transfer_preserves_customer_quantities_prices_and_zero_source_row(self):
        trip = self.trip()
        source = trip.lines.get(cd=self.cd)
        before = list(CustomerProcurementOrder.objects.values("pk", "prepayment_total", "postpayment_total"))
        self.change(trip, source, 0, target_supplier_id=self.c.pk, price_supplier_id=self.a.pk)
        source.refresh_from_db()
        target = trip.lines.get(cd=self.cd, supplier=self.c)
        self.assertEqual(source.quantity, 0)
        self.assertEqual(target.quantity, 10)
        self.assertEqual(target.supplier_price_aed_snapshot, 100)
        self.assertEqual(target.price_supplier, self.a)
        self.assertEqual(list(CustomerProcurementOrder.objects.values("pk", "prepayment_total", "postpayment_total")), before)
        self.assertFalse(OrderAdjustment.objects.exists())
        self.assertFalse(CashTransaction.objects.exists())

    def test_different_price_sources_remain_separate_lines(self):
        trip = self.trip()
        source = trip.lines.get(cd=self.cd)
        self.change(trip, source, 8, target_supplier_id=self.c.pk, price_supplier_id=self.a.pk)
        self.change(trip, source, 5, target_supplier_id=self.c.pk, price_supplier_id=self.b.pk)
        target = trip.lines.filter(cd=self.cd, supplier=self.c)
        self.assertEqual(target.count(), 2)
        self.assertCountEqual(target.values_list("quantity", "supplier_price_aed_snapshot"), [(2, Decimal("100")), (3, Decimal("110"))])

    def test_customer_chooses_reduction_and_totals_recalculate(self):
        trip = self.trip()
        source = trip.lines.get(cd=self.cd)
        self.change(trip, source, 7, action="adjust", distribution=self.distribution(source, **{str(self.first.pk): 2, str(self.second.pk): 5}))
        first_item = self.first.items.get(cd=self.cd)
        second_item = self.second.items.get(cd=self.cd)
        self.assertEqual((first_item.prepayment_quantity, first_item.postpayment_quantity), (2, 0))
        self.assertEqual((second_item.prepayment_quantity, second_item.postpayment_quantity), (1, 4))
        self.first.refresh_from_db()
        self.assertEqual(self.first.prepayment_total, first_item.prepayment_line_total)
        self.assertEqual(self.first.postpayment_total, self.first.items.get(tech=self.tech).postpayment_line_total)
        self.assertEqual(OrderAdjustment.objects.filter(order=self.first).count(), 1)

    def test_zero_client_and_supplier_rows_can_restore_original_payment_mix(self):
        trip = self.trip()
        source = trip.lines.get(cd=self.cd)
        distribution = self.distribution(source, **{str(self.first.pk): 0, str(self.second.pk): 0})
        self.change(trip, source, 0, action="adjust", distribution=distribution)
        self.assertEqual(self.first.items.get(cd=self.cd).line_total, 0)
        source.refresh_from_db()
        self.assertEqual(source.quantity, 0)
        self.change(trip, source, 10, action="adjust", distribution=self.distribution(source, **{str(self.first.pk): 5, str(self.second.pk): 5}))
        item = self.first.items.get(cd=self.cd)
        self.assertEqual((item.prepayment_quantity, item.postpayment_quantity), (2, 3))

    def test_customer_final_quantity_includes_other_suppliers(self):
        trip = self.trip()
        source = trip.lines.get(cd=self.cd)
        self.change(trip, source, 5, target_supplier_id=self.c.pk, price_supplier_id=self.a.pk)
        self.change(trip, source, 4, action="adjust", distribution=self.distribution(source, **{str(self.first.pk): 4, str(self.second.pk): 5}))
        self.assertEqual(sum(a.quantity for a in SupplierTripAllocation.objects.filter(order_item=self.first.items.get(cd=self.cd))), 4)
        self.assertEqual(self.second.items.get(cd=self.cd).postpayment_quantity, 4)

    def test_wrong_distribution_rolls_back_quantities_and_audit(self):
        trip = self.trip()
        source = trip.lines.get(cd=self.cd)
        before = list(self.first.items.values())
        with self.assertRaisesMessage(ValidationError, "Сумма долей"):
            self.change(trip, source, 3, action="adjust", distribution=self.distribution(source, **{str(self.first.pk): 2, str(self.second.pk): 2}))
        self.assertEqual(list(self.first.items.values()), before)
        self.assertFalse(OrderAdjustment.objects.exists())
        self.assertEqual(trip.changes.count(), 1)
        source.refresh_from_db()
        self.assertEqual(source.quantity, 10)

    def test_missing_or_foreign_customer_and_negative_quantity_rejected(self):
        trip = self.trip()
        source = trip.lines.get(cd=self.cd)
        for distribution in ([], [{"allocation_id": 999999, "final_quantity": 0}]):
            with self.subTest(distribution=distribution), self.assertRaises(ValidationError):
                self.change(trip, source, 0, action="adjust", distribution=distribution)
        with self.assertRaises(ValidationError):
            self.change(trip, source, -1)

    def test_invalid_target_self_transfer_and_price_from_other_product_rejected(self):
        trip = self.trip()
        source = trip.lines.get(cd=self.cd)
        for target, price in ((self.a.pk, self.a.pk), (999999, self.a.pk), (self.c.pk, 999999)):
            with self.subTest(target=target, price=price), self.assertRaises(ValidationError):
                self.change(trip, source, 9, target_supplier_id=target, price_supplier_id=price)
        self.cd_source.supplier_offers.filter(supplier=self.c).delete()
        with self.assertRaisesMessage(ValidationError, "сохранённого прайса"):
            self.change(trip, source, 9, target_supplier_id=self.b.pk, price_supplier_id=self.c.pk)

    def test_stale_revision_prevents_double_transfer(self):
        trip = self.trip()
        source = trip.lines.get(cd=self.cd)
        self.change(trip, source, 8, target_supplier_id=self.c.pk, price_supplier_id=self.a.pk)
        with self.assertRaisesMessage(ValidationError, "уже изменён"):
            change_trip_line(actor=self.user, batch_id=trip.pk, line_id=source.pk, revision=0, quantity=8,
                             action="transfer", target_supplier_id=self.c.pk, price_supplier_id=self.a.pk)
        self.assertEqual(trip.lines.get(supplier=self.c).quantity, 2)

    def test_confirmed_trip_is_read_only_and_direct_order_edit_is_blocked(self):
        trip = self.trip()
        with self.assertRaisesMessage(ValidationError, "через таблицы рейса"):
            edit_created_order(actor=self.user, order_id=self.first.pk, recipient="Изменено", comment="", lines=[{
                "price_list_item_id": self.cd_source.pk, "prepayment_quantity": 1, "postpayment_quantity": 0}])
        self.first.status = "confirmed"
        self.first.save(update_fields=("status",))
        with self.assertRaisesMessage(ValidationError, "зафиксирован"):
            self.change(trip, trip.lines.get(cd=self.cd), 9)
        response = self.client.get(reverse("orders:trip_detail", args=(trip.pk,)))
        self.assertNotContains(response, 'id="trip-editor-data"')

    def test_priority_and_price_defaults_in_dialog_data(self):
        trip = self.trip()
        response = self.client.get(reverse("orders:trip_detail", args=(trip.pk,)))
        self.assertEqual(response.status_code, 200)
        line = trip.lines.get(cd=self.cd)
        config = response.context["editor"][str(line.pk)]
        self.assertEqual(config["targets"][0]["id"], self.c.pk)
        self.assertEqual(config["price_supplier_id"], self.a.pk)
        self.assertEqual(len(config["allocations"]), 2)
        self.assertContains(response, "Скачать Excel", count=2)
        self.assertContains(response, "Получит всего")

    def test_supplier_excel_excludes_zero_rows_and_keeps_empty_workbook_valid(self):
        trip = self.trip()
        source = trip.lines.get(cd=self.cd)
        self.change(trip, source, 0, target_supplier_id=self.c.pk, price_supplier_id=self.a.pk)
        with CaptureQueriesContext(connection) as queries:
            response = self.client.get(reverse("orders:trip_export", args=(trip.pk, self.a.pk)))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response["Content-Disposition"].endswith('.xlsx"'))
        self.assertEqual(response["Content-Disposition"], f'attachment; filename="trip_{trip.pk}_A.xlsx"')
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(tuple(cell.value for cell in sheet[1]), ("Наименование", "CUSA / PPSA", "Артикул", "Цена, AED", "Количество"))
        self.assertTrue(sheet.column_dimensions["C"].hidden)
        self.assertEqual(sheet.max_row, 1)
        self.assertEqual(sheet.auto_filter.ref, "A1:E1")
        self.assertFalse(sheet.merged_cells.ranges)
        self.assertEqual(trip.lines.get(pk=source.pk).quantity, 0)
        self.assertContains(self.client.get(reverse('orders:trip_detail', args=(trip.pk,))),
                            f'data-trip-line="{source.pk}"')
        self.assertFalse(any('INSERT ' in q['sql'] or 'UPDATE ' in q['sql'] for q in queries))

    def test_supplier_excel_keeps_only_positive_lines_for_every_supplier_with_or_without_prices(self):
        trip = self.trip()
        self.change(trip, trip.lines.get(tech=self.tech), 0,
                    target_supplier_id=self.a.pk, price_supplier_id=self.b.pk)
        self.change(trip, trip.lines.get(cd=self.cd), 0,
                    target_supplier_id=self.c.pk, price_supplier_id=self.a.pk)
        from .trip_excel import supplier_order_xlsx
        before = list(trip.lines.values())
        for supplier in (self.a, self.b, self.c):
            lines = list(trip.lines.filter(supplier=supplier).select_related(
                'cd__platform', 'tech__brand', 'tech__product_type'))
            for show_prices in (True, False):
                with self.subTest(supplier=supplier.letter, show_prices=show_prices):
                    book = load_workbook(BytesIO(supplier_order_xlsx(trip, supplier, lines, show_prices=show_prices)))
                    sheet = book.active
                    data = [row for row in sheet.values if isinstance(row[-1], int)]
                    expected = [line for line in lines if line.quantity > 0]
                    self.assertCountEqual([(row[0], row[-1]) for row in data],
                        [(line.product_name_snapshot, line.quantity) for line in expected])
                    if show_prices:
                        self.assertCountEqual([(row[0], row[3]) for row in data],
                            [(line.product_name_snapshot, line.supplier_price_aed_snapshot) for line in expected])
                    self.assertTrue(sheet.column_dimensions['C'].hidden)
                    book.close()
        self.assertEqual(list(trip.lines.values()), before)

    def test_permission_checks_apply_to_services_and_direct_urls(self):
        trip = self.trip()
        viewer = User.objects.create_user("trip-viewer")
        self.client.force_login(viewer)
        for url in (reverse("orders:trip_detail", args=(trip.pk,)), reverse("orders:trip_export", args=(trip.pk, self.a.pk))):
            self.assertEqual(self.client.get(url).status_code, 403)
        with self.assertRaises(PermissionDenied):
            create_trip(actor=viewer, order_ids=[self.first.pk], creation_key=uuid4())
        for permission in ("orders.generate_supplier_orders", "partners.view_supplier_details"):
            app, codename = permission.split('.')
            viewer.user_permissions.add(Permission.objects.get(content_type__app_label=app, codename=codename))
        self.client.force_login(viewer)
        page = self.client.get(reverse("orders:trip_detail", args=(trip.pk,)))
        self.assertNotContains(page, 'id="trip-editor-data"')
        workbook = load_workbook(BytesIO(self.client.get(reverse("orders:trip_export", args=(trip.pk, self.a.pk))).content))
        self.assertEqual(workbook.active.max_column, 4)
        self.assertTrue(workbook.active.column_dimensions["C"].hidden)
        self.assertEqual(tuple(cell.value for cell in workbook.active[1]), ("Наименование", "CUSA / PPSA", "Артикул", "Количество"))
        self.assertEqual(self.client.post(reverse("orders:trip_line_update", args=(trip.pk, trip.lines.get(cd=self.cd).pk))).status_code, 403)

    def test_trip_creation_form_and_backend_methods(self):
        page = self.client.get(reverse("orders:list"))
        self.assertContains(page, "Объединить в рейс")
        response = self.client.post(reverse("orders:trip_create"), {"order_id": [self.first.pk, self.second.pk], "creation_key": str(uuid4()), "name": "Октябрь"})
        trip = SupplierOrderBatch.objects.get()
        self.assertRedirects(response, reverse("orders:trip_detail", args=(trip.pk,)))
        self.assertEqual(trip.name, "Октябрь")
        self.assertEqual(self.client.get(reverse("orders:trip_create")).status_code, 405)

    def test_http_line_update_rejects_bad_distribution_and_accepts_valid_manual_choice(self):
        trip = self.trip()
        line = trip.lines.get(cd=self.cd)
        url = reverse("orders:trip_line_update", args=(trip.pk, line.pk))
        self.assertEqual(self.client.get(url).status_code, 405)
        self.assertEqual(self.client.post(url, {"distribution": "not-json"}).status_code, 409)
        import json
        response = self.client.post(url, {"revision": trip.revision, "quantity": 8, "action": "adjust",
                  "distribution": json.dumps(self.distribution(line, **{str(self.first.pk): 3, str(self.second.pk): 5}))})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["success"])
        self.assertEqual(response.json()["revision"], 1)

    def test_receiving_reduction_keeps_linked_zero_customer_item(self):
        order = self.order("Приёмка", 0, 2)
        trip = create_trip(actor=self.user, order_ids=[order.pk], creation_key=uuid4())
        advance_order_status(actor=self.user, order_id=order.pk, next_status="confirmed")
        advance_order_status(actor=self.user, order_id=order.pk, next_status="receiving")
        item = order.items.get()
        reduce_receiving_order(actor=self.user, order_id=order.pk, comment="Товар не получен",
                               rows=[{"item_id": item.pk, "prepayment_quantity": 0, "postpayment_quantity": 0}])
        item.refresh_from_db()
        self.assertEqual(item.line_total, 0)
        self.assertEqual(trip.lines.get().quantity, 0)
        self.assertEqual(item.trip_allocations.get().quantity, 0)
        from .trips import _assert_balances
        _assert_balances(trip)


class MToFRuleTests(TestCase):
    order = classmethod(TripTests.order.__func__)
    trip = TripTests.trip

    @classmethod
    def setUpTestData(cls):
        TripTests.setUpTestData.__func__(cls)
        cls.a.letter, cls.a.name = "M", "M — исходный поставщик"
        cls.a.save(update_fields=("letter", "name"))
        cls.b.letter, cls.b.name = "F", "F — удобный поставщик"
        cls.b.save(update_fields=("letter", "name"))

    def setUp(self):
        self.client.force_login(self.user)

    def test_new_trip_assigns_m_to_f_at_m_price_without_changing_customers(self):
        before = list(CustomerProcurementOrder.objects.values("pk", "prepayment_total", "postpayment_total"))
        trip = self.trip()
        line = trip.lines.get(cd=self.cd, supplier=self.b)
        self.assertEqual((line.quantity, line.price_supplier_id, line.supplier_price_aed_snapshot), (10, self.a.pk, Decimal("100")))
        self.assertEqual(trip.lines.get(cd=self.cd, supplier=self.a).quantity, 0)
        self.assertEqual(trip.lines.get(tech=self.tech, supplier=self.b).quantity, 3)
        self.assertEqual(before, list(CustomerProcurementOrder.objects.values("pk", "prepayment_total", "postpayment_total")))
        self.assertFalse(CashTransaction.objects.exists())
        self.assertFalse(OrderAdjustment.objects.exists())

    def test_existing_trip_rule_is_atomic_and_idempotent(self):
        with patch("orders.trips.apply_m_to_f_rule"):
            trip = self.trip()
        self.assertEqual(trip.lines.get(cd=self.cd, supplier=self.a).quantity, 10)
        result = apply_m_to_f_rule(actor=self.user, batch_id=trip.pk)
        self.assertEqual(result, {"lines": 1, "units": 10})
        events = trip.changes.count()
        self.assertEqual(apply_m_to_f_rule(actor=self.user, batch_id=trip.pk), {"lines": 0, "units": 0})
        self.assertEqual(trip.changes.count(), events)
        self.assertEqual(trip.lines.get(cd=self.cd, supplier=self.b).quantity, 10)

    def test_button_returns_whole_line_to_m_at_same_price_and_keeps_customer_mix(self):
        trip = self.trip()
        line = trip.lines.get(cd=self.cd, supplier=self.b)
        page = self.client.get(reverse("orders:trip_detail", args=(trip.pk,)))
        self.assertContains(page, "Вернуть в M", count=1)
        quantities = list(self.first.items.values_list("prepayment_quantity", "postpayment_quantity"))
        url = reverse("orders:trip_line_update", args=(trip.pk, line.pk))
        payload = {"action": "return_to_m", "revision": trip.revision}
        response = self.client.post(url, payload)
        self.assertEqual(response.status_code, 200)
        line.refresh_from_db()
        returned = trip.lines.get(cd=self.cd, supplier=self.a)
        self.assertEqual((line.quantity, returned.quantity, returned.supplier_price_aed_snapshot), (0, 10, Decimal("100")))
        self.assertEqual(quantities, list(self.first.items.values_list("prepayment_quantity", "postpayment_quantity")))
        self.assertEqual(self.client.post(url, payload).status_code, 409)
        self.assertNotContains(self.client.get(reverse("orders:trip_detail", args=(trip.pk,))), "Вернуть в M")

    def test_return_rejects_native_f_row_and_unprivileged_user(self):
        trip = self.trip()
        native = trip.lines.get(tech=self.tech)
        with self.assertRaisesMessage(ValidationError, "только товар F"):
            return_line_to_m(actor=self.user, batch_id=trip.pk, line_id=native.pk, revision=trip.revision)
        viewer = User.objects.create_user("m-f-viewer")
        with self.assertRaises(PermissionDenied):
            apply_m_to_f_rule(actor=viewer, batch_id=trip.pk)

    def test_confirmed_trip_cannot_return_or_apply_rule(self):
        trip = self.trip()
        self.first.status = "confirmed"
        self.first.save(update_fields=("status",))
        line = trip.lines.get(cd=self.cd, supplier=self.b)
        with self.assertRaisesMessage(ValidationError, "зафиксирован"):
            return_line_to_m(actor=self.user, batch_id=trip.pk, line_id=line.pk, revision=trip.revision)

    def test_missing_f_does_not_leave_partially_created_trip(self):
        self.b.letter = "B"
        self.b.save(update_fields=("letter",))
        with self.assertRaisesMessage(ValidationError, "буквой F"):
            self.trip()
        self.assertFalse(SupplierOrderBatch.objects.exists())
        self.assertFalse(CustomerProcurementOrder.objects.filter(trip__isnull=False).exists())

    def test_missing_historical_m_price_rolls_back_existing_trip(self):
        with patch("orders.trips.apply_m_to_f_rule"):
            trip = self.trip()
        self.cd_source.supplier_offers.filter(supplier=self.a).delete()
        with self.assertRaisesMessage(ValidationError, "сохранённого прайса"):
            apply_m_to_f_rule(actor=self.user, batch_id=trip.pk)
        self.assertEqual(trip.lines.get(cd=self.cd, supplier=self.a).quantity, 10)
        self.assertFalse(trip.lines.filter(cd=self.cd, supplier=self.b).exists())

    def test_f_excel_contains_m_products_and_price_with_no_source_information(self):
        trip = self.trip()
        response = self.client.get(reverse("orders:trip_export", args=(trip.pk, self.b.pk)))
        book = load_workbook(BytesIO(response.content))
        sheet = book.active
        self.assertEqual(sheet.max_column, 5)
        self.assertEqual(response["Content-Disposition"], f'attachment; filename="trip_{trip.pk}_F.xlsx"')
        products = {row[2]: row for row in sheet.iter_rows(values_only=True) if row[2] in (self.cd.sku, self.tech.sku)}
        self.assertEqual(products[self.cd.sku], (self.cd.name, None, self.cd.sku, 100, 10))
        self.assertEqual(products[self.tech.sku][4], 3)
        self.assertEqual(book.sheetnames, ["Заказ поставщику"])
        self.assertEqual(sheet.sheet_state, "visible")
        self.assertFalse(book.defined_names)
        with ZipFile(BytesIO(response.content)) as archive:
            payload = "\n".join(archive.read(name).decode("utf-8") for name in archive.namelist())
        for forbidden in (self.a.name, "Источник цены", "Цена из прайса поставщика", "default_m_to_f", "Перенесено"):
            self.assertNotIn(forbidden, payload)
