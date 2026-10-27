"""Общий состав рейса и заказ поставщику из сохранённых строк."""
from dataclasses import dataclass, field
from decimal import Decimal
from io import BytesIO
from textwrap import wrap

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from catalog.nomenclature_grouping import nomenclature_product_groups
from catalog.game_codes import display_game_codes
from price.excel import LIGHT, TEAL, WHITE, safe_text


@dataclass
class TripProductRow:
    product_kind: str
    product: object
    product_name_snapshot: str
    article_snapshot: str
    supplier_price_aed_snapshot: Decimal
    quantity: int = 0
    suppliers: dict = field(default_factory=dict)

    @property
    def supplier_names(self):
        return ", ".join(supplier.name for supplier in sorted(
            self.suppliers.values(), key=lambda supplier: (supplier.letter.casefold(), supplier.pk),
        )) or "—"


def trip_order_xlsx(lines):
    """Объединяет одинаковые товары с одной ценой независимо от поставщика."""
    rows = {}
    for line in lines:
        key = (line.product_kind, line.product_id, line.product_name_snapshot,
               line.article_snapshot, line.supplier_price_aed_snapshot)
        if key not in rows:
            rows[key] = TripProductRow(
                line.product_kind, line.product, line.product_name_snapshot,
                line.article_snapshot, line.supplier_price_aed_snapshot,
            )
        rows[key].quantity += line.quantity
        if line.quantity:
            rows[key].suppliers[line.supplier_id] = line.supplier
    # Общая сортировка по товарам стабильна; разные цены одного товара идут по возрастанию.
    rows = sorted(rows.values(), key=lambda row: row.supplier_price_aed_snapshot)
    return _trip_lines_xlsx(rows, title="Товары рейса", show_prices=True, show_suppliers=True, show_total=True)


def supplier_order_xlsx(batch, supplier, lines, *, show_prices=True):
    return _trip_lines_xlsx((line for line in lines if line.quantity > 0),
                            title="Заказ поставщику", show_prices=show_prices, show_codes=True)


def _trip_lines_xlsx(lines, *, title, show_prices, show_suppliers=False, show_codes=False, show_total=False):
    book = Workbook()
    sheet = book.active
    sheet.title = title
    headers = ["Наименование"] + (["CUSA / PPSA"] if show_codes else []) + ["Артикул"]
    article_column = 3 if show_codes else 2
    price_column = article_column + 1
    if show_prices:
        headers.append("Цена, AED")
    headers.append("Количество")
    quantity_column = len(headers)
    if show_suppliers:
        headers.append("Поставщики")
    last = get_column_letter(len(headers))
    for column, title in enumerate(headers, 1):
        cell = sheet.cell(1, column, title)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 30
    lines = list(lines)
    groups = nomenclature_product_groups(
        [line for line in lines if line.product_kind == "tech"],
        [line for line in lines if line.product_kind == "cd"], product_of=lambda line: line.product,
    )
    previous_group = None
    for title, line in ((title, line) for _kind, title, group in groups for line in group):
        if title != previous_group:
            sheet.append([safe_text(title)])
            row = sheet.max_row
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
            cell = sheet.cell(row, 1)
            cell.font = Font(bold=True, color="334660")
            cell.fill = PatternFill("solid", fgColor=LIGHT)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[row].height = max(28, 15 * len(wrap(title, 80)))
            previous_group = title
        values = [safe_text(line.product_name_snapshot)]
        if show_codes:
            values.append(safe_text(display_game_codes(line.product)))
        values.append(safe_text(line.article_snapshot))
        if show_prices:
            values.append(line.supplier_price_aed_snapshot)
        values.append(line.quantity)
        if show_suppliers:
            values.append(safe_text(line.supplier_names))
        sheet.append(values)
        row = sheet.max_row
        sheet.cell(row, article_column).number_format = "@"
        if show_codes:
            sheet.cell(row, 2).number_format = '@'
        sheet.cell(row, quantity_column).number_format = "0"
        for cell in sheet[row]:
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if not line.quantity:
                cell.fill = PatternFill("solid", fgColor="FFF1D6")
        if show_prices:
            price = line.supplier_price_aed_snapshot
            sheet.cell(row, price_column).number_format = (
                "#,##0" if price == price.to_integral_value() else "#,##0.######"
            )
        sheet.row_dimensions[row].height = max(24, 15 * len(wrap(line.product_name_snapshot, 52)))
        if show_suppliers:
            sheet.row_dimensions[row].height = max(
                sheet.row_dimensions[row].height, 15 * len(wrap(line.supplier_names, 36)),
            )
    last_item_row = sheet.max_row
    if show_total and show_prices:
        total = sum((line.supplier_price_aed_snapshot * line.quantity for line in lines), Decimal('0'))
        row = sheet.max_row + 1
        sheet.cell(row, 1, 'Итого, AED')
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=article_column)
        sheet.cell(row, price_column, total)
        sheet.merge_cells(start_row=row, start_column=price_column, end_row=row, end_column=len(headers))
        for cell in sheet[row]:
            cell.font = Font(bold=True, color='334660')
            cell.fill = PatternFill('solid', fgColor=LIGHT)
            cell.alignment = Alignment(horizontal='left' if cell.column == 1 else 'right', vertical='center')
        sheet.cell(row, price_column).number_format = '#,##0.00####'
        sheet.row_dimensions[row].height = 30
    widths = ([60, 36, 24] if show_codes else [60, 24]) + ([20] if show_prices else []) + [14] + ([42] if show_suppliers else [])
    for letter, width in zip('ABCDEF', widths):
        if ord(letter) - 64 <= len(headers):
            sheet.column_dimensions[letter].width = width
    sheet.column_dimensions[get_column_letter(article_column)].hidden = True
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{last}{last_item_row}"
    sheet.print_title_rows = "1:1"
    sheet.print_area = f"A1:{last}{sheet.max_row}"
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()
