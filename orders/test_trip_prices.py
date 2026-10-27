"""Цена закупки рейса меняет расходы магазина, сохраняя цены покупателей."""
from decimal import Decimal
from copy import copy
from io import BytesIO
from unittest.mock import patch
from uuid import uuid4

from django.core.exceptions import PermissionDenied, ValidationError
from django.contrib.auth.models import Permission
from django.db import OperationalError
from django.test import Client, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import User
from catalog.models import CD, Platform
from partners.models import Supplier
from price.services import create_procurement_price_list
from pricing.models import SupplierCDPrice
from .models import CustomerProcurementOrder, SupplierTripChange
from .services import create_customer_procurement_order
from .trips import create_trip


class TripPurchasePriceTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_superuser('trip-price-admin')
        cls.supplier = Supplier.objects.create(name='Поставщик B', letter='B', priority=1,
            legal_entity='B', phone_1='1', highlight_color='#37A7BA')
        cls.other_supplier = Supplier.objects.create(name='Поставщик A', letter='A', priority=2,
            legal_entity='A', phone_1='2', highlight_color='#37A7BA')
        cls.product = CD.objects.create(cusa_ppsa_code="PPSA-00001", platform=Platform.objects.create(name='PS5'), name='Товар X', sku='000001')
        SupplierCDPrice.objects.create(supplier=cls.supplier, cd=cls.product, price=100)
        SupplierCDPrice.objects.create(supplier=cls.other_supplier, cd=cls.product, price=105)
        cls.price = create_procurement_price_list(actor=cls.user,
            exchange_rate_usdt_aed=4, exchange_rate_usdt_rub=100,
            delivery_general_usdt_per_kg=0, delivery_devices_usdt_per_kg=0,
            delivery_storage_usdt_per_kg=0, delivery_cd_rub=17)
        cls.source = cls.price.items.get()
        cls.source.prepayment_price_rub = Decimal('3500')
        cls.source.postpayment_price_rub = Decimal('3600')
        cls.source.save(update_fields=('prepayment_price_rub', 'postpayment_price_rub'))
        cls.first = cls.order('Первый покупатель', 4, 0)
        cls.second = cls.order('Второй покупатель', 0, 3)
        cls.stock = cls.order('На склад', 3, 0, purpose='stock')
        cls.trip = create_trip(actor=cls.user, order_ids=[cls.first.pk, cls.second.pk, cls.stock.pk], creation_key=uuid4())
        cls.line = cls.trip.lines.get(quantity=10)

    @classmethod
    def order(cls, recipient, pre, post, purpose='customer'):
        return create_customer_procurement_order(actor=cls.user, price_list_id=cls.price.pk,
            recipient=recipient, comment='', purpose=purpose,
            lines=[{'price_list_item_id': cls.source.pk, 'prepayment_quantity': pre, 'postpayment_quantity': post}])

    def test_preview_uses_fixed_rate_delivery_and_does_not_save(self):
        from .trip_prices import preview_trip_price
        before = self.line.supplier_price_aed_snapshot
        result = preview_trip_price(actor=self.user, batch_id=self.trip.pk, line_id=self.line.pk,
            revision=self.trip.revision, price_aed='110')
        self.line.refresh_from_db()
        self.assertEqual(self.line.supplier_price_aed_snapshot, before)
        self.assertEqual(result['quantity'], 10)
        self.assertEqual(Decimal(result['cost_delta_rub']), Decimal('2500'))
        self.assertEqual(Decimal(result['customer_profit_delta_rub']), Decimal('-1750'))
        self.assertEqual(Decimal(result['stock_cost_delta_rub']), Decimal('750'))
        self.assertEqual(len(result['allocations']), 3)
        self.assertTrue(all(Decimal(row['rate']) == 25 for row in result['allocations']))
        self.assertTrue(all(Decimal(row['new_unit_cost_rub']) == 2767 for row in result['allocations']))

    def test_change_applies_to_all_units_and_preserves_customer_snapshots(self):
        from .trip_prices import change_trip_price, order_financials
        snapshots = list(CustomerProcurementOrder.objects.filter(trip=self.trip).values(
            'pk', 'grand_total', 'prepayment_total', 'postpayment_total'))
        items = list(self.first.items.values()) + list(self.second.items.values())
        allocations = list(self.line.allocations.values())
        audit_count = SupplierTripChange.objects.filter(batch=self.trip).count()
        change_trip_price(actor=self.user, batch_id=self.trip.pk, line_id=self.line.pk,
            revision=self.trip.revision, price_aed='110', reason='Новый инвойс поставщика')
        self.line.refresh_from_db()
        self.assertEqual(self.line.supplier_price_aed_snapshot, Decimal('110'))
        self.assertEqual(self.line.quantity, 10)
        self.assertEqual(self.line.supplier_id, self.supplier.pk)
        self.assertIsNotNone(self.line.supplier_price_updated_at)
        self.assertEqual(list(self.line.allocations.values()), allocations)
        self.assertEqual(list(CustomerProcurementOrder.objects.filter(trip=self.trip).values(
            'pk', 'grand_total', 'prepayment_total', 'postpayment_total')), snapshots)
        self.assertEqual(list(self.first.items.values()) + list(self.second.items.values()), items)
        self.assertEqual(order_financials(self.first)['cost_rub'], Decimal('11068'))
        self.assertEqual(order_financials(self.first)['profit_rub'], Decimal('2932'))
        self.assertEqual(order_financials(self.second)['profit_rub'], Decimal('2499'))
        self.source.refresh_from_db()
        self.assertEqual(self.source.supplier_price_aed, Decimal('100'))
        self.assertEqual(SupplierCDPrice.objects.get(supplier=self.supplier, cd=self.product).price, Decimal('100'))
        self.assertEqual(SupplierTripChange.objects.filter(batch=self.trip).count(), audit_count + 1)

    def change(self, price='110', **kwargs):
        from .trip_prices import change_trip_price
        return change_trip_price(actor=kwargs.pop('actor', self.user), batch_id=self.trip.pk,
            line_id=self.line.pk, revision=kwargs.pop('revision', self.trip.revision), price_aed=price, **kwargs)

    def test_stock_receipt_uses_new_base_and_saved_delivery_before_price_review(self):
        from warehouse.models import Warehouse, CDWarehouseStock
        from .services import advance_order_status
        from .fulfillment_test_helpers import receive_with_test_rates as receive_stock_order
        cost_before = self.product.cost
        self.change()
        warehouse = Warehouse.objects.create(name='Склад новой закупки')
        advance_order_status(actor=self.user, order_id=self.stock.pk, next_status='confirmed')
        advance_order_status(actor=self.user, order_id=self.stock.pk, next_status='receiving')
        result = receive_stock_order(actor=self.user, order_id=self.stock.pk, warehouse_id=warehouse.pk)
        self.assertEqual(result.supply.cd_items.get().purchase_unit_cost, Decimal('2767'))
        self.assertEqual(result.supply.status, result.supply.Status.PRICE_REVIEW_REQUIRED)
        self.assertFalse(CDWarehouseStock.objects.filter(cd=self.product).exists())
        self.product.refresh_from_db()
        self.assertEqual(self.product.cost, cost_before)

    def test_zero_negative_invalid_and_overprecise_aed_are_rejected(self):
        for value in ('0', '-1', 'NaN', 'Infinity', '-Infinity', 'abc', '0.0000001', '100000000000000', '99999999999999'):
            with self.subTest(value=value), self.assertRaises(ValidationError):
                self.change(value)
        self.line.refresh_from_db()
        self.assertEqual(self.line.supplier_price_aed_snapshot, Decimal('100'))

    def test_repeat_is_noop_and_stale_different_price_cannot_overwrite(self):
        changed = self.change()
        self.assertEqual(changed.revision, self.trip.revision + 1)
        count = SupplierTripChange.objects.count()
        self.change()
        self.assertEqual(SupplierTripChange.objects.count(), count)
        with self.assertRaisesMessage(ValidationError, 'другим запросом'):
            self.change('120')
        self.line.refresh_from_db()
        self.assertEqual(self.line.supplier_price_aed_snapshot, Decimal('110'))

    def test_decrease_and_loss_are_allowed_and_audit_keeps_all_revisions(self):
        from .trip_prices import order_financials
        lowered = self.change('90')
        self.assertEqual(order_financials(self.first)['cost_rub'], Decimal('9068'))
        self.change('200', revision=lowered.revision)
        self.assertEqual(order_financials(self.first)['profit_rub'], Decimal('-6068'))
        events = SupplierTripChange.objects.filter(batch=self.trip, details__action='purchase_price').order_by('revision')
        self.assertEqual([Decimal(row.details['new_price_aed']) for row in events], [Decimal('90'), Decimal('200')])

    def test_permission_is_checked_in_service(self):
        employee = User.objects.create_user('price-forbidden')
        with self.assertRaises(PermissionDenied):
            self.change(actor=employee)

    def test_saved_rates_are_used_per_order_without_averaging(self):
        from price.models import ProcurementPriceList
        from .trip_prices import preview_trip_price
        other_price = ProcurementPriceList.objects.create(exchange_rate_aed_rub=30, created_by=self.user)
        source = copy(self.source)
        source.pk, source.price_list = None, other_price
        source.exchange_rate_aed_rub = Decimal('30')
        source.save()
        item = self.second.items.get()
        item.price_list_item = source
        item.save(update_fields=('price_list_item',))
        result = preview_trip_price(actor=self.user, batch_id=self.trip.pk, line_id=self.line.pk,
            revision=self.trip.revision, price_aed='110')
        self.assertEqual(Decimal(result['cost_delta_rub']), Decimal('2650'))
        self.assertEqual([Decimal(row['rate']) for row in result['allocations']], [Decimal('25'), Decimal('30'), Decimal('25')])

    def test_missing_rate_fails_without_partial_save_and_saved_pair_can_restore_it(self):
        from .trip_prices import preview_trip_price
        self.source.exchange_rate_aed_rub = Decimal('0')
        self.source.save(update_fields=('exchange_rate_aed_rub',))
        self.price.exchange_rate_aed_rub = Decimal('0')
        self.price.save(update_fields=('exchange_rate_aed_rub',))
        result = preview_trip_price(actor=self.user, batch_id=self.trip.pk, line_id=self.line.pk,
            revision=self.trip.revision, price_aed='110')
        self.assertEqual(Decimal(result['cost_delta_rub']), Decimal('2500'))
        self.assertIn('Сохранённая пара', result['allocations'][0]['rate_source'])
        self.price.exchange_rate_usdt_aed = self.price.exchange_rate_usdt_rub = None
        self.price.save(update_fields=('exchange_rate_usdt_aed', 'exchange_rate_usdt_rub'))
        with self.assertRaisesMessage(ValidationError, 'исторический курс'):
            self.change()
        self.line.refresh_from_db()
        self.assertEqual(self.line.supplier_price_aed_snapshot, Decimal('100'))

    def test_audit_failure_rolls_back_price_and_revision(self):
        with patch('orders.trip_prices._record_trip_change', side_effect=RuntimeError('Аудит недоступен')):
            with self.assertRaises(RuntimeError):
                self.change()
        self.line.refresh_from_db()
        self.trip.refresh_from_db()
        self.assertEqual(self.line.supplier_price_aed_snapshot, Decimal('100'))
        self.assertIsNone(self.line.supplier_price_updated_at)
        self.assertEqual(self.trip.revision, 0)

    def supply_stock(self):
        from warehouse.models import Warehouse
        from .services import advance_order_status
        from .fulfillment_test_helpers import receive_with_test_rates as receive_stock_order
        warehouse = Warehouse.objects.create(name='Склад передачи')
        advance_order_status(actor=self.user, order_id=self.stock.pk, next_status='confirmed')
        advance_order_status(actor=self.user, order_id=self.stock.pk, next_status='receiving')
        return receive_stock_order(actor=self.user, order_id=self.stock.pk, warehouse_id=warehouse.pk)

    def test_supply_link_blocks_price_and_quantity_even_if_status_was_reset(self):
        from .trips import change_trip_line
        stock = self.supply_stock()
        stock.status = stock.Status.CREATED
        stock.save(update_fields=('status',))
        with self.assertRaisesMessage(ValidationError, 'приход'):
            self.change()
        with self.assertRaisesMessage(ValidationError, 'приход'):
            change_trip_line(actor=self.user, batch_id=self.trip.pk, line_id=self.line.pk,
                revision=self.trip.revision, quantity=0, action='adjust')

    def test_completed_customer_order_blocks_price_even_without_supply(self):
        self.first.status = self.first.Status.COMPLETED
        self.first.save(update_fields=('status',))
        with self.assertRaisesMessage(ValidationError, 'финансовая история'):
            self.change()

    def test_unrelated_processed_supplier_does_not_block_customer_price_or_quantities(self):
        from .trips import change_trip_line
        change_trip_line(actor=self.user, batch_id=self.trip.pk, line_id=self.line.pk,
            revision=self.trip.revision, quantity=7, action='transfer',
            target_supplier_id=self.other_supplier.pk, price_supplier_id=self.other_supplier.pk)
        self.trip.refresh_from_db()
        other_line = self.trip.lines.get(supplier=self.other_supplier)
        self.supply_stock()
        self.change()
        self.trip.refresh_from_db()
        change_trip_line(actor=self.user, batch_id=self.trip.pk, line_id=self.line.pk,
            revision=self.trip.revision, quantity=0, action='adjust')
        other_line.refresh_from_db()
        self.assertEqual(other_line.quantity, 3)
        self.assertEqual(other_line.supplier_price_aed_snapshot, Decimal('105'))

    def test_confirmation_and_receiving_do_not_block_purchase_price(self):
        self.first.status = self.first.Status.CONFIRMED
        self.first.save(update_fields=('status',))
        self.second.status = self.second.Status.RECEIVING
        self.second.save(update_fields=('status',))
        self.change()
        self.line.refresh_from_db()
        self.assertEqual(self.line.supplier_price_aed_snapshot, Decimal('110'))

    def test_http_preview_save_replay_and_stale_edit(self):
        self.client.force_login(self.user)
        preview_url = reverse('orders:trip_price_preview', args=(self.trip.pk, self.line.pk))
        save_url = reverse('orders:trip_price_update', args=(self.trip.pk, self.line.pk))
        data = {'revision': self.trip.revision, 'price_aed': '110', 'reason': 'Последний инвойс'}
        preview = self.client.post(preview_url, data)
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(Decimal(preview.json()['cost_delta_rub']), Decimal('2500'))
        self.line.refresh_from_db()
        self.assertEqual(self.line.supplier_price_aed_snapshot, Decimal('100'))
        self.assertEqual(self.client.post(save_url, data).status_code, 200)
        self.assertEqual(self.client.post(save_url, data).status_code, 200)
        data['price_aed'] = '120'
        self.assertEqual(self.client.post(save_url, data).status_code, 409)
        with patch('orders.trip_price_views.change_trip_price', side_effect=OperationalError('database table is locked')):
            self.assertEqual(self.client.post(save_url, data).status_code, 503)
        self.line.refresh_from_db()
        self.assertEqual(self.line.supplier_price_aed_snapshot, Decimal('110'))
        self.assertEqual(self.client.get(save_url).status_code, 405)

    def test_http_permissions_and_csrf(self):
        url = reverse('orders:trip_price_update', args=(self.trip.pk, self.line.pk))
        employee = User.objects.create_user('trip-price-no-rights')
        self.client.force_login(employee)
        self.assertEqual(self.client.post(url, {'revision': 0, 'price_aed': '110'}).status_code, 403)
        csrf = Client(enforce_csrf_checks=True)
        csrf.force_login(self.user)
        self.assertEqual(csrf.post(url, {'revision': 0, 'price_aed': '110'}).status_code, 403)

    def test_explicit_price_survives_moving_stock_to_an_empty_trip(self):
        from .models import SupplierOrderBatch
        from .trips import assign_order_to_trip
        changed = self.change()
        target = SupplierOrderBatch.objects.create(price_list=self.price, created_by=self.user)
        assign_order_to_trip(actor=self.user, order_id=self.stock.pk, target_batch_id=target.pk,
            target_revision=0, source_batch_id=self.trip.pk, source_revision=changed.revision)
        new_line = target.lines.get()
        self.assertEqual(new_line.supplier_price_aed_snapshot, Decimal('110'))
        self.assertIsNotNone(new_line.supplier_price_updated_at)
        self.assertEqual(self.supply_stock().supply.cd_items.get().purchase_unit_cost, Decimal('2767'))

    def test_price_is_local_to_one_supplier_and_trip(self):
        from .trips import change_trip_line
        other_order = self.order('Другой рейс', 1, 0)
        other_trip = create_trip(actor=self.user, order_ids=[other_order.pk], creation_key=uuid4())
        change_trip_line(actor=self.user, batch_id=self.trip.pk, line_id=self.line.pk,
            revision=self.trip.revision, quantity=7, action='transfer',
            target_supplier_id=self.other_supplier.pk, price_supplier_id=self.other_supplier.pk)
        self.trip.refresh_from_db()
        self.change()
        self.assertEqual(other_trip.lines.get().supplier_price_aed_snapshot, Decimal('100'))
        self.assertEqual(self.trip.lines.get(supplier=self.other_supplier).supplier_price_aed_snapshot, Decimal('105'))

    def test_receipt_after_review_keeps_weighted_stock_cost(self):
        from supplies.services import confirm_supply_price_review
        from warehouse.models import CDWarehouseStock, Warehouse
        warehouse = Warehouse.objects.create(name='Средневзвешенный склад')
        CDWarehouseStock.objects.create(cd=self.product, warehouse=warehouse, quantity=2)
        self.product.cost = Decimal('1000')
        self.product.save(update_fields=('cost',))
        self.change()
        from .services import advance_order_status
        from .fulfillment_test_helpers import receive_with_test_rates as receive_stock_order
        advance_order_status(actor=self.user, order_id=self.stock.pk, next_status='confirmed')
        advance_order_status(actor=self.user, order_id=self.stock.pk, next_status='receiving')
        supply = receive_stock_order(actor=self.user, order_id=self.stock.pk, warehouse_id=warehouse.pk).supply
        self.assertEqual(CDWarehouseStock.objects.get(cd=self.product, warehouse=warehouse).quantity, 2)
        confirm_supply_price_review(actor=self.user, supply_id=supply.pk)
        self.assertEqual(CDWarehouseStock.objects.get(cd=self.product, warehouse=warehouse).quantity, 5)
        self.product.refresh_from_db()
        self.assertEqual(self.product.cost, Decimal('2060.20'))

    def test_same_product_and_supplier_with_different_price_sources_are_separate(self):
        from .trips import change_trip_line
        first = change_trip_line(actor=self.user, batch_id=self.trip.pk, line_id=self.line.pk,
            revision=0, quantity=8, action='transfer', target_supplier_id=self.other_supplier.pk,
            price_supplier_id=self.other_supplier.pk)
        temporary = self.trip.lines.get(supplier=self.other_supplier)
        self.other_supplier.priority = self.supplier.priority
        self.other_supplier.save(update_fields=('priority',))
        changed = change_trip_line(actor=self.user, batch_id=self.trip.pk, line_id=temporary.pk,
            revision=first.revision, quantity=0, action='transfer', target_supplier_id=self.supplier.pk,
            price_supplier_id=self.other_supplier.pk)
        self.change(revision=changed.revision)
        separate = self.trip.lines.get(supplier=self.supplier, price_supplier=self.other_supplier)
        self.assertEqual(separate.quantity, 2)
        self.assertEqual(separate.supplier_price_aed_snapshot, Decimal('105'))

    def test_regenerated_excel_uses_purchase_price_and_preserves_buyer_prices(self):
        self.change()
        self.client.force_login(self.user)
        def rows(name, *ids):
            response = self.client.get(reverse(name, args=ids))
            self.assertEqual(response.status_code, 200)
            book = load_workbook(BytesIO(response.content))
            return [row for row in book.active.iter_rows(values_only=True) if row[0] == self.product.name]
        supplier = rows('orders:trip_export', self.trip.pk, self.supplier.pk)
        full = rows('orders:trip_full_export', self.trip.pk)
        buyer = rows('orders:export', self.first.pk)
        customer_suppliers = rows('orders:export_suppliers', self.first.pk)
        stock = rows('orders:export', self.stock.pk)
        self.assertEqual(Decimal(str(supplier[0][3])), Decimal('110'))
        self.assertEqual(Decimal(str(full[0][2])), Decimal('110'))
        self.assertEqual(Decimal(str(buyer[0][3])), Decimal('3500'))
        self.assertEqual(Decimal(str(customer_suppliers[0][3])), Decimal('110'))
        self.assertEqual(Decimal(str(stock[0][2])), Decimal('2767'))

    def test_actual_supply_link_protects_old_order_edit_forms(self):
        from .services import edit_created_order, reduce_receiving_order
        stock = self.supply_stock()
        stock.trip, stock.status = None, stock.Status.CREATED
        stock.save(update_fields=('trip', 'status'))
        with self.assertRaisesMessage(ValidationError, 'приход'):
            edit_created_order(actor=self.user, order_id=stock.pk, recipient='На склад', comment='', lines=[{
                'price_list_item_id': self.source.pk, 'prepayment_quantity': 2, 'postpayment_quantity': 0}])
        stock.status = stock.Status.RECEIVING
        stock.save(update_fields=('status',))
        with self.assertRaisesMessage(ValidationError, 'приход'):
            reduce_receiving_order(actor=self.user, order_id=stock.pk, comment='Попытка изменить', rows=[])

    def test_audit_does_not_reveal_aed_to_viewer_without_price_permission(self):
        self.change()
        viewer = User.objects.create_user('trip-without-price')
        for app, code in (('orders', 'generate_supplier_orders'), ('partners', 'view_supplier_details')):
            viewer.user_permissions.add(Permission.objects.get(content_type__app_label=app, codename=code))
        self.client.force_login(viewer)
        response = self.client.get(reverse('orders:trip_detail', args=(self.trip.pk,)))
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'Цена AED:')
        self.assertNotContains(response, 'trip-price-data')
        self.assertNotContains(response, 'data-procurement-financials')
