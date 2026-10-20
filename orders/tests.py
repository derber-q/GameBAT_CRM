from decimal import Decimal
from io import BytesIO
from uuid import uuid4

from django.contrib.auth.models import Permission
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import User
from cash.models import CashRegister, CashTransaction, Safe
from catalog.models import Brand, CD, Platform, ProductType, Tech
from partners.models import Supplier
from price.services import create_procurement_price_list
from pricing.models import SupplierCDPrice, SupplierTechPrice
from warehouse.models import Warehouse

from .models import CustomerProcurementOrder, OrderAdjustment, SupplierOrderBatchLine
from .services import (
    advance_order_status,
    create_customer_procurement_order,
    create_supplier_order_batch,
    edit_created_order,
    reduce_receiving_order,
)
from .trips import create_trip


def order_product_rows(sheet):
    """Товарные строки отчёта: количество задано числом, итог в конце исключён."""
    quantity_column = sheet.max_column - 2
    return [(index, tuple(cell.value for cell in sheet[index]))
            for index in range(5, sheet.max_row)
            if isinstance(sheet.cell(index, quantity_column).value, int)]


class ProcurementOrderTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser("admin", password="StrongAdmin!123")
        self.warehouse = Warehouse.objects.create(name="Москва")
        platform = Platform.objects.create(name="PS5")
        brand = Brand.objects.create(name="Sony")
        kind = ProductType.objects.create(name="Консоли")
        self.cd = CD.objects.create(platform=platform, name="Игра", sku="CD-1", barcode="1")
        self.tech = Tech.objects.create(
            brand=brand, product_type=kind, name="Консоль", sku="T-1", barcode="2", weight_grams=1000,
        )
        self.supplier = Supplier.objects.create(
            name="Alpha", letter="A", highlight_color="#37A7BA", legal_entity="A", phone_1="1"
        )
        SupplierCDPrice.objects.create(supplier=self.supplier, cd=self.cd, price=100)
        SupplierTechPrice.objects.create(supplier=self.supplier, tech=self.tech, price=150)
        self.price_list = create_procurement_price_list(
            actor=self.user, exchange_rate_usdt_aed=4, exchange_rate_usdt_rub=80,
            delivery_general_usdt_per_kg=0, delivery_devices_usdt_per_kg=0, delivery_storage_usdt_per_kg=0, delivery_cd_rub=0,
        )
        self.cd_source = self.price_list.items.get(product_kind="cd")
        self.tech_source = self.price_list.items.get(product_kind="tech")

    def create_order(self, *, pre=2, post=3, source=None):
        source = source or self.cd_source
        return create_customer_procurement_order(
            actor=self.user, price_list_id=source.price_list_id,
            recipient="Покупатель", comment="Тест", lines=[{
                "price_list_item_id": source.pk,
                "prepayment_quantity": pre, "postpayment_quantity": post,
            }],
        )

    def test_order_copies_price_and_supplier_snapshots(self):
        order = self.create_order(pre=2, post=1)
        item = order.items.get()
        self.assertEqual(item.selected_supplier, self.supplier)
        self.assertEqual(item.supplier_price_aed_snapshot, Decimal("100"))
        self.assertEqual(item.prepayment_unit_price, Decimal("2000.00"))
        self.assertEqual(item.postpayment_unit_price, Decimal("2080.00"))
        self.assertEqual(order.prepayment_total, Decimal("4000.00"))
        self.assertEqual(order.postpayment_total, Decimal("2080.00"))
        self.assertEqual(order.grand_total, Decimal("6080.00"))

    def test_current_order_excel_has_separate_payment_rows_and_is_read_only(self):
        order = self.create_order(pre=2, post=1)
        self.client.force_login(self.user)
        with CaptureQueriesContext(connection) as queries:
            response = self.client.get(reverse("orders:export", args=(order.pk,)))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Disposition"], f'attachment; filename="zakaz_{order.pk}.xlsx"')
        self.assertIn("no-store", response["Cache-Control"])
        self.assertFalse(any(q["sql"].lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE")) for q in queries))
        self.assertEqual(sum('FROM "orders_customerprocurementorderitem"' in q["sql"] for q in queries), 1)
        book = load_workbook(BytesIO(response.content))
        sheet = book.active
        self.assertEqual(tuple(cell.value for cell in sheet[4]),
                         ("Наименование", "Оплата", "Количество", "Цена, ₽", "Сумма, ₽"))
        self.assertEqual([values for _row, values in order_product_rows(sheet)], [
            ("Игра", "Предоплата", 2, 2000, 4000), ("Игра", "Постоплата", 1, 2080, 2080),
        ])
        self.assertEqual(sheet.cell(sheet.max_row, 3).value, 3)
        self.assertEqual(sheet.cell(sheet.max_row, 5).value, 6080)
        self.assertEqual(book.sheetnames, ["Заказ"])
        self.assertNotIn(self.supplier.name, str(list(sheet.values)))
        book.close()
        order.refresh_from_db()
        self.assertEqual(order.status, order.Status.CREATED)

    def test_order_excel_refreshes_quantities_but_keeps_customer_prices(self):
        order = self.create_order(pre=2, post=3)
        self.client.force_login(self.user)
        url = reverse("orders:export", args=(order.pk,))
        self.assertEqual(self.client.get(url).status_code, 200)
        edit_created_order(actor=self.user, order_id=order.pk, recipient="Покупатель", comment="После инвойса",
                           lines=[{"price_list_item_id": self.cd_source.pk,
                                   "prepayment_quantity": 1, "postpayment_quantity": 2}])
        self.cd_source.prepayment_price_rub = Decimal("99999")
        self.cd_source.postpayment_price_rub = Decimal("99999")
        self.cd_source.save()
        book = load_workbook(BytesIO(self.client.get(url).content))
        sheet = book.active
        self.assertEqual([values for _row, values in order_product_rows(sheet)], [
            ("Игра", "Предоплата", 1, 2000, 2000), ("Игра", "Постоплата", 2, 2080, 4160),
        ])
        self.assertEqual(sheet.cell(sheet.max_row, 5).value, 6160)
        book.close()

    def test_order_excel_preserves_zero_positions_after_receiving_adjustment(self):
        order = self.create_order(pre=1, post=1)
        create_trip(actor=self.user, order_ids=[order.pk], creation_key=uuid4())
        advance_order_status(actor=self.user, order_id=order.pk, next_status=order.Status.CONFIRMED,
                             warehouse_id=self.warehouse.pk)
        advance_order_status(actor=self.user, order_id=order.pk, next_status=order.Status.RECEIVING)
        item = order.items.get()
        reduce_receiving_order(actor=self.user, order_id=order.pk, comment="Не отгружено", rows=[{
            "item_id": item.pk, "prepayment_quantity": 0, "postpayment_quantity": 0,
        }])
        self.client.force_login(self.user)
        book = load_workbook(BytesIO(self.client.get(reverse("orders:export", args=(order.pk,))).content))
        sheet = book.active
        self.assertEqual([values[2] for _row, values in order_product_rows(sheet)], [0, 0])
        self.assertEqual(sheet.cell(sheet.max_row, 5).value, 0)
        book.close()

    def test_customer_excel_requires_detail_permission_and_does_not_need_supplier_access(self):
        order = self.create_order(pre=1, post=0)
        worker = User.objects.create_user("order-reader", password="StrongWorker!123")
        worker.user_permissions.add(*Permission.objects.filter(
            content_type__app_label="orders", codename__in=("view_order_detail", "view_orders"),
        ))
        self.client.force_login(worker)
        url = reverse("orders:export", args=(order.pk,))
        self.assertEqual(self.client.get(url).status_code, 200)
        self.assertEqual(self.client.post(url).status_code, 405)
        self.assertContains(self.client.get(reverse("orders:detail", args=(order.pk,))), url)
        self.assertContains(self.client.get(reverse("orders:list")), url)

    def test_order_excel_keeps_kopecks_and_treats_names_as_text(self):
        self.cd_source.prepayment_price_rub = Decimal("1234.56")
        self.cd_source.save()
        order = self.create_order(pre=2, post=0)
        item = order.items.get()
        item.product_name_snapshot = '=HYPERLINK("https://example.invalid", "Товар")'
        item.save(update_fields=("product_name_snapshot",))
        self.client.force_login(self.user)
        book = load_workbook(BytesIO(self.client.get(reverse("orders:export", args=(order.pk,))).content))
        sheet = book.active
        row, values = order_product_rows(sheet)[0]
        self.assertEqual(sheet.cell(row, 1).data_type, "s")
        self.assertEqual(values[0], "'" + item.product_name_snapshot)
        self.assertEqual(Decimal(str(values[3])), Decimal("1234.56"))
        self.assertEqual(Decimal(str(values[4])), Decimal("2469.12"))
        self.assertEqual(Decimal(str(sheet.cell(sheet.max_row, 5).value)), Decimal("2469.12"))
        book.close()

    def test_created_order_can_add_remove_increase_and_decrease(self):
        order = self.create_order(pre=1, post=1)
        edit_created_order(
            actor=self.user, order_id=order.pk, recipient="Новый получатель", comment="Изменено",
            lines=[{
                "price_list_item_id": self.cd_source.pk,
                "prepayment_quantity": 3, "postpayment_quantity": 0,
            }, {
                "price_list_item_id": self.tech_source.pk,
                "prepayment_quantity": 0, "postpayment_quantity": 2,
            }],
        )
        order.refresh_from_db()
        self.assertEqual(order.recipient, "Новый получатель")
        self.assertEqual(order.items.count(), 2)
        self.assertEqual(order.items.get(price_list_item=self.cd_source).prepayment_quantity, 3)

    def test_confirm_prepayment_is_once_and_safe_is_untouched(self):
        order = self.create_order(pre=2, post=0)
        safe_before = Safe.objects.get(warehouse=self.warehouse).balance
        advance_order_status(
            actor=self.user, order_id=order.pk,
            next_status=CustomerProcurementOrder.Status.CONFIRMED, warehouse_id=self.warehouse.pk,
        )
        order.refresh_from_db()
        register = CashRegister.objects.get(warehouse=self.warehouse)
        self.assertEqual(register.balance, Decimal("4000.00"))
        self.assertEqual(Safe.objects.get(warehouse=self.warehouse).balance, safe_before)
        self.assertEqual(order.payment_warehouse, self.warehouse)
        self.assertEqual(CashTransaction.objects.filter(
            customer_order=order,
            operation_type=CashTransaction.OperationType.CUSTOMER_ORDER_PREPAYMENT,
        ).count(), 1)
        with self.assertRaisesMessage(ValidationError, "переход"):
            advance_order_status(
                actor=self.user, order_id=order.pk,
                next_status=CustomerProcurementOrder.Status.CONFIRMED, warehouse_id=self.warehouse.pk,
            )
        self.assertEqual(CashRegister.objects.get(warehouse=self.warehouse).balance, Decimal("4000.00"))

    def test_receiving_reduction_refunds_atomically_and_requires_comment(self):
        order = self.create_order(pre=5, post=2)
        advance_order_status(
            actor=self.user, order_id=order.pk,
            next_status=order.Status.CONFIRMED, warehouse_id=self.warehouse.pk,
        )
        advance_order_status(actor=self.user, order_id=order.pk, next_status=order.Status.RECEIVING)
        item = order.items.get()
        with self.assertRaisesMessage(ValidationError, "Комментарий"):
            reduce_receiving_order(actor=self.user, order_id=order.pk, comment="", rows=[{
                "item_id": item.pk, "prepayment_quantity": 3, "postpayment_quantity": 2,
            }])
        adjustment = reduce_receiving_order(actor=self.user, order_id=order.pk, comment="Не приехало", rows=[{
            "item_id": item.pk, "prepayment_quantity": 3, "postpayment_quantity": 1,
        }])
        order.refresh_from_db()
        item.refresh_from_db()
        self.assertEqual(item.prepayment_quantity, 3)
        self.assertEqual(item.postpayment_quantity, 1)
        self.assertEqual(adjustment.refund_amount, Decimal("4000.00"))
        self.assertEqual(CashRegister.objects.get(warehouse=self.warehouse).balance, Decimal("6000.00"))
        self.assertEqual(adjustment.changes.count(), 2)
        with self.assertRaisesMessage(ValidationError, "не изменён"):
            reduce_receiving_order(actor=self.user, order_id=order.pk, comment="Повтор", rows=[{
                "item_id": item.pk, "prepayment_quantity": 3, "postpayment_quantity": 1,
            }])
        self.assertEqual(CashTransaction.objects.filter(
            customer_order=order,
            operation_type=CashTransaction.OperationType.CUSTOMER_ORDER_REFUND,
        ).count(), 1)
        with self.assertRaisesMessage(ValidationError, "Нельзя увеличить"):
            reduce_receiving_order(actor=self.user, order_id=order.pk, comment="Назад", rows=[{
                "item_id": item.pk, "prepayment_quantity": 4, "postpayment_quantity": 1,
            }])

    def test_insufficient_cash_rolls_back_reduction(self):
        order = self.create_order(pre=5, post=0)
        advance_order_status(
            actor=self.user, order_id=order.pk,
            next_status=order.Status.CONFIRMED, warehouse_id=self.warehouse.pk,
        )
        advance_order_status(actor=self.user, order_id=order.pk, next_status=order.Status.RECEIVING)
        register = CashRegister.objects.get(warehouse=self.warehouse)
        register.balance = 1
        register.save(update_fields=("balance",))
        item = order.items.get()
        with self.assertRaisesMessage(ValidationError, "недостаточно"):
            reduce_receiving_order(actor=self.user, order_id=order.pk, comment="Возврат", rows=[{
                "item_id": item.pk, "prepayment_quantity": 4, "postpayment_quantity": 0,
            }])
        item.refresh_from_db()
        self.assertEqual(item.prepayment_quantity, 5)
        self.assertFalse(OrderAdjustment.objects.filter(order=order).exists())

    def test_postpayment_reduction_moves_no_cash_until_completed(self):
        order = self.create_order(pre=0, post=5)
        advance_order_status(actor=self.user, order_id=order.pk, next_status=order.Status.CONFIRMED)
        self.assertFalse(CashTransaction.objects.filter(customer_order=order).exists())
        advance_order_status(actor=self.user, order_id=order.pk, next_status=order.Status.RECEIVING)
        item = order.items.get()
        reduce_receiving_order(
            actor=self.user, order_id=order.pk, comment="Меньше", allow_refund=False, rows=[{
                "item_id": item.pk, "prepayment_quantity": 0, "postpayment_quantity": 3,
            }],
        )
        self.assertFalse(CashTransaction.objects.filter(customer_order=order).exists())
        order.refresh_from_db()
        advance_order_status(
            actor=self.user, order_id=order.pk,
            next_status=order.Status.COMPLETED, warehouse_id=self.warehouse.pk,
        )
        expected = Decimal("2080.00") * 3
        self.assertEqual(CashRegister.objects.get(warehouse=self.warehouse).balance, expected)
        self.assertEqual(CashTransaction.objects.filter(
            customer_order=order,
            operation_type=CashTransaction.OperationType.CUSTOMER_ORDER_POSTPAYMENT,
        ).count(), 1)
        with self.assertRaises(ValidationError):
            advance_order_status(
                actor=self.user, order_id=order.pk,
                next_status=order.Status.COMPLETED, warehouse_id=self.warehouse.pk,
            )

    def test_nonexistent_optional_payment_warehouse_is_rejected(self):
        order = self.create_order(pre=0, post=1)
        with self.assertRaisesMessage(ValidationError, "существующий склад"):
            advance_order_status(
                actor=self.user, order_id=order.pk,
                next_status=order.Status.CONFIRMED, warehouse_id=999999,
            )
        order.refresh_from_db()
        self.assertEqual(order.status, order.Status.CREATED)
        self.assertIsNone(order.payment_warehouse_id)

    def test_supplier_batch_aggregates_only_selected_orders(self):
        first = self.create_order(pre=2, post=0)
        second = self.create_order(pre=0, post=3)
        ignored = self.create_order(pre=10, post=0)
        batch = create_supplier_order_batch(actor=self.user, order_ids=[first.pk, second.pk])
        self.assertCountEqual(batch.source_orders.values_list("pk", flat=True), [first.pk, second.pk])
        line = SupplierOrderBatchLine.objects.get(batch=batch)
        self.assertEqual(line.supplier, self.supplier)
        self.assertEqual(line.cd, self.cd)
        self.assertEqual(line.quantity, 5)
        self.assertNotIn(ignored, batch.source_orders.all())

    def test_supplier_batch_uses_weighted_snapshot_when_versions_have_different_prices(self):
        old_order = self.create_order(pre=2, post=0)
        SupplierCDPrice.objects.filter(supplier=self.supplier, cd=self.cd).update(price=120)
        newer_price_list = create_procurement_price_list(
            actor=self.user, exchange_rate_usdt_aed=4, exchange_rate_usdt_rub=80,
            delivery_general_usdt_per_kg=0, delivery_devices_usdt_per_kg=0, delivery_storage_usdt_per_kg=0, delivery_cd_rub=0,
        )
        newer_source = newer_price_list.items.get(product_kind="cd")
        new_order = self.create_order(pre=3, post=0, source=newer_source)
        batch = create_supplier_order_batch(actor=self.user, order_ids=[old_order.pk, new_order.pk])
        line = SupplierOrderBatchLine.objects.get(batch=batch, cd=self.cd)
        self.assertEqual(line.quantity, 5)
        self.assertEqual(line.supplier_price_aed_snapshot, Decimal("112.000000"))

    def test_order_endpoints_enforce_permissions(self):
        worker = User.objects.create_user("worker", password="StrongWorker!123")
        self.client.force_login(worker)
        self.assertEqual(self.client.get(reverse("orders:list")).status_code, 403)
        worker.user_permissions.add(Permission.objects.get(
            content_type__app_label="orders", codename="view_orders"
        ))
        self.assertEqual(self.client.get(reverse("orders:list")).status_code, 200)
        self.assertEqual(self.client.get(reverse("orders:upload")).status_code, 403)

    def test_cash_permission_is_required_for_financial_status_transition(self):
        order = self.create_order(pre=1, post=0)
        worker = User.objects.create_user("cashless", password="StrongWorker!123")
        worker.user_permissions.add(Permission.objects.get(
            content_type__app_label="orders", codename="confirm_order"
        ))
        self.client.force_login(worker)
        response = self.client.post(reverse("orders:confirm", args=(order.pk,)), {
            "warehouse": self.warehouse.pk,
        })
        self.assertEqual(response.status_code, 403)
        order.refresh_from_db()
        self.assertEqual(order.status, order.Status.CREATED)
        self.assertFalse(CashTransaction.objects.filter(customer_order=order).exists())

    def test_all_order_endpoints_reject_missing_permissions(self):
        order = self.create_order(pre=1, post=1)
        batch = create_supplier_order_batch(actor=self.user, order_ids=[order.pk])
        worker = User.objects.create_user("restricted", password="StrongWorker!123")
        self.client.force_login(worker)
        checks = (
            ("get", reverse("orders:list")),
            ("get", reverse("orders:upload")),
            ("get", reverse("orders:preview")),
            ("get", reverse("orders:detail", args=(order.pk,))),
            ("get", reverse("orders:export", args=(order.pk,))),
            ("get", reverse("orders:edit", args=(order.pk,))),
            ("post", reverse("orders:confirm", args=(order.pk,))),
            ("post", reverse("orders:start_receiving", args=(order.pk,))),
            ("get", reverse("orders:adjust", args=(order.pk,))),
            ("post", reverse("orders:complete", args=(order.pk,))),
            ("post", reverse("orders:supplier_batch_create")),
            ("get", reverse("orders:supplier_batch_detail", args=(batch.pk,))),
        )
        for method, url in checks:
            with self.subTest(url=url):
                self.assertEqual(getattr(self.client, method)(url).status_code, 403)
