"""Группировка обычного состава заказа и отдельных таблиц рейса."""
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import User
from catalog.models import Brand, CD, GameSeries, Platform, ProductType, Tech
from partners.models import Supplier
from price.services import create_procurement_price_list
from pricing.models import SupplierCDPrice, SupplierTechPrice

from .services import create_customer_procurement_order
from .trips import change_trip_line, create_trip


def create_export_fixture(target):
    target.user = User.objects.create_superuser("export-admin", password="StrongAdmin!123")
    target.a, target.b, target.c = [Supplier.objects.create(
        name=f"Поставщик {letter}", letter=letter, highlight_color="#37A7BA", legal_entity=letter, phone_1="1", priority=index,
    ) for index, letter in enumerate("ABC", 1)]
    target.cd = CD.objects.create(name="Игра", platform=Platform.objects.create(name="PS5"),
                                  game_series=GameSeries.objects.create(name="Zelda"), sku="000123")
    target.tech = Tech.objects.create(name="Консоль", brand=Brand.objects.create(name="Sony"),
                                      product_type=ProductType.objects.create(name="Консоли"), sku="000-tech", weight_grams=1000)
    for supplier, cd_price, tech_price in ((target.a, 100, 150), (target.b, 110, 140), (target.c, 120, 160)):
        SupplierCDPrice.objects.create(supplier=supplier, cd=target.cd, price=cd_price)
        SupplierTechPrice.objects.create(supplier=supplier, tech=target.tech, price=tech_price)
    target.price = create_procurement_price_list(
        actor=target.user, exchange_rate_usdt_aed=4, exchange_rate_usdt_rub=80,
        delivery_general_usdt_per_kg=0, delivery_devices_usdt_per_kg=0,
        delivery_storage_usdt_per_kg=0, delivery_cd_rub=0,
    )
    target.cd_source = target.price.items.get(cd=target.cd)
    target.tech_source = target.price.items.get(tech=target.tech)
    orders = []
    for recipient, pre, post, tech_quantity in (("Первый", 2, 3, 1), ("Другой заказчик", 3, 4, 2)):
        orders.append(create_customer_procurement_order(
            actor=target.user, price_list_id=target.price.pk, recipient=recipient, comment="", lines=[
                {"price_list_item_id": target.cd_source.pk, "prepayment_quantity": pre, "postpayment_quantity": post},
                {"price_list_item_id": target.tech_source.pk, "prepayment_quantity": 0, "postpayment_quantity": tech_quantity},
            ],
        ))
    target.first, target.second = orders


class OrderExcelGroupingTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        create_export_fixture(cls)

    def setUp(self):
        self.client.force_login(self.user)

    @staticmethod
    def titles(sheet):
        return [row[0] for row in sheet.values if isinstance(row[0], str) and row[0].startswith(("Tech ·", "CD ·"))]

    def transfer(self, trip, source, quantity, price_supplier):
        trip.refresh_from_db()
        return change_trip_line(actor=self.user, batch_id=trip.pk, line_id=source.pk, revision=trip.revision,
                                quantity=quantity, action="transfer", target_supplier_id=self.b.pk,
                                price_supplier_id=price_supplier.pk)

    def test_order_uses_page_groups_and_keeps_customer_rows_and_totals(self):
        response = self.client.get(reverse("orders:export", args=(self.first.pk,)))
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content))
        sheet = book.active
        self.assertEqual(self.titles(sheet), ["Tech · Sony · Консоли", "CD · Платформа: PS5"])
        rows = [tuple(row) for row in sheet.iter_rows(min_row=5, max_row=sheet.max_row - 1, values_only=True)
                if isinstance(row[2], int)]
        tech_price = self.tech_source.postpayment_price_rub
        self.assertEqual(rows, [("Консоль", "Постоплата", 1, tech_price, tech_price),
                                ("Игра", "Предоплата", 2, 2000, 4000), ("Игра", "Постоплата", 3, 2080, 6240)])
        self.assertEqual(sheet.cell(sheet.max_row, 3).value, 6)
        self.assertEqual(sheet.cell(sheet.max_row, 5).value, 10240 + tech_price)
        book.close()

    def test_trip_keeps_duplicate_products_at_different_prices_in_page_order(self):
        trip = create_trip(actor=self.user, order_ids=[self.first.pk, self.second.pk], creation_key=uuid4())
        source = trip.lines.get(cd=self.cd, supplier=self.a)
        self.transfer(trip, source, 10, self.a)
        self.transfer(trip, source, 7, self.b)
        response = self.client.get(reverse("orders:trip_export", args=(trip.pk, self.b.pk)))
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content))
        sheet = book.active
        self.assertEqual(self.titles(sheet), ["Tech · Sony · Консоли", "CD · Платформа: PS5"])
        rows = [tuple(row) for row in sheet.iter_rows(min_row=2, values_only=True) if isinstance(row[4], int)]
        self.assertEqual(rows, [("Консоль", None, "000-tech", 140, 3), ("Игра", None, "000123", 100, 2), ("Игра", None, "000123", 110, 3)])
        self.assertTrue(sheet.column_dimensions["C"].hidden)
        book.close()
