"""Снимок текущего состава заказа с сохранёнными ценами покупателя."""
from decimal import Decimal
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

from core.excel_layout import wrapped_text_height
from catalog.nomenclature_grouping import nomenclature_product_groups
from price.excel import GRAY, LIGHT, TEAL, WHITE, safe_text


def order_xlsx(*, order, items):
    book = Workbook()
    sheet = book.active
    sheet.title = "Заказ"
    sheet.sheet_view.showGridLines = False
    headers = ["Наименование"]
    if not order.is_stock_order:
        headers.append("Оплата")
    headers.extend(("Количество", "Цена, ₽", "Сумма, ₽"))
    widths = (66, 14, 19, 21) if order.is_stock_order else (54, 16, 12, 16, 18)
    last_column = len(headers)
    last_letter = chr(64 + last_column)
    quantity_column = last_column - 2
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width

    def text_row(row, value, *, size=11, bold=False):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)
        cell = sheet.cell(row, 1, safe_text(value))
        cell.font = Font(name="Arial", size=size, bold=bold, color="334660")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[row].height = wrapped_text_height(
            value, sum(widths), font_name="Arial", font_size=size, bold=bold, min_height=24,
        )

    title = "Заказ на склад" if order.is_stock_order else "Заказ"
    text_row(1, f"{title} №{order.pk}", size=16, bold=True)
    text_row(2, f"Получатель: {order.recipient}")
    text_row(3, f"Сформировано: {timezone.localtime():%d.%m.%Y %H:%M:%S} · Статус: {order.get_status_display()}", size=10)
    for column, label in enumerate(headers, 1):
        cell = sheet.cell(4, column, label)
        cell.font = Font(name="Arial", size=11, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    sheet.row_dimensions[4].height = 30

    items = list(items)
    groups = nomenclature_product_groups(
        [item for item in items if item.product_kind == "tech"],
        [item for item in items if item.product_kind == "cd"], product_of=lambda item: item.product,
    )
    rows = []
    for title, item in ((title, item) for _kind, title, group in groups for item in group):
        if order.is_stock_order:
            rows.append((title, (item.product_name_snapshot, item.quantity, item.prepayment_unit_price)))
        else:
            for label, quantity, price in (
                ("Предоплата", item.prepayment_quantity, item.prepayment_unit_price),
                ("Постоплата", item.postpayment_quantity, item.postpayment_unit_price),
            ):
                if quantity or not item.quantity:
                    rows.append((title, (item.product_name_snapshot, label, quantity, price)))

    total_quantity = 0
    total_amount = Decimal("0.00")
    border = Border(bottom=Side(style="thin", color=GRAY))
    previous_group = None
    for title, values in rows:
        if title != previous_group:
            text_row(sheet.max_row + 1, title, bold=True)
            for cell in sheet[sheet.max_row]:
                cell.fill = PatternFill("solid", fgColor=LIGHT)
            previous_group = title
        quantity, price = values[-2:]
        amount = price * quantity
        sheet.append([safe_text(value) if isinstance(value, str) else value for value in (*values, amount)])
        row = sheet.max_row
        for cell in sheet[row]:
            cell.font = Font(name="Arial", size=11, color="334660")
            cell.alignment = Alignment(vertical="center", wrap_text=cell.column < quantity_column,
                                       horizontal="right" if cell.column >= quantity_column else "left")
            cell.border = border
            if not quantity:
                cell.fill = PatternFill("solid", fgColor=LIGHT)
        sheet.cell(row, quantity_column).number_format = "0"
        for column, value in ((last_column - 1, price), (last_column, amount)):
            sheet.cell(row, column).number_format = "#,##0" if value == value.to_integral_value() else "#,##0.00"
        sheet.row_dimensions[row].height = wrapped_text_height(
            values[0], widths[0], font_name="Arial", font_size=11, min_height=28,
        )
        total_quantity += quantity
        total_amount += amount
    if not rows:
        text_row(5, "В заказе нет товарных позиций.")
    last_item_row = sheet.max_row
    sheet.append(["Итого"] + [None] * (quantity_column - 2) + [total_quantity, None, total_amount])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(name="Arial", size=11, bold=True, color="334660")
        cell.fill = PatternFill("solid", fgColor=LIGHT)
        cell.alignment = Alignment(horizontal="left" if cell.column == 1 else "right", vertical="center")
    sheet.cell(sheet.max_row, last_column).number_format = "#,##0.00"
    sheet.row_dimensions[sheet.max_row].height = 30
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:{last_letter}{last_item_row}"
    sheet.print_title_rows = "4:4"
    sheet.print_area = f"A1:{last_letter}{sheet.max_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins = PageMargins(left=0.4, right=0.4, top=0.4, bottom=0.4)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()
