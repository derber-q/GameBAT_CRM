"""Закупка на наличие: сохранённая цена, общие рейсы и одноразовый приход."""
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

from django.contrib.auth.models import Permission
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import User
from cash.models import CashRegister, CashTransaction
from catalog.models import Brand, CD, Platform, ProductType, Tech
from consignment.models import CDConsignmentStock
from partners.models import Supplier, SalesPlatform
from price.excel import (HEADER_ROW, MAIN_SHEET, META_SHEET, STOCK_HEADERS, LEGACY_STOCK_HEADERS,
    generate_procurement_customer_xlsx, generate_procurement_stock_xlsx, import_procurement_order_files)
from price.services import create_procurement_price_list
from pricing.models import SupplierCDPrice, SupplierTechPrice
from sales.models import Sale
from supplies.models import Supply
from supplies.services import confirm_supply_price_review
from warehouse.models import CDWarehouseStock, TechWarehouseStock, Warehouse
from .models import CustomerProcurementOrder
from .services import (advance_order_status, create_customer_procurement_order, edit_created_order,
                       receive_stock_order, reduce_receiving_order)
from .trips import create_trip, change_trip_line, assign_order_to_trip


class StockOrderTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_superuser("stock-admin", password="StrongAdmin!123")
        cls.warehouse = Warehouse.objects.create(name="Склад приёмки")
        cls.other_warehouse = Warehouse.objects.create(name="Другой склад")
        cls.m, cls.f = [Supplier.objects.create(name=letter, letter=letter, priority=priority,
            highlight_color="#37A7BA", legal_entity=letter, phone_1="1") for letter, priority in (("M", 3), ("F", 2))]
        cls.cd = CD.objects.create(name="Игра на склад", platform=Platform.objects.create(name="PS5"),
                                   sku="000123", weight_grams=100)
        cls.tech = Tech.objects.create(name="Консоль на склад", brand=Brand.objects.create(name="Sony"),
            product_type=ProductType.objects.create(name="Консоли"), sku="000456", weight_grams=1000)
        for supplier, cd_price, tech_price in ((cls.m, 100, 150), (cls.f, 120, 170)):
            SupplierCDPrice.objects.create(supplier=supplier, cd=cls.cd, price=cd_price)
            SupplierTechPrice.objects.create(supplier=supplier, tech=cls.tech, price=tech_price)
        cls.price = create_procurement_price_list(actor=cls.user, exchange_rate_usdt_aed=4,
            exchange_rate_usdt_rub=80, delivery_general_usdt_per_kg=2, delivery_devices_usdt_per_kg=2,
            delivery_storage_usdt_per_kg=2, delivery_cd_rub=50)
        cls.cd_source = cls.price.items.get(product_kind="cd")
        cls.tech_source = cls.price.items.get(product_kind="tech")
        for source in (cls.cd_source, cls.tech_source):
            source.markup_rub = Decimal("500")
            source.prepayment_price_rub = source.stock_unit_cost + Decimal("500")
            source.postpayment_price_rub = source.prepayment_price_rub + Decimal("100")
            source.save()

    def setUp(self):
        self.client.force_login(self.user)

    def order(self, qty=3, tech=0, purpose="stock"):
        lines = [{"price_list_item_id": self.cd_source.pk, "prepayment_quantity": qty, "postpayment_quantity": 0}]
        if tech:
            lines.append({"price_list_item_id": self.tech_source.pk, "prepayment_quantity": tech, "postpayment_quantity": 0})
        return create_customer_procurement_order(actor=self.user, price_list_id=self.price.pk,
            recipient="" if purpose == "stock" else "Клиент", comment="", lines=lines, purpose=purpose)

    def receiving(self, order):
        advance_order_status(actor=self.user, order_id=order.pk, next_status="confirmed")
        advance_order_status(actor=self.user, order_id=order.pk, next_status="receiving")

    def filled_stock_file(self, qty=3, tech=0, mutate=None):
        content = generate_procurement_stock_xlsx(price_list_id=self.price.pk, actor=self.user)
        book = load_workbook(BytesIO(content))
        for row in book[META_SHEET].iter_rows(values_only=True):
            if isinstance(row[0], int):
                book[MAIN_SHEET].cell(row[0], 8, qty if row[1] == "cd" else tech)
        if mutate:
            mutate(book)
        output = BytesIO()
        book.save(output)
        book.close()
        return SimpleUploadedFile("stock.xlsx", output.getvalue())

    def employee(self, *permissions):
        employee = User.objects.create_user("stock-employee", password="StrongAdmin!123")
        for permission in permissions:
            app, code = permission.split(".")
            employee.user_permissions.add(Permission.objects.get(content_type__app_label=app, codename=code))
        return employee

    def test_export_has_only_saved_cost_delivery_and_quantity(self):
        CDWarehouseStock.objects.create(cd=self.cd, warehouse=self.warehouse, quantity=2)
        CDWarehouseStock.objects.create(cd=self.cd, warehouse=self.other_warehouse, quantity=3)
        TechWarehouseStock.objects.create(tech=self.tech, warehouse=self.warehouse, quantity=4)
        CDConsignmentStock.objects.create(cd=self.cd, warehouse=self.warehouse,
            platform=SalesPlatform.objects.create(name="На реализации"), quantity=99)
        self.cd.cost, self.tech.cost = Decimal("111.25"), Decimal("222.75")
        self.cd.save(update_fields=("cost",))
        self.tech.save(update_fields=("cost",))
        book = load_workbook(BytesIO(generate_procurement_stock_xlsx(price_list_id=self.price.pk, actor=self.user)))
        sheet = book[MAIN_SHEET]
        self.assertEqual(tuple(cell.value for cell in sheet[HEADER_ROW]), STOCK_HEADERS)
        self.assertEqual(
            [cell.value for cell in sheet[HEADER_ROW]
             if not sheet.column_dimensions[cell.column_letter].hidden],
            [STOCK_HEADERS[index] for index in (0, 1, 2, 6, 7, 8)],
        )
        self.assertEqual(book[META_SHEET].sheet_state, "veryHidden")
        rows = [row for row in book[META_SHEET].iter_rows(values_only=True) if isinstance(row[0], int)]
        self.assertEqual(len(rows), 2)
        for row, kind, _pk, article in rows:
            source = self.cd_source if kind == "cd" else self.tech_source
            self.assertEqual(sheet.cell(row, 2).value, 5 if kind == "cd" else 4)
            self.assertEqual(Decimal(str(sheet.cell(row, 3).value)),
                             Decimal("111.25") if kind == "cd" else Decimal("222.75"))
            self.assertTrue(sheet.cell(row, 2).protection.locked)
            self.assertTrue(sheet.cell(row, 3).protection.locked)
            self.assertEqual(sheet.cell(row, 4).value, article)
            self.assertTrue(article.startswith("000"))
            self.assertEqual(sheet.cell(row, 5).value, source.base_price_rub)
            self.assertEqual(sheet.cell(row, 6).value, source.delivery_rub)
            self.assertEqual(sheet.cell(row, 7).value, source.stock_unit_cost)
            self.assertEqual(sheet.cell(row, 8).value, 0)
            self.assertFalse(sheet.cell(row, 8).protection.locked)
            self.assertTrue(sheet.cell(row, 7).protection.locked)
            self.assertEqual(sheet.cell(row, 9).value, f"=G{row}*H{row}")
        text = " ".join(str(cell.value) for ws in book for row in ws for cell in row)
        self.assertNotIn("Цена предоплата", text)
        self.assertNotIn("Цена постоплата", text)
        book.close()

    def test_current_stock_order_export_uses_order_costs_and_checks_access(self):
        order = self.order(qty=3, tech=2)
        edit_created_order(actor=self.user, order_id=order.pk, recipient="На склад", comment="После инвойса", lines=[
            {"price_list_item_id": self.cd_source.pk, "prepayment_quantity": 1, "postpayment_quantity": 0},
            {"price_list_item_id": self.tech_source.pk, "prepayment_quantity": 1, "postpayment_quantity": 0},
        ])
        url = reverse("orders:export", args=(order.pk,))
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content))
        sheet = book.active
        self.assertEqual(tuple(cell.value for cell in sheet[4]),
                         ("Наименование", "Количество", "Цена, ₽", "Сумма, ₽"))
        self.assertEqual(tuple(cell.value for cell in sheet[5]), (self.cd.name, 1, 2050, 2050))
        self.assertEqual(tuple(cell.value for cell in sheet[6]), (self.tech.name, 1, 3160, 3160))
        self.assertEqual(sheet.cell(sheet.max_row, 4).value, 5210)
        book.close()
        self.assertContains(self.client.get(reverse("orders:detail", args=(order.pk,))), url)
        employee = self.employee("orders.view_order_detail", "orders.view_orders")
        self.client.force_login(employee)
        self.assertEqual(self.client.get(url).status_code, 403)
        self.assertNotContains(self.client.get(reverse("orders:list")), url)
        employee.user_permissions.add(Permission.objects.get(content_type__app_label="price", codename="view_supplier_prices"))
        self.assertEqual(self.client.get(url).status_code, 200)

    def test_import_and_creation_ignore_prices_edited_in_file(self):
        def change_prices(book):
            for row in book[META_SHEET].iter_rows(values_only=True):
                if isinstance(row[0], int):
                    for col in (2, 3, 5, 6, 7, 9):
                        book[MAIN_SHEET].cell(row[0], col, 1)
        price, lines, purpose = import_procurement_order_files(
            {"stock": self.filled_stock_file(3, 2, change_prices)}, with_purpose=True)
        order = create_customer_procurement_order(actor=self.user, price_list_id=price.pk,
            recipient="", comment="", lines=lines, purpose=purpose)
        self.assertTrue(order.is_stock_order)
        self.assertEqual(order.recipient, "На склад")
        self.assertIsNone(order.payment_warehouse_id)
        self.assertIsNone(order.supply_id)
        self.assertEqual(order.items.get(cd=self.cd).prepayment_unit_price, Decimal("2050"))
        self.assertEqual(order.items.get(tech=self.tech).prepayment_unit_price, Decimal("3160"))
        self.assertEqual(order.grand_total, Decimal("12470"))

    def test_bad_quantities_empty_and_headers_rejected(self):
        for value in (1.5, -1, "=1+1", "3", 2147483648):
            with self.subTest(value=value), self.assertRaises(ValidationError):
                import_procurement_order_files({"stock": self.filled_stock_file(value)}, with_purpose=True)
        with self.assertRaises(ValidationError):
            import_procurement_order_files({"stock": self.filled_stock_file(0)}, with_purpose=True)
        with self.assertRaisesMessage(ValidationError, "заголовки"):
            import_procurement_order_files({"stock": self.filled_stock_file(mutate=lambda b: setattr(b[MAIN_SHEET]["H8"], "value", "Изменено"))}, with_purpose=True)

    def test_stock_info_is_current_when_downloaded_and_does_not_change_saved_purchase_cost(self):
        before = load_workbook(BytesIO(generate_procurement_stock_xlsx(price_list_id=self.price.pk, actor=self.user)))
        cd_row = next(row[0] for row in before[META_SHEET].iter_rows(values_only=True)
                      if isinstance(row[0], int) and row[1] == "cd")
        self.assertEqual(before[MAIN_SHEET].cell(cd_row, 2).value, 0)
        self.assertEqual(before[MAIN_SHEET].cell(cd_row, 3).value, 0)
        before.close()
        CDWarehouseStock.objects.create(cd=self.cd, warehouse=self.warehouse, quantity=7)
        self.cd.cost = Decimal("123.45")
        self.cd.save(update_fields=("cost",))
        after = load_workbook(BytesIO(generate_procurement_stock_xlsx(price_list_id=self.price.pk, actor=self.user)))
        self.assertEqual(after[MAIN_SHEET].cell(cd_row, 2).value, 7)
        self.assertEqual(Decimal(str(after[MAIN_SHEET].cell(cd_row, 3).value)), Decimal("123.45"))
        self.assertEqual(after[MAIN_SHEET].cell(cd_row, 7).value, Decimal("2050"))
        after.close()

    def test_legacy_stock_file_without_inventory_columns_still_imports(self):
        def old_layout(book):
            sheet = book[MAIN_SHEET]
            sheet.delete_cols(2, 2)
            for column, header in enumerate(LEGACY_STOCK_HEADERS, 1):
                sheet.cell(HEADER_ROW, column, header)
        price, lines, purpose = import_procurement_order_files(
            {"stock": self.filled_stock_file(3, 2, old_layout)}, with_purpose=True)
        self.assertEqual(price.pk, self.price.pk)
        self.assertEqual(purpose, "stock")
        self.assertEqual({line["price_list_item_id"]: line["prepayment_quantity"] for line in lines},
                         {self.cd_source.pk: 3, self.tech_source.pk: 2})

    def test_documents_cannot_be_confused_or_mixed(self):
        customer = generate_procurement_customer_xlsx(price_list_id=self.price.pk, actor=self.user)
        with self.assertRaisesMessage(ValidationError, "другого типа"):
            import_procurement_order_files({"stock": SimpleUploadedFile("customer.xlsx", customer)}, with_purpose=True)
        with self.assertRaisesMessage(ValidationError, "другого типа"):
            import_procurement_order_files({"cd": self.filled_stock_file()}, with_purpose=True)
        with self.assertRaisesMessage(ValidationError, "отдельно"):
            import_procurement_order_files({"stock": self.filled_stock_file(), "cd": self.filled_stock_file()}, with_purpose=True)
        with self.assertRaises(ValidationError):
            import_procurement_order_files({"stock": self.filled_stock_file()})

    def test_stock_row_reordering_and_missing_metadata_are_rejected(self):
        def damage(book, missing_meta):
            rows = [row for row in book[META_SHEET].iter_rows() if isinstance(row[0].value, int)]
            if missing_meta:
                book[META_SHEET].delete_rows(rows[-1][0].row)
            else:
                book[MAIN_SHEET].cell(rows[0][0].value, 4, "Другой артикул")
        for missing_meta in (True, False):
            with self.subTest(missing_meta=missing_meta), self.assertRaises(ValidationError):
                import_procurement_order_files({"stock": self.filled_stock_file(mutate=lambda b: damage(b, missing_meta))},
                                               with_purpose=True)

    def test_upload_preview_create_and_edit_stock_order(self):
        response = self.client.post(reverse("orders:upload"), {"stock_file": self.filled_stock_file(3, 2)})
        self.assertRedirects(response, reverse("orders:preview"))
        preview = self.client.get(reverse("orders:preview"))
        self.assertContains(preview, "Себестоимость с доставкой")
        self.assertNotContains(preview, "Цена постоплата")
        self.client.post(reverse("orders:preview"))
        order = CustomerProcurementOrder.objects.get()
        self.assertTrue(order.is_stock_order)
        self.assertContains(self.client.get(reverse("orders:detail", args=(order.pk,))), "Заказ на склад")
        self.assertContains(self.client.get(reverse("orders:edit", args=(order.pk,))), "Себестоимость с доставкой")
        edit_created_order(actor=self.user, order_id=order.pk, recipient="Наличие", comment="",
            lines=[{"price_list_item_id": self.cd_source.pk, "prepayment_quantity": 4, "postpayment_quantity": 0}])
        self.assertEqual(order.items.get().prepayment_unit_price, Decimal("2050"))

    def test_export_button_and_backend_cost_permission(self):
        self.assertContains(self.client.get(reverse("price:imports")), reverse("price:procurement_stock_export", args=(self.price.pk,)))
        self.assertEqual(self.client.get(reverse("price:procurement_stock_export", args=(self.price.pk,))).status_code, 200)
        employee = self.employee("price.view_price_page", "price.export_procurement_price_list", "orders.create_order", "orders.view_order_detail", "orders.view_orders")
        order = self.order()
        self.client.force_login(employee)
        self.assertNotContains(self.client.get(reverse("price:imports")), "Скачать прайс для заказа на склад")
        self.assertEqual(self.client.get(reverse("price:procurement_stock_export", args=(self.price.pk,))).status_code, 403)
        self.assertEqual(self.client.post(reverse("orders:upload"), {"stock_file": self.filled_stock_file()}).status_code, 403)
        self.assertEqual(self.client.get(reverse("orders:detail", args=(order.pk,))).status_code, 403)
        self.assertNotContains(self.client.get(reverse("orders:list")), reverse("orders:detail", args=(order.pk,)))

    def test_stock_confirmation_and_receiving_need_no_cash_permissions(self):
        order = self.order()
        employee = self.employee("price.view_supplier_prices", "orders.confirm_order", "orders.start_receiving", "orders.view_order_detail")
        self.client.force_login(employee)
        self.assertContains(self.client.get(reverse("orders:detail", args=(order.pk,))), "Подтвердить закупку")
        self.client.post(reverse("orders:confirm", args=(order.pk,)))
        self.client.post(reverse("orders:start_receiving", args=(order.pk,)))
        order.refresh_from_db()
        self.assertEqual(order.status, "receiving")
        self.assertIsNone(order.payment_warehouse_id)
        self.assertFalse(CashTransaction.objects.exists())

    def test_mixed_trip_m_to_f_and_changed_supplier_price_keep_stock_cost(self):
        stock, customer = self.order(3), self.order(2, purpose="customer")
        trip = create_trip(actor=self.user, order_ids=[stock.pk, customer.pk], creation_key=uuid4())
        line = trip.lines.get(supplier=self.f)
        self.assertEqual(line.quantity, 5)
        self.assertEqual(line.price_supplier_id, self.m.pk)
        change_trip_line(actor=self.user, batch_id=trip.pk, line_id=line.pk, revision=trip.revision,
            quantity=0, action="transfer", target_supplier_id=self.m.pk, price_supplier_id=self.f.pk)
        self.receiving(stock)
        stock = receive_stock_order(actor=self.user, order_id=stock.pk, warehouse_id=self.warehouse.pk)
        item = stock.supply.cd_items.get()
        self.assertEqual(item.quantity, 3)
        self.assertEqual(item.supplier_id, self.m.pk)
        self.assertEqual(item.purchase_unit_cost, Decimal("2050"))
        self.assertIsNone(CustomerProcurementOrder.objects.get(pk=customer.pk).supply_id)
        self.assertFalse(CashTransaction.objects.exists())

    def test_stock_order_move_keeps_original_cost_and_creates_receipt_once(self):
        stock, customer = self.order(3, 2), self.order(2, purpose="customer")
        source = create_trip(actor=self.user, order_ids=[stock.pk], creation_key=uuid4())
        target = create_trip(actor=self.user, order_ids=[customer.pk], creation_key=uuid4())
        total = stock.grand_total
        assign_order_to_trip(actor=self.user, order_id=stock.pk,
            target_batch_id=target.pk, target_revision=target.revision,
            source_batch_id=source.pk, source_revision=source.revision)
        stock.refresh_from_db()
        self.assertEqual(stock.trip_id, target.pk)
        self.assertEqual(stock.grand_total, total)
        self.assertFalse(source.lines.exists())
        self.receiving(stock)
        first = receive_stock_order(actor=self.user, order_id=stock.pk, warehouse_id=self.warehouse.pk)
        second = receive_stock_order(actor=self.user, order_id=stock.pk, warehouse_id=self.warehouse.pk)
        self.assertEqual(first.supply_id, second.supply_id)
        self.assertEqual(first.supply.cd_items.get().purchase_unit_cost, Decimal("2050"))
        self.assertEqual(first.supply.tech_items.get().purchase_unit_cost, Decimal("3160"))
        self.assertEqual(first.supply.cd_items.get().supplier_id, self.f.pk)
        self.assertFalse(CashTransaction.objects.exists())
        self.assertFalse(Sale.objects.exists())

    def test_stock_increase_in_trip_uses_single_quantity_bucket(self):
        order = self.order(3)
        trip = create_trip(actor=self.user, order_ids=[order.pk], creation_key=uuid4())
        line = trip.lines.get(supplier=self.f)
        allocation = line.allocations.get()
        change_trip_line(actor=self.user, batch_id=trip.pk, line_id=line.pk, revision=trip.revision,
            quantity=5, action="adjust", distribution=[{"allocation_id": allocation.pk, "final_quantity": 5}])
        item = order.items.get()
        self.assertEqual((item.prepayment_quantity, item.postpayment_quantity), (5, 0))
        self.assertEqual(item.prepayment_unit_price, Decimal("2050"))

    def test_stock_postpayment_bucket_cannot_be_submitted_directly(self):
        lines = [{"price_list_item_id": self.cd_source.pk, "prepayment_quantity": 1, "postpayment_quantity": 2}]
        with self.assertRaisesMessage(ValidationError, "единое количество"):
            create_customer_procurement_order(actor=self.user, price_list_id=self.price.pk, recipient="",
                                              comment="", lines=lines, purpose="stock")
        order = self.order()
        with self.assertRaisesMessage(ValidationError, "единое количество"):
            edit_created_order(actor=self.user, order_id=order.pk, recipient="На склад", comment="", lines=lines)
        self.assertEqual(order.items.get().quantity, 3)

    def test_receiving_reduction_needs_no_refund_and_receipt_has_actual_quantity(self):
        order = self.order(3)
        create_trip(actor=self.user, order_ids=[order.pk], creation_key=uuid4())
        self.receiving(order)
        item = order.items.get()
        adjustment = reduce_receiving_order(actor=self.user, order_id=order.pk, comment="Приехало меньше",
            allow_refund=False, rows=[{"item_id": item.pk, "prepayment_quantity": 1, "postpayment_quantity": 0}])
        self.assertEqual(adjustment.refund_amount, 0)
        result = receive_stock_order(actor=self.user, order_id=order.pk, warehouse_id=self.warehouse.pk)
        self.assertEqual(result.supply.cd_items.get().quantity, 1)
        self.assertEqual(result.supply.grand_total, Decimal("2050"))
        self.assertFalse(CashTransaction.objects.exists())

    def test_receipt_warehouse_required_and_wrong_status_roll_back(self):
        order = self.order()
        with self.assertRaises(ValidationError):
            receive_stock_order(actor=self.user, order_id=order.pk, warehouse_id=self.warehouse.pk)
        self.receiving(order)
        for warehouse_id in (None, "", "abc", 99999):
            with self.subTest(warehouse=warehouse_id), self.assertRaises(ValidationError):
                receive_stock_order(actor=self.user, order_id=order.pk, warehouse_id=warehouse_id)
        order.refresh_from_db()
        self.assertIsNone(order.supply_id)
        self.assertEqual(order.status, "receiving")
        self.assertFalse(Supply.objects.exists())

    def test_receipt_is_idempotent_and_stock_arrives_only_once_after_price_review(self):
        order = self.order(3, 2)
        self.receiving(order)
        result = receive_stock_order(actor=self.user, order_id=order.pk, warehouse_id=self.other_warehouse.pk)
        supply = result.supply
        self.assertEqual(supply.warehouse_id, self.other_warehouse.pk)
        self.assertEqual(supply.status, Supply.Status.PRICE_REVIEW_REQUIRED)
        self.assertEqual(supply.grand_total, Decimal("12470"))
        self.assertFalse(CDWarehouseStock.objects.filter(cd=self.cd, quantity__gt=0).exists())
        again = advance_order_status(actor=self.user, order_id=order.pk, next_status="completed", warehouse_id=self.warehouse.pk)
        self.assertEqual(again.supply_id, supply.pk)
        self.assertEqual(Supply.objects.count(), 1)
        confirm_supply_price_review(actor=self.user, supply_id=supply.pk)
        confirm_supply_price_review(actor=self.user, supply_id=supply.pk)
        self.assertEqual(CDWarehouseStock.objects.get(cd=self.cd, warehouse=self.other_warehouse).quantity, 3)
        self.assertEqual(TechWarehouseStock.objects.get(tech=self.tech, warehouse=self.other_warehouse).quantity, 2)
        self.cd.refresh_from_db()
        self.assertEqual(self.cd.cost, Decimal("2050"))
        self.assertFalse(Sale.objects.exists())
        self.assertFalse(CashTransaction.objects.exists())
        self.assertTrue(all(balance == 0 for balance in CashRegister.objects.values_list("balance", flat=True)))

    def test_receipt_permission_is_checked_in_service_and_post(self):
        order = self.order()
        self.receiving(order)
        employee = self.employee("price.view_supplier_prices", "orders.complete_order")
        with self.assertRaises(PermissionDenied):
            receive_stock_order(actor=employee, order_id=order.pk, warehouse_id=self.warehouse.pk)
        self.client.force_login(employee)
        self.assertEqual(self.client.post(reverse("orders:complete", args=(order.pk,)), {"warehouse": self.warehouse.pk}).status_code, 403)
        self.assertFalse(Supply.objects.exists())

    def test_completion_post_redirects_to_existing_receipt_review(self):
        order = self.order()
        self.receiving(order)
        url = reverse("orders:complete", args=(order.pk,))
        response = self.client.post(url, {"warehouse": self.warehouse.pk})
        order.refresh_from_db()
        self.assertRedirects(response, reverse("supplies:price_review", args=(order.supply_id,)))
        self.assertRedirects(self.client.post(url, {"warehouse": self.warehouse.pk}), response.url)
        self.assertEqual(Supply.objects.count(), 1)
        detail = self.client.get(reverse("orders:detail", args=(order.pk,)))
        self.assertContains(detail, "Проверить цены и поставить на баланс")
        self.assertNotContains(detail, "Предоплата")
